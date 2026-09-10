"""读取已经完成打标的内容工作表。

该读取器只面向代表性内容筛选使用，不改变正式 Excel Import Profile。
"""

from __future__ import annotations

from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from aima_ugc.modules.analysis.representative_selection import TARGET_PLATFORMS, LabeledContent

CONTENT_SHEET_NAME = "内容"
REAL_USER_VOICE_TYPE = "真实用户发声"

_REQUIRED_HEADERS = ("平台", "内容ID", "标题", "正文", "作者")
_OPTIONAL_HEADERS = ("发布时间", "内容链接", "发声类型", "情感标签")


class LabeledContentReaderError(ValueError):
    """打标结果工作表不满足读取契约。"""


@dataclass(frozen=True, slots=True)
class LabeledContentReadSummary:
    """Excel 读取、去重和基础过滤的可观察计数。"""

    input_path: Path
    sheet_name: str
    rows_seen: int
    rows_read: int
    duplicate_rows: int
    blank_rows: int
    input_paths: tuple[Path, ...] = ()

    def __post_init__(self) -> None:
        """保持单文件调用的兼容性，同时暴露批量输入文件。"""

        if not self.input_paths:
            object.__setattr__(self, "input_paths", (self.input_path,))


def iter_labeled_contents(
    input_path: Path,
    *,
    sheet_name: str = CONTENT_SHEET_NAME,
) -> Iterator[LabeledContent]:
    """逐行读取 ``内容`` Sheet，并在输入边界完成表头和稳定 ID 校验。"""

    path = Path(input_path)
    if path.suffix.casefold() != ".xlsx":
        raise LabeledContentReaderError(f"仅支持 .xlsx 文件: {path}")
    if not path.is_file():
        raise FileNotFoundError(path)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise LabeledContentReaderError(
                f"工作簿缺少指定工作表: {sheet_name}；实际工作表: {', '.join(workbook.sheetnames)}"
            )
        worksheet = workbook[sheet_name]
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise LabeledContentReaderError(f"工作表为空: {sheet_name}") from exc

        headers = tuple(_header_name(value) for value in raw_headers)
        _validate_headers(headers, sheet_name=sheet_name)
        header_indexes = {
            header: index for index, header in enumerate(headers) if header is not None
        }

        seen_keys: set[tuple[str, str]] = set()
        for row_number, values in enumerate(rows, start=2):
            if _row_is_blank(values):
                continue
            row = _row_values(values, header_indexes)
            platform = _normalize_platform(row.get("平台"))
            content_id = _cell_text(row.get("内容ID"))
            if not platform:
                raise LabeledContentReaderError(f"第 {row_number} 行的平台不能为空")
            if not content_id:
                raise LabeledContentReaderError(f"第 {row_number} 行的内容ID不能为空")

            record = LabeledContent(
                row_number=row_number,
                platform=platform,
                content_id=content_id,
                title=_cell_text(row.get("标题")),
                text=_cell_text(row.get("正文")),
                author=_cell_text(row.get("作者")),
                published_at=_cell_text(row.get("发布时间")),
                content_url=_cell_text(row.get("内容链接")),
                voice_type=_cell_text(row.get("发声类型")),
                sentiment_label=_cell_text(row.get("情感标签")),
            )
            if record.deduplication_key in seen_keys:
                continue
            seen_keys.add(record.deduplication_key)
            yield record
    finally:
        workbook.close()


def read_labeled_contents(
    input_path: Path,
    *,
    sheet_name: str = CONTENT_SHEET_NAME,
) -> tuple[tuple[LabeledContent, ...], LabeledContentReadSummary]:
    """读取并按平台 + 内容ID 去重，返回记录和读取摘要。"""

    records: list[LabeledContent] = []
    rows_seen = 0
    duplicate_rows = 0
    blank_rows = 0

    path = Path(input_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise LabeledContentReaderError(f"工作簿缺少指定工作表: {sheet_name}")
        worksheet = workbook[sheet_name]
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise LabeledContentReaderError(f"工作表为空: {sheet_name}") from exc
        headers = tuple(_header_name(value) for value in raw_headers)
        _validate_headers(headers, sheet_name=sheet_name)
        header_indexes = {
            header: index for index, header in enumerate(headers) if header is not None
        }
        seen_keys: set[tuple[str, str]] = set()

        for row_number, values in enumerate(rows, start=2):
            if _row_is_blank(values):
                blank_rows += 1
                continue
            rows_seen += 1
            row = _row_values(values, header_indexes)
            platform = _normalize_platform(row.get("平台"))
            content_id = _cell_text(row.get("内容ID"))
            if not platform:
                raise LabeledContentReaderError(f"第 {row_number} 行的平台不能为空")
            if not content_id:
                raise LabeledContentReaderError(f"第 {row_number} 行的内容ID不能为空")
            record = LabeledContent(
                row_number=row_number,
                platform=platform,
                content_id=content_id,
                title=_cell_text(row.get("标题")),
                text=_cell_text(row.get("正文")),
                author=_cell_text(row.get("作者")),
                published_at=_cell_text(row.get("发布时间")),
                content_url=_cell_text(row.get("内容链接")),
                voice_type=_cell_text(row.get("发声类型")),
                sentiment_label=_cell_text(row.get("情感标签")),
            )
            if record.deduplication_key in seen_keys:
                duplicate_rows += 1
                continue
            seen_keys.add(record.deduplication_key)
            records.append(record)
    finally:
        workbook.close()

    return tuple(records), LabeledContentReadSummary(
        input_path=path,
        sheet_name=sheet_name,
        rows_seen=rows_seen,
        rows_read=len(records),
        duplicate_rows=duplicate_rows,
        blank_rows=blank_rows,
        input_paths=(path,),
    )


def read_labeled_content_files(
    input_paths: Sequence[Path],
    *,
    sheet_name: str = CONTENT_SHEET_NAME,
) -> tuple[tuple[LabeledContent, ...], LabeledContentReadSummary]:
    """读取多个已打标工作簿，并按平台 + 内容ID做跨文件去重。"""

    paths = tuple(Path(path) for path in input_paths)
    if not paths:
        raise LabeledContentReaderError("至少需要提供一个已打标 XLSX 文件")

    records: list[LabeledContent] = []
    seen_keys: set[tuple[str, str]] = set()
    rows_seen = 0
    duplicate_rows = 0
    blank_rows = 0

    for path in paths:
        file_records, file_summary = read_labeled_contents(path, sheet_name=sheet_name)
        rows_seen += file_summary.rows_seen
        duplicate_rows += file_summary.duplicate_rows
        blank_rows += file_summary.blank_rows
        for record in file_records:
            if record.deduplication_key in seen_keys:
                duplicate_rows += 1
                continue
            seen_keys.add(record.deduplication_key)
            records.append(record)

    return tuple(records), LabeledContentReadSummary(
        input_path=paths[0],
        input_paths=paths,
        sheet_name=sheet_name,
        rows_seen=rows_seen,
        rows_read=len(records),
        duplicate_rows=duplicate_rows,
        blank_rows=blank_rows,
    )


def _validate_headers(headers: tuple[str | None, ...], *, sheet_name: str) -> None:
    non_empty = tuple(header for header in headers if header is not None)
    duplicates = [header for header in non_empty if non_empty.count(header) > 1]
    if duplicates:
        duplicate_names = ", ".join(dict.fromkeys(duplicates))
        raise LabeledContentReaderError(f"工作表 {sheet_name} 存在重复表头: {duplicate_names}")
    missing = [header for header in _REQUIRED_HEADERS if header not in non_empty]
    if missing:
        raise LabeledContentReaderError(f"工作表 {sheet_name} 缺少必需字段: {', '.join(missing)}")


def _row_values(values: tuple[object, ...], indexes: dict[str, int]) -> dict[str, object]:
    return {
        header: values[index] if index < len(values) else None
        for header, index in indexes.items()
        if header is not None
    }


def _header_name(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).lstrip("\ufeff").strip()
    return text or None


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_platform(value: object) -> str:
    text = _cell_text(value)
    compact = "".join(text.split())
    if compact.casefold() == "douyin" or compact == "抖音":
        return "抖音"
    if compact.casefold() == "xiaohongshu" or compact == "小红书":
        return "小红书"
    return text


def _row_is_blank(values: tuple[object, ...]) -> bool:
    return all(value is None or (isinstance(value, str) and not value.strip()) for value in values)


__all__ = [
    "CONTENT_SHEET_NAME",
    "REAL_USER_VOICE_TYPE",
    "TARGET_PLATFORMS",
    "LabeledContent",
    "LabeledContentReadSummary",
    "LabeledContentReaderError",
    "iter_labeled_contents",
    "read_labeled_content_files",
    "read_labeled_contents",
]
