"""把报告图表规格写成可由飞书导入的原生 XLSX 图表工作簿。"""

from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .chart_spec import ChartSpec


@dataclass(frozen=True, slots=True)
class ChartWorkbookSummary:
    """伴随图表工作簿的可观察结果。"""

    output_path: Path
    chart_count: int
    sheet_count: int


def build_editable_chart_workbook(
    specs: tuple[ChartSpec, ...],
    output_path: Path,
) -> ChartWorkbookSummary:
    """为每个报告图表创建一个含数据和 Office 原生 Chart 的 Sheet。"""

    target = Path(output_path)
    if target.suffix.lower() != ".xlsx":
        raise ValueError("可编辑图表工作簿输出必须是 .xlsx 文件")
    if not specs:
        raise ValueError("可编辑图表工作簿至少需要一个图表")

    workbook = Workbook()
    active_sheet = workbook.active
    if active_sheet is None:
        raise OSError("新建图表工作簿缺少默认 Sheet")
    workbook.remove(active_sheet)
    for index, spec in enumerate(specs, start=1):
        _validate_spec(spec)
        sheet = workbook.create_sheet(f"图表{index:02d}")
        _write_chart_sheet(sheet, spec)

    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f".{target.stem}.tmp.xlsx")
    temporary.unlink(missing_ok=True)
    try:
        workbook.save(temporary)
        _verify_chart_workbook(temporary, expected_charts=len(specs))
        os.replace(temporary, target)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return ChartWorkbookSummary(
        output_path=target,
        chart_count=len(specs),
        sheet_count=len(specs),
    )


def _write_chart_sheet(sheet: Worksheet, spec: ChartSpec) -> None:
    sheet.cell(row=1, column=1, value="分类")
    series_names = spec.series_names or tuple(
        f"系列{index}" for index in range(1, len(spec.series) + 1)
    )
    for column, name in enumerate(series_names, start=2):
        sheet.cell(row=1, column=column, value=name)
    for row, category in enumerate(spec.categories, start=2):
        sheet.cell(row=row, column=1, value=category)
        for column, values in enumerate(spec.series, start=2):
            sheet.cell(row=row, column=column, value=values[row - 2])

    max_row = len(spec.categories) + 1
    max_column = len(spec.series) + 1
    chart: PieChart | BarChart | LineChart
    if spec.kind == "pie":
        chart = PieChart()
        data = Reference(sheet, min_col=2, min_row=1, max_row=max_row)
        labels = Reference(sheet, min_col=1, min_row=2, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
    elif spec.kind == "bar":
        chart = BarChart()
        chart.type = "bar" if spec.bar_direction == "bar" else "col"
        data = Reference(sheet, min_col=2, max_col=max_column, min_row=1, max_row=max_row)
        categories = Reference(sheet, min_col=1, min_row=2, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.y_axis.scaling.min = spec.y_min
        chart.y_axis.scaling.max = spec.y_max
    else:
        chart = LineChart()
        data = Reference(sheet, min_col=2, max_col=max_column, min_row=1, max_row=max_row)
        categories = Reference(sheet, min_col=1, min_row=2, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.y_axis.scaling.min = spec.y_min
        chart.y_axis.scaling.max = spec.y_max
    chart.title = spec.title
    chart.height = 9
    chart.width = 18
    sheet.add_chart(chart, "D2")

    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 24
    for column in range(2, max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 16


def _validate_spec(spec: ChartSpec) -> None:
    if spec.kind not in {"bar", "line", "pie"}:
        raise ValueError(f"不支持的图表类型: {spec.kind}")
    if not spec.categories or not spec.series:
        raise ValueError("图表缺少分类或数据序列")
    if any(len(values) != len(spec.categories) for values in spec.series):
        raise ValueError("图表数据序列长度与分类数量不一致")
    if spec.series_names and len(spec.series_names) != len(spec.series):
        raise ValueError("图表系列名称数量与数据序列数量不一致")


def _verify_chart_workbook(path: Path, *, expected_charts: int) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    actual_charts = 0
    for sheet in workbook.worksheets:
        actual_charts += len(getattr(sheet, "_charts", ()))
    if actual_charts != expected_charts:
        raise OSError("可编辑图表工作簿的 Office Chart 数量校验失败")
