"""在专用空数据库上用正式导入与 Replay 路径测量阶段吞吐。"""

from __future__ import annotations

import argparse
import json
import logging
import re
import shutil
import time
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from aima_ugc.adapters.persistence.postgres.content_complete import (
    PostgresCompleteContentRepository,
    PostgresCompleteExistingContentBatchItem,
)
from aima_ugc.adapters.persistence.postgres.content_contributions import (
    ContentContributionSnapshot,
    capture_content_contribution_snapshots_batch,
)
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.brand_vehicle_http import PostgresBrandVehicleHttpService
from aima_ugc.bootstrap.canonical_replay_http import PostgresCanonicalReplayHttpService
from aima_ugc.bootstrap.canonical_replay_worker import PostgresCanonicalReplayJobExecutor
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.brand_vehicle import BrandAliasCreateRequest, BrandCreateRequest
from aima_ugc.contracts.canonical import CanonicalAuthorV1, CanonicalContentV1
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.modules.identity import Principal
from aima_ugc.modules.ingestion.canonical_replay_tables import (
    canonical_replay_content_changes_table,
    canonical_replay_runs_table,
)
from aima_ugc.modules.vehicles.tables import vehicle_brands_table
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs.tables import jobs_table
from aima_ugc.platform.storage.tables import artifacts_table
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import event, func, select

_DATABASE_SUFFIX = "_canonical_replay_capacity"
_DEFAULT_DISK_BUDGET_BYTES = 512 * 1024 * 1024
_FREE_SPACE_RESERVE_BYTES = 64 * 1024 * 1024


def _require_capacity_database(name: str) -> None:
    """拒绝任何非专用容量数据库，避免基准请求全量重筛真实数据。"""

    if not name.endswith(_DATABASE_SUFFIX):
        raise ValueError(f"Replay 容量基准仅允许数据库名以 {_DATABASE_SUFFIX} 结尾")


def _fixture_xlsx(
    *,
    file_index: int,
    rows_per_file: int,
    existing_rows_per_file: int,
    nonce: str,
) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("文章")
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    for row_index in range(rows_per_file):
        external_id = f"replay-bench-{nonce}-{file_index}-{row_index}"
        keyword = "预置" if row_index < existing_rows_per_file else "星曜"
        sheet.append(
            [
                "小红书",
                f"{keyword}容量样本 {file_index}-{row_index}",
                "容量测试正文",
                "容量测试账号",
                "2026-09-11 08:00:00",
                f"https://www.xiaohongshu.com/explore/{external_id}",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def run_benchmark(
    *,
    work_dir: Path,
    file_count: int,
    rows_per_file: int,
    workers: int,
    existing_rows_per_file: int = 0,
    stable_authors: bool = False,
    scalar_stable_authors: bool = False,
    disk_budget_bytes: int = _DEFAULT_DISK_BUDGET_BYTES,
) -> dict[str, object]:
    """只测 Replay 执行窗口；输入生成与初次导入不计时。"""

    if not 1 <= file_count <= 1000:
        raise ValueError("file_count 必须在 1 到 1000 之间")
    if not 1 <= rows_per_file <= 10_000:
        raise ValueError("rows_per_file 必须在 1 到 10000 之间")
    if not 1 <= workers <= 8:
        raise ValueError("workers 必须在 1 到 8 之间")
    if not 0 <= existing_rows_per_file <= rows_per_file:
        raise ValueError("existing_rows_per_file 必须在 0 到 rows_per_file 之间")
    if scalar_stable_authors and not stable_authors:
        raise ValueError("scalar_stable_authors 要求同时启用 stable_authors")
    if disk_budget_bytes <= 0:
        raise ValueError("disk_budget_bytes 必须大于 0")
    settings = load_settings()
    _require_capacity_database(settings.db_name)
    root = work_dir.resolve()
    root.mkdir(parents=True, exist_ok=True)
    if any(root.iterdir()):
        raise ValueError("Replay 容量基准工作目录必须为空")
    estimated_bytes = file_count * (4 * 1024 * 1024 + rows_per_file * 4096)
    if estimated_bytes > disk_budget_bytes:
        raise ValueError("Replay 容量基准预计临时数据超过磁盘预算")
    free_bytes = shutil.disk_usage(root).free
    if estimated_bytes + _FREE_SPACE_RESERVE_BYTES > free_bytes:
        raise ValueError("Replay 容量基准预计临时数据超过当前可用磁盘")

    try:
        runtime = create_worker_runtime(
            settings=settings.model_copy(
                update={"data_dir": root / "data", "log_dir": root / "logs"}
            )
        )
    except BaseException:
        _cleanup_generated_outputs(root)
        raise
    completed = False
    try:
        runtime.logger.setLevel(logging.WARNING)
        _reset_capacity_database(runtime)
        with runtime.database.engine.connect() as connection:
            for table in (artifacts_table, jobs_table, vehicle_brands_table, contents_table):
                if connection.scalar(select(func.count()).select_from(table)):
                    raise ValueError("Replay 容量基准数据库必须没有业务数据")

        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        principal = Principal(
            principal_id="canonical-replay-capacity",
            display_name="Replay 容量基准",
            role="administrator",
            source="development",
        )
        brand = PostgresBrandVehicleHttpService(runtime).create_brand(
            BrandCreateRequest(
                display_name=f"Replay 容量品牌 {uuid4().hex[:8]}",
                role="owned",
                aliases=(("预置",) if existing_rows_per_file else ("导入阶段不会命中的品牌词",)),
            ),
            principal=principal,
            request_id="canonical-replay-capacity-brand",
        )
        registry = create_collection_job_registry(runtime=runtime)

        def run_one_job(index: int) -> bool:
            return create_job_worker(
                runtime=runtime,
                registry=registry,
                worker_id=f"canonical-replay-capacity-{index}",
                lease_seconds=120,
                retry_delay_seconds=0,
            ).run_once()

        nonce = uuid4().hex
        for file_index in range(file_count):
            created = client.post(
                "/api/v1/import-batches",
                files=[
                    (
                        "file",
                        (
                            f"capacity-{file_index}.xlsx",
                            _fixture_xlsx(
                                file_index=file_index,
                                rows_per_file=rows_per_file,
                                existing_rows_per_file=existing_rows_per_file,
                                nonce=nonce,
                            ),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        ),
                    ),
                    ("brand_ids", (None, str(brand.id), None)),
                ],
            )
            if created.status_code != 202 or not run_one_job(file_index):
                raise RuntimeError("基准输入导入失败")
            detail = client.get(f"/api/v1/import-batches/{created.json()['batch_id']}")
            if detail.status_code != 200 or detail.json()["status"] != "succeeded":
                raise RuntimeError("基准输入未成功产生 Canonical")
        PostgresBrandVehicleHttpService(runtime).add_alias(
            brand.id,
            BrandAliasCreateRequest(text="星曜"),
            principal=principal,
            request_id="canonical-replay-capacity-alias",
        )
        # 容量基准需要保留每个原子批次的阶段耗时，便于区分解析、Content、
        # Evidence 与账本/checkpoint 瓶颈；正式服务仍由部署日志级别控制。
        runtime.logger.setLevel(logging.DEBUG)
        for handler in runtime.logger.handlers:
            handler.setLevel(logging.DEBUG)
        created = client.post(
            "/api/v1/canonical-replays/all",
            json={"idempotency_key": f"replay-capacity-{nonce}"},
        )
        if created.status_code != 202:
            raise RuntimeError("基准全量 Replay 创建失败")
        request_id = UUID(created.json()["request_id"])
        run_count = int(created.json()["run_count"])
        statement_count = 0
        seen_inserts = 0
        ledger_inserts = 0
        statements_by_table: dict[str, int] = {}
        statement_examples: dict[str, str] = {}

        def count_sql(
            connection: object,
            cursor: object,
            statement: str,
            parameters: object,
            context: object,
            executemany: bool,
        ) -> None:
            nonlocal statement_count, seen_inserts, ledger_inserts
            del connection, cursor, parameters, context, executemany
            statement_count += 1
            first_table = re.search(r"\b(?:FROM|INTO|UPDATE)\s+([a-z_][a-z_0-9]*)", statement)
            sql_kind = statement.lstrip().split(None, 1)[0].upper()
            table_key = f"{sql_kind} {first_table.group(1) if first_table else 'other'}"
            statements_by_table[table_key] = statements_by_table.get(table_key, 0) + 1
            statement_examples.setdefault(table_key, " ".join(statement.split())[:500])
            if statement.startswith("INSERT INTO canonical_replay_seen_content"):
                seen_inserts += 1
            if statement.startswith("INSERT INTO canonical_replay_content_changes"):
                ledger_inserts += 1

        original_lineage = PostgresCanonicalReplayJobExecutor._content_with_lineage
        lineage_method_name = "_content_with_lineage"
        content_batch_method_name = "ingest_contents_with_before_snapshots_batch"
        original_content_batch = (
            PostgresCompleteContentRepository.ingest_contents_with_before_snapshots_batch
        )

        def inject_stable_author(
            session: object,
            content: CanonicalContentV1,
            **kwargs: object,
        ) -> CanonicalContentV1:
            """在正式 lineage 完成后为容量样本注入 TikHub 形态的稳定作者。"""

            mapped = original_lineage(session, content, **kwargs)  # type: ignore[arg-type]
            identity = mapped.external_content_id
            return mapped.model_copy(
                update={
                    "author": CanonicalAuthorV1(
                        external_account_id=f"capacity-author-{identity}",
                        alternate_ids={"red_id": f"capacity-red-{identity}"},
                        display_name=f"容量作者 {identity}",
                    ),
                    "observed_fields": [
                        *mapped.observed_fields,
                        "author.external_account_id",
                        "author.alternate_ids",
                        "author.display_name",
                    ],
                }
            )

        if stable_authors:
            setattr(
                PostgresCanonicalReplayJobExecutor,
                lineage_method_name,
                staticmethod(inject_stable_author),
            )

        def ingest_stable_authors_one_by_one(
            repository: PostgresCompleteContentRepository,
            entries: tuple[tuple[CanonicalContentV1, ContentContributionSnapshot], ...],
        ) -> tuple[PostgresCompleteExistingContentBatchItem, ...]:
            """仅供同机 A/B 复现旧兼容路径，不改变正式 Worker 的默认行为。"""

            completed: list[PostgresCompleteExistingContentBatchItem] = []
            for observation, before in entries:
                result = repository._ingest_content(  # noqa: SLF001
                    observation,
                    before_snapshot=before,
                )
                after = result.contribution_after
                if after is None:
                    after = capture_content_contribution_snapshots_batch(
                        repository._session,  # noqa: SLF001
                        ((observation, result.target_id),),
                    )[0]
                completed.append(
                    PostgresCompleteExistingContentBatchItem(
                        observation=observation,
                        before=before,
                        result=result,
                        contribution_after=after,
                        used_scalar_fallback=True,
                    )
                )
            return tuple(completed)

        if scalar_stable_authors:
            setattr(
                PostgresCompleteContentRepository,
                content_batch_method_name,
                ingest_stable_authors_one_by_one,
            )
        event.listen(runtime.database.engine, "before_cursor_execute", count_sql)
        started = time.perf_counter()
        try:
            with ThreadPoolExecutor(max_workers=workers) as executor:
                if not all(executor.map(run_one_job, range(file_count, file_count + run_count))):
                    raise RuntimeError("基准 Replay Job 未被领取")
        finally:
            elapsed = time.perf_counter() - started
            event.remove(runtime.database.engine, "before_cursor_execute", count_sql)
            setattr(
                PostgresCanonicalReplayJobExecutor,
                lineage_method_name,
                staticmethod(original_lineage),
            )
            setattr(
                PostgresCompleteContentRepository,
                content_batch_method_name,
                original_content_batch,
            )

        with runtime.database.engine.connect() as connection:
            statuses = (
                connection.execute(
                    select(jobs_table.c.status)
                    .select_from(
                        canonical_replay_runs_table.join(
                            jobs_table, canonical_replay_runs_table.c.job_id == jobs_table.c.id
                        )
                    )
                    .where(canonical_replay_runs_table.c.all_request_id == request_id)
                )
                .scalars()
                .all()
            )
            counters = connection.execute(
                select(
                    func.coalesce(func.sum(canonical_replay_runs_table.c.rows_seen), 0).label(
                        "rows_seen"
                    ),
                    func.coalesce(func.sum(canonical_replay_runs_table.c.rows_matched), 0).label(
                        "rows_matched"
                    ),
                    func.coalesce(func.sum(canonical_replay_runs_table.c.rows_ingested), 0).label(
                        "rows_ingested"
                    ),
                    func.coalesce(
                        func.sum(canonical_replay_runs_table.c.existing_convergence),
                        0,
                    ).label("existing_convergence"),
                ).where(canonical_replay_runs_table.c.all_request_id == request_id)
            ).one()
            ledger_count = connection.scalar(
                select(func.count())
                .select_from(canonical_replay_content_changes_table)
                .where(canonical_replay_content_changes_table.c.all_request_id == request_id)
            )
        if len(statuses) != run_count or any(status != "succeeded" for status in statuses):
            raise RuntimeError("基准 Replay 子任务未全部成功")
        expected_rows = file_count * rows_per_file
        expected_existing = file_count * existing_rows_per_file
        expected_ingested = expected_rows - expected_existing
        if ledger_count != expected_rows:
            raise RuntimeError("基准 Replay 贡献账本行数与输入不一致")
        if (
            int(counters.rows_seen) != expected_rows
            or int(counters.rows_matched) != expected_rows
            or int(counters.rows_ingested) != expected_ingested
            or int(counters.existing_convergence) != expected_existing
        ):
            raise RuntimeError("基准 Replay 持久计数与新增/既有分布不一致")
        report = {
            "schema_version": "canonical-replay-capacity.v1",
            "files": file_count,
            "rows": expected_rows,
            "existing_rows": expected_existing,
            "new_rows": expected_ingested,
            "rows_seen": int(counters.rows_seen),
            "rows_matched": int(counters.rows_matched),
            "rows_ingested": int(counters.rows_ingested),
            "existing_convergence": int(counters.existing_convergence),
            "workers": workers,
            "stable_authors": stable_authors,
            "scalar_stable_authors": scalar_stable_authors,
            "replay_runs": run_count,
            "elapsed_seconds": round(elapsed, 3),
            "rows_per_second": round(expected_rows / elapsed, 2),
            "sql_statements": statement_count,
            "seen_identity_inserts": seen_inserts,
            "contribution_ledger_inserts": ledger_inserts,
            "disk_budget_bytes": disk_budget_bytes,
            "estimated_bytes": estimated_bytes,
            "work_directory_bytes": _directory_bytes(root),
            "top_sql_tables": sorted(
                statements_by_table.items(), key=lambda item: item[1], reverse=True
            )[:15],
            "top_sql_examples": {
                key: statement_examples[key]
                for key, _ in sorted(
                    statements_by_table.items(),
                    key=lambda item: item[1],
                    reverse=True,
                )[:15]
            },
        }
        _atomic_write_json(root / "capacity_report.json", report)
        completed = True
        return report
    finally:
        try:
            runtime.close()
        finally:
            if not completed:
                _cleanup_generated_outputs(root)


def _reset_capacity_database(runtime: Any) -> None:
    """只清空后缀已校验的专用容量库，保留 Migration seed。"""

    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, accounts "
            "RESTART IDENTITY CASCADE"
        )


def _directory_bytes(root: Path) -> int:
    """统计本次专用工作目录大小；无法读取的瞬态文件按零处理。"""

    total = 0
    for path in root.rglob("*"):
        try:
            if path.is_file():
                total += path.stat().st_size
        except OSError:
            continue
    return total


def _cleanup_generated_outputs(root: Path) -> None:
    """异常时只清理由本次空目录基准创建的已知直属输出。"""

    resolved_root = root.resolve()
    for name in (
        "data",
        "logs",
        "capacity_report.json",
        "capacity_report.json.tmp",
    ):
        candidate = (resolved_root / name).resolve()
        if candidate.parent != resolved_root:
            raise RuntimeError("Replay 容量基准清理目标越出工作目录")
        if candidate.is_dir():
            shutil.rmtree(candidate)
        else:
            candidate.unlink(missing_ok=True)


def _atomic_write_json(path: Path, payload: dict[str, object]) -> None:
    """原子写入容量报告，避免中途中断留下半份 JSON。"""

    temporary = path.with_suffix(f"{path.suffix}.tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="专用空数据库上的 Canonical Replay 容量基准")
    parser.add_argument("--work-dir", type=Path, required=True)
    parser.add_argument("--files", type=int, default=1)
    parser.add_argument("--rows-per-file", type=int, default=100)
    parser.add_argument("--workers", type=int, default=1)
    parser.add_argument("--existing-rows-per-file", type=int, default=0)
    parser.add_argument("--stable-authors", action="store_true")
    parser.add_argument("--scalar-stable-authors", action="store_true")
    parser.add_argument("--disk-budget-mib", type=int, default=512)
    args = parser.parse_args()
    print(
        json.dumps(
            run_benchmark(
                work_dir=args.work_dir,
                file_count=args.files,
                rows_per_file=args.rows_per_file,
                workers=args.workers,
                existing_rows_per_file=args.existing_rows_per_file,
                stable_authors=args.stable_authors,
                scalar_stable_authors=args.scalar_stable_authors,
                disk_budget_bytes=args.disk_budget_mib * 1024 * 1024,
            ),
            ensure_ascii=False,
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
