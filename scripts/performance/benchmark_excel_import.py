"""本地 Excel 正式导入链路的有界容量基准；只允许专用空数据库。"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import statistics
import tempfile
import threading
import time
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.brand_vehicle_http import PostgresBrandVehicleHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.brand_vehicle import BrandCreateRequest
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.modules.identity import Principal
from aima_ugc.platform.config import load_settings
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import event, func, select

_DATABASE_SUFFIX = "_import_pipeline_capacity"
_DEFAULT_DISK_BUDGET_BYTES = 512 * 1024 * 1024
_FREE_SPACE_RESERVE_BYTES = 128 * 1024 * 1024
_TEMPORARY_BYTES_PATTERN = re.compile(r"\btemporary_bytes=(\d+)\b")
_TEMPORARY_PEAK_BYTES_PATTERN = re.compile(r"\btemporary_peak_bytes=(\d+)\b")


def _require_capacity_database(name: str) -> None:
    if not name.endswith(_DATABASE_SUFFIX):
        raise ValueError(f"Excel 容量基准只允许数据库名以 {_DATABASE_SUFFIX} 结尾")


def run_benchmark(
    *,
    work_dir: Path,
    row_count: int,
    runs: int,
    disk_budget_bytes: int = _DEFAULT_DISK_BUDGET_BYTES,
) -> dict[str, Any]:
    """运行至少一次完整 HTTP 上传 + Worker 导入，并输出可比较样本。"""

    if not 1 <= row_count <= 100_000:
        raise ValueError("rows 必须在 1 到 100000 之间")
    if not 1 <= runs <= 10:
        raise ValueError("runs 必须在 1 到 10 之间")
    if disk_budget_bytes <= 0:
        raise ValueError("disk_budget_bytes 必须为正数")
    root = work_dir.resolve()
    root.mkdir(parents=True, exist_ok=True)
    if any(root.iterdir()):
        raise ValueError("Excel 导入容量基准工作目录必须为空")
    estimated_bytes = 64 * 1024 * 1024 + row_count * 4096
    if estimated_bytes > disk_budget_bytes:
        raise ValueError("估算临时数据超过批准的容量基准磁盘预算")
    free_bytes = shutil.disk_usage(root).free
    if free_bytes < estimated_bytes + _FREE_SPACE_RESERVE_BYTES:
        raise RuntimeError("磁盘剩余空间不足，拒绝生成容量 Fixture")

    try:
        settings = load_settings()
        _require_capacity_database(settings.db_name)

        samples = [
            _run_once(
                settings=settings,
                run_dir=root / f"run-{run_index + 1:02d}",
                row_count=row_count,
                run_index=run_index,
            )
            for run_index in range(runs)
        ]
        p50_elapsed = statistics.median(float(sample["elapsed_seconds"]) for sample in samples)
        report: dict[str, Any] = {
            "schema_version": "excel-import-capacity.v1",
            "synthetic_fixture": True,
            "production_authorization": False,
            "input": {
                "rows": row_count,
                "runs": runs,
                "profile": "aima-monitoring-excel.v1",
            },
            "guard": {
                "disk_budget_bytes": disk_budget_bytes,
                "estimated_bytes": estimated_bytes,
                "free_bytes_before": free_bytes,
            },
            "samples": samples,
            "aggregate": {
                "elapsed_seconds_p50": p50_elapsed,
                "rows_per_second_p50": row_count / p50_elapsed,
                "sql_statements_p50": statistics.median(
                    int(sample["sql_statements"]) for sample in samples
                ),
                "sql_statements_per_1000_rows_p50": statistics.median(
                    float(sample["sql_statements_per_1000_rows"]) for sample in samples
                ),
                "temporary_bytes_max": max(int(sample["temporary_bytes"]) for sample in samples),
                "temporary_peak_bytes_max": max(
                    int(sample["temporary_peak_bytes"]) for sample in samples
                ),
                "monitored_temporary_peak_bytes_max": max(
                    int(sample["monitored_temporary_peak_bytes"]) for sample in samples
                ),
                "work_dir_bytes_max": max(int(sample["work_dir_bytes"]) for sample in samples),
            },
        }
        _atomic_write_json(root / "capacity_report.json", report)
        print(json.dumps(report["aggregate"], ensure_ascii=False, sort_keys=True))
        return report
    except BaseException:
        _cleanup_generated_outputs(
            root,
            tuple(f"run-{index + 1:02d}" for index in range(runs))
            + ("capacity_report.json", ".capacity_report.json.tmp"),
        )
        raise


def _run_once(
    *,
    settings: Any,
    run_dir: Path,
    row_count: int,
    run_index: int,
) -> dict[str, Any]:
    run_dir.mkdir()
    temporary_root = run_dir / "temporary"
    temporary_root.mkdir()
    previous_tempdir = tempfile.tempdir
    runtime = create_worker_runtime(
        settings=settings.model_copy(
            update={
                "data_dir": run_dir / "data",
                "log_dir": run_dir / "logs",
                "log_level": "INFO",
            }
        )
    )
    tempfile.tempdir = str(temporary_root)
    try:
        _reset_capacity_database(runtime)
        client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        principal = Principal(
            principal_id="excel-import-capacity",
            display_name="Excel 导入容量基准",
            role="administrator",
            source="development",
        )
        brand = PostgresBrandVehicleHttpService(runtime).create_brand(
            BrandCreateRequest(display_name="爱玛", role="owned", aliases=("爱玛",)),
            principal=principal,
            request_id=f"excel-import-capacity-brand-{run_index}",
        )
        fixture = _fixture_xlsx(row_count=row_count, nonce=uuid4().hex)
        registry = create_collection_job_registry(runtime=runtime)
        worker = create_job_worker(
            runtime=runtime,
            registry=registry,
            worker_id=f"excel-import-capacity-{run_index}",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        sql_statements = 0
        sql_by_table: dict[str, int] = {}

        def count_sql(
            connection: object,
            cursor: object,
            statement: str,
            parameters: object,
            context: object,
            executemany: bool,
        ) -> None:
            nonlocal sql_statements
            del connection, cursor, parameters, context, executemany
            sql_statements += 1
            table = re.search(r"\b(?:FROM|INTO|UPDATE)\s+([a-z_][a-z_0-9]*)", statement)
            kind = statement.lstrip().split(None, 1)[0].upper()
            key = f"{kind} {table.group(1) if table else 'other'}"
            sql_by_table[key] = sql_by_table.get(key, 0) + 1

        event.listen(runtime.database.engine, "before_cursor_execute", count_sql)
        temporary_monitor = _DirectoryPeakMonitor(temporary_root)
        temporary_monitor.start()
        started = time.perf_counter()
        try:
            created = client.post(
                "/api/v1/import-batches",
                files=[
                    (
                        "file",
                        (
                            "capacity.xlsx",
                            fixture,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        ),
                    ),
                    ("brand_ids", (None, str(brand.id), None)),
                ],
            )
            if created.status_code != 202:
                raise RuntimeError(f"创建 Excel Import 失败: {created.text}")
            if not worker.run_once():
                raise RuntimeError("Excel Import Job 未被 Worker 领取")
            batch_id = created.json()["batch_id"]
            detail = client.get(f"/api/v1/import-batches/{batch_id}")
        finally:
            elapsed_seconds = time.perf_counter() - started
            event.remove(runtime.database.engine, "before_cursor_execute", count_sql)
            temporary_monitor.stop()
        if detail.status_code != 200 or detail.json()["status"] != "succeeded":
            raise RuntimeError(f"Excel Import 未成功: {detail.text}")
        with runtime.database.engine.connect() as connection:
            content_count = int(
                connection.scalar(select(func.count()).select_from(contents_table)) or 0
            )
        if content_count != row_count:
            raise RuntimeError("Excel Import Content 数量与 Fixture 不一致")
        temporary_bytes, temporary_peak_bytes = _temporary_bytes_from_log(
            run_dir / "logs" / "worker.log"
        )
        return {
            "run": run_index + 1,
            "rows": row_count,
            "source_bytes": len(fixture),
            "elapsed_seconds": elapsed_seconds,
            "rows_per_second": row_count / elapsed_seconds,
            "sql_statements": sql_statements,
            "sql_statements_per_1000_rows": sql_statements / row_count * 1000,
            "top_sql_tables": sorted(sql_by_table.items(), key=lambda item: item[1], reverse=True)[
                :15
            ],
            "temporary_bytes": temporary_bytes,
            "temporary_peak_bytes": temporary_peak_bytes,
            "monitored_temporary_peak_bytes": temporary_monitor.peak_bytes,
            "work_dir_bytes": _directory_bytes(run_dir),
            "batch_stats": detail.json()["stats"],
        }
    finally:
        runtime.close()
        tempfile.tempdir = previous_tempdir


def _fixture_xlsx(*, row_count: int, nonce: str) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("文章")
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    for index in range(row_count):
        external_id = f"excel-capacity-{nonce}-{index}"
        sheet.append(
            [
                "小红书",
                f"爱玛本地导入容量样本 {index}",
                f"爱玛本地导入正文 {index}",
                f"容量作者 {index % 100}",
                "2026-09-24 08:00:00",
                f"https://www.xiaohongshu.com/explore/{external_id}",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _reset_capacity_database(runtime: Any) -> None:
    if not runtime.settings.db_name.endswith(_DATABASE_SUFFIX):
        raise RuntimeError("拒绝清理非专用 Excel 容量数据库")
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, accounts, "
            "audit_events RESTART IDENTITY CASCADE"
        )


def _temporary_bytes_from_log(path: Path) -> tuple[int, int]:
    if not path.is_file():
        return 0, 0
    payload = path.read_text(encoding="utf-8")
    values = [int(match.group(1)) for match in _TEMPORARY_BYTES_PATTERN.finditer(payload)]
    peaks = [int(match.group(1)) for match in _TEMPORARY_PEAK_BYTES_PATTERN.finditer(payload)]
    return max(values, default=0), max(peaks, default=0)


def _directory_bytes(path: Path) -> int:
    return sum(item.stat().st_size for item in path.rglob("*") if item.is_file())


def _cleanup_generated_outputs(root: Path, names: tuple[str, ...]) -> None:
    """异常时只清理由本次空目录基准创建的已知直属输出。"""

    resolved_root = root.resolve()
    for name in names:
        candidate = (resolved_root / name).resolve()
        if candidate.parent != resolved_root:
            raise RuntimeError("容量基准清理目标越出工作目录")
        if candidate.is_dir():
            shutil.rmtree(candidate)
        else:
            candidate.unlink(missing_ok=True)


class _DirectoryPeakMonitor:
    """短周期观测任务私有临时根；不读取文件内容。"""

    def __init__(self, root: Path) -> None:
        self._root = root
        self._stop = threading.Event()
        self._thread = threading.Thread(target=self._run, daemon=True)
        self.peak_bytes = 0

    def start(self) -> None:
        self._thread.start()

    def stop(self) -> None:
        self._stop.set()
        self._thread.join(timeout=5)
        self._sample()

    def _run(self) -> None:
        while not self._stop.wait(0.005):
            self._sample()

    def _sample(self) -> None:
        try:
            current = _directory_bytes(self._root)
        except OSError:
            return
        self.peak_bytes = max(self.peak_bytes, current)


def _atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.unlink(missing_ok=True)
    try:
        with temporary.open("w", encoding="utf-8", newline="\n") as output:
            json.dump(payload, output, ensure_ascii=False, indent=2)
            output.write("\n")
            output.flush()
            os.fsync(output.fileno())
        os.replace(temporary, path)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise


def main() -> int:
    parser = argparse.ArgumentParser(description="运行本地 Excel 导入容量基准")
    parser.add_argument("--work-dir", type=Path, required=True)
    parser.add_argument("--rows", type=int, default=1000)
    parser.add_argument("--runs", type=int, default=3)
    parser.add_argument(
        "--disk-budget-mib",
        type=int,
        default=_DEFAULT_DISK_BUDGET_BYTES // (1024 * 1024),
    )
    args = parser.parse_args()
    run_benchmark(
        work_dir=args.work_dir,
        row_count=args.rows,
        runs=args.runs,
        disk_budget_bytes=args.disk_budget_mib * 1024 * 1024,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
