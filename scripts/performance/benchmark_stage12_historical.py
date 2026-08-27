"""Stage 12 历史迁移端到端容量基准；仅允许专用容量数据库。"""

from __future__ import annotations

import argparse
import ctypes
import json
import math
import os
import platform
import shutil
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4

from aima_ugc.adapters.persistence.postgres.jobs import PostgresJobRepository
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.historical_import_http import PostgresHistoricalImportHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.modules.ingestion.historical_jobs import (
    HISTORICAL_DISCOVER_JOB_TYPE,
    HISTORICAL_IMPORT_CHUNK_JOB_TYPE,
    HISTORICAL_SNAPSHOT_JOB_TYPE,
)
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import JobHandlerResult
from aima_ugc.platform.jobs.tables import jobs_table
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import BaseModel, ConfigDict
from sqlalchemy import func, select, text

CAPACITY_SCHEMA_VERSION = "stage12-historical-capacity.v1"
_CAPACITY_DATABASE_SUFFIX = "_stage12_capacity"
_SHEET_NAME = "文章"
_HEADERS = (
    "媒体名称（中文）",
    "标题",
    "内文",
    "作者",
    "出版日期",
    "原文链接",
)
_HISTORICAL_JOB_TYPES = (
    HISTORICAL_DISCOVER_JOB_TYPE,
    HISTORICAL_SNAPSHOT_JOB_TYPE,
    HISTORICAL_IMPORT_CHUNK_JOB_TYPE,
)
_PROBE_JOB_TYPE = "stage12.capacity-ordinary-probe.v1"


class _ProbePayload(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    schema_version: Literal["stage12.capacity-ordinary-probe.v1"] = _PROBE_JOB_TYPE


def run_benchmark(
    *,
    work_dir: Path,
    row_count: int,
    rows_per_file: int,
    chunk_rows: int,
    max_in_flight: int,
) -> dict[str, Any]:
    """生成有界 XLSX Fixture，并编排生产 Campaign/Worker 取得容量证据。"""

    _require_positive(row_count, "row_count")
    _require_positive(rows_per_file, "rows_per_file")
    _require_positive(chunk_rows, "chunk_rows")
    _require_positive(max_in_flight, "max_in_flight")
    if rows_per_file > 1_048_575:
        raise ValueError("rows_per_file 不能超过 XLSX 单 Sheet 数据行上限")
    if not 100 <= chunk_rows <= 2_000:
        raise ValueError("chunk_rows 必须在 100 到 2000 之间")
    if not 1 <= max_in_flight <= 16:
        raise ValueError("max_in_flight 必须在 1 到 16 之间")

    root = Path(work_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    _require_empty(root)
    source_root = root / "input"
    source_root.mkdir()

    fixture_started = time.perf_counter()
    source_files = _write_fixture(
        source_root,
        row_count=row_count,
        rows_per_file=rows_per_file,
    )
    fixture_seconds = time.perf_counter() - fixture_started
    source_bytes = sum(path.stat().st_size for path in source_files)

    settings = load_settings().model_copy(
        update={
            "data_dir": root / "data",
            "log_dir": root / "logs",
            "historical_import_root": source_root,
            "historical_chunk_rows": chunk_rows,
            "historical_max_scan_files": max(len(source_files), 1),
            "historical_max_in_flight_jobs": max_in_flight,
            "log_level": "WARNING",
        }
    )
    if not settings.db_name.endswith(_CAPACITY_DATABASE_SUFFIX):
        raise RuntimeError(
            f"容量脚本只允许专用数据库；AIMA_DB_NAME 必须以 {_CAPACITY_DATABASE_SUFFIX} 结尾"
        )

    runtime = create_worker_runtime(settings=settings)
    try:
        _reset_capacity_database(runtime)
        before = _database_counters(runtime)
        registry = create_collection_job_registry(runtime=runtime)
        registry.register(
            job_type=_PROBE_JOB_TYPE,
            payload_version=_PROBE_JOB_TYPE,
            payload_model=_ProbePayload,
            handler=lambda payload, context: JobHandlerResult.succeeded(
                {"probe": payload.schema_version}
            ),
            retry_on_timeout=False,
        )
        worker = create_job_worker(
            runtime=runtime,
            registry=registry,
            worker_id="stage12-capacity-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
            )
        )
        pack_id = _create_keyword_pack(client)

        benchmark_started = time.perf_counter()
        cpu_started = time.process_time()
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-capacity-{uuid4()}",
                "relative_paths": [path.name for path in source_files],
                "recursive": False,
                "keyword_pack_ids": [pack_id],
            },
        )
        _require_status(created, 202, "创建容量 Campaign")
        campaign_id = created.json()["campaign_id"]

        maximum_active = _historical_active_jobs(runtime)
        maximum_lock_waiters = _lock_waiters(runtime)
        preflight_started = time.perf_counter()
        preflight_jobs, maximum_active, maximum_lock_waiters = _drain_until(
            client=client,
            runtime=runtime,
            worker=worker,
            campaign_id=campaign_id,
            terminal_statuses={"ready", "failed", "cancelled"},
            maximum_active=maximum_active,
            maximum_lock_waiters=maximum_lock_waiters,
        )
        preflight_seconds = time.perf_counter() - preflight_started
        ready = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}")
        _require_status(ready, 200, "读取预检 Campaign")
        if ready.json()["status"] != "ready":
            raise RuntimeError(f"容量 Campaign 预检失败: {ready.json()}")
        if ready.json()["total_rows"] != row_count:
            raise RuntimeError("容量 Campaign 预检行数与 Fixture 不一致")

        started = client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start")
        _require_status(started, 200, "启动容量 Campaign")
        maximum_active = max(maximum_active, _historical_active_jobs(runtime))
        probe_id = _enqueue_ordinary_probe(runtime)
        if not worker.run_once():
            raise RuntimeError("普通优先级探针 Job 未被 Worker 认领")
        ordinary_probe_passed = _job_status(runtime, probe_id) == "succeeded"
        if not ordinary_probe_passed:
            raise RuntimeError("历史低优先级 Job 抢占了普通探针 Job")

        import_started = time.perf_counter()
        import_jobs, maximum_active, maximum_lock_waiters = _drain_until(
            client=client,
            runtime=runtime,
            worker=worker,
            campaign_id=campaign_id,
            terminal_statuses={"succeeded", "partial_failed", "failed", "cancelled"},
            maximum_active=maximum_active,
            maximum_lock_waiters=maximum_lock_waiters,
        )
        import_seconds = time.perf_counter() - import_started
        elapsed_seconds = time.perf_counter() - benchmark_started
        cpu_seconds = time.process_time() - cpu_started

        response = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}")
        _require_status(response, 200, "读取容量 Campaign 终态")
        campaign = response.json()
        stats = _complete_stats(campaign.get("stats"))
        terminal_rows = sum(stats.values())
        reconciled = terminal_rows == row_count
        if campaign["status"] != "succeeded" or not reconciled:
            raise RuntimeError(f"容量 Campaign 未成功对账: {campaign}")

        after = _database_counters(runtime)
        chunk_seconds = _chunk_durations(runtime, campaign_id)
        query_latency_ms = _content_count_latency_ms(runtime)
        storage = _storage_metrics(runtime, root, source_bytes=source_bytes)
        artifact_bytes = _directory_bytes(settings.artifact_dir)
        job_count = preflight_jobs + import_jobs + 1
        report: dict[str, Any] = {
            "schema_version": CAPACITY_SCHEMA_VERSION,
            "created_at": datetime.now().astimezone().isoformat(),
            "synthetic_fixture": True,
            "production_authorization": False,
            "input": {
                "rows": row_count,
                "files": len(source_files),
                "source_bytes": source_bytes,
                "profile": "aima-monitoring-excel.v1",
                "distribution": {
                    "candidate_expected": row_count,
                    "duplicate_expected": 0,
                    "filtered_expected": 0,
                    "invalid_expected": 0,
                    "conflict_expected": 0,
                },
            },
            "configuration": {
                "chunk_rows": chunk_rows,
                "max_in_flight_jobs": max_in_flight,
                "worker_count": 1,
                "worker_lease_seconds": 120,
                "database": settings.db_name,
                "postgresql": after["postgresql"],
            },
            "runtime": {
                "python": platform.python_version(),
                "platform": platform.platform(),
                "cpu_count": os.cpu_count(),
            },
            "campaign": {
                "id": campaign_id,
                "status": campaign["status"],
                "stats": stats,
                "terminal_rows": terminal_rows,
                "reconciled": reconciled,
            },
            "measurements": {
                "fixture_seconds": fixture_seconds,
                "preflight_seconds": preflight_seconds,
                "import_seconds": import_seconds,
                "elapsed_seconds": elapsed_seconds,
                "rows_per_second": row_count / elapsed_seconds,
                "import_rows_per_second": row_count / import_seconds,
                "process_cpu_seconds": cpu_seconds,
                "process_cpu_percent_of_one_core": cpu_seconds / elapsed_seconds * 100,
                "peak_rss_bytes": _peak_rss_bytes(),
                "chunk_seconds_p50": _percentile(chunk_seconds, 0.50),
                "chunk_seconds_p95": _percentile(chunk_seconds, 0.95),
                "query_count_latency_ms": query_latency_ms,
                "wal_bytes": max(after["wal_lsn_bytes"] - before["wal_lsn_bytes"], 0),
                "temp_bytes": max(after["temp_bytes"] - before["temp_bytes"], 0),
                "deadlocks": max(after["deadlocks"] - before["deadlocks"], 0),
                "maximum_lock_waiters": maximum_lock_waiters,
            },
            "storage": {
                **storage,
                "source_bytes": source_bytes,
                "artifact_bytes": artifact_bytes,
            },
            "jobs": {
                "executed": job_count,
                "preflight_executed": preflight_jobs,
                "import_executed": import_jobs,
                "maximum_observed_active": maximum_active,
                "ordinary_job_starvation_probe": ("passed" if ordinary_probe_passed else "failed"),
            },
            "limitations": [
                "Fixture 为全新、全相关、无冲突的合成数据，主要覆盖最大全量新增存储路径。",
                "单 Worker 基准不能替代公司服务器真实硬件、真实文件分布和并发放量复测。",
                "本脚本不调用 AI，也不代表已授权或已执行生产 4000 万迁移。",
            ],
        }
        report_path = root / "capacity_report.json"
        _atomic_write_json(report_path, report)
        print(
            json.dumps(
                {
                    "report": str(report_path),
                    "rows": row_count,
                    "status": campaign["status"],
                    "rows_per_second": report["measurements"]["rows_per_second"],
                },
                ensure_ascii=False,
            )
        )
        return report
    finally:
        runtime.close()


def _write_fixture(root: Path, *, row_count: int, rows_per_file: int) -> tuple[Path, ...]:
    files: list[Path] = []
    written = 0
    file_count = math.ceil(row_count / rows_per_file)
    for file_index in range(file_count):
        path = root / f"capacity-{file_index + 1:04d}.xlsx"
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet(_SHEET_NAME)
        sheet.append(_HEADERS)
        current_rows = min(rows_per_file, row_count - written)
        for offset in range(current_rows):
            ordinal = written + offset + 1
            sheet.append(
                (
                    "小红书",
                    f"爱玛容量迁移样本 {ordinal}",
                    f"爱玛容量迁移正文 {ordinal}，用于验证有界流式导入。",
                    f"容量样本作者 {ordinal % 10_000}",
                    "2025-01-02 10:00:00",
                    f"https://www.xiaohongshu.com/explore/stage12-capacity-{ordinal}",
                )
            )
        workbook.save(path)
        files.append(path)
        written += current_rows
    return tuple(files)


def _create_keyword_pack(client: TestClient) -> str:
    created = client.post(
        "/api/v1/keyword-packs",
        json={"name": f"Stage12 容量基准 {uuid4()}"},
    )
    _require_status(created, 201, "创建容量关键词包")
    pack_id = created.json()["id"]
    added = client.post(
        f"/api/v1/keyword-packs/{pack_id}/keywords",
        json={"text": "爱玛", "priority": 10},
    )
    _require_status(added, 201, "创建容量关键词")
    return str(pack_id)


def _drain_until(
    *,
    client: TestClient,
    runtime: Any,
    worker: Any,
    campaign_id: str,
    terminal_statuses: set[str],
    maximum_active: int,
    maximum_lock_waiters: int,
) -> tuple[int, int, int]:
    executed = 0
    while True:
        current = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}")
        _require_status(current, 200, "读取容量 Campaign")
        if current.json()["status"] in terminal_statuses:
            return executed, maximum_active, maximum_lock_waiters
        maximum_active = max(maximum_active, _historical_active_jobs(runtime))
        maximum_lock_waiters = max(maximum_lock_waiters, _lock_waiters(runtime))
        if not worker.run_once():
            raise RuntimeError(f"Campaign 未到终态但 Worker 无可执行 Job: {current.json()}")
        executed += 1


def _enqueue_ordinary_probe(runtime: Any) -> Any:
    session = runtime.database.new_session()
    try:
        with session.begin():
            return (
                PostgresJobRepository(session)
                .enqueue(
                    job_type=_PROBE_JOB_TYPE,
                    payload_version=_PROBE_JOB_TYPE,
                    payload=_ProbePayload().model_dump(mode="json"),
                    internal_idempotency_key=f"stage12-capacity-probe:{uuid4()}",
                    request_id="stage12-capacity-probe",
                    priority=0,
                    max_attempts=1,
                    timeout_seconds=60,
                )
                .id
            )
    finally:
        session.close()


def _job_status(runtime: Any, job_id: Any) -> str | None:
    session = runtime.database.new_session()
    try:
        with session.begin():
            job = PostgresJobRepository(session).get(job_id)
            return job.status if job is not None else None
    finally:
        session.close()


def _historical_active_jobs(runtime: Any) -> int:
    with runtime.database.engine.begin() as connection:
        return int(
            connection.scalar(
                select(func.count())
                .select_from(jobs_table)
                .where(
                    jobs_table.c.job_type.in_(_HISTORICAL_JOB_TYPES),
                    jobs_table.c.status.in_(("queued", "running")),
                )
            )
            or 0
        )


def _lock_waiters(runtime: Any) -> int:
    with runtime.database.engine.begin() as connection:
        return int(
            connection.scalar(
                text(
                    "SELECT count(*) FROM pg_stat_activity "
                    "WHERE datname = current_database() AND wait_event_type = 'Lock'"
                )
            )
            or 0
        )


def _database_counters(runtime: Any) -> dict[str, Any]:
    with runtime.database.engine.begin() as connection:
        row = (
            connection.execute(
                text(
                    "SELECT temp_bytes, deadlocks, "
                    "pg_wal_lsn_diff(pg_current_wal_lsn(), '0/0')::bigint AS wal_lsn_bytes "
                    "FROM pg_stat_database WHERE datname = current_database()"
                )
            )
            .mappings()
            .one()
        )
        postgresql = {
            "server_version": connection.scalar(text("SHOW server_version")),
            "shared_buffers": connection.scalar(text("SHOW shared_buffers")),
            "max_connections": connection.scalar(text("SHOW max_connections")),
            "work_mem": connection.scalar(text("SHOW work_mem")),
        }
    return {
        "temp_bytes": int(row["temp_bytes"] or 0),
        "deadlocks": int(row["deadlocks"] or 0),
        "wal_lsn_bytes": int(row["wal_lsn_bytes"] or 0),
        "postgresql": postgresql,
    }


def _chunk_durations(runtime: Any, campaign_id: str) -> tuple[float, ...]:
    with runtime.database.engine.begin() as connection:
        rows = connection.execute(
            text(
                "SELECT extract(epoch FROM job.finished_at - job.started_at) AS seconds "
                "FROM jobs AS job "
                "WHERE job.job_type = :job_type AND job.status = 'succeeded' "
                "AND job.payload ->> 'chunk_item_id' IN ("
                "SELECT item.id::text FROM historical_import_campaign_items AS item "
                "WHERE item.campaign_id = CAST(:campaign_id AS uuid) "
                "AND item.item_kind = 'chunk')"
            ),
            {
                "job_type": HISTORICAL_IMPORT_CHUNK_JOB_TYPE,
                "campaign_id": campaign_id,
            },
        )
        return tuple(float(row.seconds) for row in rows if row.seconds is not None)


def _content_count_latency_ms(runtime: Any) -> float:
    started = time.perf_counter()
    with runtime.database.engine.begin() as connection:
        connection.execute(text("SELECT count(*) FROM contents")).scalar_one()
    return (time.perf_counter() - started) * 1000


def _storage_metrics(runtime: Any, root: Path, *, source_bytes: int) -> dict[str, Any]:
    relations = (
        "contents",
        "content_versions",
        "processing_import_batch_items",
        "processing_import_batch_item_conflicts",
        "historical_import_campaign_items",
        "artifacts",
        "jobs",
    )
    with runtime.database.engine.begin() as connection:
        database_bytes = int(connection.scalar(text("SELECT pg_database_size(current_database())")))
        relation_bytes = {
            relation: int(
                connection.scalar(
                    text("SELECT pg_total_relation_size(to_regclass(:relation))"),
                    {"relation": relation},
                )
                or 0
            )
            for relation in relations
        }
    disk = shutil.disk_usage(root)
    return {
        "database_bytes": database_bytes,
        "relation_bytes": relation_bytes,
        "work_dir_bytes": _directory_bytes(root),
        "disk_free_bytes": disk.free,
        "disk_total_bytes": disk.total,
        "source_bytes": source_bytes,
    }


def _reset_capacity_database(runtime: Any) -> None:
    if not runtime.settings.db_name.endswith(_CAPACITY_DATABASE_SUFFIX):
        raise RuntimeError("拒绝清理非专用容量数据库")
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts, audit_events "
            "RESTART IDENTITY CASCADE"
        )


def _complete_stats(value: object) -> dict[str, int]:
    source = value if isinstance(value, dict) else {}
    return {
        name: int(source.get(name, 0))
        for name in (
            "created",
            "filled",
            "unchanged",
            "conflict",
            "filtered",
            "duplicate",
            "invalid",
            "failed",
        )
    }


def _percentile(values: tuple[float, ...], fraction: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    index = min(math.ceil(len(ordered) * fraction) - 1, len(ordered) - 1)
    return ordered[max(index, 0)]


def _peak_rss_bytes() -> int:
    if os.name == "nt":
        return _windows_peak_rss_bytes()
    import resource

    usage = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    multiplier = 1 if platform.system() == "Darwin" else 1024
    return int(usage * multiplier)


def _windows_peak_rss_bytes() -> int:
    class ProcessMemoryCounters(ctypes.Structure):
        _fields_ = [
            ("cb", ctypes.c_ulong),
            ("PageFaultCount", ctypes.c_ulong),
            ("PeakWorkingSetSize", ctypes.c_size_t),
            ("WorkingSetSize", ctypes.c_size_t),
            ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
            ("QuotaPagedPoolUsage", ctypes.c_size_t),
            ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
            ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
            ("PagefileUsage", ctypes.c_size_t),
            ("PeakPagefileUsage", ctypes.c_size_t),
        ]

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    psapi = ctypes.WinDLL("psapi", use_last_error=True)
    get_current_process = kernel32.GetCurrentProcess
    get_current_process.restype = ctypes.c_void_p
    get_process_memory_info = psapi.GetProcessMemoryInfo
    get_process_memory_info.argtypes = [
        ctypes.c_void_p,
        ctypes.POINTER(ProcessMemoryCounters),
        ctypes.c_ulong,
    ]
    get_process_memory_info.restype = ctypes.c_int
    counters = ProcessMemoryCounters()
    counters.cb = ctypes.sizeof(counters)
    if not get_process_memory_info(
        get_current_process(),
        ctypes.byref(counters),
        ctypes.sizeof(counters),
    ):
        raise OSError(ctypes.get_last_error(), "GetProcessMemoryInfo 失败")
    return int(counters.PeakWorkingSetSize)


def _directory_bytes(path: Path) -> int:
    if not path.exists():
        return 0
    return sum(item.stat().st_size for item in path.rglob("*") if item.is_file())


def _require_status(response: Any, expected: int, action: str) -> None:
    if response.status_code != expected:
        raise RuntimeError(f"{action}失败: status={response.status_code}, body={response.text}")


def _require_empty(path: Path) -> None:
    if any(path.iterdir()):
        raise ValueError(f"容量 work_dir 必须为空目录: {path}")


def _require_positive(value: int, name: str) -> None:
    if isinstance(value, bool) or value <= 0:
        raise ValueError(f"{name} 必须是正整数")


def _atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.unlink(missing_ok=True)
    try:
        with temporary.open("w", encoding="utf-8", newline="\n") as output:
            json.dump(payload, output, ensure_ascii=False, indent=2)
            output.write("\n")
            output.flush()
            os.fsync(output.fileno())
        os.replace(temporary, path)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="运行 Stage 12 历史迁移容量基准")
    parser.add_argument("--work-dir", type=Path, required=True)
    parser.add_argument("--rows", type=int, required=True)
    parser.add_argument("--rows-per-file", type=int, default=250_000)
    parser.add_argument("--chunk-rows", type=int, default=1_000)
    parser.add_argument("--max-in-flight", type=int, default=2)
    return parser.parse_args()


def main() -> int:
    arguments = _parse_args()
    run_benchmark(
        work_dir=arguments.work_dir,
        row_count=arguments.rows,
        rows_per_file=arguments.rows_per_file,
        chunk_rows=arguments.chunk_rows,
        max_in_flight=arguments.max_in_flight,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
