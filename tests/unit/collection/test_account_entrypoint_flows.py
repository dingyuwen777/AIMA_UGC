"""通过四个平台真实公开入口验证账号、评论、Canonical 与共享文件链。"""

from __future__ import annotations

import copy
import json
from pathlib import Path
from typing import Any

import pytest
from aima_ugc.adapters.providers import tikhub_test
from aima_ugc.adapters.providers.tikhub import runtime
from aima_ugc.modules.collection.providers.transport import (
    ProviderTransportRequest,
    ProviderTransportResponse,
)
from openpyxl import load_workbook

_FIXTURES = Path(__file__).parents[2] / "fixtures/providers/tikhub"
_PLATFORMS = ["douyin", "kuaishou", "weibo", "bilibili"]


def _fixture(platform: str, name: str) -> dict[str, Any]:
    """复用已有脱敏真实形状，测试只替换业务关联与分页边界。"""
    return json.loads((_FIXTURES / platform / name).read_text(encoding="utf-8"))


def _responses(
    platform: str,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    detail = _fixture(platform, "detail.sanitized.json")
    item = runtime.extract_detail_items(platform, detail)[0]
    root_body = _fixture(platform, "comments_page1.sanitized.json")
    root = runtime.extract_comment_items(platform, root_body)[0]
    reply_name = (
        "replies_page1.sanitized.json"
        if platform in {"douyin", "bilibili"}
        else "sub_comments_page1.sanitized.json"
    )
    reply_body = _fixture(platform, reply_name)
    reply = runtime.extract_sub_comment_items(platform, reply_body)[0]
    if platform == "douyin":
        item["aweme_id"] = "100001"
        item["author"].update(uid="100002", sec_uid="MS4wLjAB" + "x" * 40)
        item["statistics"]["comment_count"] = 3
        detail["data"]["aweme_detail"] = item
        root.update(aweme_id="100001", cid="100101", reply_comment_total=1)
        root_body["data"].update(cursor=0, has_more=0, total=3)
        reply.update(aweme_id="100001", cid="100201", root_comment_id="100101", reply_id="100101")
        page1 = {"data": {"comments": [reply], "has_more": 1, "cursor": 10, "total": 1}}
        r2 = copy.deepcopy(reply)
        r2["cid"] = "100202"
        page2 = {"data": {"comments": [r2], "has_more": 0, "cursor": 20, "total": 1}}
        posts = {"data": {"aweme_list": [item], "has_more": 0, "max_cursor": 0}}
        empty = {"data": {"aweme_list": [], "has_more": 0, "max_cursor": 0}}
    elif platform == "kuaishou":
        item.update(photo_id=100001, user_id=100002, comment_count=3)
        root.update(photo_id=100001, comment_id=100101, subCommentCount=1)
        root_body["data"].update(pcursor="no_more", commentCount=3)
        reply.update(photo_id=100001, comment_id=100201, reply_to=100101)
        page1 = {"data": {"subComments": [reply], "pcursor": "reply-next"}}
        r2 = copy.deepcopy(reply)
        r2["comment_id"] = 100202
        page2 = {"data": {"subComments": [r2], "pcursor": "no_more"}}
        posts = {"data": {"photos": [item], "pcursor": "no_more"}}
        empty = {"data": {"photos": [], "pcursor": "no_more"}}
    elif platform == "weibo":
        item.update(id=100001, idstr="100001", mid="100001", comments_count=3)
        item["user"].update(id=100002, idstr="100002")
        root.update(id=100101, idstr="100101", mid="100101", rootidstr="100101", total_number=1)
        root_body["data"].pop("moreInfo")
        reply.update(id=100201, idstr="100201", mid="100201", rid="100101", rootidstr="100101")
        reply["reply_comment"].update(idstr="100101", mid="100101")
        page1 = {"data": {"data": [reply], "max_id": 100201, "max_id_type": 0}}
        r2 = copy.deepcopy(reply)
        r2.update(id=100202, idstr="100202", mid="100202")
        page2 = {"data": {"data": [r2], "max_id": 0, "max_id_type": 0}}
        posts = {"data": {"statuses": [item], "since_id": ""}}
        empty = {"data": {"statuses": [], "since_id": ""}}
    else:
        item["stat"]["reply"] = 3
        root.update(rpid=100101, rpid_str="100101", oid=100001, rcount=1, replies=[])
        root_body["data"]["data"]["cursor"].update(all_count=3, is_end=True)
        reply.update(
            rpid=100201,
            rpid_str="100201",
            oid=100001,
            root=100101,
            root_str="100101",
            parent=100101,
            parent_str="100101",
        )
        page1 = {
            "data": {
                "data": {
                    "cursor": {
                        "is_end": False,
                        "next": 1,
                        "all_count": 1,
                        "pagination_reply": {"next_offset": 1},
                    },
                    "root": {**root, "replies": [reply]},
                }
            }
        }
        r2 = copy.deepcopy(reply)
        r2.update(rpid=100202, rpid_str="100202")
        page2 = {
            "data": {
                "data": {
                    "cursor": {"is_end": True, "next": 0, "all_count": 1},
                    "root": {**root, "replies": [r2]},
                }
            }
        }
        posts = {"data": {"data": {"archives": [item]}}}
        empty = {"data": {"data": {"archives": []}}}
    return posts, detail, root_body, [page1, page2], empty


def _invoke(
    platform: str,
    tmp_path: Path,
    monkeypatch: Any,
    *,
    reply_limit: int | None = None,
    post_limit: int | None = None,
    transport_failure_delivery: str | None = None,
    account_target: Any = None,
    identity_response: dict[str, Any] | None = None,
    web_roots: dict[str, Any] | None = None,
) -> tuple[Any, dict[str, Any], list[ProviderTransportRequest]]:
    posts, detail, roots, replies, empty = _responses(platform)
    seen: list[ProviderTransportRequest] = []
    post_calls = 0

    class FakeTransport:
        """按真实 Operation path 分派，未知调用 fail closed。"""

        def __init__(self, **_: object) -> None:
            pass

        def __enter__(self) -> FakeTransport:
            return self

        def __exit__(self, *_: object) -> None:
            pass

        def send(self, request: ProviderTransportRequest) -> ProviderTransportResponse:
            nonlocal post_calls
            seen.append(request)
            path = request.path
            if "search_user" in path:
                assert identity_response is not None
                body = identity_response
            elif "post" in path and ("user" in path or "posted" in path):
                post_calls += 1
                body = posts if post_calls == 1 else empty
            elif "sub_comment" in path or "replies" in path or "reply_detail" in path:
                if "/web/" in path and platform == "kuaishou":
                    body = {"data": {"subComments": [], "pcursor": "no_more"}}
                elif replies:
                    if len(replies) == 1 and transport_failure_delivery:
                        from aima_ugc.modules.collection.providers.transport import (
                            ProviderTransportFailure,
                        )

                        if transport_failure_delivery == "unknown":
                            raise ProviderTransportFailure.unknown(
                                code="read_timeout", safe_summary="fixture 结果未知"
                            )
                        raise ProviderTransportFailure.not_sent(
                            code="connect_failed", safe_summary="fixture 未发送"
                        )
                    body = replies.pop(0)
                else:
                    raise AssertionError("回复出现预期外调用")
            elif "comment" in path:
                if "/web/" in path and platform == "kuaishou":
                    body = web_roots or {
                        "data": {"rootComments": [], "pcursor": "no_more", "commentCount": 3}
                    }
                else:
                    body = roots
            elif "video" in path or "detail" in path:
                body = detail
            else:
                raise AssertionError(f"未定义 Operation: {path}")
            return ProviderTransportResponse(
                status_code=200, external_request_id=f"fake-{len(seen)}", body=copy.deepcopy(body)
            )

    env = tmp_path / ".env"
    env.write_text(
        "TIKHUB_BASE_URL=https://api.tikhub.io\nTIKHUB_API_KEY=fixture-only-secret\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        "aima_ugc.adapters.providers.tikhub_test.operations.runner.TikHubHttpTransport",
        FakeTransport,
    )
    targets = {
        "douyin": tikhub_test.DouyinAccountTarget(sec_uid="MS4wLjAB" + "x" * 40, uid="100002"),
        "kuaishou": tikhub_test.KuaishouAccountTarget(user_id="100002"),
        "weibo": tikhub_test.WeiboAccountTarget(uid="100002"),
        "bilibili": tikhub_test.BilibiliAccountTarget(uid="100002"),
    }
    result = getattr(tikhub_test, f"run_{platform}_accounts")(
        accounts=[account_target or targets[platform]],
        start_date="2024-01-01",
        end_date="2026-12-31",
        env_file=env,
        output_root=tmp_path / "output",
        run_id="flow",
        max_account_post_pages=post_limit,
        max_comments_per_content=1,
        max_replies_per_root=1,
        max_comment_pages_per_content=3,
        max_reply_pages_per_root=reply_limit,
    )
    summary = json.loads(result.run_summary_path.read_text(encoding="utf-8"))
    return result, summary, seen


@pytest.mark.parametrize("platform", _PLATFORMS)
def test_account_public_entrypoint_keeps_replies_after_known_soft_count(
    platform: str, tmp_path: Path, monkeypatch: Any
) -> None:
    """reply_count=1 仍跟随第二页；生产 Mapper、JSONL、Excel 均保留两条回复。"""
    result, summary, seen = _invoke(platform, tmp_path, monkeypatch)
    assert result.content_count == 1, summary
    assert result.root_comment_count == 1, summary
    assert result.reply_count == 2, summary
    assert summary["status"] == "completed", summary
    assert seen[0].credential.get_secret_value() == "fixture-only-secret"
    comments = [
        json.loads(line)
        for line in (result.run_dir / "canonical/comments.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
    ]
    assert len(comments) == 3
    assert all(row["source"]["raw_artifact_id"] for row in comments)
    book = load_workbook(result.workbook_path, read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["文章"]
        rows = list(book["文章"].values)
        assert len(rows) == 4
        assert len({row[2] for row in rows[1:]}) == 3
        assert all(row[11] is None for row in rows[1:])
    finally:
        book.close()


@pytest.mark.parametrize("platform", _PLATFORMS)
def test_reply_hard_limit_cannot_report_complete_when_known_count_reached(
    platform: str, tmp_path: Path, monkeypatch: Any
) -> None:
    """第一页数量达到 reply_count 但 Provider 仍有下一页，硬上限必须 partial。"""
    result, summary, _ = _invoke(platform, tmp_path, monkeypatch, reply_limit=1)
    assert result.reply_count == 1, summary
    assert summary["status"] != "completed", summary
    assert summary["partial_content_count"] == 1, summary


def test_bilibili_post_hard_limit_is_partial(tmp_path: Path, monkeypatch: Any) -> None:
    """账号投稿接口依靠下一空页结束，技术上限不能被当作全量成功。"""
    _, summary, _ = _invoke("bilibili", tmp_path, monkeypatch, post_limit=1)
    assert summary["accounts"][0]["status"] == "partial", summary


@pytest.mark.parametrize(
    "platform,url",
    [
        ("douyin", "https://www.douyin.com/user/MS4wLjABdifferent"),
        ("kuaishou", "https://www.kuaishou.com/profile/999999"),
        ("weibo", "https://weibo.com/u/999999"),
        ("bilibili", "https://space.bilibili.com/999999"),
    ],
)
def test_account_rejects_conflicting_homepage_identity(platform: str, url: str) -> None:
    targets = {
        "douyin": {"sec_uid": "MS4wLjAB" + "x" * 40},
        "kuaishou": {"user_id": "100002"},
        "weibo": {"uid": "100002"},
        "bilibili": {"uid": "100002"},
    }
    cls = getattr(tikhub_test, f"{platform.capitalize()}AccountTarget")
    with pytest.raises(ValueError, match="不一致|冲突"):
        cls(**targets[platform], homepage_url=url)


@pytest.mark.parametrize(
    "platform,url",
    [
        ("douyin", "https://evil.invalid/user/MS4wLjABfixture"),
        ("kuaishou", "https://evil.invalid/profile/100002"),
        ("weibo", "https://evil.invalid/u/100002"),
        ("bilibili", "https://evilspace.bilibili.com/100002"),
    ],
)
def test_account_rejects_non_platform_homepage(platform: str, url: str) -> None:
    cls = getattr(tikhub_test, f"{platform.capitalize()}AccountTarget")
    with pytest.raises(ValueError, match="主页|homepage"):
        cls(homepage_url=url)


@pytest.mark.parametrize("platform", _PLATFORMS)
def test_account_unknown_publication_time_is_partial(
    platform: str, tmp_path: Path, monkeypatch: Any
) -> None:
    original = runtime.map_content

    def without_date(**kwargs: Any) -> Any:
        return original(**kwargs).model_copy(update={"published_at": None})

    monkeypatch.setattr(runtime, "map_content", without_date)
    _, summary, _ = _invoke(platform, tmp_path, monkeypatch)
    assert summary["accounts"][0]["status"] == "partial", summary
    assert summary["accounts"][0]["post_failures"][0]["error_type"] == "MissingPublicationTime"


@pytest.mark.parametrize("platform", _PLATFORMS)
def test_account_rejects_mismatched_post_author(
    platform: str, tmp_path: Path, monkeypatch: Any
) -> None:
    original = runtime.map_content

    def wrong_author(**kwargs: Any) -> Any:
        content = original(**kwargs)
        author = content.author.model_copy(
            update={"external_account_id": "999999", "alternate_ids": {"sec_uid": "wrong-sec"}}
        )
        return content.model_copy(update={"author": author})

    monkeypatch.setattr(runtime, "map_content", wrong_author)
    result, summary, seen = _invoke(platform, tmp_path, monkeypatch)
    assert result.content_count == 0, summary
    assert summary["accounts"][0]["identity_mismatches"] == 1, summary
    assert summary["accounts"][0]["status"] == "partial", summary
    assert not any("comment" in request.path for request in seen)


def test_bilibili_lean_account_feed_uses_observed_up_mid(tmp_path: Path, monkeypatch: Any) -> None:
    """真实 V2 投稿列表省略 owner，使用实际 upMid，不能把整个账号当作作者错配。"""
    original = _responses

    def lean(platform: str) -> Any:
        values = original(platform)
        if platform == "bilibili":
            item = values[0]["data"]["data"]["archives"][0]
            item.pop("owner")
            item["upMid"] = 100002
        return values

    monkeypatch.setattr(__name__ + "._responses", lean)
    result, summary, _ = _invoke("bilibili", tmp_path, monkeypatch)
    assert result.content_count == 1, summary
    assert result.root_comment_count == 1
    assert result.reply_count == 2


@pytest.mark.parametrize("platform", _PLATFORMS)
def test_labeling_workbook_reimports_distinct_comment_ids(
    platform: str, tmp_path: Path, monkeypatch: Any
) -> None:
    """使用生产 Excel reader/mapper 验证评论身份，不能只重开工作簿。"""
    from aima_ugc.adapters.providers.imports.excel_profile import get_excel_import_profile
    from aima_ugc.adapters.providers.imports.excel_reader import iter_excel_rows
    from aima_ugc.adapters.providers.imports.mapper import map_excel_row
    from aima_ugc.platform.time import beijing_now

    result, _, _ = _invoke(platform, tmp_path, monkeypatch)
    profile = get_excel_import_profile("aima-monitoring-excel.v1")
    contents = [
        map_excel_row(
            row,
            profile=profile,
            input_name="comments.xlsx",
            sheet_name="文章",
            observed_at=beijing_now(),
        )
        for row in iter_excel_rows(result.workbook_path, profile=profile)
    ]
    assert len(contents) == 3
    assert {content.external_content_id for content in contents} == {"100101", "100201", "100202"}
    assert all(content.platform == platform for content in contents)


@pytest.mark.parametrize("platform", _PLATFORMS)
def test_account_mapping_failure_keeps_good_reply_rows(
    platform: str, tmp_path: Path, monkeypatch: Any
) -> None:
    original = runtime.map_comment

    def one_bad_reply(**kwargs: Any) -> Any:
        if not kwargs["is_root"] and "page[2]" in kwargs["item_locator"]:
            raise ValueError("fixture 的第二页字段缺失")
        return original(**kwargs)

    monkeypatch.setattr(runtime, "map_comment", one_bad_reply)
    result, summary, _ = _invoke(platform, tmp_path, monkeypatch)
    assert result.content_count == 1, summary
    assert result.root_comment_count == 1
    assert result.reply_count == 1
    assert summary["status"] == "partial_success"
    book = load_workbook(result.workbook_path, read_only=True)
    try:
        assert len(list(book["文章"].values)) == 3
    finally:
        book.close()


def test_kuaishou_account_detail_keeps_public_share_identity(
    tmp_path: Path, monkeypatch: Any
) -> None:
    original = _responses

    def shared_identity(platform: str) -> Any:
        values = original(platform)
        if platform == "kuaishou":
            item = copy.deepcopy(values[0]["data"]["photos"][0])
            item["share_info"] = "photoId=3xfixturepost&shareMethod=TOKEN"
            values[0]["data"]["photos"][0] = item
        return values

    monkeypatch.setattr(__name__ + "._responses", shared_identity)
    result, summary, _ = _invoke("kuaishou", tmp_path, monkeypatch)
    assert result.content_count == 1, summary
    assert summary["status"] == "completed", summary
    rows = [
        json.loads(line)
        for line in (result.run_dir / "canonical/comments.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
    ]
    assert {row["external_content_id"] for row in rows} == {"3xfixturepost"}


@pytest.mark.parametrize("platform", _PLATFORMS)
@pytest.mark.parametrize("delivery", ["not_sent", "unknown"])
def test_account_transport_failure_keeps_successful_rows_without_retry(
    platform: str,
    delivery: str,
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    result, summary, seen = _invoke(
        platform, tmp_path, monkeypatch, transport_failure_delivery=delivery
    )
    assert result.content_count == 1, summary
    assert result.root_comment_count == 1
    assert result.reply_count == 1
    assert summary["status"] == "partial_success"
    failed = [request for request in summary["requests"] if request.get("delivery")]
    assert len(failed) == 1
    assert failed[0]["delivery"] == delivery
    assert failed[0]["billing_status"] == ("unknown" if delivery == "unknown" else "not_billable")
    assert "raw_file" not in failed[0]
    assert (
        sum(
            "sub_comment" in r.path or "replies" in r.path or "reply_detail" in r.path
            for r in seen
            if "/web/" not in r.path
        )
        == 2
    )


@pytest.mark.parametrize(
    "config",
    [
        {"user_id": "999"},
        {"eid": "3xwrong"},
        {"homepage_url": "https://www.kuaishou.com/profile/3xwrong"},
    ],
)
def test_kuaishou_search_rejects_other_configured_identity(config, tmp_path, monkeypatch):
    target = tikhub_test.KuaishouAccountTarget(kuaishou_id="handle-A", **config)
    identity = {
        "data": {
            "users": [{"userId": "100002", "kwaiId": "handle-A", "eid": "3xuser"}],
            "pcursor": "no_more",
        }
    }
    _, summary, seen = _invoke(
        "kuaishou", tmp_path, monkeypatch, account_target=target, identity_response=identity
    )
    assert summary["accounts"][0]["status"] == "failed"
    assert len(seen) == 1, [r.path for r in seen]


@pytest.mark.parametrize("missing_eid", [False, True])
def test_kuaishou_search_requires_complete_consistent_identity(missing_eid, tmp_path, monkeypatch):
    target = tikhub_test.KuaishouAccountTarget(
        kuaishou_id="handle-A",
        user_id="100002",
        eid="3xuser",
        homepage_url="https://www.kuaishou.com/profile/3xuser",
    )
    candidate = {"userId": "100002", "kwaiId": "handle-A", "eid": "3xuser"}
    if missing_eid:
        candidate.pop("eid")
    identity = {"data": {"users": [candidate], "pcursor": "no_more"}}
    _, summary, seen = _invoke(
        "kuaishou", tmp_path, monkeypatch, account_target=target, identity_response=identity
    )
    assert summary["accounts"][0]["status"] == ("failed" if missing_eid else "completed"), summary
    assert len(seen) == 1 if missing_eid else len(seen) > 1


@pytest.mark.parametrize("platform", _PLATFORMS)
def test_account_rejects_detail_author_after_correct_discovery(platform, tmp_path, monkeypatch):
    original = _responses

    def changed(name):
        posts, detail, roots, replies, empty = copy.deepcopy(original(name))
        posts = copy.deepcopy(posts)
        item = runtime.extract_detail_items(name, detail)[0]
        if name == "douyin":
            item["author"].update(uid="999", sec_uid="MS4wLjABwrong")
        elif name == "kuaishou":
            item["user_id"] = 999
        elif name == "weibo":
            item["user"].update(id=999, idstr="999")
        else:
            item["owner"]["mid"] = 999
        return posts, detail, roots, replies, empty

    monkeypatch.setattr(__import__(__name__, fromlist=["_responses"]), "_responses", changed)
    result, summary, seen = _invoke(platform, tmp_path, monkeypatch)
    rows = [
        json.loads(row)
        for row in (result.run_dir / "canonical/contents.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
    ]
    assert all(row["author"]["external_account_id"] == "100002" for row in rows), rows
    assert summary["status"] != "completed", summary
    assert not any("comment" in r.path or "replies" in r.path for r in seen)


@pytest.mark.parametrize("platform", _PLATFORMS)
@pytest.mark.parametrize("shape", ["mixed", "missing"])
def test_account_invalid_post_items_are_not_normal_exhaustion(
    platform, shape, tmp_path, monkeypatch
):
    original = _responses

    def changed(name):
        posts, detail, roots, replies, empty = original(name)
        if shape == "missing":
            posts = {"data": {}}
        else:
            items = (
                posts["data"]["aweme_list"]
                if name == "douyin"
                else posts["data"]["photos"]
                if name == "kuaishou"
                else posts["data"]["statuses"]
                if name == "weibo"
                else posts["data"]["data"]["archives"]
            )
            items.extend([{"bad": "missing-id"}, "invalid-item"])
        return posts, detail, roots, replies, empty

    monkeypatch.setattr(__import__(__name__, fromlist=["_responses"]), "_responses", changed)
    result, summary, _ = _invoke(platform, tmp_path, monkeypatch)
    assert summary["accounts"][0]["status"] != "completed", summary
    if shape == "mixed":
        assert result.content_count == 1
        assert len(summary["accounts"][0]["post_failures"]) == 2, summary
        assert all(f["item_locator"] for f in summary["accounts"][0]["post_failures"])


@pytest.mark.parametrize(
    "offset,fallback,expected",
    [(20, 99, 20), ("20", 99, 20), ("opaque", 22, 22), ("opaque", None, None), ("0", 99, None)],
)
def test_bilibili_reply_offset_uses_numeric_value_before_fallback(
    offset, fallback, expected, tmp_path, monkeypatch
):
    original = _responses

    def changed(name):
        values = original(name)
        cursor = values[3][0]["data"]["data"]["cursor"]
        cursor["pagination_reply"]["next_offset"] = offset
        if fallback is None:
            cursor.pop("next")
        else:
            cursor["next"] = fallback
        return values

    monkeypatch.setattr(__import__(__name__, fromlist=["_responses"]), "_responses", changed)
    result, summary, seen = _invoke("bilibili", tmp_path, monkeypatch)
    calls = [r for r in seen if "reply_detail" in r.path]
    if expected is None:
        assert len(calls) == 1
        assert summary["status"] != "completed"
    else:
        assert len(calls) == 2
        assert calls[1].params["next_offset"] == expected
        assert result.reply_count == 2


@pytest.mark.parametrize("app_count", [0, 1])
@pytest.mark.parametrize("web_count", [None, 1, 2])
def test_kuaishou_duplicate_web_root_can_add_reply_discovery(
    app_count, web_count, tmp_path, monkeypatch
):
    original = _responses
    values = original("kuaishou")
    web = copy.deepcopy(values[2])
    runtime.extract_comment_items("kuaishou", web)[0]["subCommentCount"] = web_count
    root = runtime.extract_comment_items("kuaishou", values[2])[0]
    root["subCommentCount"] = app_count
    monkeypatch.setattr(
        __import__(__name__, fromlist=["_responses"]),
        "_responses",
        lambda name: copy.deepcopy(values),
    )
    result, summary, seen = _invoke("kuaishou", tmp_path, monkeypatch, web_roots=web)
    assert result.root_comment_count == 1
    assert result.reply_count == 2, summary
    assert summary["status"] == "completed", summary
    assert sum("sub_comment" in r.path and "/app/" in r.path for r in seen) == 2


@pytest.mark.parametrize("platform", ["douyin", "bilibili"])
def test_account_manual_main_rejects_account_partial_with_complete_comments(
    platform, tmp_path, monkeypatch
):
    import importlib
    from types import SimpleNamespace

    module = importlib.import_module(
        f"aima_ugc.adapters.providers.tikhub_test.{platform}_accounts_test"
    )
    summary_path = tmp_path / "run_summary.json"
    summary_path.write_text(
        json.dumps(
            {
                "status": "partial_success",
                "accounts": [
                    {
                        "status": "partial",
                        "posts_discovered": 1,
                        "stop_reason": "max_account_post_pages_reached",
                    }
                ],
                "partial_content_count": 0,
            }
        ),
        encoding="utf-8",
    )
    result = SimpleNamespace(
        run_summary_path=summary_path, content_count=1, workbook_path=tmp_path / "kept.xlsx"
    )
    result.workbook_path.write_bytes(b"retained-result")
    monkeypatch.setattr(module, f"run_{platform}_accounts", lambda **kwargs: result)
    if platform == "bilibili":
        monkeypatch.setattr(module, "ACCOUNTS", [tikhub_test.BilibiliAccountTarget(uid="100002")])
    with pytest.raises(RuntimeError, match="未完成|不完整|全量"):
        module.main()
    assert result.workbook_path.read_bytes() == b"retained-result"
