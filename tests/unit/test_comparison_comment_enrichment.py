"""五平台车型共现帖子评论补采的直接行为测试。"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import UUID

import pytest
from aima_ugc.adapters.providers.imports_test.comparison_comment_enrichment import (
    enrich_comments as enrichment_module,
)
from aima_ugc.adapters.providers.imports_test.comparison_comment_enrichment.enrich_comments import (
    enrich_comparison_comments,
)
from aima_ugc.adapters.providers.imports_test.comparison_comment_enrichment.models import (
    VehiclePairCommentRecordV1,
)
from aima_ugc.adapters.providers.imports_test.vehicle_pair_filter.filter_vehicle_pairs import (
    VehicleModelMention,
    VehicleModelReference,
    VehiclePairRecordV1,
    VehiclePairReference,
)
from aima_ugc.adapters.providers.tikhub import runtime as tikhub_runtime
from aima_ugc.adapters.providers.tikhub_test.core.config import TikHubTestConfig
from aima_ugc.contracts.analysis import UnifiedContentRecordV1
from aima_ugc.contracts.canonical import (
    CanonicalCommentV1,
    CanonicalContentV1,
    CanonicalMetricsV1,
    CanonicalSourceV1,
)
from aima_ugc.contracts.platform import PLATFORM_NAMES, PlatformName
from aima_ugc.modules.collection.providers.transport import (
    ProviderTransportRequest,
    ProviderTransportResponse,
)
from openpyxl import load_workbook
from pydantic import SecretStr


class _FakeTransport:
    """按测试 handler 返回确定 Provider 响应并记录实际 Runtime 请求。"""

    def __init__(self, handler: Any) -> None:
        self.handler = handler
        self.requests: list[ProviderTransportRequest] = []

    def send(self, request: ProviderTransportRequest) -> ProviderTransportResponse:
        """记录请求后把响应生成交给测试场景。"""

        self.requests.append(request)
        return self.handler(len(self.requests), request)


def _config() -> TikHubTestConfig:
    """构造不会触及真实 Secret 的测试 Provider 配置。"""

    return TikHubTestConfig(
        base_url="https://api.tikhub.dev",
        api_key=SecretStr("test-key"),
        timeout_seconds=1,
    )


def _pair_record(
    *,
    platform: PlatformName,
    external_id: str,
    alternate_ids: dict[str, str] | None = None,
) -> VehiclePairRecordV1:
    """构造一条最小合法的第二阶段车型共现记录。"""

    observed_at = datetime(2026, 9, 12, 0, 0, tzinfo=UTC)
    content = CanonicalContentV1(
        platform=platform,
        external_content_id=external_id,
        alternate_ids=alternate_ids or {},
        content_type="video" if platform != "xiaohongshu" else "image",
        title="元宇宙与Q3对比",
        text="同一帖子同时讨论两款车型",
        observed_at=observed_at,
        observed_fields=["title", "text"],
        source=CanonicalSourceV1(
            provider_name="manual",
            source_type="test",
            source_value="comment-enrichment-test",
            observed_at=observed_at,
        ),
    )
    return VehiclePairRecordV1(
        record=UnifiedContentRecordV1(
            content=content,
            matched_keywords=["元宇宙", "Q3"],
        ),
        matched_target_models=("元宇宙",),
        matched_competitor_models=(VehicleModelReference(brand="九号", model="Q3"),),
        matched_pairs=(
            VehiclePairReference(
                target_model="元宇宙",
                competitor_brand="九号",
                competitor_model="Q3",
            ),
        ),
        model_mentions=(
            VehicleModelMention(
                brand="爱玛",
                model="元宇宙",
                fields=("title",),
                matched_aliases=("元宇宙",),
            ),
            VehicleModelMention(
                brand="九号",
                model="Q3",
                fields=("title",),
                matched_aliases=("Q3",),
            ),
        ),
    )


def _write_records(path: Path, records: list[VehiclePairRecordV1]) -> None:
    """按第二阶段真实 JSONL 形态写测试输入。"""

    path.write_text(
        "".join(f"{record.model_dump_json()}\n" for record in records),
        encoding="utf-8",
    )


def _fake_comment(
    *,
    platform: PlatformName,
    content_id: str,
    comment_id: str,
    root_comment_id: str | None,
    is_root: bool,
    reply_count: int | None,
    observed_at: datetime,
    raw_artifact_id: UUID,
    item_locator: str,
) -> CanonicalCommentV1:
    """生成满足 Canonical 关系约束的一级评论或回复。"""

    return CanonicalCommentV1(
        platform=platform,
        external_content_id=content_id,
        external_comment_id=comment_id,
        root_comment_id=comment_id if is_root else root_comment_id,
        parent_comment_id=None,
        text=f"comment-{comment_id}",
        observed_at=observed_at,
        metrics=CanonicalMetricsV1(reply_count=reply_count),
        source=CanonicalSourceV1(
            provider_name="tikhub",
            operation="test-comment",
            raw_artifact_id=raw_artifact_id,
            source_type="test",
            source_value=content_id,
            item_locator=item_locator,
            observed_at=observed_at,
        ),
        observed_fields=["root_comment_id", "parent_comment_id", "text", "metrics.reply_count"],
    )


def _install_generic_runtime(monkeypatch: pytest.MonkeyPatch) -> None:
    """保留真实 build call，只替换响应提取/Mapper/分页以隔离外网。"""

    monkeypatch.setattr(
        tikhub_runtime,
        "extract_comment_items",
        lambda platform, body: tuple(body.get("items", ())),
    )
    monkeypatch.setattr(
        tikhub_runtime,
        "extract_sub_comment_items",
        lambda platform, body: tuple(body.get("items", ())),
    )

    def map_comment(
        *,
        platform: PlatformName,
        raw: dict[str, Any],
        context: Any,
        item_locator: str,
        is_root: bool,
    ) -> CanonicalCommentV1:
        content_id = context.external_content_id
        if not isinstance(content_id, str):
            raise AssertionError("测试 Mapper 必须收到 external_content_id")
        comment_id = str(raw["id"])
        reply_count = raw.get("reply_count")
        return _fake_comment(
            platform=platform,
            content_id=content_id,
            comment_id=comment_id,
            root_comment_id=context.root_comment_id,
            is_root=is_root,
            reply_count=reply_count if isinstance(reply_count, int) else None,
            observed_at=context.observed_at,
            raw_artifact_id=context.raw_artifact_id,
            item_locator=item_locator,
        )

    monkeypatch.setattr(tikhub_runtime, "map_comment", map_comment)


def test_five_platforms_use_runtime_typed_identity_and_fetch_replies(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """统一入口应按五平台正式 typed identity 发 Comments/SubComments。"""

    _install_generic_runtime(monkeypatch)
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_sub_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )

    records = [
        _pair_record(
            platform="xiaohongshu",
            external_id="canonical-xiaohongshu",
            alternate_ids={"note_id": "typed-note-id"},
        ),
        _pair_record(
            platform="douyin",
            external_id="canonical-douyin",
            alternate_ids={"aweme_id": "typed-aweme-id"},
        ),
        _pair_record(
            platform="weibo",
            external_id="canonical-weibo",
            alternate_ids={"status_id": "5300602615631073"},
        ),
        _pair_record(
            platform="bilibili",
            external_id="canonical-bilibili",
            alternate_ids={"bv_id": "BV1typedid"},
        ),
        _pair_record(
            platform="kuaishou",
            external_id="canonical-kuaishou",
            alternate_ids={"photo_id": "typed-photo-id"},
        ),
    ]
    input_path = tmp_path / "comparison_posts.jsonl"
    _write_records(input_path, records)

    def handler(call_no: int, request: ProviderTransportRequest) -> ProviderTransportResponse:
        if call_no % 2 == 1:
            body = {"items": [{"id": f"root-{call_no}", "reply_count": 1}]}
        else:
            body = {"items": [{"id": f"reply-{call_no}", "reply_count": 0}]}
        return ProviderTransportResponse(status_code=200, body=body)

    transport = _FakeTransport(handler)
    summary = enrich_comparison_comments(
        input_path=input_path,
        output_root=tmp_path / "output",
        run_id="five-platforms",
        provider_config=_config(),
        transport=transport,
    )

    assert summary.rows_seen == 5
    assert summary.rows_complete == 5
    assert summary.root_comment_count == 5
    assert summary.reply_count == 5
    assert summary.request_count == 10
    assert summary.platform_counts == {platform: 1 for platform in PLATFORM_NAMES}

    expected_ids = (
        "typed-note-id",
        "typed-aweme-id",
        "5300602615631073",
        "BV1typedid",
        "typed-photo-id",
    )
    for index, expected_id in enumerate(expected_ids):
        platform = records[index].record.content.platform
        comment_request = transport.requests[index * 2]
        reply_request = transport.requests[index * 2 + 1]
        assert expected_id in comment_request.params.values()
        if platform == "weibo":
            assert reply_request.path.endswith("/fetch_post_sub_comments")
            assert expected_id not in reply_request.params.values()
        else:
            assert expected_id in reply_request.params.values()
        assert any(str(value).startswith("root-") for value in reply_request.params.values())

    output_lines = [
        line
        for line in summary.output_jsonl_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert len(output_lines) == 5
    for line in output_lines:
        enriched = VehiclePairCommentRecordV1.model_validate_json(line)
        assert enriched.comment_fetch.coverage == "complete"
        assert len(enriched.comments) == 2

    workbook = load_workbook(summary.workbook_path, read_only=True, data_only=True)
    try:
        assert sum(1 for _ in workbook["内容"].iter_rows()) == 6
        assert sum(1 for _ in workbook["评论"].iter_rows()) == 11
    finally:
        workbook.close()


def test_full_fetch_exceeds_old_comment_and_reply_page_limits(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """全量模式应跨过旧调试入口 20/10 页上限并严格去重。"""

    _install_generic_runtime(monkeypatch)
    comment_page = 0
    reply_page = 0

    def handler(call_no: int, request: ProviderTransportRequest) -> ProviderTransportResponse:
        nonlocal comment_page, reply_page
        if "comment_id" in request.params:
            reply_page += 1
            return ProviderTransportResponse(
                status_code=200,
                body={
                    "stage": "replies",
                    "page": reply_page,
                    "items": [{"id": f"reply-{reply_page}", "reply_count": 0}],
                    "has_more": reply_page < 11,
                },
            )
        comment_page += 1
        item = {
            "id": f"root-{comment_page}",
            "reply_count": 1 if comment_page == 1 else 0,
        }
        items = [item, dict(item)] if comment_page == 1 else [item]
        return ProviderTransportResponse(
            status_code=200,
            body={
                "stage": "comments",
                "page": comment_page,
                "items": items,
                "has_more": comment_page < 21,
            },
        )

    def advance_comments(**kwargs: Any) -> tikhub_runtime.TikHubPageAdvance:
        body = kwargs["body"]
        page = int(body["page"])
        if not body["has_more"]:
            return tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted")
        return tikhub_runtime.TikHubPageAdvance(
            {"cursor": f"comment-{page}", "index": page},
            None,
        )

    def advance_replies(**kwargs: Any) -> tikhub_runtime.TikHubPageAdvance:
        body = kwargs["body"]
        page = int(body["page"])
        if not body["has_more"]:
            return tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted")
        return tikhub_runtime.TikHubPageAdvance(
            {"cursor": f"reply-{page}", "index": page + 1},
            None,
        )

    monkeypatch.setattr(tikhub_runtime, "advance_comments", advance_comments)
    monkeypatch.setattr(tikhub_runtime, "advance_sub_comments", advance_replies)

    input_path = tmp_path / "comparison_posts.jsonl"
    _write_records(
        input_path,
        [_pair_record(platform="xiaohongshu", external_id="note-all-pages")],
    )
    transport = _FakeTransport(handler)
    summary = enrich_comparison_comments(
        input_path=input_path,
        output_root=tmp_path / "output",
        run_id="all-pages",
        provider_config=_config(),
        transport=transport,
    )

    assert comment_page == 21
    assert reply_page == 11
    assert summary.root_comment_count == 21
    assert summary.reply_count == 11
    assert summary.request_count == 32
    enriched = VehiclePairCommentRecordV1.model_validate_json(
        summary.output_jsonl_path.read_text(encoding="utf-8").strip()
    )
    assert enriched.comment_fetch.coverage == "complete"
    assert len(enriched.comments) == 32
    assert len({comment.external_comment_id for comment in enriched.comments}) == 32


def test_permanent_reply_error_publishes_partial_but_retryable_error_aborts(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """永久条目 4xx 应显式 partial；429 等运行级错误必须阻止正式发布。"""

    _install_generic_runtime(monkeypatch)
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_sub_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )
    input_path = tmp_path / "comparison_posts.jsonl"
    _write_records(
        input_path,
        [_pair_record(platform="xiaohongshu", external_id="note-errors")],
    )

    def partial_handler(
        call_no: int, request: ProviderTransportRequest
    ) -> ProviderTransportResponse:
        if call_no == 1:
            return ProviderTransportResponse(
                status_code=200,
                body={"items": [{"id": "root-error", "reply_count": 1}]},
            )
        return ProviderTransportResponse(status_code=404, body={"detail": "not found"})

    partial = enrich_comparison_comments(
        input_path=input_path,
        output_root=tmp_path / "partial-output",
        run_id="partial",
        provider_config=_config(),
        transport=_FakeTransport(partial_handler),
    )
    assert partial.rows_partial == 1
    assert partial.root_comment_count == 1
    assert partial.reply_count == 0
    assert partial.failure_count == 1
    partial_record = VehiclePairCommentRecordV1.model_validate_json(
        partial.output_jsonl_path.read_text(encoding="utf-8").strip()
    )
    assert partial_record.comment_fetch.coverage == "partial"
    assert partial_record.comment_fetch.failures[0].status_code == 404

    def retryable_handler(
        call_no: int, request: ProviderTransportRequest
    ) -> ProviderTransportResponse:
        return ProviderTransportResponse(status_code=429, body={"detail": "rate limited"})

    fatal_root = tmp_path / "fatal-output"
    with pytest.raises(RuntimeError, match="HTTP 429"):
        enrich_comparison_comments(
            input_path=input_path,
            output_root=fatal_root,
            run_id="fatal",
            provider_config=_config(),
            transport=_FakeTransport(retryable_handler),
        )
    assert not (fatal_root / "runs" / "fatal").exists()
    assert (fatal_root / ".staging-fatal").is_dir()


def test_invalid_input_fails_before_any_provider_request(tmp_path: Path) -> None:
    """坏 JSONL 必须在网络阶段开始前失败，避免产生无意义 Provider 请求。"""

    input_path = tmp_path / "comparison_posts.jsonl"
    input_path.write_text("{broken-json\n", encoding="utf-8")
    transport = _FakeTransport(
        lambda call_no, request: ProviderTransportResponse(status_code=200, body={})
    )
    output_root = tmp_path / "output"
    with pytest.raises(ValueError, match="第 1 行"):
        enrich_comparison_comments(
            input_path=input_path,
            output_root=output_root,
            run_id="invalid",
            provider_config=_config(),
            transport=transport,
        )
    assert transport.requests == []
    assert not (output_root / "runs" / "invalid").exists()


def test_weibo_hash_identity_is_unavailable_without_provider_request(tmp_path: Path) -> None:
    """缺少 status_id 的微博文章哈希不能冒充状态 ID 请求评论接口。"""

    input_path = tmp_path / "comparison_posts.jsonl"
    _write_records(
        input_path,
        [
            _pair_record(
                platform="weibo",
                external_id="25736ee699f697a90c3ef14c6447ccfc",
                alternate_ids={"source_article_id": "25736ee699f697a90c3ef14c6447ccfc"},
            )
        ],
    )
    transport = _FakeTransport(lambda call_no, request: pytest.fail("不应发送微博评论请求"))
    summary = enrich_comparison_comments(
        input_path=input_path,
        output_root=tmp_path / "output",
        run_id="unsupported-weibo-identity",
        provider_config=_config(),
        transport=transport,
    )

    assert transport.requests == []
    assert summary.rows_unavailable == 1
    assert summary.request_count == 0
    enriched = VehiclePairCommentRecordV1.model_validate_json(
        summary.output_jsonl_path.read_text(encoding="utf-8").strip()
    )
    assert enriched.comment_fetch.root_stop_reason == "unsupported_weibo_comment_identity"


def test_reply_shortfall_is_published_as_partial_instead_of_crashing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Provider 少返回已知回复时应保留结果和缺口证据，而不是触发模型异常。"""

    _install_generic_runtime(monkeypatch)
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_sub_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )

    def handler(call_no: int, request: ProviderTransportRequest) -> ProviderTransportResponse:
        if call_no == 1:
            return ProviderTransportResponse(
                status_code=200,
                body={"items": [{"id": "root-short", "reply_count": 2}]},
            )
        return ProviderTransportResponse(
            status_code=200,
            body={"items": [{"id": "reply-only", "reply_count": 0}]},
        )

    input_path = tmp_path / "comparison_posts.jsonl"
    _write_records(
        input_path,
        [_pair_record(platform="xiaohongshu", external_id="note-shortfall")],
    )
    summary = enrich_comparison_comments(
        input_path=input_path,
        output_root=tmp_path / "output",
        run_id="reply-shortfall",
        provider_config=_config(),
        transport=_FakeTransport(handler),
    )

    assert summary.rows_partial == 1
    enriched = VehiclePairCommentRecordV1.model_validate_json(
        summary.output_jsonl_path.read_text(encoding="utf-8").strip()
    )
    assert enriched.comment_fetch.coverage == "partial"
    assert len(enriched.comment_fetch.reply_shortfalls) == 1
    assert enriched.comment_fetch.reply_shortfalls[0].root_comment_id == "root-short"
    assert enriched.comment_fetch.reply_shortfalls[0].reported_count == 2
    assert enriched.comment_fetch.reply_shortfalls[0].observed_count == 1


def test_comment_identity_mismatch_is_filtered_and_recorded(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """串帖评论不得进入结果，但同页合法评论仍应保存并把覆盖降为 partial。"""

    _install_generic_runtime(monkeypatch)
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )

    def map_comment_with_observed_content(
        *,
        platform: PlatformName,
        raw: dict[str, Any],
        context: Any,
        item_locator: str,
        is_root: bool,
    ) -> CanonicalCommentV1:
        return _fake_comment(
            platform=platform,
            content_id=str(raw["content_id"]),
            comment_id=str(raw["id"]),
            root_comment_id=context.root_comment_id,
            is_root=is_root,
            reply_count=0,
            observed_at=context.observed_at,
            raw_artifact_id=context.raw_artifact_id,
            item_locator=item_locator,
        )

    monkeypatch.setattr(tikhub_runtime, "map_comment", map_comment_with_observed_content)
    input_path = tmp_path / "comparison_posts.jsonl"
    _write_records(
        input_path,
        [_pair_record(platform="xiaohongshu", external_id="note-expected")],
    )
    transport = _FakeTransport(
        lambda call_no, request: ProviderTransportResponse(
            status_code=200,
            body={
                "items": [
                    {"id": "wrong-root", "content_id": "note-other"},
                    {"id": "right-root", "content_id": "note-expected"},
                ]
            },
        )
    )
    summary = enrich_comparison_comments(
        input_path=input_path,
        output_root=tmp_path / "output",
        run_id="identity-mismatch",
        provider_config=_config(),
        transport=transport,
    )

    assert summary.rows_partial == 1
    enriched = VehiclePairCommentRecordV1.model_validate_json(
        summary.output_jsonl_path.read_text(encoding="utf-8").strip()
    )
    assert [comment.external_comment_id for comment in enriched.comments] == ["right-root"]
    assert len(enriched.comment_fetch.identity_mismatches) == 1
    mismatch = enriched.comment_fetch.identity_mismatches[0]
    assert mismatch.external_comment_id == "wrong-root"
    assert mismatch.expected_external_content_id == "note-expected"
    assert mismatch.observed_external_content_id == "note-other"


def test_completed_staging_is_preserved_after_publish_failure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """最终目录改名失败时必须保留已完成 staging，避免丢失付费抓取结果。"""

    _install_generic_runtime(monkeypatch)
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )
    input_path = tmp_path / "comparison_posts.jsonl"
    _write_records(
        input_path,
        [_pair_record(platform="xiaohongshu", external_id="note-publish")],
    )
    real_replace = enrichment_module.os.replace

    def fail_final_publish(source: object, destination: object) -> None:
        if Path(source).name == ".staging-publish-failure":
            raise PermissionError(5, "拒绝访问", str(source), str(destination))
        real_replace(source, destination)

    monkeypatch.setattr(enrichment_module.os, "replace", fail_final_publish)
    output_root = tmp_path / "output"
    with pytest.raises(PermissionError):
        enrich_comparison_comments(
            input_path=input_path,
            output_root=output_root,
            run_id="publish-failure",
            provider_config=_config(),
            transport=_FakeTransport(
                lambda call_no, request: ProviderTransportResponse(
                    status_code=200,
                    body={"items": [{"id": "root-publish", "reply_count": 0}]},
                )
            ),
        )

    staging_dir = output_root / ".staging-publish-failure"
    assert staging_dir.is_dir()
    assert (staging_dir / "run_summary.json").is_file()
    assert (staging_dir / "comparison_posts_with_comments.jsonl").is_file()


def test_incomplete_staging_resumes_strict_input_prefix_without_refetch(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """同一 run 重启应复用已落盘行，并从未完成输入继续请求且不覆盖 Raw。"""

    _install_generic_runtime(monkeypatch)
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )
    input_path = tmp_path / "comparison_posts.jsonl"
    _write_records(
        input_path,
        [
            _pair_record(platform="xiaohongshu", external_id="note-resumed"),
            _pair_record(platform="xiaohongshu", external_id="note-pending"),
        ],
    )

    first_transport = _FakeTransport(
        lambda call_no, request: (
            ProviderTransportResponse(
                status_code=200,
                body={"items": [{"id": "root-resumed", "reply_count": 0}]},
            )
            if call_no == 1
            else ProviderTransportResponse(status_code=429, body={"detail": "rate limited"})
        )
    )
    output_root = tmp_path / "output"
    with pytest.raises(RuntimeError, match="HTTP 429"):
        enrich_comparison_comments(
            input_path=input_path,
            output_root=output_root,
            run_id="resume-run",
            provider_config=_config(),
            transport=first_transport,
        )

    staging_jsonl = output_root / ".staging-resume-run" / "comparison_posts_with_comments.jsonl"
    assert staging_jsonl.read_text(encoding="utf-8").count("\n") == 1
    assert (staging_jsonl.parent / "resume_manifest.json").is_file()

    second_transport = _FakeTransport(
        lambda call_no, request: ProviderTransportResponse(
            status_code=200,
            body={"items": [{"id": "root-pending", "reply_count": 0}]},
        )
    )
    resumed = enrich_comparison_comments(
        input_path=input_path,
        output_root=output_root,
        run_id="resume-run",
        provider_config=_config(),
        transport=second_transport,
    )

    assert resumed.rows_seen == 2
    assert resumed.rows_cached == 1
    assert resumed.rows_fetched == 1
    assert len(second_transport.requests) == 1
    assert resumed.output_jsonl_path.read_text(encoding="utf-8").count("\n") == 2
    raw_names = sorted(
        path.name
        for path in (
            resumed.run_dir / "provider" / "xiaohongshu" / "runs" / "comments" / "raw"
        ).glob("*.json")
    )
    assert raw_names == [
        "0001_get_note_comments.json",
        "0002_get_note_comments.json",
        "0003_get_note_comments.json",
    ]


def test_empty_staging_rejects_changed_input_before_resume_request(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """首行失败只留下 Raw 时，也必须靠输入摘要阻止另一批数据误续跑。"""

    _install_generic_runtime(monkeypatch)
    input_path = tmp_path / "comparison_posts.jsonl"
    _write_records(
        input_path,
        [_pair_record(platform="xiaohongshu", external_id="note-original")],
    )
    output_root = tmp_path / "output"
    with pytest.raises(RuntimeError, match="HTTP 429"):
        enrich_comparison_comments(
            input_path=input_path,
            output_root=output_root,
            run_id="empty-resume",
            provider_config=_config(),
            transport=_FakeTransport(
                lambda call_no, request: ProviderTransportResponse(status_code=429, body={})
            ),
        )

    _write_records(
        input_path,
        [_pair_record(platform="xiaohongshu", external_id="note-different")],
    )
    second_transport = _FakeTransport(
        lambda call_no, request: pytest.fail("输入变化时不应发送续跑请求")
    )
    with pytest.raises(ValueError, match="续跑输入文件已变化"):
        enrich_comparison_comments(
            input_path=input_path,
            output_root=output_root,
            run_id="empty-resume",
            provider_config=_config(),
            transport=second_transport,
        )
    assert second_transport.requests == []
