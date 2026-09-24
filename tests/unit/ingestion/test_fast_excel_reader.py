"""加速 Reader 与原 openpyxl 值语义的边界回归。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from aima_ugc.adapters.providers.imports.excel_profile import (
    AIMA_MONITORING_EXCEL_V1,
    get_excel_import_profile,
)
from aima_ugc.adapters.providers.imports.excel_reader import iter_excel_rows
from openpyxl import Workbook, load_workbook


def test_fast_reader_preserves_dates_blank_rows_and_empty_strings(tmp_path: Path) -> None:
    path = tmp_path / "mixed-values.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    sheet.append(
        [
            "小红书",
            "第一条",
            "",
            "作者",
            datetime(2026, 9, 24, 8, 30),
            "https://www.xiaohongshu.com/explore/fast-reader-1",
        ]
    )
    sheet.append([None] * 6)
    sheet.append(
        [
            "小红书",
            "第二条",
            "=1+1",
            "作者",
            45924,
            "https://www.xiaohongshu.com/explore/fast-reader-2",
        ]
    )
    workbook.save(path)
    workbook.close()

    actual = tuple(
        iter_excel_rows(
            path,
            profile=get_excel_import_profile(AIMA_MONITORING_EXCEL_V1),
        )
    )
    legacy = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = legacy["文章"]
        sheet.reset_dimensions()
        raw = tuple(sheet.iter_rows(values_only=True))
    finally:
        legacy.close()

    assert tuple(item.row_number for item in actual) == (2, 4)
    assert actual[0].values["内文"] == raw[1][2]
    assert actual[0].values["出版日期"] == raw[1][4]
    assert actual[1].values["内文"] is None
    assert actual[1].values["出版日期"] == raw[3][4]


def test_fast_reader_preserves_shared_strings(tmp_path: Path) -> None:
    """Excel 常见的 sharedStrings 索引仍由已加载 Workbook 统一解释。"""

    path = tmp_path / "shared.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    sheet.append(
        [
            "小红书",
            "爱玛共享字符串",
            "正文",
            "作者",
            datetime(2026, 9, 24, 9),
            "https://www.xiaohongshu.com/explore/shared-1",
        ]
    )
    workbook.save(path)
    workbook.close()

    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    root = ElementTree.fromstring(members["xl/worksheets/sheet1.xml"])
    shared_cell = root.find(f".//{namespace}c[@r='A2']")
    assert shared_cell is not None
    shared_cell.clear()
    shared_cell.set("r", "A2")
    shared_cell.set("t", "s")
    ElementTree.SubElement(shared_cell, f"{namespace}v").text = "0"
    members["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root, encoding="utf-8")
    types_root = ElementTree.fromstring(members["[Content_Types].xml"])
    ElementTree.SubElement(
        types_root,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        {
            "PartName": "/xl/sharedStrings.xml",
            "ContentType": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
            ),
        },
    )
    members["[Content_Types].xml"] = ElementTree.tostring(types_root, encoding="utf-8")
    members["xl/sharedStrings.xml"] = (
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'count="1" uniqueCount="1"><si><t>小红书</t></si></sst>'
    ).encode()
    with ZipFile(path, "w") as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)

    legacy = load_workbook(path, read_only=True, data_only=True)
    try:
        expected = next(legacy["文章"].iter_rows(min_row=2, max_row=2, values_only=True))
    finally:
        legacy.close()
    actual = tuple(
        iter_excel_rows(path, profile=get_excel_import_profile(AIMA_MONITORING_EXCEL_V1))
    )
    assert len(actual) == 1
    assert actual[0].values["媒体名称（中文）"] == expected[0] == "小红书"


def test_fast_reader_rejects_cell_coordinate_without_column(tmp_path: Path) -> None:
    path = tmp_path / "bad-coordinate.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    sheet.append(
        [
            "小红书",
            "爱玛无效坐标",
            "正文",
            "作者",
            datetime(2026, 9, 24, 9),
            "https://www.xiaohongshu.com/explore/bad-coordinate",
        ]
    )
    workbook.save(path)
    workbook.close()
    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    root = ElementTree.fromstring(members["xl/worksheets/sheet1.xml"])
    cell = root.find(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c[@r='A2']")
    assert cell is not None
    cell.set("r", "2")
    members["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root, encoding="utf-8")
    with ZipFile(path, "w") as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    with pytest.raises(ValueError, match="坐标无效"):
        tuple(iter_excel_rows(path, profile=get_excel_import_profile(AIMA_MONITORING_EXCEL_V1)))
