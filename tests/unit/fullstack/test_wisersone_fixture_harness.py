"""网站 Fixture 和全栈辅助入口的隔离边界，不访问数据库或外部网站。"""

from __future__ import annotations

import json
import os
import runpy
import subprocess
import sys
from pathlib import Path
from uuid import uuid4

import pytest
from aima_ugc.adapters.providers.wisersone.models import ExportCancelled
from openpyxl import load_workbook


@pytest.mark.parametrize(
    ("script", "opt_in"),
    [
        ("assert_wisersone_plan_workflow.py", "AIMA_FULLSTACK_SEED"),
        ("fake_tikhub_comment_worker.py", "AIMA_FULLSTACK_FAKE_TIKHUB"),
    ],
)
def test_fullstack_entry_refuses_database_access_without_opt_in(script: str, opt_in: str) -> None:
    """无显式测试授权时必须先拒绝；非法数据库端口不得掩盖入口门禁。"""
    root = Path(__file__).resolve().parents[3]
    environ = dict(os.environ)
    environ.pop(opt_in, None)
    environ.pop("SSLKEYLOGFILE", None)
    environ["AIMA_DB_PORT"] = "1"
    completed = subprocess.run(
        [sys.executable, "-X", "utf8", f"tests/fullstack/{script}"],
        cwd=root,
        env=environ,
        capture_output=True,
        text=True,
        check=False,
        timeout=20,
    )
    assert completed.returncode != 0
    expected = "AIMA_FULLSTACK_SEED=1" if opt_in == "AIMA_FULLSTACK_SEED" else "必须显式启用"
    assert expected in completed.stderr
    assert "OperationalError" not in completed.stderr


def _exporter_type():  # type: ignore[no-untyped-def]
    root = Path(__file__).resolve().parents[3]
    return runpy.run_path(str(root / "tests/fullstack/fake_tikhub_comment_worker.py"))[
        "_WisersOneFixtureExporter"
    ]


def test_wisersone_helper_refuses_default_host_directory_before_database_access() -> None:
    """单独的测试授权不能允许下载文件落入宿主默认输入目录。"""
    root = Path(__file__).resolve().parents[3]
    environ = dict(os.environ)
    environ["AIMA_FULLSTACK_SEED"] = "1"
    environ["AIMA_DB_PORT"] = "1"
    environ.pop("AIMA_WISERSONE_INPUT_DIR", None)
    environ.pop("SSLKEYLOGFILE", None)
    completed = subprocess.run(
        [sys.executable, "-X", "utf8", "tests/fullstack/assert_wisersone_plan_workflow.py"],
        cwd=root,
        env=environ,
        capture_output=True,
        text=True,
        check=False,
        timeout=20,
    )
    assert completed.returncode != 0 and "AIMA_WISERSONE_INPUT_DIR" in completed.stderr
    assert "OperationalError" not in completed.stderr


def test_external_fixture_survives_restart_with_one_submission_and_real_xlsx(
    tmp_path: Path,
) -> None:
    """回执跨对象保留同一任务，网站响应同时给出品牌匹配和不匹配行。"""
    exporter_type = _exporter_type()
    site = exporter_type(fixture_root=tmp_path / "website")
    run_id = uuid4()
    callbacks = []
    task = site.submit(
        run_id,
        before_submit=lambda: callbacks.append("before"),
        on_submitted=lambda _task: callbacks.append("confirmed"),
    )
    assert callbacks == ["before", "confirmed"]
    restarted = exporter_type(fixture_root=site.auth_dir)
    assert restarted.saved_task(run_id) == task
    assert (
        restarted.submit(run_id, before_submit=lambda: pytest.fail("已有回执不得再次提交")) == task
    )
    destination = tmp_path / "staging" / "wisersone.xlsx"
    progress = restarted.poll(task, destination, download=False)
    assert progress.ready and progress.file is None and not destination.exists()
    result = restarted.poll(task, destination)
    assert result.ready and result.percent == 100 and result.file == destination
    book = load_workbook(destination, read_only=True)
    try:
        rows = tuple(book["文章"].iter_rows(values_only=True))
        assert len(rows) == 3
        assert rows[1][1] == f"爱玛 WisersOne 全栈 {run_id}"
        assert rows[2][1] == f"fixture-unrelated-{run_id}"
        assert rows[1][-1] != rows[2][-1]
    finally:
        book.close()
    events = [
        json.loads(line)
        for line in (site.auth_dir / f"events-{run_id.hex}.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
    ]
    assert [event["action"] for event in events] == ["submit", "poll", "download"]
    assert all(event["task_id"] == task.task_id for event in events)


def test_cancelled_fixture_does_not_submit_or_create_files(tmp_path: Path) -> None:
    """取消回调在外部发送与文件写入之前生效。"""
    site = _exporter_type()(fixture_root=tmp_path / "website")
    with pytest.raises(ExportCancelled):
        site.submit(uuid4(), cancelled=lambda: True)
    assert not site.auth_dir.exists()
