from __future__ import annotations

import base64
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from aima_ugc.platform.reporting import convert_markdown_to_docx
from aima_ugc.platform.reporting.docx_package import DocxBuilder, verify_docx
from openpyxl import load_workbook

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"

_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4//8/AAX+Av4N70a4AAAAAElFTkSuQmCC"
)


def test_docx_line_break_is_nested_in_run(tmp_path: Path) -> None:
    output_path = tmp_path / "report.docx"
    builder = DocxBuilder()
    builder.add_paragraph("第一行\n第二行")
    builder.save(output_path)
    verify_docx(output_path, expected_charts=0)

    with zipfile.ZipFile(output_path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))

    assert root.findall(f".//{{{_W}}}p/{{{_W}}}br") == []
    assert len(root.findall(f".//{{{_W}}}r/{{{_W}}}br")) == 1


def test_word_page_is_a4_landscape_with_report_margins(tmp_path: Path) -> None:
    output_path = tmp_path / "report.docx"
    builder = DocxBuilder()
    builder.add_paragraph("横向 A4")
    builder.save(output_path)

    with zipfile.ZipFile(output_path) as archive:
        document = ET.fromstring(archive.read("word/document.xml"))

    page_size = document.find(f".//{{{_W}}}sectPr/{{{_W}}}pgSz")
    assert page_size is not None
    assert page_size.get(f"{{{_W}}}w") == "16838"
    assert page_size.get(f"{{{_W}}}h") == "11906"
    assert page_size.get(f"{{{_W}}}orient") == "landscape"

    margins = document.find(f".//{{{_W}}}sectPr/{{{_W}}}pgMar")
    assert margins is not None
    for side in ("top", "right", "bottom", "left"):
        assert margins.get(f"{{{_W}}}{side}") == "850"


def test_word_chart_is_editable_office_chart_with_embedded_workbook(tmp_path: Path) -> None:
    markdown_path = tmp_path / "report.md"
    output_path = tmp_path / "report.docx"
    markdown_path.write_text(
        "# 趋势\n\n"
        "```mermaid\n"
        "xychart-beta\n"
        '    title "平台每日声量"\n'
        '    %% series ["抖音", "小红书"]\n'
        '    x-axis ["08-18", "08-19"]\n'
        '    y-axis "数量" 0 --> 10\n'
        "    line [3, 5]\n"
        "    line [2, 4]\n"
        "```\n",
        encoding="utf-8",
    )

    summary = convert_markdown_to_docx(markdown_path, output_path)
    assert summary.chart_count == 1

    with zipfile.ZipFile(output_path) as archive:
        names = set(archive.namelist())
        assert "word/charts/chart1.xml" in names
        assert "word/charts/_rels/chart1.xml.rels" in names
        assert "word/embeddings/chart1.xlsx" in names
        assert not any(name.startswith("word/media/chart-") for name in names)

        document = ET.fromstring(archive.read("word/document.xml"))
        assert len(document.findall(f".//{{{_C}}}chart")) == 1

        chart_xml_bytes = archive.read("word/charts/chart1.xml")
        chart_xml = chart_xml_bytes.decode("utf-8")
        assert "抖音" in chart_xml
        assert "小红书" in chart_xml
        chart = ET.fromstring(chart_xml_bytes)
        assert chart.find(f"./{{{_C}}}spPr/{{{_A}}}ln/{{{_A}}}noFill") is not None
        line_outlines = chart.findall(f".//{{{_C}}}lineChart/{{{_C}}}ser/{{{_C}}}spPr/{{{_A}}}ln")
        assert len(line_outlines) == 2
        assert all(outline.get("w") == "28575" for outline in line_outlines)

        data_labels = chart.find(f".//{{{_C}}}lineChart/{{{_C}}}dLbls")
        assert data_labels is not None
        assert data_labels.find(f"./{{{_C}}}showVal").get("val") == "1"
        assert data_labels.find(f"./{{{_C}}}numFmt").get("formatCode") == "#,##0"

        gridline = chart.find(
            f".//{{{_C}}}valAx/{{{_C}}}majorGridlines/{{{_C}}}spPr/{{{_A}}}ln/"
            f"{{{_A}}}solidFill/{{{_A}}}srgbClr"
        )
        assert gridline is not None
        assert gridline.get("val") == "E5E7EB"

        workbook = load_workbook(
            BytesIO(archive.read("word/embeddings/chart1.xlsx")), data_only=False
        )
        try:
            sheet = workbook.active
            assert sheet["A1"].value == "日期/分类"
            assert sheet["B1"].value == "抖音"
            assert sheet["C1"].value == "小红书"
            assert sheet["A2"].value == "08-18"
            assert sheet["B2"].value == 3
            assert sheet["C3"].value == 4
        finally:
            workbook.close()


def test_word_horizontal_bar_has_native_value_labels(tmp_path: Path) -> None:
    markdown_path = tmp_path / "report.md"
    output_path = tmp_path / "report.docx"
    markdown_path.write_text(
        "# 一级议题\n\n"
        "```mermaid\n"
        "xychart-beta\n"
        '    title "一级议题 Top 分布"\n'
        '    %% series ["数量"]\n'
        "    %% bar-direction horizontal\n"
        '    x-axis ["品牌评价", "价格与价值", "外观设计"]\n'
        '    y-axis "数量" 0 --> 40000\n'
        "    bar [36235, 6805, 5678]\n"
        "```\n",
        encoding="utf-8",
    )

    convert_markdown_to_docx(markdown_path, output_path)

    with zipfile.ZipFile(output_path) as archive:
        chart = ET.fromstring(archive.read("word/charts/chart1.xml"))

    direction = chart.find(f".//{{{_C}}}barChart/{{{_C}}}barDir")
    assert direction is not None
    assert direction.get("val") == "bar"
    labels = chart.find(f".//{{{_C}}}barChart/{{{_C}}}dLbls")
    assert labels is not None
    assert labels.find(f"./{{{_C}}}showVal").get("val") == "1"
    assert labels.find(f"./{{{_C}}}dLblPos").get("val") == "outEnd"
    assert labels.find(f"./{{{_C}}}numFmt").get("formatCode") == "#,##0"


def test_word_pie_charts_have_no_outer_border_and_show_two_decimal_percentages(
    tmp_path: Path,
) -> None:
    markdown_path = tmp_path / "report.md"
    output_path = tmp_path / "report.docx"
    markdown_path.write_text(
        "# 情感结构\n\n"
        "```mermaid\n"
        "pie showData\n"
        '    title "情感结构"\n'
        '    "正面" : 10\n'
        '    "负面" : 2\n'
        "```\n\n"
        "```mermaid\n"
        "pie showData\n"
        '    title "平台分布"\n'
        '    "抖音" : 7\n'
        '    "小红书" : 3\n'
        "```\n",
        encoding="utf-8",
    )

    convert_markdown_to_docx(markdown_path, output_path)

    with zipfile.ZipFile(output_path) as archive:
        charts = [ET.fromstring(archive.read(f"word/charts/chart{index}.xml")) for index in (1, 2)]
    for chart in charts:
        assert chart.find(f"./{{{_C}}}spPr/{{{_A}}}ln/{{{_A}}}noFill") is not None
        number_format = chart.find(f".//{{{_C}}}pieChart/{{{_C}}}dLbls/{{{_C}}}numFmt")
        assert number_format is not None
        assert number_format.get("formatCode") == "0.00%"
        assert number_format.get("sourceLinked") == "0"


def test_word_table_uses_editorial_report_style(tmp_path: Path) -> None:
    output_path = tmp_path / "report.docx"
    builder = DocxBuilder()
    builder.add_table(
        ("情感标签", "内容量", "内容占比"),
        (("中性", "28,886", "65.31%"), ("正面", "14,710", "33.26%")),
    )
    builder.save(output_path)

    with zipfile.ZipFile(output_path) as archive:
        document = ET.fromstring(archive.read("word/document.xml"))

    table = document.find(f".//{{{_W}}}tbl")
    assert table is not None
    tbl_width = table.find(f"./{{{_W}}}tblPr/{{{_W}}}tblW")
    assert tbl_width is not None
    assert tbl_width.get(f"{{{_W}}}type") == "pct"
    assert tbl_width.get(f"{{{_W}}}w") == "5000"
    assert table.find(f"./{{{_W}}}tblPr/{{{_W}}}tblLayout").get(f"{{{_W}}}type") == "fixed"
    assert table.find(f"./{{{_W}}}tblPr/{{{_W}}}tblCellMar") is not None

    grid_columns = table.findall(f"./{{{_W}}}tblGrid/{{{_W}}}gridCol")
    assert len(grid_columns) == 3
    assert sum(int(column.get(f"{{{_W}}}w", "0")) for column in grid_columns) == 15138

    borders = table.find(f"./{{{_W}}}tblPr/{{{_W}}}tblBorders")
    assert borders is not None
    for side in ("left", "right", "insideV"):
        border = borders.find(f"./{{{_W}}}{side}")
        assert border is not None
        assert border.get(f"{{{_W}}}val") == "nil"
    inside_h = borders.find(f"./{{{_W}}}insideH")
    assert inside_h is not None
    assert inside_h.get(f"{{{_W}}}color") == "E5E7EB"

    rows = table.findall(f"./{{{_W}}}tr")
    header_shadings = rows[0].findall(f".//{{{_W}}}shd")
    assert all(shading.get(f"{{{_W}}}fill") != "1F4E78" for shading in header_shadings)
    assert rows[0].find(f"./{{{_W}}}trPr/{{{_W}}}tblHeader") is not None
    assert all(row.find(f"./{{{_W}}}trPr/{{{_W}}}cantSplit") is not None for row in rows)

    numeric_cell_alignment = (
        rows[1].findall(f"./{{{_W}}}tc")[1].find(f"./{{{_W}}}p/{{{_W}}}pPr/{{{_W}}}jc")
    )
    assert numeric_cell_alignment is not None
    assert numeric_cell_alignment.get(f"{{{_W}}}val") == "right"


def test_ranking_metadata_renders_editable_native_ranking(tmp_path: Path) -> None:
    markdown_path = tmp_path / "report.md"
    output_path = tmp_path / "report.docx"
    markdown_path.write_text(
        "# 一级议题\n\n"
        "<!-- aima:table-style=ranking -->\n"
        "| 一级标签 | 标签对数量 | 标签对占比 |\n"
        "| --- | ---: | ---: |\n"
        "| 品牌评价 | 36,235 | 57.96% |\n"
        "| 价格与价值 | 6,805 | 10.89% |\n",
        encoding="utf-8",
    )

    summary = convert_markdown_to_docx(markdown_path, output_path)
    assert summary.table_count >= 1

    with zipfile.ZipFile(output_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        document = ET.fromstring(archive.read("word/document.xml"))
        rels = ET.fromstring(archive.read("word/_rels/document.xml.rels"))

    for text in ("01", "品牌评价", "36,235", "57.96%", "02", "价格与价值"):
        assert text in document_xml
    assert len(document.findall(f".//{{{_W}}}tbl")) >= 3
    fills = {
        node.get(f"{{{_W}}}fill")
        for node in document.findall(f".//{{{_W}}}shd")
        if node.get(f"{{{_W}}}fill")
    }
    assert "2F6BCA" in fills
    assert "E8EEF7" in fills
    assert document.find(f".//{{{_W}}}keepNext") is not None
    assert not any(
        (relationship.get("Type") or "").endswith("/image")
        for relationship in rels.findall(f"./{{{_PKG_REL}}}Relationship")
    )


def test_markdown_png_is_packaged_as_docx_media(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    image_path = assets / "primary_topics_wordcloud.png"
    image_path.write_bytes(_PNG_1X1)
    markdown_path = tmp_path / "report.md"
    output_path = tmp_path / "report.docx"
    markdown_path.write_text(
        "# 一级议题词云\n\n![一级议题词云](assets/primary_topics_wordcloud.png)\n",
        encoding="utf-8",
    )

    summary = convert_markdown_to_docx(markdown_path, output_path)
    assert summary.image_count == 1

    with zipfile.ZipFile(output_path) as archive:
        names = set(archive.namelist())
        media = sorted(name for name in names if name.startswith("word/media/"))
        assert media == ["word/media/image1.png"]
        assert archive.read(media[0]) == _PNG_1X1
        content_types = archive.read("[Content_Types].xml").decode("utf-8")
        assert 'Extension="png"' in content_types
        document = ET.fromstring(archive.read("word/document.xml"))
        rels = ET.fromstring(archive.read("word/_rels/document.xml.rels"))

    blip = document.find(f".//{{{_A}}}blip")
    assert blip is not None
    image_rid = blip.get(f"{{{_R}}}embed")
    image_rels = [
        relationship
        for relationship in rels.findall(f"./{{{_PKG_REL}}}Relationship")
        if (relationship.get("Type") or "").endswith("/image")
    ]
    assert len(image_rels) == 1
    assert image_rels[0].get("Id") == image_rid
    assert image_rels[0].get("Target") == "media/image1.png"
