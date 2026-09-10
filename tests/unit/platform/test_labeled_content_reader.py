from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from aima_ugc.adapters.providers.imports import (
    LabeledContentReaderError,
    read_labeled_content_files,
    read_labeled_contents,
)
from openpyxl import Workbook


def _write_workbook(path: Path, *, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    content = workbook.active
    content.title = "内容"
    content.append(headers)
    for row in rows:
        content.append(row)
    detail = workbook.create_sheet("标签明细")
    detail.append(["内容ID", "平台"])
    detail.append(["should-not-be-read", "抖音"])
    workbook.create_sheet("评论")
    workbook.save(path)


def test_reads_content_sheet_and_deduplicates_platform_content_id(tmp_path: Path) -> None:
    path = tmp_path / "labeled.xlsx"
    headers = ["平台", "内容ID", "标题", "正文", "作者", "发布时间", "内容链接"]
    _write_workbook(
        path,
        headers=headers,
        rows=[
            [
                "抖音",
                "same",
                "标题",
                "正文",
                "用户",
                "2026-09-01 10:00:00",
                "https://example.test/1",
            ],
            [
                "抖音",
                "same",
                "重复",
                "重复",
                "用户",
                "2026-09-01 10:01:00",
                "https://example.test/2",
            ],
            ["小红书", "same", "另一个平台", "正文", "用户", datetime(2026, 9, 1, 11), None],
        ],
    )

    records, summary = read_labeled_contents(path)

    assert [record.content_id for record in records] == ["same", "same"]
    assert [record.platform for record in records] == ["抖音", "小红书"]
    assert records[1].published_at == "2026-09-01 11:00:00"
    assert summary.rows_seen == 3
    assert summary.rows_read == 2
    assert summary.duplicate_rows == 1


def test_requires_core_headers_and_explicit_content_sheet(tmp_path: Path) -> None:
    path = tmp_path / "invalid.xlsx"
    _write_workbook(
        path,
        headers=["平台", "标题", "正文", "作者"],
        rows=[["抖音", "标题", "正文", "用户"]],
    )

    with pytest.raises(LabeledContentReaderError, match="内容ID"):
        read_labeled_contents(path)

    with pytest.raises(LabeledContentReaderError, match="工作表为空|指定工作表"):
        read_labeled_contents(path, sheet_name="评论")


def test_reads_multiple_labeled_workbooks_and_deduplicates_across_files(tmp_path: Path) -> None:
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    headers = ["平台", "内容ID", "标题", "正文", "作者"]
    _write_workbook(
        first,
        headers=headers,
        rows=[
            ["抖音", "same", "第一份", "正文", "用户"],
            ["抖音", "first", "第一条", "正文", "用户"],
        ],
    )
    _write_workbook(
        second,
        headers=headers,
        rows=[
            ["抖音", "same", "重复", "重复正文", "用户"],
            ["小红书", "second", "第二条", "正文", "用户"],
        ],
    )

    records, summary = read_labeled_content_files((first, second))

    assert [record.content_id for record in records] == ["same", "first", "second"]
    assert summary.input_paths == (first, second)
    assert summary.rows_seen == 4
    assert summary.rows_read == 3
    assert summary.duplicate_rows == 1


def test_multiple_labeled_workbooks_require_at_least_one_file() -> None:
    with pytest.raises(LabeledContentReaderError, match="至少需要提供一个"):
        read_labeled_content_files(())


def test_missing_content_id_row_fails_closed(tmp_path: Path) -> None:
    path = tmp_path / "invalid-row.xlsx"
    _write_workbook(
        path,
        headers=["平台", "内容ID", "标题", "正文", "作者"],
        rows=[["抖音", None, "标题", "正文", "用户"]],
    )

    with pytest.raises(LabeledContentReaderError, match="内容ID不能为空"):
        read_labeled_contents(path)
