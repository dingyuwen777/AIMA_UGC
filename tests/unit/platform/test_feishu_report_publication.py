from __future__ import annotations

import hashlib
import json
import uuid
from collections.abc import Mapping
from datetime import date
from pathlib import Path

import httpx
import pytest
from aima_ugc.adapters.feishu import (
    FeishuApiError,
    FeishuReportPublisher,
    FeishuReportPublisherConfig,
    load_feishu_report_publisher_config,
)
from aima_ugc.adapters.feishu.report_publisher import _ImportResult, _table_cell_text_block
from aima_ugc.bootstrap import feishu_report_publication as publication_module
from aima_ugc.bootstrap.feishu_report_publication import publish_all_report_to_feishu
from aima_ugc.platform.config import PlatformSettings
from aima_ugc.platform.reporting import (
    ChartSpec,
    build_editable_chart_workbook,
    build_feishu_native_document,
)
from aima_ugc.platform.reporting.chart_png import render_chart_png
from aima_ugc.platform.reporting.feishu_native_document import (
    FeishuNativeBlock,
    FeishuNativeDocument,
)
from openpyxl import load_workbook
from pydantic import SecretStr


def _spec() -> ChartSpec:
    return ChartSpec(
        kind="bar",
        title="平台声量",
        categories=("抖音", "小红书"),
        series=((12.0, 8.0),),
        series_names=("内容量",),
    )


def _report_markdown(tmp_path: Path) -> Path:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "wordcloud.png").write_bytes(render_chart_png(_spec())[0])
    path = tmp_path / "report.md"
    path.write_text(
        "# AIMA 舆情报告\n\n正文保留在原生文档中。\n\n"
        "| 平台 | 内容量 |\n| --- | --- |\n| 抖音 | 12 |\n| 小红书 | 8 |\n\n"
        "```mermaid\nxychart-beta\n```\n\n"
        "![一级议题词云](assets/wordcloud.png)\n",
        encoding="utf-8",
    )
    return path


def test_representative_link_cell_projects_blue_clickable_text() -> None:
    payload = _table_cell_text_block(
        "[爱玛售后服务不错](https://example.com/post)",
        bold=False,
    )

    text_run = payload["text"]["elements"][0]["text_run"]
    assert text_run["content"] == "爱玛售后服务不错"
    assert text_run["text_element_style"]["link"] == {"url": "https://example.com/post"}


def test_native_document_parser_keeps_text_table_chart_and_image_order(tmp_path: Path) -> None:
    document = build_feishu_native_document(_report_markdown(tmp_path), (_spec(),))

    assert [block.kind for block in document.blocks] == [
        "heading",
        "paragraph",
        "table",
        "chart",
        "image",
    ]
    assert document.blocks[2].rows == (("平台", "内容量"), ("抖音", "12"), ("小红书", "8"))
    assert document.blocks[3].chart_index == 1
    assert document.blocks[4].image_path == tmp_path / "assets" / "wordcloud.png"


def test_native_document_parser_pivots_compact_daily_table_like_word(tmp_path: Path) -> None:
    markdown_path = tmp_path / "report.md"
    markdown_path.write_text(
        "<!-- aima:table-style=compact-daily -->\n"
        "| 日期 | 平台 | 数量 |\n"
        "| --- | --- | --- |\n"
        "| 2026-08-13 | 抖音 | 4,083 |\n"
        "| 2026-08-13 | 小红书 | 860 |\n"
        "| 2026-08-14 | 抖音 | 4,345 |\n"
        "| 2026-08-14 | 小红书 | 848 |\n\n"
        "```mermaid\nxychart-beta\n```\n",
        encoding="utf-8",
    )

    document = build_feishu_native_document(markdown_path, (_spec(),))

    assert [block.kind for block in document.blocks] == ["table", "chart"]
    assert document.blocks[0].rows == (
        ("日期", "抖音", "小红书"),
        ("2026-08-13", "4,083", "860"),
        ("2026-08-14", "4,345", "848"),
    )


def test_native_document_parser_normalizes_empty_table_cells(tmp_path: Path) -> None:
    markdown_path = tmp_path / "report.md"
    markdown_path.write_text(
        "| 指标 | 数量 |\n| --- | --- |\n| 暂无数据 |  |\n\n```mermaid\nxychart-beta\n```\n",
        encoding="utf-8",
    )

    document = build_feishu_native_document(markdown_path, (_spec(),))

    assert document.blocks[0].rows == (("指标", "数量"), ("暂无数据", "—"))


def test_native_document_parser_replaces_representative_section_with_bitable(
    tmp_path: Path,
) -> None:
    markdown_path = tmp_path / "report.md"
    markdown_path.write_text(
        "# AIMA 舆情报告\n\n"
        "## 6. 代表性评论与关联页面\n\n"
        "### 6.1 抖音正面评价\n\n"
        "| 原文链接 | 评论内容 | 处理建议 |\n"
        "| --- | --- | --- |\n"
        "| 打开原文 |  | 建议跟进 |\n\n"
        "## 7. 后续说明\n\n"
        "```mermaid\nxychart-beta\n```\n",
        encoding="utf-8",
    )

    document = build_feishu_native_document(
        markdown_path,
        (_spec(),),
        embed_representative_bitable=True,
    )

    assert [block.kind for block in document.blocks] == [
        "heading",
        "heading",
        "bitable",
        "heading",
        "chart",
    ]
    assert document.blocks[1].text == "6. 代表性评论与关联页面"
    assert document.blocks[3].text == "7. 后续说明"


def test_report_publisher_writes_real_bitable_block_payload() -> None:
    requests: list[dict[str, object]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/open-apis/auth/v3/tenant_access_token/internal":
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "tenant-token"})
        assert request.url.path == "/open-apis/docx/v1/documents/doc/blocks/doc/children"
        assert request.headers["Authorization"] == "Bearer tenant-token"
        requests.append(json.loads(request.content))
        return httpx.Response(
            200,
            json={
                "code": 0,
                "data": {
                    "children": [
                        {
                            "block_id": "bitable-block",
                            "block_type": 18,
                            "bitable": {"token": "bascnExample_tblExample"},
                        }
                    ]
                },
            },
        )

    client = httpx.Client(
        base_url="https://open.feishu.cn",
        transport=httpx.MockTransport(handler),
    )
    publisher = FeishuReportPublisher(
        FeishuReportPublisherConfig(
            app_id="app-id",
            app_secret=SecretStr("app-secret"),
            folder_token="folder-token",
        ),
        client=client,
    )
    try:
        embedded = publisher._append_bitable_block(  # noqa: SLF001
            document_token="doc",
            idempotency_source="test-bitable",
        )
    finally:
        client.close()

    assert requests == [
        {
            "index": -1,
            "children": [
                {
                    "block_type": 18,
                    "bitable": {"view_type": 1},
                }
            ],
        }
    ]
    assert embedded.token == "bascnExample_tblExample"
    assert embedded.app_token == "bascnExample"
    assert embedded.table_id == "tblExample"


def test_report_publisher_creates_embedded_bitable_in_document_order() -> None:
    requests: list[dict[str, object]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/open-apis/auth/v3/tenant_access_token/internal":
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "tenant-token"})
        assert request.url.path == "/open-apis/docx/v1/documents/doc/blocks/doc/children"
        body = json.loads(request.content)
        requests.append(body)
        if body["children"] == [{"block_type": 18, "bitable": {"view_type": 1}}]:
            return httpx.Response(
                200,
                json={
                    "code": 0,
                    "data": {
                        "children": [
                            {
                                "block_id": "bitable-block",
                                "block_type": 18,
                                "bitable": {"token": "bascnExample_tblTarget"},
                            }
                        ]
                    },
                },
            )
        return httpx.Response(200, json={"code": 0, "data": {}})

    client = httpx.Client(
        base_url="https://open.feishu.cn",
        transport=httpx.MockTransport(handler),
    )
    publisher = FeishuReportPublisher(
        FeishuReportPublisherConfig(
            app_id="app-id",
            app_secret=SecretStr("app-secret"),
            folder_token="folder-token",
        ),
        client=client,
    )
    document = FeishuNativeDocument(
        markdown_path=Path("report.md"),
        blocks=(
            FeishuNativeBlock(kind="heading", text="6. 代表性评论与关联页面", level=2),
            FeishuNativeBlock(kind="bitable"),
        ),
        chart_specs=(),
    )
    try:
        embedded = publisher._write_native_document(  # noqa: SLF001
            native_document=_ImportResult("doc", "https://feishu.example/doc", ()),
            document=document,
            sheet_url=None,
            idempotency_source="test-embedded-bitable",
        )
    finally:
        client.close()

    assert requests[1]["children"] == [{"block_type": 18, "bitable": {"view_type": 1}}]
    assert embedded is not None
    assert embedded.token == "bascnExample_tblTarget"


def test_docx_client_tokens_are_stable_uuid_v4() -> None:
    from aima_ugc.adapters.feishu.report_publisher import _stable_client_token

    first = _stable_client_token("same-operation")
    second = _stable_client_token("same-operation")

    assert first == second
    assert uuid.UUID(first).version == 4


def test_report_publisher_rejects_bitable_response_without_token() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/open-apis/auth/v3/tenant_access_token/internal":
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "tenant-token"})
        return httpx.Response(
            200,
            json={"code": 0, "data": {"children": [{"block_type": 18, "bitable": {}}]}},
        )

    client = httpx.Client(
        base_url="https://open.feishu.cn",
        transport=httpx.MockTransport(handler),
    )
    publisher = FeishuReportPublisher(
        FeishuReportPublisherConfig(
            app_id="app-id",
            app_secret=SecretStr("app-secret"),
            folder_token="folder-token",
        ),
        client=client,
    )
    with pytest.raises(FeishuApiError, match="响应缺少 token"):
        publisher._append_bitable_block(  # noqa: SLF001
            document_token="doc",
            idempotency_source="test-bitable-compat",
        )
    client.close()


def test_native_document_parser_keeps_word_ranking_limits_for_wordcloud_sections(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "wordcloud.png").write_bytes(render_chart_png(_spec())[0])
    markdown_path = tmp_path / "report.md"
    ranking_rows = "\n".join(
        f"| 标签{index:02d} | {100 - index} | {index}.00% |" for index in range(1, 13)
    )
    markdown_path.write_text(
        "<!-- aima:layout=ranking-image-top8 -->\n"
        "<!-- aima:table-style=ranking -->\n"
        "| 二级标签 | 标签对数量 | 标签对占比 |\n"
        "| --- | --- | --- |\n"
        f"{ranking_rows}\n\n"
        "![二级议题词云](assets/wordcloud.png)\n\n"
        "```mermaid\nxychart-beta\n```\n",
        encoding="utf-8",
    )

    document = build_feishu_native_document(markdown_path, (_spec(),))

    assert [block.kind for block in document.blocks] == ["paragraph", "table", "image", "chart"]
    assert document.blocks[0].text == "Top 8"
    assert document.blocks[1].rows[0] == ("排名", "二级标签", "标签对数量", "标签对占比")
    assert len(document.blocks[1].rows) == 9
    assert document.blocks[1].rows[-1] == ("08", "标签08", "92", "8.00%")


def test_build_editable_chart_workbook_keeps_data_and_native_charts(tmp_path: Path) -> None:
    output_path = tmp_path / "report-charts.xlsx"
    summary = build_editable_chart_workbook((_spec(),), output_path)

    assert summary.chart_count == 1
    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["图表01"]
    assert len(workbook["图表01"]._charts) == 1  # noqa: SLF001
    assert workbook["图表01"]["A2"].value == "抖音"


def test_report_publisher_creates_native_document_without_docx_import(
    tmp_path: Path,
) -> None:
    word_path = tmp_path / "report.docx"
    markdown_path = _report_markdown(tmp_path)
    word_bytes = b"PK\x03\x04original-office-chart-docx"
    word_path.write_bytes(word_bytes)
    chart_workbook_path = tmp_path / "report-charts.xlsx"
    build_editable_chart_workbook((_spec(),), chart_workbook_path)
    requests: list[tuple[str, str, bytes]] = []
    upload_count = 0
    image_block_count = 0
    uploaded_image_parents: list[str] = []
    replaced_images: list[dict[str, object]] = []
    deleted_files: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal upload_count, image_block_count
        body = request.content
        requests.append((request.method, request.url.path, body))
        if request.url.path == "/open-apis/auth/v3/tenant_access_token/internal":
            assert b"app-secret-value" in body
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "tenant-token"})
        assert request.headers["Authorization"] == "Bearer tenant-token"
        if request.url.path == "/open-apis/drive/v1/files/upload_all":
            upload_count += 1
            if upload_count == 1:
                assert word_bytes in body
            else:
                assert b"xl/workbook.xml" in body
            return httpx.Response(
                200, json={"code": 0, "data": {"file_token": f"file-{upload_count}"}}
            )
        if request.url.path == "/open-apis/drive/v1/import_tasks":
            payload = json.loads(body)
            assert payload["file_token"] == "file-2"
            assert payload["type"] == "sheet"
            return httpx.Response(200, json={"code": 0, "data": {"ticket": "sheet-ticket"}})
        if request.url.path == "/open-apis/drive/v1/import_tasks/sheet-ticket":
            return httpx.Response(
                200,
                json={
                    "code": 0,
                    "data": {
                        "result": {
                            "job_status": 0,
                            "token": "native-sheet",
                            "url": "https://tenant.feishu.cn/sheets/native-sheet",
                        }
                    },
                },
            )
        if request.url.path == "/open-apis/drive/v1/files/file-2" and request.method == "DELETE":
            assert request.url.params["type"] == "file"
            deleted_files.append(request.url.path.rsplit("/", 1)[-1])
            return httpx.Response(200, json={"code": 0, "data": {}})
        if request.url.path == "/open-apis/docx/v1/documents":
            assert b"folder-token" in body
            return httpx.Response(
                200,
                json={
                    "code": 0,
                    "data": {
                        "document": {
                            "document_id": "native-doc",
                            "url": "https://tenant.feishu.cn/docx/native-doc",
                        }
                    },
                },
            )
        if request.url.path == "/open-apis/drive/v1/medias/upload_all":
            parent_node = (
                request.content.split(b'name="parent_node"\r\n\r\n', 1)[1]
                .split(b"\r\n", 1)[0]
                .decode("utf-8")
            )
            uploaded_image_parents.append(parent_node)
            return httpx.Response(200, json={"code": 0, "data": {"file_token": "image-token"}})
        if (
            request.url.path
            == "/open-apis/docx/v1/documents/native-doc/blocks/native-doc/descendant"
        ):
            payload = json.loads(body)
            assert len(payload["children_id"]) == 1
            table = payload["descendants"][0]
            assert table["block_type"] == 31
            assert table["table"]["property"]["row_size"] == 3
            assert table["table"]["property"]["column_size"] == 2
            assert payload["descendants"][1]["block_type"] == 32
            assert payload["descendants"][2]["text"]["elements"][0]["text_run"]["content"] == "平台"
            return httpx.Response(200, json={"code": 0, "data": {}})
        if request.url.path == "/open-apis/docx/v1/documents/native-doc/blocks/native-doc/children":
            payload = json.loads(body)
            if payload["children"] == [{"block_type": 27, "image": {}}]:
                image_block_count += 1
                return httpx.Response(
                    200,
                    json={
                        "code": 0,
                        "data": {"children": [{"block_id": f"image-block-{image_block_count}"}]},
                    },
                )
            return httpx.Response(200, json={"code": 0, "data": {}})
        if request.url.path.startswith(
            "/open-apis/docx/v1/documents/native-doc/blocks/image-block-"
        ):
            payload = json.loads(body)
            replaced_images.append(payload["replace_image"])
            return httpx.Response(200, json={"code": 0, "data": {}})
        raise AssertionError(f"unexpected request: {request.method} {request.url}")

    publisher = FeishuReportPublisher(
        FeishuReportPublisherConfig(
            app_id="app-id",
            app_secret=SecretStr("app-secret-value"),
            folder_token="folder-token",
            poll_interval_seconds=0,
        ),
        client=httpx.Client(
            base_url="https://open.feishu.cn", transport=httpx.MockTransport(handler)
        ),
        sleep=lambda _seconds: None,
    )
    summary = publisher.publish(
        word_path=word_path,
        markdown_path=markdown_path,
        chart_specs=(_spec(),),
        chart_workbook_path=chart_workbook_path,
        title="AIMA 舆情报告",
    )

    assert summary.native_document_token == "native-doc"
    assert summary.word_file_token == "file-1"
    assert summary.word_sha256 == hashlib.sha256(word_bytes).hexdigest()
    assert summary.editable_chart_sheet_url == "https://tenant.feishu.cn/sheets/native-sheet"
    assert summary.chart_workbook_file_token is None
    assert deleted_files == ["file-2"]
    paths = [path for _method, path, _body in requests]
    assert "/open-apis/docx/v1/documents" in paths
    assert "/open-apis/drive/v1/import_tasks" in paths
    block_payloads = [
        json.loads(body) for _method, path, body in requests if path.endswith("/children")
    ]
    native_content = next(
        payload
        for payload in block_payloads
        if any("heading1" in child for child in payload["children"])
    )
    heading = native_content["children"][0]
    assert heading["heading1"]["elements"][0]["text_run"]["content"] == "AIMA 舆情报告"
    assert "text" not in heading
    assert any(
        "编辑图表数据（图表01）" in json.dumps(payload, ensure_ascii=False)
        for payload in block_payloads
    )
    assert any(
        "原始 Word 下载" in json.dumps(payload, ensure_ascii=False) for payload in block_payloads
    )
    assert uploaded_image_parents == ["image-block-1", "image-block-2"]
    assert len(replaced_images) == 2
    assert all(image["token"] == "image-token" for image in replaced_images)
    assert all(image["width"] > 0 and image["height"] > 0 for image in replaced_images)


def test_report_publisher_fails_closed_and_redacts_app_secret(tmp_path: Path) -> None:
    word_path = tmp_path / "report.docx"
    word_path.write_bytes(b"docx")
    markdown_path = _report_markdown(tmp_path)
    publisher = FeishuReportPublisher(
        FeishuReportPublisherConfig(
            app_id="app-id", app_secret=SecretStr("app-secret-value"), folder_token="folder-token"
        ),
        client=httpx.Client(
            base_url="https://open.feishu.cn",
            transport=httpx.MockTransport(
                lambda _request: httpx.Response(
                    400, json={"code": 10003, "msg": "invalid app-secret-value"}
                )
            ),
        ),
        sleep=lambda _seconds: None,
    )
    with pytest.raises(FeishuApiError) as error:
        publisher.publish(
            word_path=word_path,
            markdown_path=markdown_path,
            chart_specs=(_spec(),),
            chart_workbook_path=None,
            title="报告",
        )

    assert "app-secret-value" not in str(error.value)
    assert "[REDACTED]" in str(error.value)
    assert error.value.status_code == 400
    assert error.value.api_code == 10003
    assert error.value.http_method == "POST"
    assert error.value.endpoint == "/open-apis/auth/v3/tenant_access_token/internal"


def test_load_feishu_config_is_opt_in_and_reads_secret_from_approved_root(tmp_path: Path) -> None:
    secret_root = tmp_path / "secrets"
    secret_file = secret_root / "feishu" / "report_app_secret"
    secret_file.parent.mkdir(parents=True)
    secret_file.write_text("secret-from-file\n", encoding="utf-8")
    assert load_feishu_report_publisher_config({}, secret_root=secret_root) is None
    environ: Mapping[str, str] = {
        "AIMA_FEISHU_REPORT_ENABLED": "true",
        "AIMA_FEISHU_APP_ID": "app-id",
        "AIMA_FEISHU_FOLDER_TOKEN": "folder-token",
        "AIMA_FEISHU_APP_SECRET_REF": "feishu/report_app_secret",
    }
    config = load_feishu_report_publisher_config(environ, secret_root=secret_root)
    assert config is not None
    assert config.app_secret.get_secret_value() == "secret-from-file"


def test_publish_all_dry_run_generates_report_without_loading_or_calling_feishu(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "current.xlsx"
    previous = tmp_path / "previous.xlsx"
    source.write_bytes(b"current")
    previous.write_bytes(b"previous")
    settings = PlatformSettings(
        data_dir=tmp_path / "data",
        log_dir=tmp_path / "logs",
        secret_dir=tmp_path / "secrets",
        llm_base_url="https://llm.example/v1",
        llm_model="fake-model",
    )
    rows = (object(), object())
    monkeypatch.setattr(
        publication_module,
        "prepare_representative_report",
        lambda **_kwargs: type(
            "Preparation",
            (),
            {"rows": rows, "selection_run": type("Selection", (), {"selected": ()})()},
        )(),
    )
    captured: dict[str, object] = {}

    def fake_generate_report(**kwargs: object):
        captured.update(kwargs)
        return publication_module.ReportGenerationSummary(
            source_excel_path=source,
            template_path=tmp_path / "template.docx",
            markdown_path=tmp_path / "report.md",
            word_path=tmp_path / "report.docx",
            content_rows=3,
            label_rows=2,
            comment_rows=1,
            start_date="2026-09-01",
            end_date="2026-09-09",
            word_chart_count=1,
        )

    monkeypatch.setattr(publication_module, "generate_excel_report", fake_generate_report)
    monkeypatch.setattr(
        publication_module,
        "_load_publisher_config",
        lambda *_args, **_kwargs: pytest.fail("Dry Run 不应读取飞书发布配置"),
    )

    result = publish_all_report_to_feishu(
        input_path=source,
        previous_input_path=previous,
        output_dir=tmp_path / "output",
        report_date_range=(date(2026, 9, 1), date(2026, 9, 9)),
        settings=settings,
        environ={},
        dry_run=True,
    )

    assert result.dry_run is True
    assert result.publication is None
    assert result.representative_count == 2
    assert captured["representative_rows"] == rows


def test_publish_all_real_run_creates_embedded_bitable_before_sync(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "current.xlsx"
    previous = tmp_path / "previous.xlsx"
    source.write_bytes(b"current")
    previous.write_bytes(b"previous")
    settings = PlatformSettings(
        data_dir=tmp_path / "data",
        log_dir=tmp_path / "logs",
        secret_dir=tmp_path / "secrets",
    )
    rows = (object(),)
    monkeypatch.setattr(
        publication_module,
        "prepare_representative_report",
        lambda **_kwargs: type(
            "Preparation",
            (),
            {"rows": rows, "selection_run": type("Selection", (), {"selected": ("selected",)})()},
        )(),
    )
    monkeypatch.setattr(
        publication_module,
        "generate_excel_report",
        lambda **_kwargs: publication_module.ReportGenerationSummary(
            source_excel_path=source,
            template_path=tmp_path / "template.docx",
            markdown_path=tmp_path / "report.md",
            word_path=tmp_path / "report.docx",
            content_rows=3,
            label_rows=2,
            comment_rows=1,
            start_date="2026-09-01",
            end_date="2026-09-09",
            word_chart_count=1,
        ),
    )
    monkeypatch.setattr(
        publication_module,
        "_load_publisher_config",
        lambda *_args, **_kwargs: object(),
    )
    captured: dict[str, object] = {}
    events: list[str] = []

    def fake_sync(**kwargs: object):
        events.append("sync")
        captured.update(kwargs)
        return type(
            "Sync",
            (),
            {
                "target_bitable_block_token": "bitable-token",
                "target_table_url": "https://feishu.example/table",
                "target_table_name": "代表性内容",
                "created_count": 1,
                "updated_count": 0,
                "verified_count": 1,
            },
        )()

    monkeypatch.setattr(publication_module, "publish_selected_representatives_to_feishu", fake_sync)

    def fake_publish(report: object, config: object, **kwargs: object):
        del report, config
        events.append("publish")
        captured["publish_kwargs"] = kwargs
        return type(
            "Publication",
            (),
            {
                "native_document_token": "doc-token",
                "native_document_url": "https://feishu.example/doc",
                "editable_chart_sheet_url": "https://feishu.example/sheet",
                "representative_bitable_token": "bascnEmbedded_tblEmbedded",
            },
        )()

    monkeypatch.setattr(
        publication_module,
        "_publish_generated_report",
        fake_publish,
    )

    result = publish_all_report_to_feishu(
        input_path=source,
        previous_input_path=previous,
        output_dir=tmp_path / "output",
        report_date_range=(date(2026, 9, 1), date(2026, 9, 9)),
        settings=settings,
        environ={},
        dry_run=False,
    )

    assert result.dry_run is False
    assert result.representative_sync is not None
    assert result.publication is not None
    assert captured["settings"] is settings
    assert captured["selected"] == ("selected",)
    assert captured["report_rows"] == rows
    assert captured["target_bitable_block_token"] == "bascnEmbedded_tblEmbedded"
    assert captured["target_document_token"] == "doc-token"
    assert captured["target_document_url"] == "https://feishu.example/doc"
    assert captured["publish_kwargs"] == {"embed_representative_bitable": True}
    assert events == ["publish", "sync"]
