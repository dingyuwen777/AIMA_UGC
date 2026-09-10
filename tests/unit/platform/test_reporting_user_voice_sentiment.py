from __future__ import annotations

from datetime import datetime
from pathlib import Path

from aima_ugc.platform.reporting import generate_excel_report
from openpyxl import Workbook


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    content = workbook.active
    content.title = "内容"
    content.append(
        (
            "内容ID",
            "平台",
            "发布时间",
            "命中关键词",
            "发声类型",
            "情感标签",
            "一级标签",
            "二级标签",
        )
    )
    rows = (
        (
            "c1",
            "抖音",
            "2026-09-01 10:00:00",
            "爱玛",
            "真实用户发声",
            "正面",
            "用户体验",
            "骑行体验",
        ),
        (
            "c2",
            "抖音",
            "2026-09-01 11:00:00",
            "爱玛",
            "营销推广发声",
            "负面",
            "营销议题",
            "推广内容",
        ),
        (
            "c3",
            "小红书",
            "2026-09-01 12:00:00",
            "爱玛",
            "真实用户发声",
            "中性",
            "用户咨询",
            "购买意愿",
        ),
        (
            "c4",
            "微博",
            "2026-09-01 13:00:00",
            "爱玛",
            "媒体机构发声",
            "混合",
            "媒体议题",
            "资讯内容",
        ),
    )
    for row in rows:
        content.append(row)

    labels = workbook.create_sheet("标签明细")
    labels.append(
        (
            "内容ID",
            "平台",
            "发布时间",
            "发声类型",
            "情感标签",
            "一级标签",
            "二级标签",
        )
    )
    for row in rows:
        labels.append((row[0], row[1], row[2], row[4], row[5], row[6], row[7]))

    comments = workbook.create_sheet("评论")
    comments.append(("平台", "评论时间"))
    comments.append(("抖音", "2026-09-01 14:00:00"))
    workbook.save(path)
    workbook.close()


def test_report_sentiment_only_counts_real_user_voice(tmp_path: Path) -> None:
    workbook_path = tmp_path / "labeled_data.xlsx"
    _make_workbook(workbook_path)

    summary = generate_excel_report(
        input_path=workbook_path,
        output_dir=tmp_path / "reports",
        generated_at=datetime(2026, 9, 2, 10, 30, 0),
    )
    markdown = summary.markdown_path.read_text(encoding="utf-8")

    # 内容总量、平台声量和整体议题仍包含全部四条内容。
    assert "| 内容声量 | 4 |" in markdown
    assert "| 抖音 | 2 | 50.00% | 1 |" in markdown
    assert "| 营销议题 | 1 |" in markdown

    # 情感统计只包含两条真实用户发声，比例分母也限定为这两条。
    assert "| 正面 | 1 | 50.00% |" in markdown
    assert "| 中性 | 1 | 50.00% |" in markdown
    assert "| 负面 |" not in markdown.split("### 3.1 总体情感结构", 1)[1].split("### 3.2", 1)[0]
    assert "| 混合 |" not in markdown.split("### 3.1 总体情感结构", 1)[1].split("### 3.2", 1)[0]

    sentiment_platform = markdown.split("### 2.2 平台与情感对比", 1)[1].split("### 2.3", 1)[0]
    assert "| 抖音 | 1 | 0 | 1 | 0.00% |" in sentiment_platform
    assert "| 小红书 | 0 | 1 | 1 | 0.00% |" in sentiment_platform
    assert "| 微博 | 0 | 0 | 0 | 0.00% |" in sentiment_platform

    positive_topics = markdown.split("### 3.3 正面一级议题", 1)[1].split("### 3.4", 1)[0]
    negative_topics = markdown.split("### 3.5 负面一级议题", 1)[1].split("### 3.6", 1)[0]
    assert "用户体验" in positive_topics
    assert "营销议题" not in negative_topics
