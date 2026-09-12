"""车型共现筛选统一 Excel 导出的直接行为测试。"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

import pytest
import aima_ugc.adapters.providers.imports_test.vehicle_pair_filter.filter_vehicle_pairs as vehicle_pair_filter_module
from aima_ugc.contracts.analysis import UnifiedContentRecordV1
from aima_ugc.contracts.canonical import CanonicalContentV1, CanonicalSourceV1
from openpyxl import load_workbook


def _record(*, external_id: str, title: str, text: str) -> UnifiedContentRecordV1:
    """构造一个能命中多个目标/竞品车型的合法帖子。"""

    observed_at = datetime(2026, 9, 12, 10, 0, tzinfo=UTC)
    return UnifiedContentRecordV1(
        content=CanonicalContentV1(
            platform="xiaohongshu",
            external_content_id=external_id,
            content_type="note",
            title=title,
            text=text,
            observed_at=observed_at,
            observed_fields=["title", "text"],
            source=CanonicalSourceV1(
                provider_name="manual",
                source_type="test",
                source_value="vehicle-pair-excel-test",
                observed_at=observed_at,
            ),
        ),
        matched_keywords=["爱玛", "九号"],
    )


def _write_catalog(path: Path) -> None:
    """写入包含两个目标车型和两个竞品车型的测试目录。"""

    path.write_text(
        json.dumps(
            {
                "schema_version": "vehicle-pair-catalog.v1",
                "target_brand": "爱玛",
                "brands": [
                    {
                        "name": "爱玛",
                        "models": [
                            {"name": "元宇宙", "aliases": []},
                            {"name": "墩墩", "aliases": []},
                        ],
                    },
                    {
                        "name": "雅迪",
                        "models": [{"name": "莱茵", "aliases": []}],
                    },
                    {
                        "name": "九号",
                        "models": [{"name": "Q3", "aliases": []}],
                    },
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def _write_input(path: Path) -> None:
    """写入一篇会形成四个 matched_pairs、但仍只应导出一行的帖子。"""

    record = _record(
        external_id="excel-post-001",
        title="元宇宙和墩墩怎么选",
        text="莱茵和Q3也都试过",
    )
    path.write_text(f"{record.model_dump_json()}\n", encoding="utf-8")


def test_vehicle_pair_filter_exports_unified_excel_with_brand_vehicle_columns(tmp_path: Path) -> None:
    """成功筛选应在同一 run 生成一帖一行、含品牌车型列的统一 Excel。"""

    input_path = tmp_path / "contents.jsonl"
    catalog_path = tmp_path / "vehicle_catalog.json"
    _write_input(input_path)
    _write_catalog(catalog_path)

    summary = vehicle_pair_filter_module.filter_vehicle_pairs(
        input_path=input_path,
        catalog_path=catalog_path,
        output_root=tmp_path / "output",
        run_id="excel-export",
    )

    assert summary.output_path.is_file()
    assert summary.workbook_path.is_file()
    assert summary.run_summary_path.is_file()
    assert summary.rows_with_cross_brand_pair == 1
    assert len([line for line in summary.output_path.read_bytes().splitlines() if line.strip()]) == 1

    workbook = load_workbook(summary.workbook_path, read_only=True, data_only=True)
    try:
        content_rows = list(workbook["内容"].iter_rows(values_only=True))
        label_rows = list(workbook["标签明细"].iter_rows(values_only=True))
        comment_rows = list(workbook["评论"].iter_rows(values_only=True))
    finally:
        workbook.close()

    assert len(content_rows) == 2
    headers = tuple(str(value) for value in content_rows[0])
    values = dict(zip(headers, content_rows[1], strict=True))
    assert values["平台"] == "小红书"
    assert values["内容ID"] == "excel-post-001"
    assert values["标题"] == "元宇宙和墩墩怎么选"
    assert values["品牌"] == "爱玛；雅迪；九号"
    assert values["品牌角色"] == "自有品牌；竞品品牌；竞品品牌"
    assert values["竞品范围"] == "混合品牌"
    assert values["车型"] == "元宇宙；墩墩；莱茵；Q3"
    assert len(label_rows) == 1
    assert len(comment_rows) == 1

    run_summary = json.loads(summary.run_summary_path.read_text(encoding="utf-8"))
    assert run_summary["outputs"]["comparison_posts"] == str(summary.output_path)
    assert run_summary["outputs"]["comparison_posts_excel"] == str(summary.workbook_path)


def test_vehicle_pair_filter_removes_run_when_excel_export_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """JSONL 成功后若共享 Excel Exporter 失败，本次 run 也必须整体清理。"""

    input_path = tmp_path / "contents.jsonl"
    catalog_path = tmp_path / "vehicle_catalog.json"
    output_root = tmp_path / "output"
    _write_input(input_path)
    _write_catalog(catalog_path)

    def fail_export(*args: object, **kwargs: object) -> None:
        """模拟共享 Excel Exporter 在发布前失败。"""

        raise RuntimeError("simulated excel export failure")

    monkeypatch.setattr(vehicle_pair_filter_module, "export_unified_data_excel", fail_export)

    with pytest.raises(RuntimeError, match="simulated excel export failure"):
        vehicle_pair_filter_module.filter_vehicle_pairs(
            input_path=input_path,
            catalog_path=catalog_path,
            output_root=output_root,
            run_id="excel-failure",
        )

    assert not (output_root / "runs" / "excel-failure").exists()
