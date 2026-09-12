"""验证本地历史自动接管后，三阶段只处理新增数据且评论缓存不重复请求。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from aima_ugc.adapters.providers.imports_test.comparison_comment_enrichment.enrich_comments import (
    enrich_comparison_comments,
)
from aima_ugc.adapters.providers.imports_test.comparison_comment_enrichment.models import (
    VehiclePairCommentRecordV1,
)
from aima_ugc.adapters.providers.imports_test.monitoring_excel_filter.process_directory import (
    process_directory,
)
from aima_ugc.adapters.providers.imports_test.vehicle_pair_filter.filter_vehicle_pairs import (
    filter_vehicle_pairs,
)
from aima_ugc.adapters.providers.tikhub import runtime as tikhub_runtime
from aima_ugc.adapters.providers.tikhub_test.core.config import TikHubTestConfig
from aima_ugc.contracts.canonical import (
    CanonicalCommentV1,
    CanonicalMetricsV1,
    CanonicalSourceV1,
)
from aima_ugc.modules.collection.providers.transport import (
    ProviderTransportRequest,
    ProviderTransportResponse,
)
from openpyxl import Workbook, load_workbook
from pydantic import SecretStr

_HEADERS = (
    "媒体名称（中文）",
    "文章编号",
    "标题",
    "内文",
    "作者",
    "出版日期",
    "原文链接",
    "粉丝数",
)


class _OneRootTransport:
    """每个 cache miss 返回一个无回复一级评论，并记录实际请求。"""

    def __init__(self) -> None:
        self.requests: list[ProviderTransportRequest] = []

    def send(self, request: ProviderTransportRequest) -> ProviderTransportResponse:
        """记录请求并返回一个确定根评论。"""

        self.requests.append(request)
        return ProviderTransportResponse(
            status_code=200,
            body={"items": [{"id": f"root-{len(self.requests)}", "reply_count": 0}]},
        )


class _NoProviderAllowed:
    """用于证明全量 current 再跑时所有帖子均命中本地缓存。"""

    def send(self, request: ProviderTransportRequest) -> ProviderTransportResponse:
        """生产代码一旦尝试请求 Provider 就让测试立即失败。"""

        raise AssertionError(f"历史缓存帖子不应再次请求 Provider: {request.path}")


def _config() -> TikHubTestConfig:
    """构造不接触真实 Secret 的测试 Provider 配置。"""

    return TikHubTestConfig(
        base_url="https://api.tikhub.dev",
        api_key=SecretStr("test-key"),
        timeout_seconds=1,
    )


def _write_workbook(path: Path, rows: list[tuple[object, ...]]) -> None:
    """写入符合 Monitoring Excel Profile 的真实 XLSX fixture。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "文章"
    worksheet.append(list(_HEADERS))
    for row in rows:
        worksheet.append(list(row))
    workbook.save(path)
    workbook.close()


def _row(*, source_id: str, note_id: str) -> tuple[object, ...]:
    """构造一条同时包含目标和竞品车型的小红书帖子。"""

    return (
        "小红书",
        source_id,
        "元宇宙实际体验",
        "Q3也试过",
        "测试作者",
        datetime(2026, 9, 12, 12, 0),
        f"https://www.xiaohongshu.com/explore/{note_id}",
        100,
    )


def _write_catalog(path: Path) -> None:
    """写入全链路测试需要的最小爱玛/九号车型目录。"""

    path.write_text(
        """{
  "schema_version": "vehicle-pair-catalog.v1",
  "target_brand": "爱玛",
  "brands": [
    {"name": "爱玛", "models": [{"name": "元宇宙", "aliases": []}]},
    {"name": "九号", "models": [{"name": "Q3", "aliases": []}]}
  ]
}
""",
        encoding="utf-8",
    )


def _install_comment_runtime(monkeypatch: pytest.MonkeyPatch) -> None:
    """保留正式 request builder，只隔离 Provider 响应提取/Mapper/分页。"""

    monkeypatch.setattr(
        tikhub_runtime,
        "extract_comment_items",
        lambda platform, body: tuple(body.get("items", ())),
    )
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )

    def map_comment(
        *,
        platform: str,
        raw: dict[str, Any],
        context: Any,
        item_locator: str,
        is_root: bool,
    ) -> CanonicalCommentV1:
        if not is_root or not isinstance(context.external_content_id, str):
            raise AssertionError("全链回归只产生一级评论")
        comment_id = str(raw["id"])
        return CanonicalCommentV1(
            platform=platform,
            external_content_id=context.external_content_id,
            external_comment_id=comment_id,
            root_comment_id=comment_id,
            parent_comment_id=None,
            text=f"comment-{context.external_content_id}",
            observed_at=context.observed_at,
            metrics=CanonicalMetricsV1(reply_count=0),
            source=CanonicalSourceV1(
                provider_name="tikhub",
                operation=context.operation,
                provider_request_id=context.provider_request_id,
                provider_attempt_id=context.provider_attempt_id,
                raw_artifact_id=context.raw_artifact_id,
                source_type=context.source_type,
                source_value=context.source_value,
                item_locator=item_locator,
                observed_at=context.observed_at,
            ),
            observed_fields=["root_comment_id", "parent_comment_id", "text", "metrics.reply_count"],
        )

    monkeypatch.setattr(tikhub_runtime, "map_comment", map_comment)


def test_existing_three_stage_history_bootstraps_then_only_new_data_is_processed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """真实文件链应自动接管旧 output；第二轮只处理新增，累计评论再次运行零请求。"""

    _install_comment_runtime(monkeypatch)
    input_dir = tmp_path / "input"
    keyword_pack = tmp_path / "keyword_pack.txt"
    keyword_pack.write_text("爱玛\n元宇宙\nQ3\n", encoding="utf-8")
    catalog = tmp_path / "vehicle_catalog.json"
    _write_catalog(catalog)

    monitoring_root = tmp_path / "monitoring-output"
    pair_root = tmp_path / "pair-output"
    comment_root = tmp_path / "comment-output"

    _write_workbook(input_dir / "old.xlsx", [_row(source_id="old-source", note_id="old-note")])
    stage_one_old = process_directory(
        input_dir=input_dir,
        output_root=monitoring_root,
        keyword_pack_file=keyword_pack,
        sheet_name="文章",
        run_id="stage-one-old",
    )
    stage_two_old = filter_vehicle_pairs(
        input_path=stage_one_old.deduplicated_path,
        catalog_path=catalog,
        output_root=pair_root,
        run_id="stage-two-old",
    )
    old_transport = _OneRootTransport()
    stage_three_old = enrich_comparison_comments(
        input_path=stage_two_old.output_path,
        output_root=comment_root,
        run_id="stage-three-old",
        provider_config=_config(),
        transport=old_transport,
    )
    assert stage_three_old.rows_fetched == 1
    assert stage_three_old.request_count == 1

    # 模拟用户从旧代码升级：真实 runs 留在本机，只删除新版本才有的 state/current。
    for root in (monitoring_root, pair_root, comment_root):
        for name in ("state", "current"):
            target = root / name
            if target.exists():
                import shutil

                shutil.rmtree(target)

    _write_workbook(
        input_dir / "new.xlsx",
        [
            _row(source_id="old-duplicate", note_id="old-note"),
            _row(source_id="new-source", note_id="new-note"),
        ],
    )

    stage_one_new = process_directory(
        input_dir=input_dir,
        output_root=monitoring_root,
        keyword_pack_file=keyword_pack,
        sheet_name="文章",
        run_id="stage-one-new",
    )
    assert stage_one_new.files_processed == 1
    assert stage_one_new.files_skipped_unchanged == 1
    assert stage_one_new.historical_duplicates_removed == 1
    assert stage_one_new.rows_after_deduplication == 1

    stage_two_new = filter_vehicle_pairs(
        input_path=stage_one_new.deduplicated_path,
        catalog_path=catalog,
        output_root=pair_root,
        run_id="stage-two-new",
    )
    assert stage_two_new.rows_seen == 1
    assert stage_two_new.rows_with_cross_brand_pair == 1

    new_transport = _OneRootTransport()
    stage_three_new = enrich_comparison_comments(
        input_path=stage_two_new.output_path,
        output_root=comment_root,
        run_id="stage-three-new",
        provider_config=_config(),
        transport=new_transport,
    )
    assert stage_three_new.rows_cached == 0
    assert stage_three_new.rows_fetched == 1
    assert stage_three_new.request_count == 1

    current_records = [
        VehiclePairCommentRecordV1.model_validate_json(line)
        for line in stage_three_new.current_jsonl_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert {item.record.record.content.external_content_id for item in current_records} == {
        "old-note",
        "new-note",
    }

    # 将 Stage 2 累计 current 再喂给 Stage 3；历史与新增都已缓存，因此严格 0 Provider 请求。
    all_cached = enrich_comparison_comments(
        input_path=stage_two_new.current_output_path,
        output_root=comment_root,
        run_id="stage-three-all-cached",
        provider_config=_config(),
        transport=_NoProviderAllowed(),
    )
    assert all_cached.rows_seen == 2
    assert all_cached.rows_cached == 2
    assert all_cached.rows_fetched == 0
    assert all_cached.request_count == 0

    workbook = load_workbook(all_cached.current_workbook_path, read_only=True, data_only=True)
    try:
        assert sum(1 for _ in workbook["内容"].iter_rows()) == 3
        assert sum(1 for _ in workbook["评论"].iter_rows()) == 3
    finally:
        workbook.close()
