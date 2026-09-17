"""验证三阶段离线链自动接管本地旧 output，并且后续只处理增量。"""

from __future__ import annotations

import json
import shutil
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from aima_ugc.adapters.providers.imports_test.comparison_comment_enrichment.enrich_comments import (
    enrich_comparison_comments,
)
from aima_ugc.adapters.providers.imports_test.comparison_comment_enrichment.models import (
    CommentFetchCoverageV1,
    VehiclePairCommentRecordV1,
)
from aima_ugc.adapters.providers.imports_test.monitoring_excel_filter.process_directory import (
    process_directory,
)
from aima_ugc.adapters.providers.imports_test.vehicle_pair_filter.filter_vehicle_pairs import (
    VehicleModelMention,
    VehicleModelReference,
    VehiclePairRecordV1,
    VehiclePairReference,
    filter_vehicle_pairs,
)
from aima_ugc.adapters.providers.tikhub import runtime as tikhub_runtime
from aima_ugc.adapters.providers.tikhub_test.core.config import TikHubTestConfig
from aima_ugc.contracts.analysis import UnifiedContentRecordV1
from aima_ugc.contracts.canonical import (
    CanonicalCommentV1,
    CanonicalContentV1,
    CanonicalMetricsV1,
    CanonicalSourceV1,
)
from aima_ugc.modules.collection.providers.transport import (
    ProviderTransportRequest,
    ProviderTransportResponse,
)
from openpyxl import Workbook, load_workbook
from pydantic import SecretStr

_HEADERS = (
    "媒体名称（中文）",
    "文章编号",
    "标题",
    "内文",
    "作者",
    "出版日期",
    "原文链接",
    "粉丝数",
)


class _OneCommentTransport:
    """每篇帖子只返回一个无回复根评论，并记录真实 Runtime 请求。"""

    def __init__(self) -> None:
        self.requests: list[ProviderTransportRequest] = []

    def send(self, request: ProviderTransportRequest) -> ProviderTransportResponse:
        """记录请求并返回一个根评论。"""

        self.requests.append(request)
        return ProviderTransportResponse(
            status_code=200,
            body={"items": [{"id": f"root-{len(self.requests)}", "reply_count": 0}]},
        )


class _FailIfCalledTransport:
    """用于证明 cache hit 时不会产生任何 Provider 请求。"""

    def send(self, request: ProviderTransportRequest) -> ProviderTransportResponse:
        """只要生产代码尝试访问 Provider 就立即让测试失败。"""

        raise AssertionError(f"cached content 不应再次请求 Provider: {request.path}")


def _provider_config() -> TikHubTestConfig:
    """构造不会访问真实 Secret 的测试配置。"""

    return TikHubTestConfig(
        base_url="https://api.tikhub.dev",
        api_key=SecretStr("test-key"),
        timeout_seconds=1,
    )


def _write_workbook(path: Path, rows: list[tuple[object, ...]]) -> None:
    """写入符合 Monitoring Excel Profile 的真实 XLSX fixture。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "文章"
    worksheet.append(list(_HEADS := _HEADERS))
    for row in rows:
        worksheet.append(list(row))
    workbook.save(path)
    workbook.close()
    assert _HEADS == _HEADERS


def _excel_row(*, article_id: str, title: str, url: str) -> tuple[object, ...]:
    """构造一条小红书监测数据。"""

    return (
        "小红书",
        article_id,
        title,
        "Q3也试过",
        "测试作者",
        datetime(2026, 9, 12, 12, 0),
        url,
        100,
    )


def _content_record(*, external_id: str) -> UnifiedContentRecordV1:
    """构造 Stage 2 可识别的最小统一帖子。"""

    observed_at = datetime(2026, 9, 12, 0, 0, tzinfo=UTC)
    return UnifiedContentRecordV1(
        content=CanonicalContentV1(
            platform="xiaohongshu",
            external_content_id=external_id,
            content_type="image",
            title="元宇宙实际体验",
            text="Q3也试过",
            observed_at=observed_at,
            observed_fields=["title", "text"],
            source=CanonicalSourceV1(
                provider_name="manual",
                source_type="test",
                source_value="incremental-test",
                observed_at=observed_at,
            ),
        ),
        matched_keywords=["元宇宙", "Q3"],
    )


def _write_content_jsonl(path: Path, records: list[UnifiedContentRecordV1]) -> None:
    """写 Stage 2 输入 JSONL。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "".join(f"{record.model_dump_json()}\n" for record in records),
        encoding="utf-8",
    )


def _write_catalog(path: Path) -> None:
    """写最小爱玛/九号车型目录。"""

    path.write_text(
        json.dumps(
            {
                "schema_version": "vehicle-pair-catalog.v1",
                "target_brand": "爱玛",
                "brands": [
                    {"name": "爱玛", "models": [{"name": "元宇宙", "aliases": []}]},
                    {"name": "九号", "models": [{"name": "Q3", "aliases": []}]},
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def _pair_record(*, external_id: str, target_model: str = "元宇宙") -> VehiclePairRecordV1:
    """构造可改变 pair metadata、但 content identity 保持不变的 Stage 3 输入。"""

    content_record = _content_record(external_id=external_id)
    return VehiclePairRecordV1(
        record=content_record,
        matched_target_models=(target_model,),
        matched_competitor_models=(VehicleModelReference(brand="九号", model="Q3"),),
        matched_pairs=(
            VehiclePairReference(
                target_model=target_model,
                competitor_brand="九号",
                competitor_model="Q3",
            ),
        ),
        model_mentions=(
            VehicleModelMention(
                brand="爱玛",
                model=target_model,
                fields=("title",),
                matched_aliases=(target_model,),
            ),
            VehicleModelMention(
                brand="九号",
                model="Q3",
                fields=("text",),
                matched_aliases=("Q3",),
            ),
        ),
    )


def _write_pair_jsonl(path: Path, records: list[VehiclePairRecordV1]) -> None:
    """写 Stage 3 输入 JSONL。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "".join(f"{record.model_dump_json()}\n" for record in records),
        encoding="utf-8",
    )


def _install_simple_comment_runtime(monkeypatch: pytest.MonkeyPatch) -> None:
    """保留正式 request builder，只隔离 Provider 响应提取/Mapper/分页。"""

    monkeypatch.setattr(
        tikhub_runtime,
        "extract_comment_items",
        lambda platform, body: tuple(body.get("items", ())),
    )
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )

    def map_comment(
        *,
        platform: str,
        raw: dict[str, Any],
        context: Any,
        item_locator: str,
        is_root: bool,
    ) -> CanonicalCommentV1:
        if not is_root or not isinstance(context.external_content_id, str):
            raise AssertionError("测试只产生一级评论")
        comment_id = str(raw["id"])
        return CanonicalCommentV1(
            platform=platform,
            external_content_id=context.external_content_id,
            external_comment_id=comment_id,
            root_comment_id=comment_id,
            parent_comment_id=None,
            text="历史评论",
            observed_at=context.observed_at,
            metrics=CanonicalMetricsV1(reply_count=0),
            source=CanonicalSourceV1(
                provider_name="tikhub",
                operation=context.operation,
                provider_request_id=context.provider_request_id,
                provider_attempt_id=context.provider_attempt_id,
                raw_artifact_id=context.raw_artifact_id,
                source_type=context.source_type,
                source_value=context.source_value,
                item_locator=item_locator,
                observed_at=context.observed_at,
            ),
            observed_fields=["root_comment_id", "parent_comment_id", "text", "metrics.reply_count"],
        )

    monkeypatch.setattr(tikhub_runtime, "map_comment", map_comment)


def test_stage_one_bootstraps_local_run_and_only_processes_new_excel(tmp_path: Path) -> None:
    """删除 state 模拟升级后，旧 run 应自动成为基线，新 run 只打开新增 Excel。"""

    input_dir = tmp_path / "input"
    old_file = input_dir / "old.xlsx"
    _write_workbook(
        old_file,
        [
            _excel_row(
                article_id="old-source",
                title="爱玛元宇宙体验",
                url="https://www.xiaohongshu.com/explore/old-note",
            )
        ],
    )
    keyword_pack = tmp_path / "keyword_pack.txt"
    keyword_pack.write_text("爱玛\n元宇宙\nQ3\n", encoding="utf-8")
    output_root = tmp_path / "monitoring-output"

    legacy = process_directory(
        input_dir=input_dir,
        output_root=output_root,
        keyword_pack_file=keyword_pack,
        sheet_name="文章",
        run_id="legacy",
    )
    assert legacy.rows_after_deduplication == 1
    shutil.rmtree(output_root / "state")
    shutil.rmtree(output_root / "current")

    new_file = input_dir / "new.xlsx"
    _write_workbook(
        new_file,
        [
            _excel_row(
                article_id="duplicate-source",
                title="爱玛元宇宙重复",
                url="https://www.xiaohongshu.com/explore/old-note",
            ),
            _excel_row(
                article_id="new-source",
                title="爱玛元宇宙新增",
                url="https://www.xiaohongshu.com/explore/new-note",
            ),
        ],
    )

    incremental = process_directory(
        input_dir=input_dir,
        output_root=output_root,
        keyword_pack_file=keyword_pack,
        sheet_name="文章",
        run_id="incremental",
    )
    assert incremental.input_file_count == 2
    assert incremental.files_processed == 1
    assert incremental.files_skipped_unchanged == 1
    assert incremental.rows_seen == 2
    assert incremental.historical_duplicates_removed == 1
    assert incremental.rows_after_deduplication == 1
    record = UnifiedContentRecordV1.model_validate_json(
        incremental.deduplicated_path.read_text(encoding="utf-8").strip()
    )
    assert record.content.external_content_id == "new-note"
    assert legacy.run_dir.is_dir()

    _write_workbook(
        old_file,
        [
            _excel_row(
                article_id="changed-source",
                title="爱玛元宇宙被改过",
                url="https://www.xiaohongshu.com/explore/changed-note",
            )
        ],
    )
    with pytest.raises(ValueError, match="已处理 Excel 内容发生变化"):
        process_directory(
            input_dir=input_dir,
            output_root=output_root,
            keyword_pack_file=keyword_pack,
            sheet_name="文章",
            run_id="changed-old-file",
        )


def test_stage_one_signature_and_stage_two_catalog_drift_fail_closed(tmp_path: Path) -> None:
    """过滤词包或车型目录变化时不得静默沿用旧增量 state。"""

    input_dir = tmp_path / "input"
    _write_workbook(
        input_dir / "data.xlsx",
        [
            _excel_row(
                article_id="source",
                title="爱玛元宇宙体验",
                url="https://www.xiaohongshu.com/explore/pair-note",
            )
        ],
    )
    keyword_pack = tmp_path / "keyword_pack.txt"
    keyword_pack.write_text("爱玛\n元宇宙\nQ3\n", encoding="utf-8")
    monitoring_root = tmp_path / "monitoring-output"
    first = process_directory(
        input_dir=input_dir,
        output_root=monitoring_root,
        keyword_pack_file=keyword_pack,
        sheet_name="文章",
        run_id="first",
    )
    keyword_pack.write_text("爱玛\n元宇宙\nQ3\n九号\n", encoding="utf-8")
    with pytest.raises(ValueError, match="增量配置已变化"):
        process_directory(
            input_dir=input_dir,
            output_root=monitoring_root,
            keyword_pack_file=keyword_pack,
            sheet_name="文章",
            run_id="keyword-drift",
        )

    catalog = tmp_path / "vehicle_catalog.json"
    _write_catalog(catalog)
    pair_root = tmp_path / "pair-output"
    pair_first = filter_vehicle_pairs(
        input_path=first.deduplicated_path,
        catalog_path=catalog,
        output_root=pair_root,
        run_id="pair-first",
    )
    assert pair_first.rows_with_cross_brand_pair == 1
    shutil.rmtree(pair_root / "state")
    shutil.rmtree(pair_root / "current")

    cached = filter_vehicle_pairs(
        input_path=first.deduplicated_path,
        catalog_path=catalog,
        output_root=pair_root,
        run_id="pair-cached",
    )
    assert cached.input_skipped_cached is True
    assert cached.rows_seen == 0
    assert cached.output_path.read_text(encoding="utf-8") == ""
    assert cached.current_output_path.read_text(encoding="utf-8").count("\n") == 1

    payload = json.loads(catalog.read_text(encoding="utf-8"))
    payload["brands"][1]["models"][0]["aliases"] = ["九号Q3"]
    catalog.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ValueError, match="vehicle_catalog.json 已变化"):
        filter_vehicle_pairs(
            input_path=first.deduplicated_path,
            catalog_path=catalog,
            output_root=pair_root,
            run_id="catalog-drift",
        )


def test_stage_three_bootstrap_reuses_complete_and_partial_without_provider(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """旧 complete/partial 都应自动进入 cache；pair metadata 更新也不能触发评论重抓。"""

    _install_simple_comment_runtime(monkeypatch)
    comment_root = tmp_path / "comment-output"
    historical_input = tmp_path / "historical-pairs.jsonl"
    _write_pair_jsonl(historical_input, [_pair_record(external_id="cached-complete")])
    first_transport = _OneCommentTransport()
    historical = enrich_comparison_comments(
        input_path=historical_input,
        output_root=comment_root,
        run_id="legacy-complete",
        provider_config=_provider_config(),
        transport=first_transport,
    )
    assert historical.request_count == 1
    assert len(first_transport.requests) == 1

    partial_pair = _pair_record(external_id="cached-partial")
    partial_record = VehiclePairCommentRecordV1(
        record=partial_pair,
        comments=(),
        comment_fetch=CommentFetchCoverageV1(
            coverage="partial",
            root_comment_count=0,
            reply_count=0,
            request_count=1,
            root_stop_reason="legacy-partial",
        ),
    )
    partial_run = comment_root / "runs" / "legacy-partial"
    partial_run.mkdir(parents=True)
    partial_jsonl = partial_run / "comparison_posts_with_comments.jsonl"
    partial_jsonl.write_text(f"{partial_record.model_dump_json()}\n", encoding="utf-8")
    (partial_run / "run_summary.json").write_text(
        json.dumps(
            {
                "schema_version": "comparison-comment-enrichment-run.v1",
                "run_id": "legacy-partial",
                "outputs": {"jsonl": str(partial_jsonl)},
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    shutil.rmtree(comment_root / "state")
    shutil.rmtree(comment_root / "current")

    rerun_input = tmp_path / "rerun-pairs.jsonl"
    _write_pair_jsonl(
        rerun_input,
        [
            _pair_record(external_id="cached-complete", target_model="墩墩"),
            partial_pair,
        ],
    )
    reused = enrich_comparison_comments(
        input_path=rerun_input,
        output_root=comment_root,
        run_id="reuse-cache",
        provider_config=_provider_config(),
        transport=_FailIfCalledTransport(),
    )
    assert reused.rows_seen == 2
    assert reused.rows_cached == 2
    assert reused.rows_fetched == 0
    assert reused.request_count == 0
    assert reused.rows_complete == 1
    assert reused.rows_partial == 1

    records = [
        VehiclePairCommentRecordV1.model_validate_json(line)
        for line in reused.output_jsonl_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    by_id = {item.record.record.content.external_content_id: item for item in records}
    complete = by_id["cached-complete"]
    assert complete.record.matched_target_models == ("墩墩",)
    assert len(complete.comments) == 1
    assert complete.comments[0].text == "历史评论"
    assert by_id["cached-partial"].comment_fetch.coverage == "partial"

    current_records = [
        VehiclePairCommentRecordV1.model_validate_json(line)
        for line in reused.current_jsonl_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert len(current_records) == 2
    workbook = load_workbook(reused.current_workbook_path, read_only=True, data_only=True)
    try:
        assert sum(1 for _ in workbook["内容"].iter_rows()) == 3
        assert sum(1 for _ in workbook["评论"].iter_rows()) == 2
    finally:
        workbook.close()
