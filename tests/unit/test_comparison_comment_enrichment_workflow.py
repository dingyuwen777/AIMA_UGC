"""验证 Excel 到车型共现帖子全量评论 JSONL/Excel 的三阶段文件工作流。"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import SecretStr
from aima_ugc.adapters.providers.imports_test.comparison_comment_enrichment.enrich_comments import (
    enrich_comparison_comments,
)
from aima_ugc.adapters.providers.imports_test.comparison_comment_enrichment.models import (
    VehiclePairCommentRecordV1,
)
from aima_ugc.adapters.providers.imports_test.monitoring_excel_filter.process_directory import (
    process_directory as process_monitoring_directory,
)
from aima_ugc.adapters.providers.imports_test.vehicle_pair_filter.filter_vehicle_pairs import (
    filter_vehicle_pairs,
)
from aima_ugc.adapters.providers.tikhub import runtime as tikhub_runtime
from aima_ugc.adapters.providers.tikhub_test.core.config import TikHubTestConfig
from aima_ugc.contracts.canonical import CanonicalCommentV1, CanonicalMetricsV1, CanonicalSourceV1
from aima_ugc.contracts.platform import PLATFORM_NAMES, PlatformName
from aima_ugc.modules.collection.providers.transport import (
    ProviderTransportRequest,
    ProviderTransportResponse,
)

_HEADERS = (
    "媒体名称（中文）",
    "文章编号",
    "标题",
    "内文",
    "作者",
    "出版日期",
    "原文链接",
    "粉丝数",
)


class _WorkflowTransport:
    """为三阶段工作流返回每帖一个根评论的隔离 Provider。"""

    def __init__(self) -> None:
        self.requests: list[ProviderTransportRequest] = []

    def send(self, request: ProviderTransportRequest) -> ProviderTransportResponse:
        """记录真实 Runtime 请求并返回可由测试 Mapper 消费的根评论。"""

        self.requests.append(request)
        return ProviderTransportResponse(
            status_code=200,
            body={"items": [{"id": f"root-{len(self.requests)}", "reply_count": 0}]},
        )


def _write_monitoring_workbook(path: Path) -> None:
    """生成五平台车型共现帖、一个单车型帖和一个非目标平台帖。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "文章"
    worksheet.append(list(_HEADERS))
    rows = (
        (
            "小红书",
            "xiaohongshu-pair",
            "元宇宙实际体验",
            "Q3也试过",
            "测试作者",
            datetime(2026, 6, 15, 12, 0),
            "https://www.xiaohongshu.com/explore/xiaohongshu-pair",
            100,
        ),
        (
            "抖音",
            "douyin-pair",
            "墩墩实际体验",
            "莱茵也试过",
            "测试作者",
            datetime(2026, 6, 15, 12, 1),
            "https://www.douyin.com/video/douyin-pair",
            100,
        ),
        (
            "新浪微博",
            "weibo-pair",
            "元宇宙实际体验",
            "Y果冻也试过",
            "测试作者",
            datetime(2026, 6, 15, 12, 2),
            "https://weibo.com/1/weibo-pair",
            100,
        ),
        (
            "哔哩哔哩",
            "bilibili-pair",
            "墩墩实际体验",
            "QZ1也试过",
            "测试作者",
            datetime(2026, 6, 15, 12, 3),
            "https://www.bilibili.com/video/BV1testpair",
            100,
        ),
        (
            "快手",
            "kuaishou-pair",
            "元宇宙实际体验",
            "Mo1也试过",
            "测试作者",
            datetime(2026, 6, 15, 12, 4),
            "https://www.kuaishou.com/short-video/kuaishou-pair",
            100,
        ),
        (
            "小红书",
            "xiaohongshu-target-only",
            "元宇宙单车体验",
            "这里只聊一个爱玛车型",
            "测试作者",
            datetime(2026, 6, 15, 12, 5),
            "https://www.xiaohongshu.com/explore/xiaohongshu-target-only",
            100,
        ),
        (
            "微信",
            "wechat-pair",
            "元宇宙和Q3",
            "非目标平台应在第一阶段跳过",
            "测试作者",
            datetime(2026, 6, 15, 12, 6),
            "https://example.com/wechat-pair",
            100,
        ),
    )
    for row in rows:
        worksheet.append(list(row))
    workbook.save(path)
    workbook.close()


def _write_vehicle_catalog(path: Path) -> None:
    """写入三阶段工作流需要的目标品牌和竞品车型目录。"""

    path.write_text(
        json.dumps(
            {
                "schema_version": "vehicle-pair-catalog.v1",
                "target_brand": "爱玛",
                "brands": [
                    {
                        "name": "爱玛",
                        "models": [
                            {"name": "元宇宙", "aliases": []},
                            {"name": "墩墩", "aliases": []},
                        ],
                    },
                    {"name": "雅迪", "models": [{"name": "莱茵", "aliases": []}]},
                    {
                        "name": "九号",
                        "models": [
                            {"name": "Q3", "aliases": []},
                            {"name": "QZ1", "aliases": []},
                        ],
                    },
                    {"name": "小牛", "models": [{"name": "Y果冻", "aliases": []}]},
                    {"name": "极核", "models": [{"name": "Mo1", "aliases": []}]},
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def _install_workflow_runtime(monkeypatch: pytest.MonkeyPatch) -> None:
    """使用正式 build call，仅隔离外部响应解析以验证真实三阶段文件接线。"""

    monkeypatch.setattr(
        tikhub_runtime,
        "extract_comment_items",
        lambda platform, body: tuple(body.get("items", ())),
    )
    monkeypatch.setattr(
        tikhub_runtime,
        "advance_comments",
        lambda **kwargs: tikhub_runtime.TikHubPageAdvance(None, "provider_exhausted"),
    )

    def map_comment(
        *,
        platform: PlatformName,
        raw: dict[str, Any],
        context: Any,
        item_locator: str,
        is_root: bool,
    ) -> CanonicalCommentV1:
        if not is_root or not isinstance(context.external_content_id, str):
            raise AssertionError("工作流 fixture 只应产生一级评论")
        comment_id = str(raw["id"])
        return CanonicalCommentV1(
            platform=platform,
            external_content_id=context.external_content_id,
            external_comment_id=comment_id,
            root_comment_id=comment_id,
            parent_comment_id=None,
            text=f"真实工作流评论 {platform}",
            observed_at=context.observed_at,
            metrics=CanonicalMetricsV1(reply_count=0),
            source=CanonicalSourceV1(
                provider_name="tikhub",
                operation=context.operation,
                provider_request_id=context.provider_request_id,
                provider_attempt_id=context.provider_attempt_id,
                raw_artifact_id=context.raw_artifact_id,
                source_type=context.source_type,
                source_value=context.source_value,
                item_locator=item_locator,
                observed_at=context.observed_at,
            ),
            observed_fields=["root_comment_id", "parent_comment_id", "text", "metrics.reply_count"],
        )

    monkeypatch.setattr(tikhub_runtime, "map_comment", map_comment)


def test_three_stage_workflow_outputs_unified_comments_jsonl_and_excel(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Runner 应真实执行 XLSX → 宽筛去重 → 车型共现 → 五平台评论 JSONL/Excel。"""

    input_dir = tmp_path / "input"
    _write_monitoring_workbook(input_dir / "2026-06-15_sample.xlsx")
    keyword_pack = tmp_path / "keyword_pack.txt"
    keyword_pack.write_text(
        "元宇宙\n墩墩\nQ3\nQZ1\n莱茵\nY果冻\nMo1\n",
        encoding="utf-8",
    )
    first = process_monitoring_directory(
        input_dir=input_dir,
        output_root=tmp_path / "monitoring-output",
        keyword_pack_file=keyword_pack,
        sheet_name="文章",
        run_id="stage-one",
    )
    assert first.rows_seen == 7
    assert first.rows_supported_platform == 6
    assert first.rows_skipped_platform_unmapped == 1
    assert first.rows_after_deduplication == 6

    catalog_path = tmp_path / "vehicle_catalog.json"
    _write_vehicle_catalog(catalog_path)
    second = filter_vehicle_pairs(
        input_path=first.deduplicated_path,
        catalog_path=catalog_path,
        output_root=tmp_path / "pair-output",
        run_id="stage-two",
    )
    assert second.rows_seen == 6
    assert second.rows_with_cross_brand_pair == 5
    assert second.rows_filtered_out == 1

    _install_workflow_runtime(monkeypatch)
    transport = _WorkflowTransport()
    third = enrich_comparison_comments(
        input_path=second.output_path,
        output_root=tmp_path / "comment-output",
        run_id="stage-three",
        provider_config=TikHubTestConfig(
            base_url="https://api.tikhub.dev",
            api_key=SecretStr("test-key"),
            timeout_seconds=1,
        ),
        transport=transport,
    )

    assert third.rows_seen == 5
    assert third.rows_complete == 5
    assert third.rows_partial == 0
    assert third.rows_unavailable == 0
    assert third.root_comment_count == 5
    assert third.reply_count == 0
    assert third.request_count == 5
    assert third.platform_counts == {platform: 1 for platform in PLATFORM_NAMES}
    assert len(transport.requests) == 5

    enriched_records = [
        VehiclePairCommentRecordV1.model_validate_json(line)
        for line in third.output_jsonl_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert len(enriched_records) == 5
    assert {item.record.record.content.platform for item in enriched_records} == set(PLATFORM_NAMES)
    assert all(len(item.comments) == 1 for item in enriched_records)
    assert all(item.comment_fetch.coverage == "complete" for item in enriched_records)

    workbook = load_workbook(third.workbook_path, read_only=True, data_only=True)
    try:
        assert workbook["内容"].max_row == 6
        assert workbook["评论"].max_row == 6
    finally:
        workbook.close()

    for platform in PLATFORM_NAMES:
        raw_dir = third.run_dir / "provider" / platform / "runs" / "comments" / "raw"
        assert raw_dir.is_dir()
        assert len(list(raw_dir.glob("*.json"))) == 1
