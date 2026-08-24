"""AI relevant → 人工 irrelevant → 撤销的 PostgreSQL 审计闭环。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from aima_ugc.bootstrap.analysis_worker import PostgresContentAnalysisJobExecutor
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.content_http import PostgresContentHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.http import (
    ContentAnalysisSubmitRequest,
    ContentListQuery,
    ContentTargetSelection,
)
from aima_ugc.contracts.relevance_review import ContentRelevanceReviewRequest
from aima_ugc.modules.analysis import (
    CONTENT_LABELING_PROMPT_PATH,
    ContentLabelingService,
    FakeContentLabelingLLM,
    FrozenPromptTaxonomyLoader,
    PromptTaxonomyLoader,
)
from aima_ugc.modules.analysis.content_analysis_job import (
    ContentAnalysisJobHandler,
    register_content_analysis_job,
)
from aima_ugc.modules.analysis.relevance_review import ContentRelevanceReviewConflict
from aima_ugc.modules.analysis.relevance_review_tables import (
    analysis_content_relevance_reviews_table,
)
from aima_ugc.modules.analysis.tables import analysis_content_results_table
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import JobRegistry
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    sheet.append(
        [
            "小红书",
            "爱玛 Q7 相关性人工纠偏",
            "AI 认为与爱玛相关，但人工需要把它排除出业务相关集合。",
            "用户乙",
            "2026-08-20 11:00:00",
            "https://www.xiaohongshu.com/explore/manual-review-relevant-content",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _seed_import(client: TestClient) -> None:
    pack = client.post(
        "/api/v1/keyword-packs",
        json={"name": f"双向人工复核 {uuid4()}"},
    )
    assert pack.status_code == 201
    keyword = client.post(
        f"/api/v1/keyword-packs/{pack.json()['id']}/keywords",
        json={"text": "爱玛", "priority": 10},
    )
    assert keyword.status_code == 201
    uploaded = client.post(
        "/api/v1/import-batches",
        files=[
            (
                "file",
                (
                    "bidirectional-review.xlsx",
                    _xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("keyword_pack_ids", (None, pack.json()["id"])),
        ],
    )
    assert uploaded.status_code == 202


def _analysis_registry(runtime) -> JobRegistry:  # type: ignore[no-untyped-def]
    taxonomy = PromptTaxonomyLoader(CONTENT_LABELING_PROMPT_PATH).load()
    service = ContentLabelingService(
        prompt_loader=FrozenPromptTaxonomyLoader(taxonomy),
        llm=FakeContentLabelingLLM(
            responses=[
                '{"items":[{"item_no":1,"relevance":"relevant",'
                '"voice_type":"user_voice","sentiment":"中性",'
                '"labels":[{"primary_label":"骑行性能","secondary_label":"舒适性"}]}]}'
            ]
        ),
    )
    executor = PostgresContentAnalysisJobExecutor(
        runtime,
        service_factory=lambda: (service, lambda: None),
    )
    registry = JobRegistry()
    register_content_analysis_job(registry, ContentAnalysisJobHandler(executor))
    return registry


def test_manual_irrelevant_override_and_undo_are_append_only_and_preserve_ai_result(
    tmp_path: Path,
) -> None:
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "llm_base_url": "https://fake.example/v1",
            "llm_provider_name": "fake",
            "llm_model": "fake-content-labeler-v1",
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts RESTART IDENTITY CASCADE"
        )
    try:
        import_client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        _seed_import(import_client)
        import_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="bidirectional-review-import",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert import_worker.run_once() is True

        with runtime.database.engine.begin() as connection:
            content_id = connection.scalar(select(contents_table.c.id))
        assert content_id is not None

        content_service = PostgresContentHttpService(
            runtime,
            cursor_signing_secret=b"bidirectional-review-cursor-key-32bytes",
        )
        created = content_service.create_analysis(
            ContentAnalysisSubmitRequest(
                targets=ContentTargetSelection(scope="selected", content_ids=(content_id,))
            ),
            request_id="bidirectional-review-analysis",
        )
        analysis_worker = create_job_worker(
            runtime=runtime,
            registry=_analysis_registry(runtime),
            worker_id="bidirectional-review-analysis",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert analysis_worker.run_once() is True
        assert content_service.get_analysis_job(created.job_id).status == "succeeded"

        default_item = content_service.list_contents(ContentListQuery()).items[0]
        assert default_item.id == content_id
        assert default_item.analysis.relevance == "relevant"
        assert default_item.effective_relevance == "relevant"
        assert default_item.relevance_source == "ai"
        assert content_service.list_contents(ContentListQuery(relevance="irrelevant")).items == ()

        excluded = content_service.review_relevance(
            ContentRelevanceReviewRequest(content_ids=(content_id,), decision="irrelevant"),
            request_id="manual-exclude",
        )
        assert excluded.requested_count == 1
        assert excluded.changed_count == 1
        assert excluded.unchanged_count == 0
        assert content_service.list_contents(ContentListQuery()).items == ()
        irrelevant_item = content_service.list_contents(
            ContentListQuery(relevance="irrelevant")
        ).items[0]
        assert irrelevant_item.id == content_id
        assert irrelevant_item.analysis.relevance == "relevant"
        assert irrelevant_item.effective_relevance == "irrelevant"
        assert irrelevant_item.relevance_source == "manual_review"

        repeated = content_service.review_relevance(
            ContentRelevanceReviewRequest(content_ids=(content_id,), decision="irrelevant"),
            request_id="manual-exclude-repeat",
        )
        assert repeated.changed_count == 0
        assert repeated.unchanged_count == 1

        with pytest.raises(ContentRelevanceReviewConflict):
            content_service.review_relevance(
                ContentRelevanceReviewRequest(content_ids=(content_id,), decision="relevant"),
                request_id="manual-direct-reverse",
            )
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(
                    select(func.count()).select_from(analysis_content_relevance_reviews_table)
                )
                == 1
            )

        # 模型身份变化后当前 AI Result 变 stale，但活动人工覆盖仍必须可识别并可撤销。
        runtime.settings = runtime.settings.model_copy(update={"llm_model": "stale-model"})
        stale_override = content_service.list_contents(
            ContentListQuery(relevance="irrelevant")
        ).items[0]
        assert stale_override.analysis.status == "stale"
        assert stale_override.analysis.relevance is None
        assert stale_override.effective_relevance == "irrelevant"
        assert stale_override.relevance_source == "manual_review"

        undone = content_service.review_relevance(
            ContentRelevanceReviewRequest(content_ids=(content_id,), decision="inherit_ai"),
            request_id="manual-undo",
        )
        assert undone.changed_count == 1
        assert undone.unchanged_count == 0
        inherited = content_service.list_contents(ContentListQuery()).items[0]
        assert inherited.id == content_id
        assert inherited.analysis.status == "stale"
        assert inherited.effective_relevance is None
        assert inherited.relevance_source is None
        assert content_service.list_contents(ContentListQuery(relevance="irrelevant")).items == ()

        undo_again = content_service.review_relevance(
            ContentRelevanceReviewRequest(content_ids=(content_id,), decision="inherit_ai"),
            request_id="manual-undo-repeat",
        )
        assert undo_again.changed_count == 0
        assert undo_again.unchanged_count == 1

        with runtime.database.engine.begin() as connection:
            raw_result = connection.execute(
                select(
                    analysis_content_results_table.c.id,
                    analysis_content_results_table.c.relevance,
                    analysis_content_results_table.c.sentiment,
                ).where(analysis_content_results_table.c.content_id == content_id)
            ).one()
            events = connection.execute(
                select(
                    analysis_content_relevance_reviews_table.c.review_no,
                    analysis_content_relevance_reviews_table.c.decision,
                    analysis_content_relevance_reviews_table.c.analysis_result_id,
                )
                .where(analysis_content_relevance_reviews_table.c.content_id == content_id)
                .order_by(analysis_content_relevance_reviews_table.c.review_no)
            ).all()
        assert raw_result.relevance == "relevant"
        assert raw_result.sentiment == "中性"
        assert [(row.review_no, row.decision) for row in events] == [
            (1, "irrelevant"),
            (2, "inherit_ai"),
        ]
        assert all(row.analysis_result_id == raw_result.id for row in events)
    finally:
        runtime.close()
