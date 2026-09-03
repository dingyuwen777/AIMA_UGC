"""CHG-314 第一阶段 Review 发现的跨 Contract / PostgreSQL 回归。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

from aima_ugc.adapters.persistence.postgres.analysis import PostgresAnalysisRepository
from aima_ugc.adapters.persistence.postgres.analysis_schemes import (
    PostgresAnalysisSchemeRepository,
)
from aima_ugc.bootstrap.analysis_worker import (
    PostgresContentAnalysisJobExecutor,
    PostgresContentAnalysisPlanJobExecutor,
    create_analysis_job_terminal_callback,
)
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.content_http import PostgresContentHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.modules.analysis import (
    ContentLabelingService,
    FakeContentLabelingLLM,
    FrozenPromptTaxonomyLoader,
)
from aima_ugc.modules.analysis.content_analysis_job import (
    ContentAnalysisJobHandler,
    ContentAnalysisPlanJobHandler,
    register_content_analysis_job,
)
from aima_ugc.modules.analysis.schemes import prompt_taxonomy_from_version
from aima_ugc.modules.analysis.tables import (
    analysis_content_requests_table,
    analysis_content_run_targets_table,
)
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.modules.system.tables import audit_events_table, keyword_packs_table
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import JobRegistry
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select


def _xlsx() -> bytes:
    """生成三条固定 Content 的最小 Excel Fixture。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    for index in range(3):
        sheet.append(
            [
                "小红书",
                f"CHG314 Review {index}",
                f"第 {index} 条固定测试内容",
                f"用户 {index}",
                f"2026-09-03 08:0{index}:00",
                f"https://www.xiaohongshu.com/explore/chg314-review-{index}",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _runtime(tmp_path: Path):  # type: ignore[no-untyped-def]
    """创建启用 Fake LLM 身份的隔离测试 Runtime，并清空共享 PostgreSQL 事实。"""

    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "llm_base_url": "https://fake.example/v1",
            "llm_provider_name": "fake",
            "llm_model": "fake-content-labeler-v1",
            "analysis_run_shard_size": 1,
            "analysis_run_max_in_flight_jobs": 2,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts, audit_events "
            "RESTART IDENTITY CASCADE"
        )
    return runtime


def _client(runtime) -> TestClient:  # type: ignore[no-untyped-def]
    """构造同时使用正式 Import 与 Content Service 的测试 HTTP Client。"""

    return TestClient(
        create_app(
            import_service=PostgresImportHttpService(runtime),
            content_service=PostgresContentHttpService(
                runtime,
                cursor_signing_secret=b"chg314-review-content-cursor-key",
            ),
        )
    )


def _seed_contents(client: TestClient, runtime) -> tuple[UUID, ...]:  # type: ignore[no-untyped-def]
    """通过正式 Excel Import + Worker 链写入三条 Content，并返回稳定排序后的 ID。"""

    pack = client.post("/api/v1/keyword-packs", json={"name": f"CHG314 Review {uuid4()}"})
    assert pack.status_code == 201
    keyword = client.post(
        f"/api/v1/keyword-packs/{pack.json()['id']}/keywords",
        json={"text": "CHG314", "priority": 10},
    )
    assert keyword.status_code == 201
    uploaded = client.post(
        "/api/v1/import-batches",
        files=[
            (
                "file",
                (
                    "chg314-review.xlsx",
                    _xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("keyword_pack_ids", (None, pack.json()["id"])),
        ],
    )
    assert uploaded.status_code == 202
    worker = create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id="chg314-review-import",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    assert worker.run_once() is True
    with runtime.database.engine.begin() as connection:
        return tuple(
            connection.execute(select(contents_table.c.id).order_by(contents_table.c.id))
            .scalars()
            .all()
        )


def _analysis_registry(runtime, *, raw_response: str) -> JobRegistry:  # type: ignore[no-untyped-def]
    """构造复用正式 Planner/Sharding 的 Fake LLM Analysis Registry。"""

    session = runtime.database.new_session()
    try:
        with session.begin():
            version = PostgresAnalysisSchemeRepository(session).get_active_version()
            assert version is not None
            taxonomy = prompt_taxonomy_from_version(version)
    finally:
        session.close()
    service = ContentLabelingService(
        prompt_loader=FrozenPromptTaxonomyLoader(taxonomy),
        llm=FakeContentLabelingLLM(responses=[raw_response] * 8),
    )
    registry = JobRegistry()
    callback = create_analysis_job_terminal_callback(runtime)
    register_content_analysis_job(
        registry,
        ContentAnalysisJobHandler(
            PostgresContentAnalysisJobExecutor(
                runtime,
                service_factory=lambda: (service, lambda: None),
            )
        ),
        terminal_callback=callback,
        planner_handler=ContentAnalysisPlanJobHandler(
            PostgresContentAnalysisPlanJobExecutor(runtime)
        ),
        planner_terminal_callback=callback,
    )
    return registry


def _irrelevant_response() -> str:
    """返回符合 V3 Contract 的固定“不相关”模型结果。"""

    return (
        '{"items":[{"item_no":1,"relevance":"irrelevant",'
        '"voice_type":"真实用户发声","sentiment":null,"labels":[]}]}'
    )


def _relevant_response() -> str:
    """返回符合当前 Taxonomy 的固定“相关”模型结果。"""

    return (
        '{"items":[{"item_no":1,"relevance":"relevant",'
        '"voice_type":"真实用户发声","sentiment":"负面",'
        '"labels":[{"primary_label":"骑行性能","secondary_label":"舒适性"}]}]}'
    )


def _run_selected_analysis(
    client: TestClient,
    runtime,
    content_id: UUID,
    *,
    raw_response: str,
) -> None:  # type: ignore[no-untyped-def]
    """对一条 Content 运行真实 Planner + Shard，并等待 Worker 队列清空。"""

    targets = {"scope": "selected", "content_ids": [str(content_id)]}
    preview = client.post("/api/v1/analysis/content-runs/preview", json={"targets": targets})
    assert preview.status_code == 200
    created = client.post(
        "/api/v1/analysis/content-runs",
        json={
            "client_idempotency_key": f"chg314-review-selected-{uuid4()}",
            "targets": targets,
            "expected_target_count": 1,
            "expected_configuration_hash": preview.json()["configuration_hash"],
            "run_intent": "manual_reanalysis",
        },
    )
    assert created.status_code == 202
    worker = create_job_worker(
        runtime=runtime,
        registry=_analysis_registry(runtime, raw_response=raw_response),
        worker_id=f"chg314-review-analysis-{uuid4()}",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    completed = 0
    while worker.run_once():
        completed += 1
        if completed > 4:
            raise AssertionError("单条 Analysis Run 超出预期有界 Job 数")
    run = client.get(f"/api/v1/analysis/content-runs/{created.json()['run_id']}")
    assert run.status_code == 200
    assert run.json()["status"] == "succeeded"


def test_atomic_keyword_pack_create_returns_committed_version(tmp_path: Path) -> None:
    """原子创建返回值与审计必须反映关键词写入后的真实 Pack Version。"""

    runtime = _runtime(tmp_path)
    try:
        client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        created = client.post(
            "/api/v1/keyword-packs",
            json={
                "name": f"CHG314 Atomic {uuid4()}",
                "keywords": [
                    {"text": f"关键词-A-{uuid4()}", "priority": 100},
                    {"text": f"关键词-B-{uuid4()}", "priority": 80},
                ],
            },
        )
        assert created.status_code == 201
        pack_id = UUID(created.json()["id"])
        assert created.json()["version"] == 3
        with runtime.database.engine.begin() as connection:
            assert connection.scalar(
                select(keyword_packs_table.c.version).where(keyword_packs_table.c.id == pack_id)
            ) == created.json()["version"]
            safe_detail = connection.scalar(
                select(audit_events_table.c.safe_detail)
                .where(
                    audit_events_table.c.event_type == "keyword_pack_created",
                    audit_events_table.c.object_id == str(pack_id),
                )
                .order_by(audit_events_table.c.created_at.desc())
                .limit(1)
            )
            assert safe_detail is not None
            assert safe_detail["version"] == created.json()["version"]
    finally:
        runtime.close()


def test_legacy_empty_query_run_remains_query_scope(tmp_path: Path) -> None:
    """历史兼容 query 空筛选不能因为新增公开 all Scope 被重新解释。"""

    runtime = _runtime(tmp_path)
    try:
        client = _client(runtime)
        assert len(_seed_contents(client, runtime)) == 3
        created = client.post(
            "/api/v1/content-analysis-requests",
            json={"targets": {"scope": "query", "filters": {}}},
        )
        assert created.status_code == 202
        run = client.get(f"/api/v1/analysis/content-runs/{created.json()['run_id']}")
        assert run.status_code == 200
        assert run.json()["scope"] == "query"
    finally:
        runtime.close()


def test_all_scope_preview_includes_current_irrelevant_content(tmp_path: Path) -> None:
    """公开 all 必须覆盖全部 Current Content，不能继承声音广场默认隐藏不相关内容。"""

    runtime = _runtime(tmp_path)
    try:
        client = _client(runtime)
        content_ids = _seed_contents(client, runtime)
        assert len(content_ids) == 3
        _run_selected_analysis(
            client,
            runtime,
            content_ids[0],
            raw_response=_irrelevant_response(),
        )
        irrelevant = client.get("/api/v1/contents?relevance=irrelevant&limit=10")
        assert irrelevant.status_code == 200
        assert len(irrelevant.json()["items"]) == 1

        preview = client.post(
            "/api/v1/analysis/content-runs/preview",
            json={"targets": {"scope": "all"}},
        )
        assert preview.status_code == 200
        assert preview.json()["target_count"] == 3
    finally:
        runtime.close()


def test_all_scope_planner_does_not_use_unbounded_freeze(
    tmp_path: Path,
    monkeypatch,
) -> None:  # type: ignore[no-untyped-def]
    """全量 Planner 不得回退到一次事务冻结整个目标集合的旧入口。"""

    runtime = _runtime(tmp_path)
    try:
        client = _client(runtime)
        assert len(_seed_contents(client, runtime)) == 3
        preview = client.post(
            "/api/v1/analysis/content-runs/preview",
            json={"targets": {"scope": "all"}},
        )
        assert preview.status_code == 200
        created = client.post(
            "/api/v1/analysis/content-runs",
            json={
                "client_idempotency_key": f"chg314-review-all-bounded-{uuid4()}",
                "targets": {"scope": "all"},
                "expected_target_count": 3,
                "expected_configuration_hash": preview.json()["configuration_hash"],
                "run_intent": "manual_reanalysis",
            },
        )
        assert created.status_code == 202

        def reject_unbounded_freeze(self, *, run_id, target_statement):  # type: ignore[no-untyped-def]
            """若 all 仍调用旧全量 INSERT...SELECT，则明确让回归失败。"""

            del self, run_id, target_statement
            raise AssertionError("all Scope 不得使用单事务全量 freeze_run_targets")

        monkeypatch.setattr(
            PostgresAnalysisRepository,
            "freeze_run_targets",
            reject_unbounded_freeze,
        )
        worker = create_job_worker(
            runtime=runtime,
            registry=_analysis_registry(runtime, raw_response=_relevant_response()),
            worker_id="chg314-review-all-bounded",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert worker.run_once() is True
        run_id = UUID(created.json()["run_id"])
        with runtime.database.engine.begin() as connection:
            assert connection.scalar(
                select(func.count()).where(analysis_content_run_targets_table.c.run_id == run_id)
            ) == 3
            assert connection.scalar(
                select(func.count()).where(analysis_content_requests_table.c.run_id == run_id)
            ) == 2
    finally:
        runtime.close()
