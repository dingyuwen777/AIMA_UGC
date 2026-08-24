"""AI irrelevant → 人工 relevant 的 PostgreSQL 复核闭环。"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from aima_ugc.bootstrap.analysis_worker import PostgresContentAnalysisJobExecutor
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.content_http import PostgresContentHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.reporting_http import PostgresReportingHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.http import (
    ContentAnalysisSubmitRequest,
    ContentFilterSnapshot,
    ContentListQuery,
    ContentTargetSelection,
    DataExportSubmitRequest,
)
from aima_ugc.contracts.relevance_review import ContentRelevanceReviewRequest
from aima_ugc.modules.analysis import (
    CONTENT_LABELING_PROMPT_PATH,
    ContentLabelingService,
    FakeContentLabelingLLM,
    FrozenPromptTaxonomyLoader,
    PromptTaxonomyLoader,
)
from aima_ugc.modules.analysis.content_analysis_job import (
    ContentAnalysisJobHandler,
    register_content_analysis_job,
)
from aima_ugc.modules.analysis.relevance_review import ContentRelevanceReviewConflict
from aima_ugc.modules.analysis.relevance_review_tables import (
    analysis_content_relevance_reviews_table,
)
from aima_ugc.modules.analysis.tables import analysis_content_results_table
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import JobRegistry
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select


def _xlsx(*, text_suffix: str = "") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    sheet.append(
        [
            "小红书",
            "爱玛 Q7 人工复核",
            f"第一条内容{text_suffix}",
            "用户甲",
            "2026-08-20 10:00:00",
            "https://www.xiaohongshu.com/explore/manual-review-content-1",
        ]
    )
    sheet.append(
        [
            "小红书",
            "爱玛其他内容",
            f"第二条内容{text_suffix}",
            "用户乙",
            "2026-08-20 11:00:00",
            "https://www.xiaohongshu.com/explore/manual-review-content-2",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _seed_import(client: TestClient, *, text_suffix: str = "") -> str:
    pack = client.post(
        "/api/v1/keyword-packs",
        json={"name": f"人工复核相关性 {uuid4()}"},
    )
    assert pack.status_code == 201
    keyword = client.post(
        f"/api/v1/keyword-packs/{pack.json()['id']}/keywords",
        json={"text": "爱玛", "priority": 10},
    )
    assert keyword.status_code == 201
    uploaded = client.post(
        "/api/v1/import-batches",
        files=[
            (
                "file",
                (
                    "manual-review.xlsx",
                    _xlsx(text_suffix=text_suffix),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("keyword_pack_ids", (None, pack.json()["id"])),
        ],
    )
    assert uploaded.status_code == 202
    return str(uploaded.json()["batch_id"])


def _analysis_registry(runtime, response: str) -> JobRegistry:  # type: ignore[no-untyped-def]
    taxonomy = PromptTaxonomyLoader(CONTENT_LABELING_PROMPT_PATH).load()
    service = ContentLabelingService(
        prompt_loader=FrozenPromptTaxonomyLoader(taxonomy),
        llm=FakeContentLabelingLLM(responses=[response]),
    )
    executor = PostgresContentAnalysisJobExecutor(
        runtime,
        service_factory=lambda: (service, lambda: None),
    )
    registry = JobRegistry()
    register_content_analysis_job(registry, ContentAnalysisJobHandler(executor))
    return registry


def _irrelevant_response() -> str:
    return (
        '{"items":[{"item_no":1,"relevance":"irrelevant",'
        '"voice_type":"media_information","sentiment":null,"labels":[]}]}'
    )


def test_manual_relevance_review_preserves_ai_result_and_drives_business_queries(
    tmp_path: Path,
) -> None:
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "llm_base_url": "https://fake.example/v1",
            "llm_provider_name": "fake",
            "llm_model": "fake-content-labeler-v1",
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts RESTART IDENTITY CASCADE"
        )
    try:
        import_client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        _seed_import(import_client)
        import_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="manual-review-import",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert import_worker.run_once() is True

        with runtime.database.engine.begin() as connection:
            content_ids = tuple(
                connection.execute(
                    select(contents_table.c.id).order_by(contents_table.c.published_at)
                ).scalars()
            )
        assert len(content_ids) == 2

        content_service = PostgresContentHttpService(
            runtime,
            cursor_signing_secret=b"manual-review-content-cursor-key-32-bytes-minimum",
        )
        created = content_service.create_analysis(
            ContentAnalysisSubmitRequest(
                targets=ContentTargetSelection(scope="selected", content_ids=(content_ids[0],))
            ),
            request_id="manual-review-analysis",
        )
        analysis_worker = create_job_worker(
            runtime=runtime,
            registry=_analysis_registry(runtime, _irrelevant_response()),
            worker_id="manual-review-analysis",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert analysis_worker.run_once() is True
        assert content_service.get_analysis_job(created.job_id).status == "succeeded"

        assert content_ids[0] not in {
            item.id for item in content_service.list_contents(ContentListQuery()).items
        }
        irrelevant = content_service.list_contents(ContentListQuery(relevance="irrelevant"))
        assert [item.id for item in irrelevant.items] == [content_ids[0]]
        assert irrelevant.items[0].analysis.relevance == "irrelevant"

        # 混入一个未打标 Content 时整批失败，不能悄悄只复核部分选择。
        with pytest.raises(ContentRelevanceReviewConflict):
            content_service.review_relevance(
                ContentRelevanceReviewRequest(content_ids=content_ids),
                request_id="manual-review-atomic",
            )
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(
                    select(func.count()).select_from(analysis_content_relevance_reviews_table)
                )
                == 0
            )

        reviewed = content_service.review_relevance(
            ContentRelevanceReviewRequest(content_ids=(content_ids[0],)),
            request_id="manual-review-single",
        )
        assert reviewed.requested_count == 1
        assert reviewed.reviewed_count == 1
        assert reviewed.already_reviewed_count == 0

        repeated = content_service.review_relevance(
            ContentRelevanceReviewRequest(content_ids=(content_ids[0],)),
            request_id="manual-review-repeat",
        )
        assert repeated.reviewed_count == 0
        assert repeated.already_reviewed_count == 1

        default_ids = {item.id for item in content_service.list_contents(ContentListQuery()).items}
        assert content_ids[0] in default_ids
        relevant = content_service.list_contents(ContentListQuery(relevance="relevant"))
        assert [item.id for item in relevant.items] == [content_ids[0]]
        assert content_service.list_contents(ContentListQuery(relevance="irrelevant")).items == ()

        # Public Analysis 仍然是模型原始判断；人工决定只改变有效业务相关性。
        detail = content_service.get_content(content_ids[0])
        assert detail.analysis.relevance == "irrelevant"
        assert detail.analysis.sentiment is None
        assert detail.analysis.labels == ()
        with runtime.database.engine.begin() as connection:
            result = connection.execute(
                select(
                    analysis_content_results_table.c.relevance,
                    analysis_content_results_table.c.sentiment,
                ).where(analysis_content_results_table.c.content_id == content_ids[0])
            ).one()
            assert result.relevance == "irrelevant"
            assert result.sentiment is None

        # 查询型 Analysis / Export 都复用同一 effective relevance，不另造过滤规则。
        query_analysis = content_service.create_analysis(
            ContentAnalysisSubmitRequest(
                targets=ContentTargetSelection(
                    scope="query",
                    filters=ContentFilterSnapshot(relevance="relevant"),
                )
            ),
            request_id="manual-review-query-analysis",
        )
        assert query_analysis.target_count == 1
        export = PostgresReportingHttpService(runtime).create_export(
            DataExportSubmitRequest(
                targets=ContentTargetSelection(
                    scope="query",
                    filters=ContentFilterSnapshot(relevance="relevant"),
                )
            ),
            request_id="manual-review-query-export",
        )
        assert export.target_count == 1

        # 相同外部 Content 出现新正文版本后，旧人工判断继续留作审计但不套到 V2。
        updated_batch_id = _seed_import(import_client, text_suffix="（更新）")
        for _ in range(10):
            batch = import_client.get(f"/api/v1/import-batches/{updated_batch_id}")
            assert batch.status_code == 200
            batch_status = batch.json()["status"]
            if batch_status == "succeeded":
                break
            assert batch_status in {"queued", "running"}
            assert import_worker.run_once() is True
        else:
            pytest.fail("更新版 Import Batch 未在测试预算内进入 succeeded")

        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(
                    select(contents_table.c.current_version).where(
                        contents_table.c.id == content_ids[0]
                    )
                )
                == 2
            )
            review_row = connection.execute(
                select(
                    analysis_content_relevance_reviews_table.c.content_version,
                    analysis_content_relevance_reviews_table.c.decision,
                    analysis_content_relevance_reviews_table.c.reviewed_at,
                ).where(analysis_content_relevance_reviews_table.c.content_id == content_ids[0])
            ).one()
            assert review_row.content_version == 1
            assert review_row.decision == "relevant"
            assert isinstance(review_row.reviewed_at, datetime)
            assert review_row.reviewed_at.tzinfo is not None

        assert content_service.list_contents(ContentListQuery(relevance="relevant")).items == ()
        with pytest.raises(ContentRelevanceReviewConflict):
            content_service.review_relevance(
                ContentRelevanceReviewRequest(content_ids=(content_ids[0],)),
                request_id="manual-review-version-2",
            )
    finally:
        runtime.close()
