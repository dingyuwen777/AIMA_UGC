"""Stage 5 Brand/Vehicle 查询、目标冻结与 Excel 导出 PostgreSQL 纵切。"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

from aima_ugc.adapters.persistence.postgres.brand_vehicle import (
    PostgresBrandVehicleRepository,
)
from aima_ugc.adapters.persistence.postgres.reporting import PostgresDataExportRepository
from aima_ugc.adapters.persistence.postgres.vehicles import PostgresVehicleCatalogRepository
from aima_ugc.bootstrap.administration_http import PostgresAdministrationHttpService
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.brand_vehicle_http import PostgresBrandVehicleHttpService
from aima_ugc.bootstrap.content_http import PostgresContentHttpService
from aima_ugc.bootstrap.export_worker import PostgresDataExportJobExecutor
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.product_http import PostgresProductHttpService
from aima_ugc.bootstrap.reporting_http import PostgresReportingHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.administration import VehicleModelCreateRequest
from aima_ugc.contracts.brand_vehicle import BrandCreateRequest
from aima_ugc.contracts.http import (
    ContentAnalysisSubmitRequest,
    ContentCountRequest,
    ContentFilterSnapshot,
    ContentListQuery,
    ContentTargetSelection,
    DataExportSubmitRequest,
)
from aima_ugc.modules.analysis.tables import analysis_content_run_targets_table
from aima_ugc.modules.content.tables import content_versions_table, contents_table
from aima_ugc.modules.identity import Principal
from aima_ugc.modules.reporting.data_export_job import (
    DataExportJobHandler,
    register_data_export_job,
)
from aima_ugc.modules.reporting.tables import reporting_data_export_items_table
from aima_ugc.modules.vehicles.tables import (
    content_brand_evidence_table,
    content_vehicle_evidence_table,
)
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import JobRegistry
from aima_ugc.platform.time import beijing_now
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import insert, select, update


def _principal() -> Principal:
    return Principal(
        principal_id="stage5-admin",
        display_name="Stage5 管理员",
        role="administrator",
        source="development",
    )


def _xlsx(rows: tuple[tuple[str, str], ...]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    for index, (title, url_key) in enumerate(rows):
        sheet.append(
            [
                "小红书",
                title,
                "Stage 5 品牌车型查询测试正文",
                f"Stage5 作者{index}",
                f"2026-09-10 1{index}:00:00",
                f"https://www.xiaohongshu.com/explore/{url_key}",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _import(runtime, workbook: bytes, *, worker_id: str) -> None:  # type: ignore[no-untyped-def]
    client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
    response = client.post(
        "/api/v1/import-batches",
        files={
            "file": (
                "stage5.xlsx",
                workbook,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 202
    worker = create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id=worker_id,
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    assert worker.run_once() is True


def test_brand_vehicle_filters_share_targets_and_export_frozen_version(tmp_path: Path) -> None:
    """五种竞品范围、Brand/Vehicle AND、四消费者与冻结导出保持同一事实。"""

    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "llm_base_url": "https://fake.example/v1",
            "llm_provider_name": "fake",
            "llm_model": "fake-stage5-labeler",
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts, vehicle_brands, "
            "vehicle_models RESTART IDENTITY CASCADE"
        )
    try:
        principal = _principal()
        brand_service = PostgresBrandVehicleHttpService(runtime)
        vehicle_service = PostgresAdministrationHttpService(runtime)
        owned = brand_service.create_brand(
            BrandCreateRequest(
                code="AIMA-STAGE5",
                display_name="爱玛 Stage5",
                role="owned",
                aliases=("爱玛舞台",),
            ),
            principal=principal,
            request_id="stage5-owned",
        )
        competitor = brand_service.create_brand(
            BrandCreateRequest(
                code="COMP-STAGE5",
                display_name="竞品 Stage5",
                role="competitor",
                aliases=("竞品舞台",),
            ),
            principal=principal,
            request_id="stage5-competitor",
        )
        other = brand_service.create_brand(
            BrandCreateRequest(
                code="OTHER-STAGE5",
                display_name="其他 Stage5",
                role="other",
                aliases=("其他舞台",),
            ),
            principal=principal,
            request_id="stage5-other",
        )
        source_vehicle = vehicle_service.create_vehicle_model(
            VehicleModelCreateRequest(
                code="OLD-STAGE5",
                display_name="旧车型 Stage5",
                brand_id=owned.id,
                aliases=("旧车型舞台",),
            ),
            principal=principal,
            request_id="stage5-source-vehicle",
        )
        target_vehicle = vehicle_service.create_vehicle_model(
            VehicleModelCreateRequest(
                code="NEW-STAGE5",
                display_name="新车型 Stage5",
                brand_id=owned.id,
                aliases=("新车型舞台",),
            ),
            principal=principal,
            request_id="stage5-target-vehicle",
        )

        rows = (
            ("爱玛舞台 旧车型舞台", "stage5-owned-vehicle"),
            ("竞品舞台", "stage5-competitor"),
            ("爱玛舞台 竞品舞台", "stage5-mixed"),
            ("其他舞台", "stage5-other"),
            ("爱玛舞台 历史未分类", "stage5-none"),
        )
        _import(runtime, _xlsx(rows), worker_id="stage5-import")

        with runtime.database.engine.begin() as connection:
            title_rows = connection.execute(
                select(contents_table.c.id, content_versions_table.c.title)
                .join(
                    content_versions_table,
                    (content_versions_table.c.content_id == contents_table.c.id)
                    & (
                        content_versions_table.c.version_no
                        == contents_table.c.current_version
                    ),
                )
            ).all()
            content_by_title = {str(title): UUID(str(content_id)) for content_id, title in title_rows}
            none_id = content_by_title["爱玛舞台 历史未分类"]
            connection.execute(
                update(content_brand_evidence_table)
                .where(content_brand_evidence_table.c.content_id == none_id)
                .values(is_active=False)
            )
            connection.execute(
                update(content_vehicle_evidence_table)
                .where(content_vehicle_evidence_table.c.content_id == none_id)
                .values(is_active=False)
            )

        session = runtime.database.new_session()
        try:
            with session.begin():
                PostgresVehicleCatalogRepository(session).merge_model(
                    source_vehicle.id,
                    target_vehicle.id,
                    actor_ref=principal.principal_id,
                )
        finally:
            session.close()

        content_service = PostgresContentHttpService(
            runtime,
            cursor_signing_secret=b"stage5-query-cursor-secret-is-32-bytes",
        )
        scope_titles = {
            scope: {
                item.title
                for item in content_service.list_contents(
                    ContentListQuery(competition_scopes=(scope,), limit=100)
                ).items
            }
            for scope in (
                "owned_only",
                "competitor_only",
                "mixed",
                "other_only",
                "none_detected",
            )
        }
        assert scope_titles == {
            "owned_only": {"爱玛舞台 旧车型舞台"},
            "competitor_only": {"竞品舞台"},
            "mixed": {"爱玛舞台 竞品舞台"},
            "other_only": {"其他舞台"},
            "none_detected": {"爱玛舞台 历史未分类"},
        }

        mixed = content_service.list_contents(
            ContentListQuery(brand_ids=(owned.id, competitor.id), competition_scopes=("mixed",))
        ).items[0]
        assert [(brand.display_name, brand.role) for brand in mixed.brands] == [
            ("爱玛 Stage5", "owned"),
            ("竞品 Stage5", "competitor"),
        ]
        assert mixed.competition_scope == "mixed"

        shared_filter = ContentFilterSnapshot(
            brand_ids=(owned.id,),
            vehicle_model_ids=(target_vehicle.id,),
            competition_scopes=("owned_only",),
        )
        listed = content_service.list_contents(
            ContentListQuery.model_validate(shared_filter.model_dump())
        )
        assert len(listed.items) == 1
        projected = listed.items[0]
        assert projected.title == "爱玛舞台 旧车型舞台"
        assert projected.competition_scope == "owned_only"
        assert projected.brands[0].display_name == "爱玛 Stage5"
        assert {evidence.source for evidence in projected.brands[0].evidences} == {
            "alias_match",
            "vehicle_match",
        }
        assert projected.vehicles[0].vehicle_model_id == target_vehicle.id
        assert projected.vehicles[0].display_name == "新车型 Stage5"
        assert projected.vehicles[0].brand is not None
        assert projected.vehicles[0].brand.id == owned.id

        count = PostgresProductHttpService(runtime).count_contents(
            ContentCountRequest(filters=shared_filter, count_mode="exact", exact_limit=100)
        )
        assert count.count == 1
        assert count.count_kind == "exact"

        analysis = content_service.create_analysis(
            ContentAnalysisSubmitRequest(
                targets=ContentTargetSelection(scope="query", filters=shared_filter)
            ),
            request_id="stage5-analysis-targets",
        )
        assert analysis.run_id is not None

        reporting = PostgresReportingHttpService(runtime)
        export = reporting.create_export(
            DataExportSubmitRequest(
                targets=ContentTargetSelection(scope="query", filters=shared_filter),
                columns=(
                    "matched_keywords",
                    "brands",
                    "brand_roles",
                    "competition_scope",
                    "vehicles",
                ),
            ),
            request_id="stage5-export-targets",
            actor_ref="user:stage5",
        )
        expected_ids = {projected.id}
        with runtime.database.engine.begin() as connection:
            analysis_ids = set(
                connection.scalars(
                    select(analysis_content_run_targets_table.c.content_id).where(
                        analysis_content_run_targets_table.c.run_id == analysis.run_id
                    )
                )
            )
            export_ids = set(
                connection.scalars(
                    select(reporting_data_export_items_table.c.content_id).where(
                        reporting_data_export_items_table.c.export_id == export.export_id
                    )
                )
            )
        assert analysis.target_count == export.target_count == count.count == len(listed.items) == 1
        assert analysis_ids == export_ids == expected_ids

        # 模拟目标冻结后 Content Current 前进到 v2；Export 必须继续读冻结的 v1 Evidence。
        session = runtime.database.new_session()
        try:
            with session.begin():
                old_version = (
                    session.execute(
                        select(content_versions_table).where(
                            content_versions_table.c.content_id == projected.id,
                            content_versions_table.c.version_no == 1,
                        )
                    )
                    .mappings()
                    .one()
                )
                next_version = dict(old_version)
                next_version.update(
                    {
                        "id": uuid4(),
                        "version_no": 2,
                        "title": "竞品舞台 更新版本",
                        "observed_at": datetime(2026, 9, 10, 12, 0, tzinfo=UTC),
                    }
                )
                session.execute(insert(content_versions_table).values(**next_version))
                session.execute(
                    update(contents_table)
                    .where(contents_table.c.id == projected.id)
                    .values(current_version=2, updated_at=beijing_now())
                )
                PostgresBrandVehicleRepository(session).replace_manual_brand_evidence(
                    content_id=projected.id,
                    content_version=2,
                    brand_ids=(competitor.id,),
                    unlock_existing=False,
                    actor_ref=principal.principal_id,
                )
        finally:
            session.close()

        current = content_service.get_content(projected.id)
        assert current.content_version == 2
        assert [brand.id for brand in current.brands] == [competitor.id]
        assert current.competition_scope == "competitor_only"

        export_registry = JobRegistry()
        register_data_export_job(
            export_registry,
            DataExportJobHandler(PostgresDataExportJobExecutor(runtime)),
        )
        export_worker = create_job_worker(
            runtime=runtime,
            registry=export_registry,
            worker_id="stage5-export",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert export_worker.run_once() is True
        downloaded = reporting.download_export(export.export_id)
        workbook = load_workbook(BytesIO(b"".join(download.chunks)), read_only=True, data_only=True)
        try:
            assert list(workbook["内容"].iter_rows(values_only=True)) == [
                ("命中关键词", "品牌", "品牌角色", "竞品范围", "车型"),
                (None, "爱玛 Stage5", "自有品牌", "仅自有品牌", "新车型 Stage5"),
            ]
        finally:
            workbook.close()

        session = runtime.database.new_session()
        try:
            with session.begin():
                frozen_page = PostgresDataExportRepository(session).load_page(
                    export.export_id,
                    after_ordinal=-1,
                    limit=10,
                )
        finally:
            session.close()
        assert frozen_page[0][1].content.brands == ("爱玛 Stage5",)
        assert frozen_page[0][1].content.competition_scope == "owned_only"
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts, vehicle_brands, "
                "vehicle_models RESTART IDENTITY CASCADE"
            )
        runtime.close()
