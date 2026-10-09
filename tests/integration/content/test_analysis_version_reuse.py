"""补采等价输入与正式导入共用的 Analysis 版本复用回归。"""

from io import BytesIO

import pytest
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.content_http import PostgresContentHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.http import ContentAnalysisSubmitRequest, ContentTargetSelection
from aima_ugc.modules.analysis.tables import (
    analysis_content_results_table,
    analysis_content_runs_table,
)
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.platform.config import load_settings
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from tests.integration.content.test_bidirectional_relevance_review import _seed_import, _xlsx
from tests.integration.content.test_stage12_analysis_runs import _analysis_registry
from tests.integration.stage3_brand_support import stage3_filter_brand_id


@pytest.fixture
def reuse_runtime(tmp_path):  # type: ignore[no-untyped-def]
    """独占任务测试数据库，模型只使用正式 Worker 中的 Fake。"""
    runtime = create_worker_runtime(
        settings=load_settings().model_copy(
            update={
                "data_dir": tmp_path / "data",
                "log_dir": tmp_path / "logs",
                "llm_base_url": "https://fake.example/v1",
                "llm_provider_name": "fake",
                "llm_model": "fake-content-labeler-v1",
            }
        )
    )
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE jobs, artifacts, keyword_packs, accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        _seed_import(client, runtime)
        _import_worker(runtime)
        service = PostgresContentHttpService(runtime, cursor_signing_secret=b"reuse-test" * 4)
        with runtime.database.engine.begin() as connection:
            content_id = connection.scalar(select(contents_table.c.id))
        service.create_analysis(
            ContentAnalysisSubmitRequest(
                targets=ContentTargetSelection(
                    scope="selected",
                    content_ids=(content_id,),
                )
            ),
            request_id="reuse-initial",
        )
        _analyze(runtime)
        yield runtime, service, content_id
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE jobs, artifacts, keyword_packs, accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def _import_worker(runtime):  # type: ignore[no-untyped-def]
    """通过实际导入 Job 完成一次入库。"""
    worker = create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id="reuse-import",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    assert worker.run_once()


def _analyze(runtime, *, sentiment="中性"):  # type: ignore[no-untyped-def]
    """Planner 与 Shard 均执行生产实现，不直接伪造成功 Result。"""
    worker = create_job_worker(
        runtime=runtime,
        registry=_analysis_registry(runtime, sentiment=sentiment),
        worker_id="reuse-analysis",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    assert worker.run_once()
    assert worker.run_once()


def _observe(runtime, *, text=None, title=None, author=None, published="2026-08-21 11:00:00"):  # type: ignore[no-untyped-def]
    """相同内容身份经正式 Excel 入口观察变化，覆盖相关入库编排。"""
    workbook = load_workbook(BytesIO(_xlsx()))
    sheet = workbook.active
    if text is not None:
        sheet["C2"] = text
    if title is not None:
        sheet["B2"] = title
    if author is not None:
        sheet["D2"] = author
    sheet["E2"] = published
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
    brand = stage3_filter_brand_id(runtime, alias="爱玛")
    response = client.post(
        "/api/v1/import-batches",
        files=[
            (
                "file",
                (
                    "reuse.xlsx",
                    output.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("brand_ids", (None, brand)),
        ],
    )
    assert response.status_code == 202, response.text
    _import_worker(runtime)


def test_non_ai_version_change_reuses_real_result_without_new_run(reuse_runtime):  # type: ignore[no-untyped-def]
    """发布时间变更形成新版本，却不改变实际模型输入或新增模型执行。"""
    runtime, service, content_id = reuse_runtime
    before = service.get_content(content_id)
    with runtime.database.engine.begin() as connection:
        result_id = connection.scalar(select(analysis_content_results_table.c.id))
        run_count = connection.scalar(select(func.count()).select_from(analysis_content_runs_table))
    assert before.analysis.status == "completed"
    _observe(runtime)
    from aima_ugc.adapters.persistence.postgres.analysis_reuse import (
        PostgresAnalysisReuseRepository,
    )
    from aima_ugc.modules.analysis.tables import analysis_content_version_reuses_table

    with runtime.database.new_session() as session:
        plans = PostgresAnalysisReuseRepository(session).plan_reuses(
            ((content_id, before.content_version + 1),)
        )
        relations = tuple(session.execute(select(analysis_content_version_reuses_table)).mappings())
        assert relations, plans
    after = service.get_content(content_id)
    assert after.content_version == before.content_version + 1
    assert after.analysis.status == "completed"
    assert after.analysis.analyzed_at == before.analysis.analyzed_at
    with runtime.database.engine.begin() as connection:
        assert (
            connection.scalar(select(func.count()).select_from(analysis_content_results_table)) == 1
        )
        assert connection.scalar(select(analysis_content_results_table.c.id)) == result_id
        assert (
            connection.scalar(select(func.count()).select_from(analysis_content_runs_table))
            == run_count
        )


@pytest.mark.parametrize("field", ("title", "text", "author"))
def test_one_actual_input_field_change_stays_stale(reuse_runtime, field):  # type: ignore[no-untyped-def]
    """任一实际模型字段变化都不能继承，且不自动投放模型 Job。"""
    runtime, service, content_id = reuse_runtime
    _observe(runtime, **{field: "爱玛 新的实际输入"})
    assert service.get_content(content_id).analysis.status == "stale"
    with runtime.database.engine.begin() as connection:
        assert connection.scalar(select(func.count()).select_from(analysis_content_runs_table)) == 1


@pytest.mark.parametrize(
    "schema",
    (
        "content-label-analysis.v1",
        "content-label-analysis.v2",
        "content-label-analysis.v3",
        "unknown.v99",
    ),
)
def test_hash_proof_is_independent_of_supported_output_version(reuse_runtime, schema):  # type: ignore[no-untyped-def]
    """已支持输出协议只要实际六字段 Hash 可证实便兼容；未知协议保守拒绝。"""
    from sqlalchemy import update

    runtime, service, content_id = reuse_runtime
    with runtime.database.engine.begin() as connection:
        connection.execute(update(analysis_content_results_table).values(schema_version=schema))
    _observe(runtime)
    expected = "stale" if schema == "unknown.v99" else "completed"
    assert service.get_content(content_id).analysis.status == expected


def test_nonadjacent_equal_input_reuses_original_success(reuse_runtime):  # type: ignore[no-untyped-def]
    """A→B→A 直接引用最初成功的真实结果，不创建递归副本。"""
    from aima_ugc.modules.analysis.tables import analysis_content_version_reuses_table

    runtime, service, content_id = reuse_runtime
    _observe(runtime, text="爱玛 B 输入")
    assert service.get_content(content_id).analysis.status == "stale"
    _observe(runtime, published="2026-08-22 11:00:00")
    after = service.get_content(content_id)
    assert after.content_version == 3
    assert after.analysis.status == "completed"
    with runtime.database.engine.begin() as connection:
        reuse = connection.execute(select(analysis_content_version_reuses_table)).mappings().one()
        assert reuse["target_content_version"] == 3
        assert reuse["source_content_version"] == 1


def _review(service, content_id, version, **changes):  # type: ignore[no-untyped-def]
    """使用正式人工复核入口验证锁定及审计写入。"""
    from aima_ugc.contracts.product import ContentAnalysisManualReviewRequest

    return service.review_analysis(
        content_id,
        ContentAnalysisManualReviewRequest(content_version=version, **changes),
        request_id="reuse-manual",
        actor_ref="local-tester",
    )


def test_inherited_manual_locks_require_unlock_and_tombstones_survive(reuse_runtime):  # type: ignore[no-untyped-def]
    """继承锁仍阻止隐式覆盖；当前版本解锁后下一版本不会复活旧锁。"""
    from aima_ugc.modules.content.http import ContentAnalysisRunConflict

    runtime, service, content_id = reuse_runtime
    _review(service, content_id, 1, sentiment="负面")
    _observe(runtime)
    after = service.get_content(content_id)
    assert after.analysis.sentiment == "负面"
    assert "sentiment" in after.analysis.manual_locked_dimensions
    with pytest.raises(ContentAnalysisRunConflict):
        _review(service, content_id, 2, sentiment="正面")
    _review(service, content_id, 2, unlock_dimensions=("sentiment",))
    assert service.get_content(content_id).analysis.sentiment == "中性"
    _observe(runtime, published="2026-08-22 11:00:00")
    after = service.get_content(content_id)
    assert after.analysis.sentiment == "中性"
    assert after.analysis.manual_locked_dimensions == ()


def test_inherited_relevance_can_return_to_ai_without_resurrection(reuse_runtime):  # type: ignore[no-untyped-def]
    """新版本 inherit_ai 创建自己的撤销事实，后续等价版本继承撤销。"""
    from aima_ugc.contracts.relevance_review import ContentRelevanceReviewRequest
    from aima_ugc.modules.analysis.relevance_review_tables import (
        analysis_content_relevance_reviews_table,
    )

    runtime, service, content_id = reuse_runtime

    def review(decision):  # type: ignore[no-untyped-def]
        return service.review_relevance(
            ContentRelevanceReviewRequest(
                content_ids=(content_id,),
                decision=decision,
            ),
            request_id="reuse-relevance",
        )

    assert review("irrelevant").changed_count == 1
    _observe(runtime)
    assert service.get_content(content_id).effective_relevance == "irrelevant"
    assert review("inherit_ai").changed_count == 1
    assert service.get_content(content_id).effective_relevance == "relevant"
    _observe(runtime, published="2026-08-22 11:00:00")
    assert service.get_content(content_id).effective_relevance == "relevant"
    with runtime.database.engine.begin() as connection:
        events = (
            connection.execute(
                select(analysis_content_relevance_reviews_table).order_by(
                    analysis_content_relevance_reviews_table.c.content_version
                )
            )
            .mappings()
            .all()
        )
        assert [(row["content_version"], row["decision"], row["review_no"]) for row in events] == [
            (1, "irrelevant", 1),
            (2, "inherit_ai", 1),
        ]


def test_explicit_reanalysis_executes_and_direct_result_keeps_inherited_lock(reuse_runtime):  # type: ignore[no-untyped-def]
    """主动重新打标正常调用 Fake 模型，由直接新结果优先且不丢人工锁。"""
    runtime, service, content_id = reuse_runtime
    _review(service, content_id, 1, sentiment="负面")
    _observe(runtime)
    service.create_analysis(
        ContentAnalysisSubmitRequest(
            targets=ContentTargetSelection(
                scope="selected",
                content_ids=(content_id,),
            )
        ),
        request_id="reuse-manual-reanalysis",
    )
    _analyze(runtime, sentiment="正面")
    assert service.get_content(content_id).analysis.sentiment == "负面"
    with runtime.database.engine.begin() as connection:
        assert (
            connection.scalar(select(func.count()).select_from(analysis_content_results_table)) == 2
        )
        direct = (
            connection.execute(
                select(analysis_content_results_table).where(
                    analysis_content_results_table.c.content_version == 2
                )
            )
            .mappings()
            .one()
        )
        assert direct["sentiment"] == "正面"
    _review(service, content_id, 2, unlock_dimensions=("sentiment",))
    assert service.get_content(content_id).analysis.sentiment == "正面"


def test_frozen_export_and_report_use_target_version_and_inherited_manual(reuse_runtime):  # type: ignore[no-untyped-def]
    """冻结旧版本在 Current 后续改变后仍能读取合法来源和当时的人工效果。"""
    from aima_ugc.adapters.persistence.postgres.report_runs import PostgresReportRepository
    from aima_ugc.adapters.persistence.postgres.reporting import PostgresDataExportRepository

    runtime, service, content_id = reuse_runtime
    _review(service, content_id, 1, sentiment="负面")
    _observe(runtime)
    _observe(runtime, text="爱玛 已改变的 Current", published="2026-08-22 11:00:00")
    assert service.get_content(content_id).analysis.status == "stale"
    with runtime.database.new_session() as session:
        exported = PostgresDataExportRepository(session)._analysis_by_content({content_id: 2})
        assert exported[content_id].sentiment == "负面"
        basis = PostgresReportRepository(session)._analysis_basis({content_id: 2})
        assert basis[content_id]["content_version"] == 1
        assert basis[content_id]["target_content_version"] == 2
        assert basis[content_id]["manual_override_version"] == 1


def _supplement(runtime, content_id):  # type: ignore[no-untyped-def]
    """真实 Fence/Attempt 下用稀疏详情补采同一已分析内容。"""
    from datetime import timedelta

    from aima_ugc.adapters.persistence.postgres.analysis import canonical_analysis_content_from_row
    from aima_ugc.adapters.persistence.postgres.analysis_reuse import (
        PostgresAnalysisReuseRepository,
    )
    from aima_ugc.adapters.persistence.postgres.brand_vehicle import PostgresBrandVehicleRepository
    from aima_ugc.adapters.persistence.postgres.collection_content import (
        PostgresFencedCollectionIngestionWriter,
    )
    from aima_ugc.modules.vehicles.brand_vehicle import BrandVehicleResolver
    from aima_ugc.platform.time import beijing_now

    from tests.integration.collection.test_collection_content_runtime import (
        _canonical,
        _create_live_source,
    )

    source = _create_live_source(
        runtime.database, source_value="analysis-reuse-supplement", enrichment=True
    )
    with runtime.database.new_session() as session, session.begin():
        current = session.scalar(
            select(contents_table.c.current_version).where(contents_table.c.id == content_id)
        )
        original = canonical_analysis_content_from_row(
            PostgresAnalysisReuseRepository(session)._version_inputs((content_id,))[
                (content_id, current)
            ]
        )
        catalog = PostgresBrandVehicleRepository(session).snapshot(brand_ids=None)
    observation = original.model_copy(
        update={
            "title": None,
            "text": None,
            "author": None,
            "published_at": original.published_at + timedelta(days=1),
            "observed_at": beijing_now(),
            "observed_fields": ["published_at"],
            "source": _canonical(source).source,
        }
    )
    resolution = BrandVehicleResolver().resolve(
        catalog, title=None, raw_text=None, transcript_text=None
    )
    return PostgresFencedCollectionIngestionWriter(runtime.database.new_session).ingest_content(
        canonical=observation,
        fence=source.fence,
        expected_content_id=content_id,
        brand_vehicle_snapshot=catalog,
        brand_vehicle_resolution=resolution,
    )


def test_supplement_reuse_exception_rolls_back_content_evidence_and_relation(
    reuse_runtime, monkeypatch
):  # type: ignore[no-untyped-def]
    """已有 Content/Evidence 后，复用关联写入后的异常让整个补采事务原子回滚。"""
    from aima_ugc.adapters.persistence.postgres.analysis_reuse import (
        PostgresAnalysisReuseRepository,
    )
    from aima_ugc.modules.analysis.tables import analysis_content_version_reuses_table
    from aima_ugc.modules.content.tables import content_versions_table
    from aima_ugc.modules.vehicles.tables import content_brand_evidence_table

    runtime, service, content_id = reuse_runtime
    before = service.get_content(content_id)
    with runtime.database.engine.begin() as connection:
        original_evidence = connection.execute(
            select(content_brand_evidence_table).where(
                content_brand_evidence_table.c.content_id == content_id
            )
        ).all()
    original = PostgresAnalysisReuseRepository.converge_reuses

    def fail_after_reuse(self, pairs):  # type: ignore[no-untyped-def]
        original(self, pairs)
        raise RuntimeError("reuse-rollback-probe")

    monkeypatch.setattr(PostgresAnalysisReuseRepository, "converge_reuses", fail_after_reuse)
    with pytest.raises(RuntimeError, match="reuse-rollback-probe"):
        _supplement(runtime, content_id)
    assert service.get_content(content_id) == before
    with runtime.database.engine.begin() as connection:
        assert (
            connection.scalar(
                select(func.count())
                .select_from(content_versions_table)
                .where(content_versions_table.c.content_id == content_id)
            )
            == 1
        )
        assert (
            connection.scalar(
                select(func.count()).select_from(analysis_content_version_reuses_table)
            )
            == 0
        )
        assert (
            connection.execute(
                select(content_brand_evidence_table).where(
                    content_brand_evidence_table.c.content_id == content_id
                )
            ).all()
            == original_evidence
        )


def _wait_content_lock(runtime):  # type: ignore[no-untyped-def]
    """观察 PostgreSQL 真实锁等待，避免把线程调度延迟当成并发证据。"""
    from time import monotonic, sleep

    from sqlalchemy import text

    deadline = monotonic() + 10
    while monotonic() < deadline:
        with runtime.database.engine.connect() as connection:
            if connection.scalar(
                text(
                    "SELECT count(*) FROM pg_stat_activity WHERE datname=current_database() "
                    "AND wait_event_type='Lock' AND query ILIKE '%contents%'"
                )
            ):
                return
        sleep(0.01)
    pytest.fail("未观察到 Content 行锁等待")


def test_ai_persistence_content_lock_serializes_supplement(reuse_runtime, monkeypatch):  # type: ignore[no-untyped-def]
    """AI 已锁住版本时补采必须等待提交，再引用刚成功的新结果。"""
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event

    from aima_ugc.adapters.persistence.postgres.analysis_batch import (
        PostgresAnalysisBatchRepository,
    )

    runtime, service, content_id = reuse_runtime
    service.create_analysis(
        ContentAnalysisSubmitRequest(
            targets=ContentTargetSelection(scope="selected", content_ids=(content_id,))
        ),
        request_id="concurrent-reanalysis",
    )
    worker = create_job_worker(
        runtime=runtime,
        registry=_analysis_registry(runtime, sentiment="正面"),
        worker_id="reuse-concurrent-ai",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    assert worker.run_once()
    locked, release = Event(), Event()
    original = PostgresAnalysisBatchRepository._current_versions

    def pause_after_lock(self, ids):  # type: ignore[no-untyped-def]
        versions = original(self, ids)
        locked.set()
        assert release.wait(10)
        return versions

    monkeypatch.setattr(PostgresAnalysisBatchRepository, "_current_versions", pause_after_lock)
    with ThreadPoolExecutor(max_workers=2) as pool:
        ai = pool.submit(worker.run_once)
        assert locked.wait(10)
        supplement = pool.submit(_supplement, runtime, content_id)
        try:
            _wait_content_lock(runtime)
            assert not supplement.done()
        finally:
            release.set()
        assert ai.result(timeout=15)
        assert supplement.result(timeout=15).version_no == 2
    after = service.get_content(content_id)
    assert after.analysis.status == "completed"
    assert after.analysis.sentiment == "正面"


def test_manual_content_lock_serializes_supplement_and_preserves_new_lock(
    reuse_runtime, monkeypatch
):  # type: ignore[no-untyped-def]
    """人工审核与补采实际并发，补采读到提交后的人工锁并保留其来源。"""
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event

    from aima_ugc.adapters.persistence.postgres.analysis_manual_reviews import (
        PostgresAnalysisManualReviewRepository,
    )

    runtime, service, content_id = reuse_runtime
    written, release = Event(), Event()
    original = PostgresAnalysisManualReviewRepository.review

    def pause_after_review(self, **kwargs):  # type: ignore[no-untyped-def]
        row = original(self, **kwargs)
        written.set()
        assert release.wait(10)
        return row

    monkeypatch.setattr(PostgresAnalysisManualReviewRepository, "review", pause_after_review)
    with ThreadPoolExecutor(max_workers=2) as pool:
        manual = pool.submit(_review, service, content_id, 1, sentiment="负面")
        assert written.wait(10)
        supplement = pool.submit(_supplement, runtime, content_id)
        try:
            _wait_content_lock(runtime)
            assert not supplement.done()
        finally:
            release.set()
        manual.result(timeout=15)
        assert supplement.result(timeout=15).version_no == 2
    assert service.get_content(content_id).analysis.sentiment == "负面"
    assert "sentiment" in service.get_content(content_id).analysis.manual_locked_dimensions


def test_all_current_consumers_share_reuse_and_manual_values(reuse_runtime):  # type: ignore[no-untyped-def]
    """投影、兼容读取、Count、Target、导出及工作台都采用同一继承人工结果。"""
    from datetime import UTC, datetime

    from aima_ugc.adapters.persistence.postgres.analysis_schemes import (
        PostgresAnalysisSchemeRepository,
    )
    from aima_ugc.adapters.persistence.postgres.content_queries import (
        PostgresContentQueryRepository,
    )
    from aima_ugc.adapters.persistence.postgres.reporting import PostgresDataExportRepository
    from aima_ugc.adapters.persistence.postgres.workbench import PostgresWorkbenchRepository
    from aima_ugc.contracts.http import ContentFilterSnapshot
    from aima_ugc.contracts.workbench import WorkbenchQuery
    from aima_ugc.modules.analysis.schemes import prompt_taxonomy_from_version
    from aima_ugc.modules.content.query import ContentReadQuery
    from aima_ugc.modules.content.read_model_tables import voice_plaza_projection_state_table
    from sqlalchemy import update

    runtime, service, content_id = reuse_runtime
    with runtime.database.new_session() as session:
        version = PostgresAnalysisSchemeRepository(session).get_active_version()
        taxonomy = prompt_taxonomy_from_version(version)
    primary = next(iter(taxonomy.labels))
    secondary = taxonomy.labels[primary][0]
    voice = taxonomy.voice_types[0]
    _review(
        service,
        content_id,
        1,
        sentiment="负面",
        voice_type=voice,
        labels=({"primary_label": primary, "secondary_label": secondary},),
    )
    _observe(runtime)
    with runtime.database.new_session() as session, session.begin():
        reader = PostgresContentQueryRepository(session, analysis_identity=None)
        filters = ContentFilterSnapshot(
            sentiment="负面",
            voice_type=voice,
            primary_label=primary,
            secondary_label=secondary,
            analysis_status="completed",
        )
        legacy, _ = reader._legacy_base_statement(filters)
        rows = tuple(session.execute(legacy).mappings())
        assert len(rows) == 1 and rows[0]["id"] == content_id
        session.execute(update(voice_plaza_projection_state_table).values(status="ready"))
        projected = reader.list_contents(ContentReadQuery(filters=filters, position=None, limit=20))
        assert len(projected) == 1
        assert projected[0].analysis.sentiment == "负面"
        assert projected[0].analysis.manual_locked_dimensions == (
            "voice_type",
            "sentiment",
            "labels",
        )
        assert reader.projection_count(filters) == 1
        assert reader.count_filtered_analysis_targets(filters) == 1
        assert (primary, secondary) in reader.list_filter_values().label_pairs
        exported = PostgresDataExportRepository(session)._analysis_by_content({content_id: 2})[
            content_id
        ]
        assert (exported.sentiment, exported.voice_type) == ("负面", voice)
        assert exported.label_pairs[0].secondary_label == secondary
        workbench = PostgresWorkbenchRepository(session)
        stream = workbench.stream_rows(
            active_scheme_version_id=version.id,
            query=WorkbenchQuery(sentiments=("负面",)),
            start_at=datetime(2026, 8, 1, tzinfo=UTC),
            end_at=datetime(2026, 9, 1, tzinfo=UTC),
        )
        assert len(stream) == 1 and stream[0]["effective_sentiment"] == "负面"
        snapshot = workbench.trend_snapshot(
            active_scheme_version_id=version.id,
            query=WorkbenchQuery(sentiments=("负面",)),
            previous_start_at=datetime(2026, 7, 1, tzinfo=UTC),
            current_start_at=datetime(2026, 8, 1, tzinfo=UTC),
            end_at=datetime(2026, 9, 1, tzinfo=UTC),
        )
        assert snapshot["current_summary"]["analyzed_count"] == 1


def test_source_snapshot_must_prove_persisted_hash_and_apply_is_idempotent(reuse_runtime):  # type: ignore[no-untyped-def]
    """仅让旧 Hash 恰好等于目标不足以复用；可信来源的 Dry-run 零写且修复幂等。"""
    from aima_ugc.adapters.persistence.postgres.analysis_reuse import (
        PostgresAnalysisReuseRepository,
    )
    from aima_ugc.modules.analysis.tables import analysis_content_version_reuses_table
    from sqlalchemy import delete, update

    runtime, service, content_id = reuse_runtime
    _observe(runtime)
    with runtime.database.new_session() as session, session.begin():
        session.execute(delete(analysis_content_version_reuses_table))
        owner = PostgresAnalysisReuseRepository(session)
        pairs = ((content_id, 2),)
        plan = owner.plan_reuses(pairs)[0]
        assert plan.reason == "equivalent"
        assert (
            session.scalar(select(func.count()).select_from(analysis_content_version_reuses_table))
            == 0
        )
        assert owner.converge_reuses(pairs)[0].status == "inserted"
        first = session.execute(select(analysis_content_version_reuses_table)).mappings().one()
        assert owner.converge_reuses(pairs)[0].status == "existing"
        assert (
            session.execute(select(analysis_content_version_reuses_table)).mappings().one() == first
        )
    _observe(runtime, text="爱玛 当前输入 B", published="2026-08-22 11:00:00")
    with runtime.database.new_session() as session, session.begin():
        owner = PostgresAnalysisReuseRepository(session)
        target_hash = owner.plan_reuses(((content_id, 3),))[0].input_hash
        session.execute(update(analysis_content_results_table).values(input_hash=target_hash))
        plan = owner.plan_reuses(((content_id, 3),))[0]
        assert plan.reason == "unknown_protocol"
        assert owner.converge_reuses(((content_id, 3),))[0].status == "rejected"
    assert service.get_content(content_id).analysis.status == "stale"


def test_late_frozen_ai_work_does_not_overwrite_reused_current(reuse_runtime):  # type: ignore[no-untyped-def]
    """补采先切版时，旧冻结 AI 请求仍保留 stale 审计且不能覆盖 Current。"""
    from aima_ugc.modules.analysis.tables import analysis_content_request_items_table

    runtime, service, content_id = reuse_runtime
    service.create_analysis(
        ContentAnalysisSubmitRequest(
            targets=ContentTargetSelection(scope="selected", content_ids=(content_id,))
        ),
        request_id="late-frozen-reanalysis",
    )
    _supplement(runtime, content_id)
    _analyze(runtime, sentiment="正面")
    assert service.get_content(content_id).analysis.sentiment == "中性"
    with runtime.database.engine.begin() as connection:
        assert (
            connection.scalar(select(func.count()).select_from(analysis_content_results_table)) == 1
        )
        assert (
            "stale"
            in connection.scalars(select(analysis_content_request_items_table.c.status)).all()
        )


def test_collection_irrelevant_filter_uses_reused_raw_ai_not_manual_overlay(reuse_runtime):  # type: ignore[no-untyped-def]
    """补采目标的既有原始 AI irrelevant 口径仍保留，同时支持正式复用来源。"""
    from aima_ugc.adapters.persistence.postgres.collection_targets import (
        PostgresCollectionTargetReader,
    )
    from aima_ugc.contracts.relevance_review import ContentRelevanceReviewRequest

    from tests.integration.content.test_manual_relevance_review import (
        _analysis_registry as irrelevant_registry,
    )
    from tests.integration.content.test_manual_relevance_review import (
        _irrelevant_response,
    )

    runtime, service, content_id = reuse_runtime
    service.create_analysis(
        ContentAnalysisSubmitRequest(
            targets=ContentTargetSelection(scope="selected", content_ids=(content_id,))
        ),
        request_id="irrelevant-reanalysis",
    )
    worker = create_job_worker(
        runtime=runtime,
        registry=irrelevant_registry(runtime, _irrelevant_response()),
        worker_id="reuse-irrelevant-ai",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    assert worker.run_once()
    assert worker.run_once()
    service.review_relevance(
        ContentRelevanceReviewRequest(content_ids=(content_id,), decision="relevant"),
        request_id="relevant-overlay",
    )
    _observe(runtime)
    assert service.get_content(content_id).effective_relevance == "relevant"
    with runtime.database.new_session() as session:
        assert PostgresCollectionTargetReader(session)._latest_current_irrelevant_content_ids(
            (content_id,)
        ) == {content_id}


def test_reuse_preflight_batches_targets_without_n_plus_one(reuse_runtime):  # type: ignore[no-untyped-def]
    """一百条显式目标共用固定数量的只读查询，不逐内容读取历史或调用模型。"""
    from aima_ugc.adapters.persistence.postgres.analysis_reuse import (
        PostgresAnalysisReuseRepository,
    )
    from sqlalchemy import event

    runtime, _service, original_id = reuse_runtime
    workbook = load_workbook(BytesIO(_xlsx()))
    sheet = workbook.active
    template = [cell.value for cell in sheet[2]]
    sheet.delete_rows(2)
    for ordinal in range(100):
        row = list(template)
        row[5] = f"https://www.xiaohongshu.com/explore/reuse-batch-{ordinal}"
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
    brand = stage3_filter_brand_id(runtime, alias="爱玛")
    response = client.post(
        "/api/v1/import-batches",
        files=[
            (
                "file",
                (
                    "bounded.xlsx",
                    output.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("brand_ids", (None, brand)),
        ],
    )
    assert response.status_code == 202, response.text
    _import_worker(runtime)
    with runtime.database.new_session() as session:
        pairs = tuple(
            (row[0], row[1])
            for row in session.execute(
                select(contents_table.c.id, contents_table.c.current_version).where(
                    contents_table.c.id != original_id
                )
            )
        )
        assert len(pairs) == 100
        statements = []

        def capture(_connection, _cursor, statement, _parameters, _context, _executemany):  # type: ignore[no-untyped-def]
            statements.append(statement)

        event.listen(runtime.database.engine, "before_cursor_execute", capture)
        try:
            plans = PostgresAnalysisReuseRepository(session).plan_reuses(pairs)
        finally:
            event.remove(runtime.database.engine, "before_cursor_execute", capture)
        assert len(plans) == 100
        assert {plan.reason for plan in plans} == {"no_success"}
        assert len(statements) <= 8
        assert all(statement.lstrip().startswith("SELECT") for statement in statements)
        assert session.scalar(select(func.count()).select_from(analysis_content_runs_table)) == 1


def test_legacy_non_model_url_format_does_not_add_an_equivalence_condition(reuse_runtime):  # type: ignore[no-untyped-def]
    """历史非模型 URL 格式不扩大六字段 Hash 协议。"""
    from aima_ugc.modules.content.tables import content_versions_table
    from sqlalchemy import update

    runtime, service, content_id = reuse_runtime
    with runtime.database.engine.begin() as connection:
        connection.execute(
            update(content_versions_table)
            .where(content_versions_table.c.content_id == content_id)
            .values(share_url="legacy-invalid-url")
        )
    _observe(runtime)
    assert service.get_content(content_id).analysis.status == "completed"
