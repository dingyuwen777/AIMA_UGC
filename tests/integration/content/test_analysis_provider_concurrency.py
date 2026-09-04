"""数据库 Provider 驱动的正式 Analysis 并发与批量持久化集成回归。"""

from __future__ import annotations

import json
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
from threading import Event, Lock, Thread
from time import monotonic
from uuid import UUID, uuid4

import pytest
from aima_ugc.adapters.llm import OpenAICompatibleContentLabelingLLM
from aima_ugc.adapters.llm.openai_compatible import OpenAICompatibleLLMError
from aima_ugc.adapters.persistence.postgres.analysis import PostgresAnalysisRepository
from aima_ugc.adapters.persistence.postgres.system import PostgresProviderConfigRepository
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.content_http import PostgresContentHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.modules.analysis import (
    ContentLabelingLLMRequest,
    ContentLabelingLLMResponse,
)
from aima_ugc.modules.analysis.tables import (
    analysis_content_request_items_table,
    analysis_content_requests_table,
    analysis_content_results_table,
    analysis_content_runs_table,
)
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.modules.system.models import ProviderConfig
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import LeaseLostError
from aima_ugc.platform.jobs.tables import jobs_table
from aima_ugc.platform.jobs.worker import JobExecutionContext
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select, text, update


class _ConcurrentFakeLLM:
    """请求阻塞到指定并发峰值后统一释放。"""

    provider_name = "fake-db"
    model_name = "fake-content-labeler-v1"

    def __init__(self, *, max_concurrency: int, fail_parallel_call: int | None = None) -> None:
        """创建线程安全 Fake；可让 第 N 个调用返回非重试 Transport 错误。"""

        self._max_concurrency = max_concurrency
        self._fail_parallel_call = fail_parallel_call
        self._release = Event()
        self._lock = Lock()
        self._calls = 0
        self._parallel_calls = 0
        self._active = 0
        self.peak_active = 0
        self.item_sizes: list[int] = []

    def close(self) -> None:
        """Fake 不持有 HTTP 连接，由正式执行器按 Adapter 生命周期关闭。"""

    def request_metrics(self) -> dict[str, int]:
        """Fake 没有真实 HTTP，不伪造物理请求统计。"""

        return {}

    def complete(self, request: ContentLabelingLLMRequest) -> ContentLabelingLLMResponse:
        """记录请求粒度和实际并发；首批也必须真实并发。"""

        with self._lock:
            self._calls += 1
            self.item_sizes.append(len(request.items))
            self._parallel_calls += 1
            parallel_call_no = self._parallel_calls
            self._active += 1
            self.peak_active = max(self.peak_active, self._active)
            if self.peak_active >= self._max_concurrency:
                self._release.set()

        assert self._release.wait(timeout=5)
        try:
            if parallel_call_no == self._fail_parallel_call:
                raise OpenAICompatibleLLMError(
                    "fake bad request",
                    error_code="http_400",
                    retryable=False,
                    status_code=400,
                )
            return ContentLabelingLLMResponse(raw_text=_valid_response())
        finally:
            with self._lock:
                self._active -= 1


def _valid_response() -> str:
    """返回当前测试 Scheme 已有的稳定 V3 合法结果。"""

    return (
        '{"items":[{"item_no":1,"relevance":"relevant",'
        '"voice_type":"真实用户发声","sentiment":"负面",'
        '"labels":[{"primary_label":"骑行性能","secondary_label":"舒适性"}]}]}'
    )


def _xlsx(row_count: int = 8) -> bytes:
    """生成足以观察并发的固定 Excel Content Fixture。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    for index in range(row_count):
        sheet.append(
            [
                "小红书",
                f"Provider Concurrency {index}",
                f"爱玛第 {index} 条并发测试内容",
                f"用户 {index}",
                f"2026-09-04 01:{index:02d}:00",
                f"https://www.xiaohongshu.com/explore/provider-concurrency-{index}",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _runtime(tmp_path: Path, *, max_concurrency: int = 4):  # type: ignore[no-untyped-def]
    """创建不依赖真实 LLM Secret 的隔离 PostgreSQL Runtime。"""

    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "analysis_run_max_in_flight_jobs": 2,
            "external_secret_dir": tmp_path / "secrets",
            "llm_base_url": "https://fake.example/v1",
            "llm_provider_name": "fake-db",
            "llm_model": "fake-content-labeler-v1",
            "llm_max_connections": max_concurrency,
            "llm_validation_retries": 0,
        }
    )
    for secret_ref in ("llm_api_key", "providers/tests/fake.key"):
        secret_path = settings.external_secret_root / secret_ref
        secret_path.parent.mkdir(parents=True, exist_ok=True)
        secret_path.write_text("integration-fake-key", encoding="utf-8")
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts, audit_events, "
            "provider_configs RESTART IDENTITY CASCADE"
        )

    def cleanup_provider_configs() -> None:
        """测试结束时清除 Provider，避免污染同一 CI 数据库的环境配置回归。"""

        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql("TRUNCATE TABLE provider_configs RESTART IDENTITY CASCADE")

    runtime.add_resource_closer(cleanup_provider_configs)
    return runtime


def _seed_provider(runtime, *, max_concurrency: int = 4) -> None:  # type: ignore[no-untyped-def]
    """直接通过 System Owner Repository 写入正式默认 LLM Provider 配置。"""

    session = runtime.database.new_session()
    try:
        with session.begin():
            PostgresProviderConfigRepository(session).create(
                ProviderConfig(
                    id=uuid4(),
                    provider="fake-db",
                    provider_kind="llm",
                    display_name="并发集成 Fake",
                    base_url="https://fake.example/v1",
                    model="fake-content-labeler-v1",
                    secret_ref="providers/tests/fake.key",
                    timeout_seconds=45,
                    max_retries=0,
                    max_concurrency=max_concurrency,
                    max_rps=None,
                    is_default=True,
                    enabled=True,
                )
            )
    finally:
        session.close()


def _client(runtime) -> TestClient:  # type: ignore[no-untyped-def]
    """使用正式 Import/Content Service 构造测试 HTTP Client。"""

    return TestClient(
        create_app(
            import_service=PostgresImportHttpService(runtime),
            content_service=PostgresContentHttpService(
                runtime,
                cursor_signing_secret=b"provider-concurrency-content-cursor-key",
            ),
        )
    )


def _seed_contents(client: TestClient, runtime, *, row_count: int = 8) -> tuple[UUID, ...]:  # type: ignore[no-untyped-def]
    """通过正式 Excel Import + Worker 链写入指定数量的 Content。"""

    pack = client.post("/api/v1/keyword-packs", json={"name": f"Concurrency {uuid4()}"})
    assert pack.status_code == 201
    keyword = client.post(
        f"/api/v1/keyword-packs/{pack.json()['id']}/keywords",
        json={"text": "爱玛", "priority": 10},
    )
    assert keyword.status_code == 201
    uploaded = client.post(
        "/api/v1/import-batches",
        files=[
            (
                "file",
                (
                    "analysis-provider-concurrency.xlsx",
                    _xlsx(row_count),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("keyword_pack_ids", (None, pack.json()["id"])),
        ],
    )
    assert uploaded.status_code == 202
    worker = create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id="analysis-provider-concurrency-import",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    assert worker.run_once() is True
    with runtime.database.engine.begin() as connection:
        return tuple(
            connection.execute(select(contents_table.c.id).order_by(contents_table.c.id))
            .scalars()
            .all()
        )


def _create_run(
    client: TestClient,
    content_ids: tuple[UUID, ...],
    *,
    key: str,
    shard_size: int = 80,
    shard_count: int = 1,
) -> str:
    """从公开 Preview/Create 入口创建 selected Analysis Run 并返回 Run ID。"""

    targets = {"scope": "selected", "content_ids": [str(item) for item in content_ids]}
    preview = client.post("/api/v1/analysis/content-runs/preview", json={"targets": targets})
    assert preview.status_code == 200
    assert preview.json()["shard_size"] == shard_size
    assert preview.json()["shard_count"] == shard_count
    created = client.post(
        "/api/v1/analysis/content-runs",
        json={
            "client_idempotency_key": key,
            "targets": targets,
            "expected_target_count": len(content_ids),
            "expected_configuration_hash": preview.json()["configuration_hash"],
            "run_intent": "manual_reanalysis",
        },
    )
    assert created.status_code == 202
    return str(created.json()["run_id"])


def _drain(runtime, llm: _ConcurrentFakeLLM, *, worker_id: str) -> int:  # type: ignore[no-untyped-def]
    """使用正式 Registry 和冻结配置执行，仅替换外部模型 Adapter。"""

    worker = create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id=worker_id,
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    count = 0
    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(
            "aima_ugc.bootstrap.analysis_concurrent_worker.OpenAICompatibleContentLabelingLLM",
            lambda **_kwargs: llm,
        )
        while worker.run_once():
            count += 1
            if count > 4:
                raise AssertionError("Provider 并发测试出现异常 Job 扩张")
    return count


@pytest.mark.parametrize("config_source", ["database", "environment"])
def test_provider_drives_concurrency_shard_and_batch_persistence(
    tmp_path: Path, config_source: str
) -> None:
    """两种配置来源都自动分片并由正式装配达到配置并发，八条结果完整落库。"""

    runtime = _runtime(tmp_path)
    try:
        if config_source == "database":
            _seed_provider(runtime, max_concurrency=4)
        client = _client(runtime)
        content_ids = _seed_contents(client, runtime)
        assert len(content_ids) == 8
        run_id = _create_run(client, content_ids, key="provider-concurrency-success")

        llm = _ConcurrentFakeLLM(max_concurrency=4)
        assert _drain(runtime, llm, worker_id="provider-concurrency-success") == 2
        assert llm.peak_active == 4
        assert llm.item_sizes == [1] * 8

        run = client.get(f"/api/v1/analysis/content-runs/{run_id}")
        assert run.status_code == 200
        assert run.json()["status"] == "succeeded"
        assert run.json()["shard_size"] == 80
        assert run.json()["stats"] == {
            "pending": 0,
            "succeeded": 8,
            "failed": 0,
            "stale": 0,
            "cancelled": 0,
        }
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(select(func.count()).select_from(analysis_content_results_table))
                == 8
            )
            snapshot = connection.execute(
                select(analysis_content_runs_table.c.runtime_config_snapshot).where(
                    analysis_content_runs_table.c.id == UUID(run_id)
                )
            ).scalar_one()
            assert snapshot["max_concurrency"] == 4
    finally:
        runtime.close()


def test_parallel_transport_error_only_fails_one_content(tmp_path: Path) -> None:
    """单条非重试 Transport 错误只终结对应 Content，不重跑整个 Shard。"""

    runtime = _runtime(tmp_path)
    try:
        _seed_provider(runtime, max_concurrency=4)
        client = _client(runtime)
        content_ids = _seed_contents(client, runtime)
        run_id = _create_run(client, content_ids, key="provider-concurrency-one-failure")

        llm = _ConcurrentFakeLLM(max_concurrency=4, fail_parallel_call=2)
        assert _drain(runtime, llm, worker_id="provider-concurrency-failure") == 2
        assert llm.peak_active == 4
        assert llm.item_sizes == [1] * 8

        run = client.get(f"/api/v1/analysis/content-runs/{run_id}")
        assert run.status_code == 200
        assert run.json()["status"] == "partial_failed"
        assert run.json()["stats"] == {
            "pending": 0,
            "succeeded": 7,
            "failed": 1,
            "stale": 0,
            "cancelled": 0,
        }
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(select(func.count()).select_from(analysis_content_results_table))
                == 7
            )
    finally:
        runtime.close()


def test_automatic_shards_refill_bounded_job_window(tmp_path: Path) -> None:
    """41 条内容自动拆成 3 个 Shard，超过双 Job 窗口后仍全部完成且无重复请求。"""

    runtime = _runtime(tmp_path, max_concurrency=1)
    try:
        client = _client(runtime)
        content_ids = _seed_contents(client, runtime, row_count=41)
        run_id = _create_run(
            client, content_ids, key="automatic-multiple-shards", shard_size=20, shard_count=3
        )
        planner_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="automatic-multiple-shards-plan",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert planner_worker.run_once() is True
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(
                    select(func.count()).where(
                        analysis_content_requests_table.c.run_id == UUID(run_id)
                    )
                )
                == 2
            )
        llm = _ConcurrentFakeLLM(max_concurrency=1)
        assert _drain(runtime, llm, worker_id="automatic-multiple-shards") == 3
        assert llm.item_sizes == [1] * 41
        run = client.get(f"/api/v1/analysis/content-runs/{run_id}")
        assert run.status_code == 200
        assert run.json()["status"] == "succeeded"
        assert run.json()["stats"]["succeeded"] == 41
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(select(func.count()).select_from(analysis_content_results_table))
                == 41
            )
    finally:
        runtime.close()


def test_concurrent_executor_preserves_existing_frozen_single_item_shards(tmp_path: Path) -> None:
    """已有 Run 的单条分片快照仍由新执行器完成，不按当前配置重划分历史目标。"""

    runtime = _runtime(tmp_path)
    try:
        client = _client(runtime)
        content_ids = _seed_contents(client, runtime, row_count=3)
        run_id = _create_run(client, content_ids, key="existing-frozen-single-item-shards")
        # 构造旧版本已创建但尚未执行的冻结分片事实，仅作用于隔离测试库。
        with runtime.database.engine.begin() as connection:
            connection.execute(
                update(analysis_content_runs_table)
                .where(analysis_content_runs_table.c.id == UUID(run_id))
                .values(shard_size=1, shard_count=3)
            )
        llm = _ConcurrentFakeLLM(max_concurrency=1)
        assert _drain(runtime, llm, worker_id="existing-frozen-shards") == 4
        run = client.get(f"/api/v1/analysis/content-runs/{run_id}")
        assert run.status_code == 200
        assert run.json()["shard_size"] == 1
        assert len(run.json()["shards"]) == 3
        assert run.json()["status"] == "succeeded"
        assert run.json()["stats"]["succeeded"] == 3
        assert llm.item_sizes == [1] * 3
    finally:
        runtime.close()


@contextmanager
def _controlled_http(*, expected: int, block_all: bool = False, status: int = 200):
    """真实本地 HTTP 服务：首请求或全部请求由测试显式释放，不调用付费模型。"""

    arrived = Event()
    release = Event()
    lock = Lock()
    bodies: list[dict] = []
    active = 0
    peak = 0

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *_args) -> None:
            pass

        def do_POST(self) -> None:
            nonlocal active, peak
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            with lock:
                bodies.append(body)
                ordinal = len(bodies)
                active += 1
                peak = max(peak, active)
                if ordinal == expected:
                    arrived.set()
            try:
                if ordinal == 1 or block_all:
                    assert release.wait(10)
                data = json.dumps(
                    {"choices": [{"message": {"content": _valid_response()}}]}
                ).encode()
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)
            finally:
                with lock:
                    active -= 1

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    adapters: list[OpenAICompatibleContentLabelingLLM] = []

    def factory(**kwargs):
        kwargs["base_url"] = f"http://127.0.0.1:{server.server_port}/v1"
        adapter = OpenAICompatibleContentLabelingLLM(**kwargs)
        adapters.append(adapter)
        return adapter

    try:
        with pytest.MonkeyPatch.context() as patch:
            patch.setattr(
                "aima_ugc.bootstrap.analysis_concurrent_worker.OpenAICompatibleContentLabelingLLM",
                factory,
            )
            yield arrived, release, bodies, adapters
    finally:
        release.set()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def _planned_run(tmp_path: Path, *, rows: int = 2):
    """完成真实 Import 和 Planner，返回等待执行的 Shard。"""

    runtime = _runtime(tmp_path, max_concurrency=2)
    client = _client(runtime)
    ids = _seed_contents(client, runtime, row_count=rows)
    run_id = _create_run(client, ids, key=str(uuid4()), shard_size=40)
    worker = create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id="streaming-test",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    assert worker.run_once()
    return runtime, client, run_id, worker


def test_small_run_persists_first_result_without_waiting_for_batch_timer(tmp_path: Path) -> None:
    """两条并发中一条仍等待模型时，另一条立即可读，不再等一秒批量计时器。"""

    runtime, client, run_id, worker = _planned_run(tmp_path)
    try:
        with _controlled_http(expected=2) as (arrived, release, _bodies, adapters):
            with ThreadPoolExecutor(max_workers=1) as pool:
                future = pool.submit(worker.run_once)
                try:
                    assert arrived.wait(4)
                    deadline = monotonic() + 0.75
                    while True:
                        state = client.get(f"/api/v1/analysis/content-runs/{run_id}").json()
                        if state["stats"]["succeeded"] == 1:
                            break
                        assert monotonic() < deadline, "少量已完成结果仍等待批量计时器"
                        Event().wait(0.02)
                    assert not future.done()
                    assert state["stats"]["pending"] == 1
                finally:
                    release.set()
                assert future.result(timeout=4)
            assert adapters[0].request_metrics()["http_requests"] == 2
        assert client.get(f"/api/v1/analysis/content-runs/{run_id}").json()["status"] == "succeeded"
    finally:
        runtime.close()


def test_real_http_streams_across_pages_and_flushes_before_slow_tail(tmp_path: Path) -> None:
    """一个慢请求不能阻止跨过多个 2C 页，其他结果在慢请求返回前可从 API 读取。"""

    runtime, client, run_id, worker = _planned_run(tmp_path, rows=12)
    try:
        with _controlled_http(expected=12) as (arrived, release, bodies, adapters):
            with ThreadPoolExecutor(max_workers=1) as pool:
                future = pool.submit(worker.run_once)
                try:
                    assert arrived.wait(4), "后续页被首个慢请求挡住"
                    deadline = monotonic() + 3
                    while True:
                        state = client.get(f"/api/v1/analysis/content-runs/{run_id}").json()
                        if state["stats"]["succeeded"] == 11:
                            break
                        assert monotonic() < deadline, "完成结果没有按时间落库"
                        Event().wait(0.05)
                    assert not future.done()
                    assert state["stats"]["pending"] == 1
                    assert all(body["model"] == "fake-content-labeler-v1" for body in bodies)
                    assert all(
                        body["response_format"] == {"type": "json_object"} for body in bodies
                    )
                    assert len({body["messages"][0]["content"] for body in bodies}) == 1
                finally:
                    release.set()
                assert future.result(timeout=4)
            assert adapters[0].request_metrics()["http_peak_active"] == 2
            assert adapters[0].request_metrics()["http_requests"] == 12
        state = client.get(f"/api/v1/analysis/content-runs/{run_id}").json()
        assert state["status"] == "succeeded"
        assert state["stats"]["succeeded"] == 12
    finally:
        runtime.close()


@pytest.mark.parametrize("stop_kind", ["cancel", "lease", "deadline"])
def test_stop_during_http_prevents_retries_and_stale_writes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    stop_kind: str,
) -> None:
    """真实 HTTP 等待期间发现取消/失权后，503 不重试、成功回包也不能越权写入。"""

    runtime, client, run_id, worker = _planned_run(tmp_path)
    observed = Event()
    original = JobExecutionContext.cancel_requested

    def observe(context):
        try:
            cancelled = original(context)
        except LeaseLostError:
            observed.set()
            raise
        if cancelled:
            observed.set()
        return cancelled

    monkeypatch.setattr(JobExecutionContext, "cancel_requested", observe)
    try:
        with _controlled_http(
            expected=2, block_all=True, status=200 if stop_kind == "lease" else 503
        ) as (
            arrived,
            release,
            bodies,
            _adapters,
        ):
            with ThreadPoolExecutor(max_workers=1) as pool:
                future = pool.submit(worker.run_once)
                try:
                    assert arrived.wait(3), "两条数据未从首批并发发送"
                    if stop_kind == "cancel":
                        response = client.post(f"/api/v1/analysis/content-runs/{run_id}/cancel")
                        assert response.status_code == 200
                    else:
                        with runtime.database.engine.begin() as connection:
                            column = (
                                "lease_expires_at"
                                if stop_kind == "lease"
                                else "attempt_deadline_at"
                            )
                            expiration = {column: text("clock_timestamp() - interval '1 second'")}
                            if stop_kind == "deadline":
                                expiration.update(
                                    {
                                        "attempt_started_at": text(
                                            "clock_timestamp() - interval '2 seconds'"
                                        ),
                                        "lease_expires_at": text(
                                            "clock_timestamp() - interval '1 second'"
                                        ),
                                    }
                                )
                            # 仅在隔离测试库模拟进程失权，不修改生产时钟或 Fence 逻辑。
                            connection.execute(
                                update(jobs_table)
                                .where(
                                    jobs_table.c.id.in_(
                                        select(analysis_content_requests_table.c.job_id).where(
                                            analysis_content_requests_table.c.run_id == UUID(run_id)
                                        )
                                    )
                                )
                                .values(expiration)
                            )
                    assert observed.wait(2), "等待 HTTP 时没有检查执行控制状态"
                finally:
                    release.set()
                if stop_kind == "cancel":
                    assert future.result(timeout=3)
                else:
                    with pytest.raises(LeaseLostError):
                        future.result(timeout=3)
            assert len(bodies) == 2
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(select(func.count()).select_from(analysis_content_results_table))
                == 0
            )
        if stop_kind == "cancel":
            state = client.get(f"/api/v1/analysis/content-runs/{run_id}").json()
            assert state["status"] == "cancelled"
            assert state["stats"]["cancelled"] == 2
        elif stop_kind == "lease":
            # 同一 Request 由新 Attempt 接管，只有新的执行者能够完成写入。
            llm = _ConcurrentFakeLLM(max_concurrency=2)
            assert _drain(runtime, llm, worker_id="recovered") == 1
            state = client.get(f"/api/v1/analysis/content-runs/{run_id}").json()
            assert state["stats"]["succeeded"] == 2
    finally:
        runtime.close()


def test_entire_stale_page_does_not_finish_remaining_items(tmp_path: Path) -> None:
    """前一页全部内容过期时仍扫描后页，不产生 pending 非零的成功 Job。"""

    runtime, client, run_id, _worker = _planned_run(tmp_path, rows=6)
    try:
        with runtime.database.engine.begin() as connection:
            first_page = select(analysis_content_request_items_table.c.content_id).where(
                analysis_content_request_items_table.c.request_id.in_(
                    select(analysis_content_requests_table.c.id).where(
                        analysis_content_requests_table.c.run_id == UUID(run_id),
                    ),
                ),
                analysis_content_request_items_table.c.ordinal < 4,
            )
            connection.execute(
                update(contents_table)
                .where(
                    contents_table.c.id.in_(first_page),
                )
                .values(current_version=2)
            )
        llm = _ConcurrentFakeLLM(max_concurrency=2)
        assert _drain(runtime, llm, worker_id="all-stale-page") == 1
        state = client.get(f"/api/v1/analysis/content-runs/{run_id}").json()
        assert state["stats"]["pending"] == 0
        assert state["stats"]["stale"] == 4
        assert state["stats"]["succeeded"] == 2
        assert len(llm.item_sizes) == 2
    finally:
        runtime.close()


def test_cancel_before_start_sends_no_http(tmp_path: Path) -> None:
    """排队期间取消，不创建模型 Client，也不发送任何请求。"""

    runtime, client, run_id, worker = _planned_run(tmp_path)
    try:
        with _controlled_http(expected=1) as (_arrived, _release, bodies, adapters):
            assert client.post(f"/api/v1/analysis/content-runs/{run_id}/cancel").status_code == 200
            worker.run_once()
            assert not bodies
            assert not adapters
        assert client.get(f"/api/v1/analysis/content-runs/{run_id}").json()["status"] == "cancelled"
    finally:
        runtime.close()


def test_systemic_http_failure_stops_refilling(tmp_path: Path) -> None:
    """整批认证错误最多发送当前初始窗口，不逐条重复触发 401。"""

    runtime, client, run_id, worker = _planned_run(tmp_path, rows=10)
    try:
        with _controlled_http(expected=2, block_all=True, status=401) as (
            arrived,
            release,
            bodies,
            _adapters,
        ):
            with ThreadPoolExecutor(max_workers=1) as pool:
                future = pool.submit(worker.run_once)
                try:
                    assert arrived.wait(3)
                finally:
                    release.set()
                assert future.result(timeout=3)
            assert len(bodies) == 2
        state = client.get(f"/api/v1/analysis/content-runs/{run_id}").json()
        assert state["status"] == "failed"
        assert state["stats"]["failed"] == 10
    finally:
        runtime.close()


def test_active_run_remains_visible_outside_recent_limit(tmp_path: Path) -> None:
    """最近历史的数量限制不允许隐藏仍在运行的旧任务。"""

    runtime, client, old_run_id, _worker = _planned_run(tmp_path)
    try:
        with runtime.database.engine.begin() as connection:
            ids = tuple(connection.execute(select(contents_table.c.id)).scalars())
        new_run_id = _create_run(client, ids, key=str(uuid4()), shard_size=40)
        session = runtime.database.new_session()
        try:
            with session.begin():
                rows = PostgresAnalysisRepository(session).list_runs(limit=1)
                assert [str(row["id"]) for row in rows] == [new_run_id, old_run_id]
        finally:
            session.close()
    finally:
        runtime.close()
