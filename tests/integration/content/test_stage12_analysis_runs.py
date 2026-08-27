"""Stage 12C Analysis Run 的 PostgreSQL/Job/HTTP Golden Path。"""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from aima_ugc.bootstrap.analysis_worker import (
    PostgresContentAnalysisJobExecutor,
    PostgresContentAnalysisPlanJobExecutor,
    create_analysis_job_terminal_callback,
)
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.content_http import PostgresContentHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.modules.analysis import (
    CONTENT_LABELING_PROMPT_PATH,
    ContentLabelingService,
    FakeContentLabelingLLM,
    FrozenPromptTaxonomyLoader,
    PromptTaxonomyLoader,
)
from aima_ugc.modules.analysis.content_analysis_job import (
    ContentAnalysisJobHandler,
    ContentAnalysisPlanJobHandler,
    register_content_analysis_job,
)
from aima_ugc.modules.analysis.tables import (
    analysis_content_requests_table,
    analysis_content_results_table,
    analysis_content_run_targets_table,
    analysis_content_runs_table,
)
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.modules.system.tables import audit_events_table
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import JobRegistry
from aima_ugc.platform.jobs.tables import jobs_table
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select, update


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    for index in range(3):
        sheet.append(
            [
                "小红书",
                f"爱玛 Stage12 Run {index}",
                f"第 {index} 条固定测试内容",
                f"用户 {index}",
                f"2026-08-26 10:0{index}:00",
                f"https://www.xiaohongshu.com/explore/stage12-run-{index}",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _seed_contents(client: TestClient) -> None:
    pack = client.post("/api/v1/keyword-packs", json={"name": f"Stage12 Run {uuid4()}"})
    assert pack.status_code == 201
    keyword = client.post(
        f"/api/v1/keyword-packs/{pack.json()['id']}/keywords",
        json={"text": "爱玛", "priority": 10},
    )
    assert keyword.status_code == 201
    uploaded = client.post(
        "/api/v1/import-batches",
        files=[
            (
                "file",
                (
                    "stage12-analysis.xlsx",
                    _xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("keyword_pack_ids", (None, pack.json()["id"])),
        ],
    )
    assert uploaded.status_code == 202


def _response(sentiment: str) -> str:
    return (
        '{"items":[{"item_no":1,"relevance":"relevant",'
        '"voice_type":"user_voice","sentiment":"'
        + sentiment
        + '","labels":[{"primary_label":"骑行性能",'
        '"secondary_label":"舒适性"}]}]}'
    )


def _analysis_registry(runtime, *, sentiment: str) -> JobRegistry:  # type: ignore[no-untyped-def]
    taxonomy = PromptTaxonomyLoader(CONTENT_LABELING_PROMPT_PATH).load()
    service = ContentLabelingService(
        prompt_loader=FrozenPromptTaxonomyLoader(taxonomy),
        llm=FakeContentLabelingLLM(responses=[_response(sentiment)] * 3),
    )
    registry = JobRegistry()
    callback = create_analysis_job_terminal_callback(runtime)
    register_content_analysis_job(
        registry,
        ContentAnalysisJobHandler(
            PostgresContentAnalysisJobExecutor(
                runtime,
                service_factory=lambda: (service, lambda: None),
            )
        ),
        terminal_callback=callback,
        planner_handler=ContentAnalysisPlanJobHandler(
            PostgresContentAnalysisPlanJobExecutor(runtime)
        ),
        planner_terminal_callback=callback,
    )
    return registry


def _drain_analysis(runtime, *, sentiment: str, worker_id: str) -> int:  # type: ignore[no-untyped-def]
    worker = create_job_worker(
        runtime=runtime,
        registry=_analysis_registry(runtime, sentiment=sentiment),
        worker_id=worker_id,
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    completed = 0
    while worker.run_once():
        completed += 1
        if completed > 10:
            raise AssertionError("Analysis Run 超出预期有界 Job 数")
    return completed


def test_analysis_runs_freeze_targets_bound_shards_and_keep_run_order_current(
    tmp_path: Path,
) -> None:
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "llm_base_url": "https://fake.example/v1",
            "llm_provider_name": "fake",
            "llm_model": "fake-content-labeler-v1",
            "analysis_run_shard_size": 1,
            "analysis_run_max_in_flight_jobs": 2,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts, audit_events "
            "RESTART IDENTITY CASCADE"
        )
    try:
        content_service = PostgresContentHttpService(
            runtime,
            cursor_signing_secret=b"stage12-analysis-cursor-key-32-bytes",
        )
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                content_service=content_service,
            )
        )
        _seed_contents(client)
        import_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-analysis-import",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert import_worker.run_once() is True
        with runtime.database.engine.begin() as connection:
            content_ids = tuple(
                connection.execute(select(contents_table.c.id).order_by(contents_table.c.id))
                .scalars()
                .all()
            )
        assert len(content_ids) == 3

        targets = {"scope": "selected", "content_ids": [str(item) for item in content_ids]}
        preview = client.post(
            "/api/v1/analysis/content-runs/preview",
            json={"targets": targets},
        )
        assert preview.status_code == 200
        assert preview.json()["target_count"] == 3
        assert preview.json()["shard_count"] == 3
        assert preview.json()["shard_size"] == 1
        assert preview.json()["cost_estimate_available"] is False

        create_body = {
            "client_idempotency_key": "stage12-analysis-run-1",
            "targets": targets,
            "expected_target_count": 3,
            "expected_configuration_hash": preview.json()["configuration_hash"],
            "run_intent": "manual_reanalysis",
        }
        created = client.post("/api/v1/analysis/content-runs", json=create_body)
        assert created.status_code == 202
        run_id = created.json()["run_id"]
        repeated = client.post("/api/v1/analysis/content-runs", json=create_body)
        assert repeated.status_code == 202
        assert repeated.json() == created.json()

        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(select(func.count()).select_from(analysis_content_runs_table))
                == 1
            )
            assert (
                connection.scalar(
                    select(func.count()).select_from(analysis_content_run_targets_table)
                )
                == 0
            )
            assert (
                connection.scalar(select(func.count()).select_from(analysis_content_requests_table))
                == 0
            )
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(jobs_table)
                    .where(jobs_table.c.job_type.like("analysis.%"))
                )
                == 1
            )
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(audit_events_table)
                    .where(
                        audit_events_table.c.event_type == "analysis_run_created",
                        audit_events_table.c.object_id == run_id,
                    )
                )
                == 1
            )

        assert _drain_analysis(runtime, sentiment="负面", worker_id="stage12-run-1") == 4
        run = client.get(f"/api/v1/analysis/content-runs/{run_id}")
        assert run.status_code == 200
        assert run.json()["status"] == "succeeded"
        assert run.json()["stats"] == {
            "pending": 0,
            "succeeded": 3,
            "failed": 0,
            "stale": 0,
            "cancelled": 0,
        }
        assert len(run.json()["shards"]) == 3
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(
                    select(func.count()).select_from(analysis_content_run_targets_table)
                )
                == 3
            )

        preview_2 = client.post(
            "/api/v1/analysis/content-runs/preview",
            json={"targets": targets},
        ).json()
        created_2 = client.post(
            "/api/v1/analysis/content-runs",
            json={
                **create_body,
                "client_idempotency_key": "stage12-analysis-run-2",
                "expected_configuration_hash": preview_2["configuration_hash"],
            },
        )
        assert created_2.status_code == 202
        assert _drain_analysis(runtime, sentiment="正面", worker_id="stage12-run-2") == 4
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(select(func.count()).select_from(analysis_content_results_table))
                == 6
            )
        current = content_service.get_content(content_ids[0])
        assert current.analysis.sentiment == "正面"
        assert str(current.analysis.latest_run_id) == created_2.json()["run_id"]
        assert current.analysis.latest_run_status == "succeeded"

        created_3 = client.post(
            "/api/v1/analysis/content-runs",
            json={
                **create_body,
                "client_idempotency_key": "stage12-analysis-run-3",
                "expected_configuration_hash": preview_2["configuration_hash"],
            },
        )
        assert created_3.status_code == 202
        planner_worker = create_job_worker(
            runtime=runtime,
            registry=_analysis_registry(runtime, sentiment="中性"),
            worker_id="stage12-run-3-planner",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert planner_worker.run_once() is True
        cancelled = client.post(
            f"/api/v1/analysis/content-runs/{created_3.json()['run_id']}/cancel"
        )
        assert cancelled.status_code == 200
        assert cancelled.json()["status"] == "cancelled"
        assert cancelled.json()["stats"]["cancelled"] == 3
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(audit_events_table)
                    .where(
                        audit_events_table.c.event_type == "analysis_run_cancel_requested",
                        audit_events_table.c.object_id == created_3.json()["run_id"],
                    )
                )
                == 1
            )
        after_cancel = content_service.get_content(content_ids[0])
        assert after_cancel.analysis.sentiment == "正面"
        assert str(after_cancel.analysis.latest_run_id) == created_3.json()["run_id"]
        assert after_cancel.analysis.latest_run_status == "cancelled"

        conflict = client.post(
            "/api/v1/analysis/content-runs",
            json={
                **create_body,
                "targets": {
                    "scope": "selected",
                    "content_ids": [str(content_ids[0])],
                },
                "expected_target_count": 1,
            },
        )
        assert conflict.status_code == 409
        assert conflict.json()["errors"][0]["code"] == "content_analysis_run_conflict"
    finally:
        runtime.close()


def test_analysis_planner_rolls_back_when_frozen_selection_count_changed(tmp_path: Path) -> None:
    """Run 的选择快照异常变化时 Planner 必须失败，不能留下部分冻结目标。"""

    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "llm_base_url": "https://fake.example/v1",
            "llm_provider_name": "fake",
            "llm_model": "fake-content-labeler-v1",
            "analysis_run_shard_size": 1,
            "analysis_run_max_in_flight_jobs": 2,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts, audit_events "
            "RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                content_service=PostgresContentHttpService(
                    runtime,
                    cursor_signing_secret=b"stage12-analysis-cursor-key-32-bytes",
                ),
            )
        )
        _seed_contents(client)
        import_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-analysis-changed-import",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert import_worker.run_once() is True
        with runtime.database.engine.begin() as connection:
            content_ids = tuple(connection.scalars(select(contents_table.c.id)).all())
        targets = {"scope": "selected", "content_ids": [str(item) for item in content_ids]}
        preview = client.post(
            "/api/v1/analysis/content-runs/preview",
            json={"targets": targets},
        ).json()
        created = client.post(
            "/api/v1/analysis/content-runs",
            json={
                "client_idempotency_key": "stage12-analysis-target-changed",
                "targets": targets,
                "expected_target_count": 3,
                "expected_configuration_hash": preview["configuration_hash"],
                "run_intent": "manual_reanalysis",
            },
        )
        assert created.status_code == 202
        with runtime.database.engine.begin() as connection:
            connection.execute(
                update(analysis_content_runs_table)
                .where(analysis_content_runs_table.c.id == UUID(created.json()["run_id"]))
                .values(
                    filter_snapshot={
                        "scope": "selected",
                        "content_ids": [str(item) for item in content_ids[1:]],
                    }
                )
            )

        worker = create_job_worker(
            runtime=runtime,
            registry=_analysis_registry(runtime, sentiment="负面"),
            worker_id="stage12-analysis-target-changed",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert worker.run_once() is True
        run = client.get(f"/api/v1/analysis/content-runs/{created.json()['run_id']}").json()
        assert run["status"] == "failed"
        assert run["error_code"] == "content_analysis_target_changed"
        assert run["stats"]["failed"] == 3
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(
                    select(func.count()).select_from(analysis_content_run_targets_table)
                )
                == 0
            )
            assert (
                connection.scalar(select(func.count()).select_from(analysis_content_requests_table))
                == 0
            )
    finally:
        runtime.close()


@pytest.mark.parametrize("legacy_backfill", (False, True))
def test_analysis_run_runtime_configuration_policy(
    tmp_path: Path,
    legacy_backfill: bool,
) -> None:
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "llm_base_url": "https://fake.example/v1",
            "llm_provider_name": "fake",
            "llm_model": "fake-content-labeler-v1",
            "analysis_run_shard_size": 1,
            "analysis_run_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts, audit_events "
            "RESTART IDENTITY CASCADE"
        )
    drifted_runtime = None
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                content_service=PostgresContentHttpService(
                    runtime,
                    cursor_signing_secret=b"stage12-analysis-cursor-key-32-bytes",
                ),
            )
        )
        _seed_contents(client)
        import_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-analysis-drift-import",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert import_worker.run_once() is True
        with runtime.database.engine.begin() as connection:
            content_id = connection.scalar(select(contents_table.c.id).limit(1))
        targets = {"scope": "selected", "content_ids": [str(content_id)]}
        preview = client.post(
            "/api/v1/analysis/content-runs/preview",
            json={"targets": targets},
        ).json()
        created = client.post(
            "/api/v1/analysis/content-runs",
            json={
                "client_idempotency_key": "stage12-analysis-drift",
                "targets": targets,
                "expected_target_count": 1,
                "expected_configuration_hash": preview["configuration_hash"],
                "run_intent": "manual_reanalysis",
            },
        )
        assert created.status_code == 202
        if legacy_backfill:
            run_id = UUID(created.json()["run_id"])
            with runtime.database.engine.begin() as connection:
                connection.execute(
                    update(analysis_content_runs_table)
                    .where(analysis_content_runs_table.c.id == run_id)
                    .values(
                        client_idempotency_key=f"legacy-request:{run_id}",
                        generation_config={},
                        generation_config_hash=hashlib.sha256(b"{}").hexdigest(),
                    )
                )

        drifted_runtime = create_worker_runtime(
            settings=settings.model_copy(update={"llm_model": "fake-content-labeler-v2"})
        )
        taxonomy = PromptTaxonomyLoader(CONTENT_LABELING_PROMPT_PATH).load()
        fake_llm = FakeContentLabelingLLM(
            responses=[_response("负面")],
            model_name="fake-content-labeler-v2",
        )
        service = ContentLabelingService(
            prompt_loader=FrozenPromptTaxonomyLoader(taxonomy),
            llm=fake_llm,
        )
        callback = create_analysis_job_terminal_callback(drifted_runtime)
        registry = JobRegistry()
        register_content_analysis_job(
            registry,
            ContentAnalysisJobHandler(
                PostgresContentAnalysisJobExecutor(
                    drifted_runtime,
                    service_factory=lambda: (service, lambda: None),
                )
            ),
            terminal_callback=callback,
            planner_handler=ContentAnalysisPlanJobHandler(
                PostgresContentAnalysisPlanJobExecutor(drifted_runtime)
            ),
            planner_terminal_callback=callback,
        )
        worker = create_job_worker(
            runtime=drifted_runtime,
            registry=registry,
            worker_id="stage12-analysis-drift-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert worker.run_once() is True
        assert worker.run_once() is True

        assert len(fake_llm.calls) == (1 if legacy_backfill else 0)
        run = client.get(f"/api/v1/analysis/content-runs/{created.json()['run_id']}").json()
        assert run["status"] == ("succeeded" if legacy_backfill else "failed")
        assert run["stats"]["succeeded"] == (1 if legacy_backfill else 0)
        assert run["stats"]["failed"] == (0 if legacy_backfill else 1)
        with runtime.database.engine.begin() as connection:
            assert connection.scalar(
                select(func.count()).select_from(analysis_content_results_table)
            ) == (1 if legacy_backfill else 0)
    finally:
        if drifted_runtime is not None:
            drifted_runtime.close()
        runtime.close()
