from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.modules.ingestion.tables import processing_import_batches_table
from aima_ugc.platform.config import load_settings
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(
        [
            "媒体名称（中文）",
            "标题",
            "内文",
            "作者",
            "出版日期",
            "原文链接",
        ]
    )
    sheet.append(
        [
            "小红书",
            "爱玛新品发布",
            "普通正文",
            "用户甲",
            "2026-08-20 10:00:00",
            "https://www.xiaohongshu.com/explore/multi-pack-1",
        ]
    )
    sheet.append(
        [
            "小红书",
            "试驾体验",
            "黑翼骑起来很稳",
            "用户乙",
            "2026-08-20 11:00:00",
            "https://www.xiaohongshu.com/explore/multi-pack-2",
        ]
    )
    sheet.append(
        [
            "小红书",
            "周末出游",
            "天气很好",
            "用户丙",
            "2026-08-20 12:00:00",
            "https://www.xiaohongshu.com/explore/multi-pack-3",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _create_pack(client: TestClient, *, keyword: str) -> str:
    response = client.post(
        "/api/v1/keyword-packs",
        json={"name": f"多词包导入 {keyword} {uuid4()}"},
    )
    assert response.status_code == 201
    pack_id = response.json()["id"]
    response = client.post(
        f"/api/v1/keyword-packs/{pack_id}/keywords",
        json={"text": keyword, "priority": 10},
    )
    assert response.status_code == 201
    return pack_id


def test_excel_import_uses_union_of_multiple_selected_keyword_packs(tmp_path) -> None:
    settings = load_settings().model_copy(
        update={"data_dir": tmp_path / "data", "log_dir": tmp_path / "logs"}
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        brand_pack_id = _create_pack(client, keyword="爱玛")
        model_pack_id = _create_pack(client, keyword="黑翼")

        created = client.post(
            "/api/v1/import-batches",
            files=[
                (
                    "file",
                    (
                        "multi-pack.xlsx",
                        _xlsx(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
                ("keyword_pack_ids", (None, brand_pack_id)),
                ("keyword_pack_ids", (None, model_pack_id)),
            ],
        )
        assert created.status_code == 202

        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="multi-keyword-pack-import-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert worker.run_once() is True
        assert worker.run_once() is False

        batch = client.get(f"/api/v1/import-batches/{created.json()['batch_id']}")
        assert batch.status_code == 200
        assert batch.json()["status"] == "succeeded"
        assert batch.json()["stats"] == {
            "rows_seen": 3,
            "rows_matched": 2,
            "rows_filtered_out": 1,
            "duplicates_removed": 0,
            "rows_ingested": 2,
            "rows_rejected": 0,
        }

        with runtime.database.engine.begin() as connection:
            external_ids = set(
                connection.scalars(select(contents_table.c.external_content_id)).all()
            )
            persisted_batch = (
                connection.execute(select(processing_import_batches_table)).mappings().one()
            )
        assert external_ids == {"multi-pack-1", "multi-pack-2"}
        selection = persisted_batch["stats"]["keyword_selection"]
        assert set(selection["effective_keywords"]) == {"爱玛", "黑翼"}
        assert len(selection["keyword_packs"]) == 2
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()
