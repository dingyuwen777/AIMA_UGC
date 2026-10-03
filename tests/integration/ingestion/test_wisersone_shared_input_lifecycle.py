"""公开 server_path 的新消费者与 Wise 原 XLSX 清理共享同一物理生命周期。"""

from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from pathlib import Path
from threading import Event
from typing import Any
from uuid import UUID, uuid4

import pytest
from aima_ugc.adapters.persistence.postgres.artifact_metadata import (
    PostgresArtifactMetadataRepository,
)
from aima_ugc.bootstrap.wisersone_cleanup import cleanup_wisersone_files
from aima_ugc.modules.ingestion.historical_directory import HistoricalDirectoryBrowser
from aima_ugc.modules.ingestion.historical_tables import (
    historical_import_campaign_items_table as items,
)
from aima_ugc.modules.ingestion.historical_tables import (
    historical_import_campaigns_table as campaigns,
)
from aima_ugc.modules.ingestion.wisersone_tables import wisersone_downloads_table as downloads
from aima_ugc.platform.jobs.tables import jobs_table as jobs
from aima_ugc.platform.storage.service import ArtifactService
from aima_ugc.platform.time import beijing_now
from openpyxl import load_workbook
from sqlalchemy import func, select

from .test_wisersone_review_boundaries import _canonical_state
from .test_wisersone_review_regressions import _settle
from .test_wisersone_workflow import _create
from .test_wisersone_workflow import workflow as workflow

_SERVER = "/api/v1/data-import-campaigns/server"
_KINDS = ("file", "parent", "multi", "root_recursive")


def _finished_native(workflow: Any) -> tuple[dict[str, Any], Path, Any]:
    """正式下载、入库后不再运行 Worker；通过 housekeeping 时间参数满足七天条件。"""
    runtime, _, client, brand, new_worker, tick = workflow
    created = _create(client, brand)
    result = _settle(client, new_worker(), tick, created["id"], {"succeeded"})
    published = runtime.settings.wisersone_input_dir / result["id"] / "wisersone_last24h.xlsx"
    assert published.is_file() and result["campaign_id"]
    cleanup_at = beijing_now() + timedelta(days=8)
    with runtime.database.new_session() as session:
        native = (
            session.execute(select(campaigns).where(campaigns.c.id == UUID(result["campaign_id"])))
            .mappings()
            .one()
        )
        assert native["status"] == "succeeded" and native["finished_at"] < cleanup_at - timedelta(
            days=7
        )
    return result, published, cleanup_at


def _body(workflow: Any, native: dict[str, Any], kind: str) -> dict[str, Any]:
    runtime = workflow[0]
    exact = f"wisersone/{native['id']}/wisersone_last24h.xlsx"
    (runtime.settings.historical_import_root / "manual-empty").mkdir(exist_ok=True)
    selected = {
        "file": [exact],
        "parent": [f"wisersone/{native['id']}"],
        "multi": ["manual-empty", exact],
        "root_recursive": [""],
    }[kind]
    return {
        "client_idempotency_key": f"shared-input:{uuid4()}",
        "relative_paths": selected,
        "recursive": kind == "root_recursive",
        "brand_ids": [workflow[3]],
        "ingestion_policy": "standard_observation",
    }


def _create_consumer(workflow: Any, body: dict[str, Any]) -> dict[str, Any]:
    response = workflow[2].post(_SERVER, json=body)
    assert response.status_code == 202, response.text
    consumer = response.json()
    with workflow[0].database.new_session() as session:
        row = (
            session.execute(
                select(campaigns).where(campaigns.c.id == UUID(consumer["campaign_id"]))
            )
            .mappings()
            .one()
        )
        assert row["source_kind"] == "server_path" and row["status"] == "discovering"
        assert row["profile_snapshot"]["relative_paths"] == body["relative_paths"]
        assert row["client_idempotency_key"] == body["client_idempotency_key"]
        assert (
            session.scalar(
                select(func.count()).select_from(items).where(items.c.campaign_id == row["id"])
            )
            == 0
        )
    return consumer


def _source_rows(runtime: Any, consumer: dict[str, Any]) -> list[Any]:
    with runtime.database.new_session() as session:
        return list(
            session.execute(
                select(items).where(
                    items.c.campaign_id == UUID(consumer["campaign_id"]),
                    items.c.item_kind == "source_file",
                )
            ).mappings()
        )


def _wait_campaign(workflow: Any, consumer: dict[str, Any], expected: str) -> dict[str, Any]:
    _, _, client, _, new_worker, tick = workflow
    worker = new_worker()
    for _ in range(20):
        response = client.get(f"/api/v1/data-import-campaigns/{consumer['campaign_id']}")
        assert response.status_code == 200
        row = response.json()
        if row["status"] == expected:
            return row
        assert row["status"] not in {"failed", "partial_failed", "cancelled"}, row
        tick(worker)
    pytest.fail(f"公开 Campaign 未收敛到 {expected}：{row}")


def _assert_independent_source(
    workflow: Any, consumer: dict[str, Any], source_bytes: bytes
) -> None:
    rows = _source_rows(workflow[0], consumer)
    assert len(rows) == 1 and rows[0]["artifact_id"] is not None
    with workflow[0].database.new_session() as session:
        artifact = PostgresArtifactMetadataRepository(session).get(rows[0]["artifact_id"])
        assert artifact is not None
        assert workflow[0].artifact_store.read(artifact.storage_key) == source_bytes
    assert sum(count for _, count in _canonical_state(workflow[0], consumer).values()) == 1


def _import_after_original_removed(workflow: Any, consumer: dict[str, Any], canonical: Any) -> None:
    response = workflow[2].post(f"/api/v1/data-import-campaigns/{consumer['campaign_id']}/start")
    assert response.status_code == 200, response.text
    _wait_campaign(workflow, consumer, "succeeded")
    assert _canonical_state(workflow[0], consumer) == canonical


def _complete_then_expire(
    workflow: Any,
    consumer: dict[str, Any],
    canonical: Any,
    published: Path,
    cleanup_at: Any,
) -> None:
    """不强制独立 Source 后立即释放保留期；所有消费者终态满七天后验证合法清理。"""
    _import_after_original_removed(workflow, consumer, canonical)
    assert cleanup_wisersone_files(workflow[0], now=cleanup_at) == 1
    assert not published.exists()
    assert _canonical_state(workflow[0], consumer) == canonical


def _assert_atomic_rejection(workflow: Any, body: dict[str, Any], expected: int) -> None:
    response = workflow[2].post(_SERVER, json=body)
    assert response.status_code == expected, response.text
    request_id = response.headers["x-request-id"]
    with workflow[0].database.new_session() as session:
        assert (
            session.scalar(
                select(func.count())
                .select_from(campaigns)
                .where(
                    campaigns.c.client_idempotency_key == body["client_idempotency_key"],
                )
            )
            == 0
        )
        assert (
            session.scalar(
                select(func.count()).select_from(jobs).where(jobs.c.request_id == request_id)
            )
            == 0
        )


@pytest.mark.parametrize("kind", _KINDS)
@pytest.mark.parametrize("phase", ["discovering", "snapshotting"])
def test_new_public_consumer_protects_original_until_source_is_frozen(
    workflow: Any,
    kind: str,
    phase: str,
) -> None:
    runtime, _, client, _, new_worker, tick = workflow
    native, published, cleanup_at = _finished_native(workflow)
    original_bytes = published.read_bytes()
    consumer = _create_consumer(workflow, _body(workflow, native, kind))
    # HTTP 仅保存 discovering；Worker 明确停在发现前或 Source 尚未绑定的边界。
    if phase == "snapshotting":
        tick(new_worker())
        discovered = client.get(f"/api/v1/data-import-campaigns/{consumer['campaign_id']}")
        assert discovered.json()["status"] == phase, discovered.text
        rows = _source_rows(runtime, consumer)
        assert len(rows) == 1 and rows[0]["artifact_id"] is None
    assert (
        client.get(f"/api/v1/data-import-campaigns/{consumer['campaign_id']}").json()["status"]
        == phase
    )
    assert cleanup_wisersone_files(runtime, now=cleanup_at) == 0
    assert published.read_bytes() == original_bytes
    _wait_campaign(workflow, consumer, "ready")
    _assert_independent_source(workflow, consumer, original_bytes)
    canonical = _canonical_state(runtime, consumer)
    _complete_then_expire(workflow, consumer, canonical, published, cleanup_at)


@pytest.mark.parametrize("kind", _KINDS)
def test_claimed_cleanup_atomically_rejects_new_public_consumer(
    workflow: Any,
    monkeypatch: pytest.MonkeyPatch,
    kind: str,
) -> None:
    runtime = workflow[0]
    native, published, cleanup_at = _finished_native(workflow)
    body = _body(workflow, native, kind)
    entered, release = Event(), Event()
    original = Path.unlink

    def blocked_unlink(path: Path, *args: Any, **kwargs: Any) -> None:
        if path == published:
            entered.set()
            assert release.wait(15)
        original(path, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", blocked_unlink)
    with ThreadPoolExecutor(max_workers=1) as pool:
        deleting = pool.submit(cleanup_wisersone_files, runtime, now=cleanup_at)
        try:
            assert entered.wait(15) and published.exists()
            with runtime.database.new_session() as session:
                assert (
                    session.scalar(
                        select(downloads.c.file_delete_pending_at).where(
                            downloads.c.id == UUID(native["id"])
                        )
                    )
                    is not None
                )
            _assert_atomic_rejection(workflow, body, 409)
        finally:
            release.set()
        assert deleting.result(timeout=15) == 1
    assert not published.exists()


@pytest.mark.parametrize("boundary", ["source_copy", "source_recheck"])
def test_source_copy_and_recheck_keep_original_protected_then_import_independently(
    workflow: Any,
    monkeypatch: pytest.MonkeyPatch,
    boundary: str,
) -> None:
    runtime, _, _, _, new_worker, tick = workflow
    native, published, cleanup_at = _finished_native(workflow)
    source_bytes = published.read_bytes()
    consumer = _create_consumer(workflow, _body(workflow, native, "file"))
    tick(new_worker())
    assert _source_rows(runtime, consumer)[0]["artifact_id"] is None
    entered, release = Event(), Event()
    original_store = ArtifactService.store_stream
    original_describe = HistoricalDirectoryBrowser.describe
    descriptions = 0

    def copying(self: Any, **options: Any) -> Any:
        if (
            options["kind"] == "historical-import.source"
            and Path(options["source"].name) == published
        ):
            entered.set()
            assert release.wait(15)
        return original_store(self, **options)

    def rechecking(self: Any, relative_path: str) -> Any:
        nonlocal descriptions
        if relative_path == f"wisersone/{native['id']}/wisersone_last24h.xlsx":
            descriptions += 1
            if descriptions == 2:
                entered.set()
                assert release.wait(15)
        return original_describe(self, relative_path)

    if boundary == "source_copy":
        monkeypatch.setattr(ArtifactService, "store_stream", copying)
    else:
        monkeypatch.setattr(HistoricalDirectoryBrowser, "describe", rechecking)
    with ThreadPoolExecutor(max_workers=1) as pool:
        snapshot = pool.submit(tick, new_worker())
        try:
            assert entered.wait(15)
            assert _source_rows(runtime, consumer)[0]["artifact_id"] is None
            cleanup_error: OSError | None = None
            removed: int | None = None
            try:
                removed = cleanup_wisersone_files(runtime, now=cleanup_at)
            except OSError as exc:
                # Windows 可阻止删除打开的源文件；这不能替代数据库中的生命周期保护。
                cleanup_error = exc
            with runtime.database.new_session() as session:
                pending = session.scalar(
                    select(downloads.c.file_delete_pending_at).where(
                        downloads.c.id == UUID(native["id"])
                    )
                )
            assert pending is None, f"活动 Source 复制/复核期间错误认领清理：{cleanup_error}"
            assert removed == 0, cleanup_error
            assert published.read_bytes() == source_bytes
        finally:
            release.set()
        snapshot.result(timeout=15)
    monkeypatch.setattr(ArtifactService, "store_stream", original_store)
    monkeypatch.setattr(HistoricalDirectoryBrowser, "describe", original_describe)
    _wait_campaign(workflow, consumer, "ready")
    _assert_independent_source(workflow, consumer, source_bytes)
    canonical = _canonical_state(runtime, consumer)
    _complete_then_expire(workflow, consumer, canonical, published, cleanup_at)


def test_staging_is_rejected_without_campaign_or_job(workflow: Any) -> None:
    runtime = workflow[0]
    native, published, _ = _finished_native(workflow)
    stage = runtime.settings.wisersone_input_dir / ".staging" / native["id"] / published.name
    stage.parent.mkdir(parents=True, exist_ok=True)
    stage.write_bytes(published.read_bytes())
    body = _body(workflow, native, "file")
    body["relative_paths"] = [f"wisersone/.staging/{native['id']}/{published.name}"]
    _assert_atomic_rejection(workflow, body, 400)
    listed = workflow[2].get(
        "/api/v1/data-import-sources/server/directories", params={"relative_path": "wisersone"}
    )
    assert listed.status_code == 200
    assert all(".staging" not in item["relative_path"] for item in listed.json()["items"])


@pytest.mark.parametrize("namespace", ["same_named_directory", "manual"])
def test_user_input_does_not_protect_or_redirect_unrelated_managed_file(
    workflow: Any,
    namespace: str,
) -> None:
    runtime = workflow[0]
    native, published, cleanup_at = _finished_native(workflow)
    relative = (
        f"wisersone/{native['id']}/{published.name}"
        if namespace == "same_named_directory"
        else f"manual/{published.name}"
    )
    user_file = runtime.settings.historical_import_root / relative
    user_file.parent.mkdir(parents=True, exist_ok=True)
    user_file.write_bytes(published.read_bytes())
    book = load_workbook(user_file)
    try:
        book["文章"].cell(2, 2).value = f"爱玛人工输入 {namespace}"
        book["文章"].cell(2, 6).value = f"https://www.xiaohongshu.com/explore/manual-{namespace}"
        book.save(user_file)
    finally:
        book.close()
    user_bytes = user_file.read_bytes()
    assert user_bytes != published.read_bytes()
    body = _body(workflow, native, "file")
    body["relative_paths"] = [relative]
    consumer = _create_consumer(workflow, body)
    assert cleanup_wisersone_files(runtime, now=cleanup_at) == 1
    assert not published.exists() and user_file.read_bytes() == user_bytes
    _wait_campaign(workflow, consumer, "ready")
    _assert_independent_source(workflow, consumer, user_bytes)
    _import_after_original_removed(workflow, consumer, _canonical_state(runtime, consumer))
