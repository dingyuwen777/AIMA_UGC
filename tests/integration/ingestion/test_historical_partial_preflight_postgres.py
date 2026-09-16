"""历史批量导入 Source File 预检容错的真实 PostgreSQL/API/Worker 回归。"""

from __future__ import annotations

from collections.abc import Iterator
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

import pytest
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.brand_vehicle_http import PostgresBrandVehicleHttpService
from aima_ugc.bootstrap.historical_import_http import PostgresHistoricalImportHttpService
from aima_ugc.bootstrap.runtime import PlatformRuntime
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.brand_vehicle import BrandCreateRequest
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.modules.identity import Principal
from aima_ugc.platform.config import load_settings
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select

_HEADERS = [
    "媒体名称（中文）",
    "标题",
    "内文",
    "作者",
    "出版日期",
    "原文链接",
]


def _workbook_bytes(rows: list[list[object]]) -> bytes:
    """生成符合历史导入 Profile 的测试 XLSX。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(_HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _good_xlsx(*, suffix: str) -> bytes:
    """生成一条能命中爱玛 Brand Filter 的合法历史数据。"""

    return _workbook_bytes(
        [
            [
                "小红书",
                f"爱玛历史导入 {suffix}",
                "正常数据应继续完成导入",
                "测试账号",
                "2025-01-02 10:00:00",
                f"https://www.xiaohongshu.com/explore/{suffix}",
            ]
        ]
    )


def _empty_xlsx() -> bytes:
    """生成只有合法表头、没有数据行的 XLSX。"""

    return _workbook_bytes([])


def _bad_xlsx() -> bytes:
    """生成可打开但不符合当前 Excel Profile 的 XLSX。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["错误列"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _principal() -> Principal:
    """构造集成测试使用的管理员身份。"""

    return Principal(
        principal_id="historical-partial-preflight-admin",
        display_name="历史导入预检管理员",
        role="administrator",
        source="development",
    )


def _create_brand(runtime: PlatformRuntime) -> str:
    """创建能够命中合法测试行的爱玛品牌目录。"""

    brand = PostgresBrandVehicleHttpService(runtime).create_brand(
        BrandCreateRequest(
            code=f"AIMA-PARTIAL-PREFLIGHT-{uuid4()}",
            display_name="爱玛",
            role="owned",
            aliases=("爱玛",),
        ),
        principal=_principal(),
        request_id="historical-partial-preflight-brand",
    )
    return str(brand.id)


def _drain(worker: Any, *, maximum: int = 30) -> int:
    """在测试预算内同步执行当前持久 Job 队列。"""

    executed = 0
    for _ in range(maximum):
        if not worker.run_once():
            break
        executed += 1
    return executed


def _reset_database(runtime: PlatformRuntime) -> None:
    """清理当前集成测试拥有的数据库事实，保持用例彼此隔离。"""

    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )


@pytest.fixture
def historical_environment(tmp_path: Path) -> Iterator[tuple[PlatformRuntime, TestClient, Any, str, Path]]:
    """建立真实 PostgreSQL/API/Worker 历史导入测试环境。"""

    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 4,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    _reset_database(runtime)
    client = TestClient(
        create_app(historical_import_service=PostgresHistoricalImportHttpService(runtime))
    )
    worker = create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id="historical-partial-preflight-worker",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    brand_id = _create_brand(runtime)
    try:
        yield runtime, client, worker, brand_id, historical_root
    finally:
        client.close()
        _reset_database(runtime)
        runtime.close()


def _create_campaign(client: TestClient, *, brand_id: str, relative_paths: list[str]) -> str:
    """创建使用当前历史补空语义的服务器目录 Campaign。"""

    response = client.post(
        "/api/v1/historical-import-campaigns",
        json={
            "client_idempotency_key": f"historical-partial-preflight-{uuid4()}",
            "relative_paths": relative_paths,
            "recursive": False,
            "brand_ids": [brand_id],
        },
    )
    assert response.status_code == 202
    return str(response.json()["campaign_id"])


def _campaign(client: TestClient, campaign_id: str) -> dict[str, Any]:
    """读取 Campaign 当前公开状态。"""

    response = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}")
    assert response.status_code == 200
    return response.json()


def _items(client: TestClient, campaign_id: str) -> list[dict[str, Any]]:
    """读取 Campaign 当前有限明细。"""

    response = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}/items")
    assert response.status_code == 200
    return response.json()["items"]


def _content_count(runtime: PlatformRuntime) -> int:
    """读取已经进入 Content Owner 的业务内容数量。"""

    with runtime.database.engine.begin() as connection:
        return int(connection.scalar(select(func.count()).select_from(contents_table)) or 0)


def test_good_and_empty_sources_import_good_data(
    historical_environment: tuple[PlatformRuntime, TestClient, Any, str, Path],
) -> None:
    """空文件应成功跳过，不能阻止同批正常文件导入。"""

    runtime, client, worker, brand_id, root = historical_environment
    (root / "good.xlsx").write_bytes(_good_xlsx(suffix="good-empty"))
    (root / "empty.xlsx").write_bytes(_empty_xlsx())
    campaign_id = _create_campaign(
        client,
        brand_id=brand_id,
        relative_paths=["good.xlsx", "empty.xlsx"],
    )

    assert _drain(worker) == 3
    ready = _campaign(client, campaign_id)
    assert ready["status"] == "ready"
    assert ready["can_start"] is True
    assert ready["ready_item_count"] == 2
    assert ready["total_rows"] == 1
    empty_source = next(item for item in _items(client, campaign_id) if item["relative_path"] == "empty.xlsx")
    assert empty_source["item_kind"] == "source_file"
    assert empty_source["status"] == "succeeded"
    assert empty_source["row_count"] == 0
    assert empty_source["error_code"] == "historical_source_empty"
    assert empty_source["stats"]["warning_code"] == "historical_source_empty"

    started = client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start")
    assert started.status_code == 200
    assert _drain(worker) == 1
    completed = _campaign(client, campaign_id)
    assert completed["status"] == "succeeded"
    assert _content_count(runtime) == 1


def test_good_and_bad_sources_finish_partial_failed(
    historical_environment: tuple[PlatformRuntime, TestClient, Any, str, Path],
) -> None:
    """少量坏文件应被隔离，正常文件仍可写入并最终标记部分失败。"""

    runtime, client, worker, brand_id, root = historical_environment
    (root / "good.xlsx").write_bytes(_good_xlsx(suffix="good-bad"))
    (root / "bad.xlsx").write_bytes(_bad_xlsx())
    campaign_id = _create_campaign(
        client,
        brand_id=brand_id,
        relative_paths=["good.xlsx", "bad.xlsx"],
    )

    assert _drain(worker) == 3
    ready = _campaign(client, campaign_id)
    assert ready["status"] == "ready"
    assert ready["can_start"] is True
    assert ready["ready_item_count"] == 2
    bad_source = next(item for item in _items(client, campaign_id) if item["relative_path"] == "bad.xlsx")
    assert bad_source["status"] == "failed"
    assert bad_source["error_code"] == "historical_snapshot_invalid"

    started = client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start")
    assert started.status_code == 200
    assert _drain(worker) == 1
    completed = _campaign(client, campaign_id)
    assert completed["status"] == "partial_failed"
    assert completed["failed_chunk_count"] == 0
    assert _content_count(runtime) == 1


def test_all_empty_sources_finish_successfully_without_writes(
    historical_environment: tuple[PlatformRuntime, TestClient, Any, str, Path],
) -> None:
    """全部为空文件时应无写入成功结束，而不是制造失败任务。"""

    runtime, client, worker, brand_id, root = historical_environment
    (root / "empty-a.xlsx").write_bytes(_empty_xlsx())
    (root / "empty-b.xlsx").write_bytes(_empty_xlsx())
    campaign_id = _create_campaign(
        client,
        brand_id=brand_id,
        relative_paths=["empty-a.xlsx", "empty-b.xlsx"],
    )

    assert _drain(worker) == 3
    completed = _campaign(client, campaign_id)
    assert completed["status"] == "succeeded"
    assert completed["can_start"] is False
    assert completed["ready_item_count"] == 2
    assert completed["total_rows"] == 0
    assert _content_count(runtime) == 0
    sources = [item for item in _items(client, campaign_id) if item["item_kind"] == "source_file"]
    assert {item["status"] for item in sources} == {"succeeded"}
    assert {item["error_code"] for item in sources} == {"historical_source_empty"}


def test_all_bad_sources_remain_failed(
    historical_environment: tuple[PlatformRuntime, TestClient, Any, str, Path],
) -> None:
    """全部文件都真正无效时仍应失败关闭且不能开始导入。"""

    runtime, client, worker, brand_id, root = historical_environment
    (root / "bad-a.xlsx").write_bytes(_bad_xlsx())
    (root / "bad-b.xlsx").write_bytes(_bad_xlsx())
    campaign_id = _create_campaign(
        client,
        brand_id=brand_id,
        relative_paths=["bad-a.xlsx", "bad-b.xlsx"],
    )

    assert _drain(worker) == 3
    failed = _campaign(client, campaign_id)
    assert failed["status"] == "failed"
    assert failed["can_start"] is False
    assert failed["ready_item_count"] == 2
    assert failed["total_rows"] == 0
    assert _content_count(runtime) == 0
    sources = [item for item in _items(client, campaign_id) if item["item_kind"] == "source_file"]
    assert {item["status"] for item in sources} == {"failed"}
    assert {item["error_code"] for item in sources} == {"historical_snapshot_invalid"}
