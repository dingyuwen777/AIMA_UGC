"""Data Import Campaign 撤销的 PostgreSQL 来源安全集成测试。"""

from __future__ import annotations

import time
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from aima_ugc.adapters.persistence.postgres.content_queries import (
    PostgresContentQueryRepository,
)
from aima_ugc.adapters.persistence.postgres.import_revocation_lifecycle import (
    PostgresImportRevocationLifecycleRepository,
)
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.brand_vehicle_http import PostgresBrandVehicleHttpService
from aima_ugc.bootstrap.canonical_replay_http import PostgresCanonicalReplayHttpService
from aima_ugc.bootstrap.collection_http import PostgresCollectionHttpService
from aima_ugc.bootstrap.historical_import_http import PostgresHistoricalImportHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.import_revocation_http import PostgresImportRevocationHttpService
from aima_ugc.bootstrap.import_revocation_worker import PostgresImportRevocationJobExecutor
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.brand_vehicle import BrandAliasCreateRequest
from aima_ugc.contracts.http import CollectionRuntimeListQuery, ContentFilterSnapshot
from aima_ugc.contracts.lifecycle import DataImportRevokeRequest
from aima_ugc.modules.collection.tables import (
    provider_request_attempts_table,
    provider_requests_table,
)
from aima_ugc.modules.content.contribution_tables import content_source_contributions_table
from aima_ugc.modules.content.query import ContentReadQuery
from aima_ugc.modules.content.read_model_tables import (
    voice_plaza_content_projection_table,
    voice_plaza_filter_catalog_entries_table,
)
from aima_ugc.modules.content.tables import content_versions_table, contents_table
from aima_ugc.modules.identity import Principal
from aima_ugc.modules.ingestion.reversal_shard_tables import reversal_shards_table
from aima_ugc.modules.ingestion.revocation_tables import (
    historical_import_campaign_revocations_table,
    historical_import_revocation_content_versions_table,
    historical_import_revocation_requests_table,
)
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs.tables import jobs_table
from aima_ugc.platform.storage.tables import artifacts_table
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import and_, delete, event, func, insert, select, update
from sqlalchemy.exc import DBAPIError, SQLAlchemyError

from tests.integration.stage3_brand_support import stage3_filter_brand_id


def _xlsx(rows: tuple[tuple[str, str, str | None], ...]) -> bytes:
    """生成最小 AIMA Excel；每行使用独立小红书 URL 作为稳定身份。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(
        [
            "媒体名称（中文）",
            "标题",
            "内文",
            "作者",
            "出版日期",
            "原文链接",
        ]
    )
    for title, content_key, text in rows:
        sheet.append(
            [
                "小红书",
                title,
                text,
                "官方账号",
                "2025-01-02 10:00:00",
                f"https://www.xiaohongshu.com/explore/{content_key}",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _drain(worker, *, maximum: int = 20) -> int:
    """同步执行当前测试已经入队的有限 Job，避免依赖后台调度时间。"""

    executed = 0
    for _ in range(maximum):
        if not worker.run_once():
            break
        executed += 1
    return executed


def _runtime(tmp_path: Path, historical_root: Path, *, chunk_rows: int = 100):
    """建立隔离 Runtime，并使用单 Chunk 便于精确断言生命周期结果。"""

    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": chunk_rows,
            "historical_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, accounts "
            "RESTART IDENTITY CASCADE"
        )
    return runtime


def _client(runtime) -> TestClient:
    """复用正式导入 Service 建立 API Client；撤销执行另走真实 Application Service。"""

    return TestClient(
        create_app(
            historical_import_service=PostgresHistoricalImportHttpService(runtime),
            import_service=PostgresImportHttpService(runtime),
            canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
        )
    )


def _cleanup(runtime) -> None:
    """清空当前 Integration 数据并释放 Runtime。"""

    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, accounts "
            "RESTART IDENTITY CASCADE"
        )
    runtime.close()


def test_large_import_revocation_uses_durable_content_shards_and_exact_accounting(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """父 Job 终态失败后由新 Job 接管已提交分片，计数仍严格相等。"""

    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    row_count = 10_000
    (historical_root / "large.xlsx").write_bytes(
        _xlsx(
            tuple(
                (f"星曜分片撤销 {index}", f"formal-import-reversal-{uuid4()}-{index}", "正文")
                for index in range(row_count)
            )
        )
    )
    runtime = _runtime(tmp_path, historical_root, chunk_rows=2000)
    try:
        client = _client(runtime)
        brand_id = stage3_filter_brand_id(runtime, alias="暂不命中的测试词")
        setup_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="large-import-reversal-setup",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"formal-import-reversal-{uuid4()}",
                "relative_paths": ["large.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
                "ingestion_policy": "standard_observation",
            },
        )
        assert created.status_code == 202
        campaign_id = UUID(created.json()["campaign_id"])
        assert _drain(setup_worker, maximum=20) >= 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )
        assert _drain(setup_worker, maximum=30) >= 1
        PostgresBrandVehicleHttpService(runtime).add_alias(
            UUID(brand_id),
            BrandAliasCreateRequest(text="星曜"),
            principal=Principal(
                principal_id="formal-reversal",
                display_name="分片撤销测试",
                role="administrator",
                source="development",
            ),
            request_id="formal-reversal-alias",
        )
        replay = client.post(
            "/api/v1/canonical-replays/all",
            json={"idempotency_key": f"formal-reversal-replay-{uuid4()}"},
        )
        assert replay.status_code == 202
        assert _drain(setup_worker, maximum=20) >= 1
        service = PostgresImportRevocationHttpService(
            runtime.database.new_session, runtime.artifact_store
        )
        assert service.preview(campaign_id).impact.affected_content_count == row_count
        queued = service.revoke(
            campaign_id,
            DataImportRevokeRequest(reason="正式分片撤销回归"),
            actor_ref="integration-admin",
            request_id="formal-reversal-request",
        )
        assert queued.job_id is not None
        first_job_id = queued.job_id
        original_process = PostgresImportRevocationJobExecutor.process_shard
        fault_injected = False

        def fail_parent_after_commit(self, shard_id, *, fence, context):  # type: ignore[no-untyped-def]
            nonlocal fault_injected
            processed = original_process(self, shard_id, fence=fence, context=context)
            if fence.job_id == first_job_id and not fault_injected:
                fault_injected = True
                raise ValueError("模拟父 Job 首片提交后终态失败")
            return processed

        monkeypatch.setattr(
            PostgresImportRevocationJobExecutor, "process_shard", fail_parent_after_commit
        )
        failed_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="large-import-reversal-failed-parent",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert failed_worker.run_once()
        assert fault_injected
        with runtime.database.engine.connect() as connection:
            assert (
                connection.scalar(
                    select(jobs_table.c.status).where(jobs_table.c.id == first_job_id)
                )
                == "failed"
            )
        parent_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="large-import-reversal-parent",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        child_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="large-import-reversal-child",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        child_worker.run_once()  # 已取消子 Job 可能在请求取消时直接结清。
        retried = service.revoke(
            campaign_id,
            DataImportRevokeRequest(reason="从分片断点继续撤销"),
            actor_ref="integration-admin",
            request_id="formal-reversal-retry",
        )
        assert retried.job_id is not None and retried.job_id != first_job_id
        with ThreadPoolExecutor(max_workers=1) as pool:
            parent_result = pool.submit(parent_worker.run_once)
            deadline = time.monotonic() + 120
            while time.monotonic() < deadline:
                with runtime.database.engine.connect() as connection:
                    status = connection.scalar(
                        select(jobs_table.c.status).where(jobs_table.c.id == retried.job_id)
                    )
                if status == "running":
                    break
                time.sleep(0.01)
            else:
                pytest.fail("导入撤销父 Job 未进入运行状态")
            while not parent_result.done() and time.monotonic() < deadline:
                if not child_worker.run_once():
                    time.sleep(0.02)
            assert parent_result.result(timeout=1) is True
        with runtime.database.engine.connect() as connection:
            units = (
                connection.execute(
                    select(reversal_shards_table).where(
                        reversal_shards_table.c.import_campaign_id == campaign_id
                    )
                )
                .mappings()
                .all()
            )
            actual = connection.scalar(
                select(func.count())
                .select_from(historical_import_revocation_content_versions_table)
                .where(
                    historical_import_revocation_content_versions_table.c.campaign_id == campaign_id
                )
            )
            request = (
                connection.execute(
                    select(historical_import_revocation_requests_table).where(
                        historical_import_revocation_requests_table.c.campaign_id == campaign_id
                    )
                )
                .mappings()
                .one()
            )
        assert len(units) == 4
        assert all(unit["status"] == "succeeded" for unit in units)
        assert sum(unit["processed_content_count"] for unit in units) == row_count
        assert actual == row_count
        assert request["status"] == "succeeded"
        assert request["recomputed_content_count"] == row_count
    finally:
        _cleanup(runtime)


def test_revocation_preview_counts_replay_contributions_after_filtered_import(
    tmp_path: Path,
) -> None:
    """原导入行全过滤时，后续重筛的同源贡献仍应进入撤销预览与进度。"""

    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "filtered.xlsx").write_bytes(
        _xlsx(
            (
                ("星曜重筛第一条", f"replay-revocation-{uuid4()}", "星曜正文"),
                ("星曜重筛第二条", f"replay-revocation-{uuid4()}", "星曜正文"),
            )
        )
    )
    runtime = _runtime(tmp_path, historical_root)
    try:
        client = _client(runtime)
        brand_id = stage3_filter_brand_id(runtime, alias="暂不命中的测试词")
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="replay-revocation-integration-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"filtered-replay-revocation-{uuid4()}",
                "relative_paths": ["filtered.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
                "ingestion_policy": "standard_observation",
            },
        )
        assert created.status_code == 202
        campaign_id = UUID(created.json()["campaign_id"])
        assert _drain(worker) == 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )
        assert _drain(worker) >= 1
        imported = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert imported["status"] == "succeeded"
        assert imported["stats"]["filtered"] == 2
        assert imported["stats"]["created"] == 0

        PostgresBrandVehicleHttpService(runtime).add_alias(
            UUID(brand_id),
            BrandAliasCreateRequest(text="星曜"),
            principal=Principal(
                principal_id="replay-revocation-integration",
                display_name="重筛撤销测试",
                role="administrator",
                source="development",
            ),
            request_id="replay-revocation-alias",
        )
        replay = client.post(
            "/api/v1/canonical-replays/all",
            json={"idempotency_key": f"filtered-replay-{uuid4()}"},
        )
        assert replay.status_code == 202
        assert _drain(worker) >= 1

        service = PostgresImportRevocationHttpService(
            runtime.database.new_session, runtime.artifact_store
        )
        preview = service.preview(campaign_id)
        assert preview.eligible is True
        assert preview.impact.affected_content_count == 2
        assert preview.impact.hidden_content_count == 2
        assert preview.impact.retained_shared_content_count == 0
        assert preview.impact.unreversible_content_count == 0

        queued = service.revoke(
            campaign_id,
            DataImportRevokeRequest(reason="全过滤后重筛贡献撤销回归"),
            actor_ref="integration-admin",
            request_id="replay-revocation-request",
        )
        assert queued.status == "queued"
        assert _drain(worker) == 1
        completed = service.preview(campaign_id)
        assert completed.status == "succeeded"
        assert completed.recomputed_content_count == 2
        assert completed.impact.affected_content_count == 2
        runtime_row = (
            PostgresCollectionHttpService(runtime, cursor_signing_secret=b"r" * 32)
            .list_runtime_runs(CollectionRuntimeListQuery(record_types=("data_import_campaign",)))
            .items[0]
        )
        assert runtime_row.import_stats is not None
        assert runtime_row.import_stats.rows_seen == 2
        assert runtime_row.import_stats.rows_filtered_out == 2
        assert runtime_row.revocation_recomputed_content_count == 2
    finally:
        _cleanup(runtime)


def test_revocation_batches_common_contributions_without_per_content_sql(
    tmp_path: Path,
) -> None:
    """批量历史导入撤销保持每条 Version/来源追溯，同时限制数据库往返。"""

    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "batch.xlsx").write_bytes(
        _xlsx(
            tuple(
                (f"爱玛批量撤销 {index}", f"revocation-batch-{uuid4()}-{index}", "正文")
                for index in range(101)
            )
        )
    )
    runtime = _runtime(tmp_path, historical_root)
    try:
        client = _client(runtime)
        brand_id = stage3_filter_brand_id(runtime)
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="revocation-batch-integration-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"revocation-batch-{uuid4()}",
                "relative_paths": ["batch.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
                "ingestion_policy": "standard_observation",
            },
        )
        assert created.status_code == 202
        campaign_id = UUID(created.json()["campaign_id"])
        assert _drain(worker) == 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )
        assert _drain(worker) == 2
        service = PostgresImportRevocationHttpService(
            runtime.database.new_session, runtime.artifact_store
        )
        assert service.preview(campaign_id).impact.affected_content_count == 101
        with runtime.database.new_session() as session:
            lifecycle = PostgresImportRevocationLifecycleRepository(session)
            first, checkpoint = lifecycle.next_campaign_contribution_batch(
                campaign_id, after_content_id=None, content_limit=1
            )
            second, _ = lifecycle.next_campaign_contribution_batch(
                campaign_id, after_content_id=checkpoint, content_limit=1
            )
            assert len(first) == len(second) == 1
            assert first[0]["content_id"] != second[0]["content_id"]
            duplicate = dict(first[0])
            duplicate["id"] = uuid4()
            duplicate["source_item_key"] = uuid4().hex + uuid4().hex
            savepoint = session.begin_nested()
            try:
                session.execute(insert(content_source_contributions_table).values(**duplicate))
                complete_group, _ = lifecycle.next_campaign_contribution_batch(
                    campaign_id, after_content_id=None, content_limit=1
                )
                assert len(complete_group) == 2
                assert {row["content_id"] for row in complete_group} == {first[0]["content_id"]}
            finally:
                savepoint.rollback()
        statement_count = 0
        per_content_updates = 0
        per_content_inserts = 0

        def count_sql(
            connection: object,
            cursor: object,
            statement: str,
            parameters: object,
            context: object,
            executemany: bool,
        ) -> None:
            nonlocal statement_count, per_content_updates, per_content_inserts
            del connection, cursor, parameters, context
            statement_count += 1
            if statement.lstrip().startswith("UPDATE contents") and executemany:
                per_content_updates += 1
            if (
                statement.lstrip().startswith(
                    (
                        "INSERT INTO content_versions",
                        "INSERT INTO historical_import_revocation_content_versions",
                    )
                )
                and executemany
            ):
                per_content_inserts += 1

        event.listen(runtime.database.engine, "before_cursor_execute", count_sql)
        try:
            revoked = service.revoke(
                campaign_id,
                DataImportRevokeRequest(reason="批量撤销回归"),
                actor_ref="integration-admin",
                request_id="revocation-batch-request",
            )
            assert revoked.status == "queued"
            assert revoked.job_id is not None
            assert _drain(worker) == 1
        finally:
            event.remove(runtime.database.engine, "before_cursor_execute", count_sql)
        assert revoked.impact.hidden_content_count == 101
        assert statement_count < 100
        assert per_content_updates == 0
        assert per_content_inserts == 0
        with runtime.database.engine.connect() as connection:
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(historical_import_revocation_content_versions_table)
                    .where(
                        historical_import_revocation_content_versions_table.c.campaign_id
                        == campaign_id
                    )
                )
                == 101
            )
            assert (
                connection.scalar(select(func.count()).select_from(content_versions_table)) == 202
            )
    finally:
        _cleanup(runtime)


def test_revocation_hides_exclusive_content_and_retains_shared_content(tmp_path: Path) -> None:
    """撤销 Campaign 会重组 Current；共享 Content 保留，独占 Content 退出后续业务目标。"""

    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    shared_key = f"revocation-shared-{uuid4()}"
    exclusive_key = f"revocation-exclusive-{uuid4()}"
    (historical_root / "campaign.xlsx").write_bytes(
        _xlsx(
            (
                ("爱玛共享来源", shared_key, "共享正文"),
                ("爱玛独占来源", exclusive_key, "独占正文"),
            )
        )
    )
    runtime = _runtime(tmp_path, historical_root)
    try:
        client = _client(runtime)
        brand_id = stage3_filter_brand_id(runtime)
        baseline = client.post(
            "/api/v1/import-batches",
            files=[
                (
                    "file",
                    (
                        "baseline.xlsx",
                        _xlsx((("爱玛共享来源", shared_key, "共享正文"),)),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
                ("brand_ids", (None, brand_id)),
            ],
        )
        assert baseline.status_code == 202

        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="revocation-integration-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker) == 1

        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"revocation-{uuid4()}",
                "relative_paths": ["campaign.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        assert created.status_code == 202
        campaign_id = UUID(created.json()["campaign_id"])
        assert _drain(worker) == 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )
        assert _drain(worker) == 1
        completed = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert completed["status"] == "succeeded"

        before_session = runtime.database.new_session()
        try:
            with before_session.begin():
                before_repository = PostgresContentQueryRepository(
                    before_session,
                    analysis_identity=None,
                )
                before = before_repository.list_contents(
                    ContentReadQuery(
                        filters=ContentFilterSnapshot(),
                        position=None,
                        limit=20,
                    )
                )
                assert len(before) == 2
                by_title = {record.title: record for record in before}
                shared_id = by_title["爱玛共享来源"].id
                exclusive_id = by_title["爱玛独占来源"].id
                assert before_repository.count_all_analysis_targets() == 2
                assert dict(
                    before_session.execute(
                        select(
                            voice_plaza_content_projection_table.c.content_id,
                            voice_plaza_content_projection_table.c.is_visible,
                        ).where(
                            voice_plaza_content_projection_table.c.content_id.in_(
                                (shared_id, exclusive_id)
                            )
                        )
                    ).all()
                ) == {shared_id: True, exclusive_id: True}
        finally:
            before_session.close()

        service = PostgresImportRevocationHttpService(
            runtime.database.new_session,
            runtime.artifact_store,
        )
        preview = service.preview(campaign_id)
        assert preview.eligible is True
        assert preview.already_revoked is False
        assert preview.ineligible_reason is None
        assert preview.impact.affected_content_count == 2
        assert preview.impact.hidden_content_count == 1
        assert preview.impact.retained_shared_content_count == 1
        assert preview.impact.unreversible_content_count == 0

        first = service.revoke(
            campaign_id,
            DataImportRevokeRequest(reason="集成测试撤销"),
            actor_ref="integration-admin",
            request_id="revocation-request-1",
        )
        second = service.revoke(
            campaign_id,
            DataImportRevokeRequest(reason="重复点击"),
            actor_ref="integration-admin",
            request_id="revocation-request-2",
        )
        assert first.status == "queued"
        assert second.status == "queued"
        assert first.job_id == second.job_id
        assert _drain(worker) == 1
        completed_revocation = service.revoke(
            campaign_id,
            DataImportRevokeRequest(reason="完成后重复点击"),
            actor_ref="integration-admin",
            request_id="revocation-request-3",
        )
        assert completed_revocation.status == "succeeded"
        assert completed_revocation.already_revoked is True
        assert completed_revocation.revoked_at is not None
        assert second.impact == first.impact

        after_session = runtime.database.new_session()
        try:
            with after_session.begin():
                after_repository = PostgresContentQueryRepository(
                    after_session,
                    analysis_identity=None,
                )
                after = after_repository.list_contents(
                    ContentReadQuery(
                        filters=ContentFilterSnapshot(),
                        position=None,
                        limit=20,
                    )
                )
                assert [record.id for record in after] == [shared_id]
                assert after_repository.get_content(shared_id) is not None
                assert after_repository.get_content(exclusive_id) is None
                assert after_repository.count_all_analysis_targets() == 1
                assert dict(
                    after_session.execute(
                        select(
                            voice_plaza_content_projection_table.c.content_id,
                            voice_plaza_content_projection_table.c.is_visible,
                        ).where(
                            voice_plaza_content_projection_table.c.content_id.in_(
                                (shared_id, exclusive_id)
                            )
                        )
                    ).all()
                ) == {shared_id: True, exclusive_id: False}
                assert (
                    after_session.scalar(
                        select(func.count())
                        .select_from(voice_plaza_filter_catalog_entries_table)
                        .where(
                            voice_plaza_filter_catalog_entries_table.c.content_id == exclusive_id
                        )
                    )
                    == 0
                )
                selected = after_repository.freeze_targets(content_ids=(shared_id, exclusive_id))
                assert [target.content_id for target in selected] == [shared_id]
                assert (
                    after_session.scalar(
                        select(func.count())
                        .select_from(historical_import_campaign_revocations_table)
                        .where(
                            historical_import_campaign_revocations_table.c.campaign_id
                            == campaign_id
                        )
                    )
                    == 1
                )
                assert (
                    after_session.scalar(
                        select(func.count())
                        .select_from(historical_import_revocation_content_versions_table)
                        .where(
                            historical_import_revocation_content_versions_table.c.campaign_id
                            == campaign_id
                        )
                    )
                    == 2
                )
                lifecycle_rows = tuple(
                    after_session.execute(
                        select(
                            contents_table.c.id,
                            provider_requests_table.c.provider,
                            provider_requests_table.c.operation,
                            artifacts_table.c.kind,
                            artifacts_table.c.storage_status,
                        )
                        .select_from(
                            contents_table.join(
                                content_versions_table,
                                and_(
                                    content_versions_table.c.content_id == contents_table.c.id,
                                    content_versions_table.c.version_no
                                    == contents_table.c.current_version,
                                ),
                            )
                            .join(
                                provider_request_attempts_table,
                                provider_request_attempts_table.c.id
                                == content_versions_table.c.provider_attempt_id,
                            )
                            .join(
                                provider_requests_table,
                                provider_requests_table.c.id
                                == provider_request_attempts_table.c.provider_request_id,
                            )
                            .join(
                                artifacts_table,
                                artifacts_table.c.id == content_versions_table.c.raw_artifact_id,
                            )
                        )
                        .where(contents_table.c.id.in_((shared_id, exclusive_id)))
                    ).mappings()
                )
                assert len(lifecycle_rows) == 2
                assert {row["provider"] for row in lifecycle_rows} == {"imports"}
                assert {row["operation"] for row in lifecycle_rows} == {"data_import_revoke"}
                assert {row["kind"] for row in lifecycle_rows} == {"provider-raw"}
                assert {row["storage_status"] for row in lifecycle_rows} == {"linked"}

                guarded_mutations = (
                    update(content_source_contributions_table).values(
                        observed_at=content_source_contributions_table.c.observed_at
                    ),
                    delete(historical_import_campaign_revocations_table).where(
                        historical_import_campaign_revocations_table.c.campaign_id == campaign_id
                    ),
                    update(historical_import_revocation_content_versions_table)
                    .where(
                        historical_import_revocation_content_versions_table.c.campaign_id
                        == campaign_id
                    )
                    .values(
                        created_at=historical_import_revocation_content_versions_table.c.created_at
                    ),
                )
                for mutation in guarded_mutations:
                    with pytest.raises(DBAPIError), after_session.begin_nested():
                        after_session.execute(mutation)
        finally:
            after_session.close()
    finally:
        _cleanup(runtime)


def test_revocation_restores_field_filled_by_historical_fill_only(tmp_path: Path) -> None:
    """历史补空贡献被撤销后，共享 Content 恢复为导入前的空字段而不是删除 Content。"""

    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    shared_key = f"revocation-fill-{uuid4()}"
    (historical_root / "fill.xlsx").write_bytes(
        _xlsx((("爱玛补空内容", shared_key, "历史补空正文"),))
    )
    runtime = _runtime(tmp_path, historical_root)
    try:
        client = _client(runtime)
        brand_id = stage3_filter_brand_id(runtime)
        baseline = client.post(
            "/api/v1/import-batches",
            files=[
                (
                    "file",
                    (
                        "baseline-empty.xlsx",
                        _xlsx((("爱玛补空内容", shared_key, None),)),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
                ("brand_ids", (None, brand_id)),
            ],
        )
        assert baseline.status_code == 202
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="revocation-fill-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker) == 1

        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"revocation-fill-{uuid4()}",
                "relative_paths": ["fill.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        campaign_id = UUID(created.json()["campaign_id"])
        assert _drain(worker) == 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )
        assert _drain(worker) == 1

        session = runtime.database.new_session()
        try:
            with session.begin():
                before = PostgresContentQueryRepository(
                    session,
                    analysis_identity=None,
                ).list_contents(
                    ContentReadQuery(
                        filters=ContentFilterSnapshot(),
                        position=None,
                        limit=20,
                    )
                )
                assert len(before) == 1
                content_id = before[0].id
                assert before[0].text == "历史补空正文"
                version_before_revoke = session.scalar(
                    select(contents_table.c.current_version).where(
                        contents_table.c.id == content_id
                    )
                )
        finally:
            session.close()

        service = PostgresImportRevocationHttpService(
            runtime.database.new_session,
            runtime.artifact_store,
        )
        preview = service.preview(campaign_id)
        assert preview.eligible is True
        assert preview.impact.unreversible_content_count == 0
        result = service.revoke(
            campaign_id,
            DataImportRevokeRequest(reason="撤销历史补空"),
            actor_ref="integration-admin",
            request_id="revocation-fill-request",
        )
        assert result.already_revoked is False
        assert result.status == "queued"
        assert result.impact.retained_shared_content_count == 1
        assert _drain(worker) == 1
        assert service.preview(campaign_id).status == "succeeded"

        session = runtime.database.new_session()
        try:
            with session.begin():
                after_repository = PostgresContentQueryRepository(
                    session,
                    analysis_identity=None,
                )
                after = after_repository.list_contents(
                    ContentReadQuery(
                        filters=ContentFilterSnapshot(),
                        position=None,
                        limit=20,
                    )
                )
                assert len(after) == 1
                assert after[0].id == content_id
                assert after[0].text is None
                version_after_revoke = session.scalar(
                    select(contents_table.c.current_version).where(
                        contents_table.c.id == content_id
                    )
                )
                assert version_after_revoke == version_before_revoke + 1
                assert (
                    session.scalar(
                        select(func.count())
                        .select_from(historical_import_revocation_content_versions_table)
                        .where(
                            historical_import_revocation_content_versions_table.c.campaign_id
                            == campaign_id
                        )
                    )
                    == 1
                )
        finally:
            session.close()
    finally:
        _cleanup(runtime)


def test_revocation_retries_from_committed_batch_checkpoint(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """第二批瞬时失败后，同一请求从已提交断点继续且不重复创建 Version。"""

    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "resume.xlsx").write_bytes(
        _xlsx(
            tuple(
                (f"爱玛断点撤销 {index}", f"revocation-resume-{uuid4()}-{index}", "正文")
                for index in range(501)
            )
        )
    )
    runtime = _runtime(tmp_path, historical_root)
    try:
        client = _client(runtime)
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="revocation-resume-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"revocation-resume-{uuid4()}",
                "relative_paths": ["resume.xlsx"],
                "recursive": False,
                "brand_ids": [stage3_filter_brand_id(runtime)],
                "ingestion_policy": "standard_observation",
            },
        )
        campaign_id = UUID(created.json()["campaign_id"])
        assert _drain(worker) == 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )
        assert _drain(worker) == 6
        service = PostgresImportRevocationHttpService(
            runtime.database.new_session, runtime.artifact_store
        )
        requested = service.revoke(
            campaign_id,
            DataImportRevokeRequest(reason="断点恢复验证"),
            actor_ref="integration-admin",
            request_id="revocation-resume-request",
        )
        assert requested.status == "queued"
        original = PostgresImportRevocationJobExecutor._apply_batch
        calls = 0

        def fail_second_batch(self: PostgresImportRevocationJobExecutor, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise SQLAlchemyError("injected transient failure")
            return original(self, **kwargs)

        with monkeypatch.context() as patch:
            patch.setattr(PostgresImportRevocationJobExecutor, "_apply_batch", fail_second_batch)
            assert worker.run_once()
        with runtime.database.engine.connect() as connection:
            checkpoint = (
                connection.execute(
                    select(historical_import_revocation_requests_table).where(
                        historical_import_revocation_requests_table.c.campaign_id == campaign_id
                    )
                )
                .mappings()
                .one()
            )
            assert checkpoint["recomputed_content_count"] == 500
            assert checkpoint["checkpoint_content_id"] is not None
        assert worker.run_once()
        completed = service.preview(campaign_id)
        assert completed.status == "succeeded"
        assert completed.recomputed_content_count == 501
        with runtime.database.engine.connect() as connection:
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(historical_import_revocation_content_versions_table)
                    .where(
                        historical_import_revocation_content_versions_table.c.campaign_id
                        == campaign_id
                    )
                )
                == 501
            )
    finally:
        _cleanup(runtime)
