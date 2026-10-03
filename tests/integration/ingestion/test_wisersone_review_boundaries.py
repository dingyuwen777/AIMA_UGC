"""独立验证恢复准入、取消终态投影与文件清理的真实 PostgreSQL 边界。"""

from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from pathlib import Path
from threading import Event
from typing import Any
from uuid import UUID

import pytest
from aima_ugc.adapters.persistence.postgres.artifact_metadata import (
    PostgresArtifactMetadataRepository,
)
from aima_ugc.adapters.persistence.postgres.historical_cancellation import (
    PostgresHistoricalCancellationRepository,
)
from aima_ugc.adapters.persistence.postgres.historical_import import (
    PostgresHistoricalImportRepository,
)
from aima_ugc.bootstrap import wisersone_http
from aima_ugc.bootstrap.wisersone_cleanup import cleanup_wisersone_files
from aima_ugc.bootstrap.wisersone_worker import PostgresWisersOneJobExecutor
from aima_ugc.modules.ingestion.historical_jobs import (
    HISTORICAL_DISCOVER_JOB_TYPE,
    HISTORICAL_IMPORT_CHUNK_JOB_TYPE,
    HISTORICAL_SNAPSHOT_JOB_TYPE,
)
from aima_ugc.modules.ingestion.historical_tables import (
    historical_import_campaign_items_table as items,
)
from aima_ugc.modules.ingestion.historical_tables import (
    historical_import_campaigns_table as campaigns,
)
from aima_ugc.modules.ingestion.historical_tables import (
    processing_import_batch_items_table as ledger,
)
from aima_ugc.modules.ingestion.wisersone_jobs import WISERSONE_JOB_TYPE
from aima_ugc.modules.ingestion.wisersone_tables import wisersone_downloads_table as downloads
from aima_ugc.platform.jobs.tables import jobs_table as jobs
from aima_ugc.platform.storage.canonical import CanonicalArtifactReader
from aima_ugc.platform.storage.tables import canonical_artifact_links_table as canonical_links
from aima_ugc.platform.time import beijing_now
from openpyxl import Workbook, load_workbook
from sqlalchemy import Text, cast, func, or_, select, update

from .test_wisersone_review_regressions import _campaign, _settle
from .test_wisersone_workflow import workflow as workflow


def _owned_jobs(runtime: Any, created: dict[str, Any]) -> list[Any]:
    """所有队列断言只关联本次 Download、Campaign 和它的来源项。"""
    campaign_id = UUID(created["campaign_id"])
    item_ids = select(cast(items.c.id, Text)).where(items.c.campaign_id == campaign_id)
    with runtime.database.new_session() as session:
        return list(
            session.execute(
                select(jobs)
                .where(
                    or_(
                        jobs.c.payload["download_id"].astext == created["id"],
                        jobs.c.payload["campaign_id"].astext == created["campaign_id"],
                        jobs.c.payload["campaign_item_id"].astext.in_(item_ids),
                        jobs.c.payload["chunk_item_id"].astext.in_(item_ids),
                    )
                )
                .order_by(jobs.c.created_at)
            ).mappings()
        )


def _queued(runtime: Any, created: dict[str, Any], job_type: str) -> UUID:
    queued = [
        job
        for job in _owned_jobs(runtime, created)
        if job["job_type"] == job_type and job["status"] == "queued"
    ]
    assert queued, (job_type, _owned_jobs(runtime, created))
    return queued[0]["id"]


def _run_job(workflow: Any, created: dict[str, Any], job_id: UUID, worker: Any = None) -> None:
    """只控制本次测试队列时钟；领取、Fence、Handler 和回调仍使用正式 Worker。"""
    runtime = workflow[0]
    owned = tuple(job["id"] for job in _owned_jobs(runtime, created) if job["status"] == "queued")
    with runtime.database.engine.begin() as connection:
        connection.execute(
            update(jobs)
            .where(jobs.c.id.in_(owned))
            .values(available_at=beijing_now() + timedelta(hours=1))
        )
        connection.execute(
            update(jobs)
            .where(jobs.c.id == job_id)
            .values(available_at=beijing_now() - timedelta(seconds=1))
        )
    assert (worker or workflow[4]()).run_once()


def _campaign_row(runtime: Any, created: dict[str, Any]) -> dict[str, Any]:
    with runtime.database.new_session() as session:
        return dict(
            session.execute(select(campaigns).where(campaigns.c.id == UUID(created["campaign_id"])))
            .mappings()
            .one()
        )


def _ready(workflow: Any) -> dict[str, Any]:
    runtime, _, _, _, created = _campaign(workflow)
    for job_type in (HISTORICAL_DISCOVER_JOB_TYPE, HISTORICAL_SNAPSHOT_JOB_TYPE):
        _run_job(workflow, created, _queued(runtime, created, job_type))
    assert _campaign_row(runtime, created)["status"] == "ready"
    return created


def _canonical_state(runtime: Any, created: dict[str, Any]) -> dict[UUID, tuple[bytes, int]]:
    with runtime.database.new_session() as session:
        ids = tuple(
            session.scalars(
                select(canonical_links.c.artifact_id)
                .join(
                    items,
                    items.c.id == canonical_links.c.historical_import_campaign_item_id,
                )
                .where(items.c.campaign_id == UUID(created["campaign_id"]))
            )
        )
        result = {}
        repository = PostgresArtifactMetadataRepository(session)
        for artifact_id in ids:
            artifact = repository.get(artifact_id)
            assert artifact is not None
            count = len(tuple(CanonicalArtifactReader(store=runtime.artifact_store).read(artifact)))
            result[artifact_id] = (runtime.artifact_store.read(artifact.storage_key), count)
        assert result
        return result


def _scoped_created_count(runtime: Any, created: dict[str, Any]) -> int:
    with runtime.database.new_session() as session:
        return session.scalar(
            select(func.count())
            .select_from(ledger)
            .join(
                items,
                items.c.id == ledger.c.campaign_item_id,
            )
            .where(
                items.c.campaign_id == UUID(created["campaign_id"]), ledger.c.outcome == "created"
            )
        )


def _running_fixture(workflow: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    """生产 Owner 完成第一分块后仍有第二分块，得到真实的 running Campaign。"""
    runtime, website = workflow[:2]
    runtime.settings = runtime.settings.model_copy(update={"historical_chunk_rows": 100})
    original = website.poll

    def two_chunks(task: Any, destination: Path, **options: Any) -> Any:
        result = original(task, destination, **options)
        if result.file is not None:
            book = load_workbook(result.file)
            try:
                sheet = book["文章"]
                for index in range(100):
                    sheet.append(
                        [
                            "小红书",
                            f"爱玛分块 {index}",
                            "爱玛",
                            "测试",
                            "2026-10-03 10:00:00",
                            f"https://www.xiaohongshu.com/explore/{task.task_id}-{index}",
                        ]
                    )
                book.save(result.file)
            finally:
                book.close()
        return result

    monkeypatch.setattr(website, "poll", two_chunks)


@pytest.mark.parametrize("campaign_state", ["ready", "running", "succeeded"])
def test_failed_monitor_resumes_existing_campaign_without_failed_chunks(
    workflow: Any,
    monkeypatch: pytest.MonkeyPatch,
    campaign_state: str,
) -> None:
    runtime, website, client, _, new_worker, tick = workflow
    if campaign_state == "running":
        _running_fixture(workflow, monkeypatch)
    created = _ready(workflow)
    if campaign_state != "ready":
        _run_job(workflow, created, _queued(runtime, created, WISERSONE_JOB_TYPE))
        _run_job(workflow, created, _queued(runtime, created, HISTORICAL_IMPORT_CHUNK_JOB_TYPE))
    assert _campaign_row(runtime, created)["status"] == campaign_state
    canonical = _canonical_state(runtime, created)
    monitor = _queued(runtime, created, WISERSONE_JOB_TYPE)
    original = PostgresWisersOneJobExecutor._campaign

    def failed_monitor(self: Any, payload: Any, context: Any, campaign_id: UUID) -> Any:
        if str(campaign_id) == created["campaign_id"]:
            raise OSError("模拟父监控读取临时失败，导入 Owner 继续保留已有事实")
        return original(self, payload, context, campaign_id)

    monkeypatch.setattr(PostgresWisersOneJobExecutor, "_campaign", failed_monitor)
    _run_job(workflow, created, monitor)
    monkeypatch.setattr(PostgresWisersOneJobExecutor, "_campaign", original)
    assert client.get(f"/api/v1/wisersone-downloads/{created['id']}").json()["status"] == "failed"
    assert _campaign_row(runtime, created)["status"] == campaign_state
    with runtime.database.new_session() as session:
        assert not PostgresHistoricalImportRepository(session).has_failed_chunks(
            UUID(created["campaign_id"])
        )
    before_jobs = len(_owned_jobs(runtime, created))
    response = client.post(f"/api/v1/wisersone-downloads/{created['id']}/retry")
    assert response.status_code == 200, response.text
    assert response.json()["campaign_id"] == created["campaign_id"]
    assert len(_owned_jobs(runtime, created)) == before_jobs + 1
    assert client.post(f"/api/v1/wisersone-downloads/{created['id']}/retry").status_code == 409
    assert len(_owned_jobs(runtime, created)) == before_jobs + 1
    result = _settle(client, new_worker(), tick, created["id"], {"succeeded"})
    assert result["campaign_id"] == created["campaign_id"] and website.submissions == 1
    assert _canonical_state(runtime, created) == canonical
    assert _scoped_created_count(runtime, created) == (101 if campaign_state == "running" else 1)


def test_permanently_invalid_snapshot_cannot_retry_or_export_again(
    workflow: Any,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, website, client, _, new_worker, tick = workflow
    original = website.poll

    def invalid_input(task: Any, destination: Path, **options: Any) -> Any:
        result = original(task, destination, **options)
        if result.file is not None:
            book = Workbook()
            sheet = book.active
            assert sheet is not None
            sheet.append(["无法识别的表头"])
            sheet.append(["永久无效输入"])
            try:
                book.save(result.file)
            finally:
                book.close()
        return result

    monkeypatch.setattr(website, "poll", invalid_input)
    _, _, worker, _, created = _campaign(workflow)
    _settle(client, worker, tick, created["id"], {"failed"})
    campaign = _campaign_row(runtime, created)
    assert campaign["status"] == "failed"
    with runtime.database.new_session() as session:
        errors = session.scalars(
            select(items.c.error_code).where(
                items.c.campaign_id == UUID(created["campaign_id"]),
                items.c.item_kind == "source_file",
            )
        ).all()
        assert errors == ["historical_snapshot_invalid"]
    before = _owned_jobs(runtime, created)
    response = client.post(f"/api/v1/wisersone-downloads/{created['id']}/retry")
    assert response.status_code == 409, response.text
    assert _campaign_row(runtime, created) == campaign
    assert [job["id"] for job in _owned_jobs(runtime, created)] == [job["id"] for job in before]
    assert website.submissions == 1 and new_worker().run_once() is False


def test_cancel_waits_for_last_chunk_commit_and_projects_actual_success(
    workflow: Any,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, website, client, _, new_worker, tick = workflow
    created = _ready(workflow)
    _run_job(workflow, created, _queued(runtime, created, WISERSONE_JOB_TYPE))
    chunk = _queued(runtime, created, HISTORICAL_IMPORT_CHUNK_JOB_TYPE)
    completed, release, cancel_entered = Event(), Event(), Event()
    original = PostgresHistoricalImportRepository.refresh_batch_and_campaign
    original_gate = wisersone_http.lock_historical_campaign_cancel_gate

    def completed_before_commit(self: Any, *, campaign_id: UUID, batch_id: UUID) -> str:
        status = original(self, campaign_id=campaign_id, batch_id=batch_id)
        if str(campaign_id) == created["campaign_id"]:
            assert status == "succeeded"
            completed.set()
            assert release.wait(15), "最后分块事务未获得释放"
        return status

    def cancel_gate(session: Any, campaign_id: UUID, *, shared: bool) -> None:
        if str(campaign_id) == created["campaign_id"] and not shared:
            cancel_entered.set()
        original_gate(session, campaign_id, shared=shared)

    monkeypatch.setattr(
        PostgresHistoricalImportRepository, "refresh_batch_and_campaign", completed_before_commit
    )
    monkeypatch.setattr(wisersone_http, "lock_historical_campaign_cancel_gate", cancel_gate)
    with ThreadPoolExecutor(max_workers=2) as pool:
        writing = pool.submit(_run_job, workflow, created, chunk, new_worker())
        try:
            assert completed.wait(15)
            cancelling = pool.submit(
                client.post, f"/api/v1/wisersone-downloads/{created['id']}/cancel"
            )
            assert cancel_entered.wait(15)
            assert not cancelling.done(), "取消必须等待仍持有共享取消门的最后写事务"
        finally:
            release.set()
        writing.result(timeout=15)
        response = cancelling.result(timeout=15)
        assert response.status_code == 200, response.text
        assert response.json()["status"] == "cancelling"
    result = _settle(client, new_worker(), tick, created["id"], {"succeeded"})
    assert result["campaign_id"] == created["campaign_id"]
    assert _campaign_row(runtime, created)["status"] == "succeeded"
    assert _scoped_created_count(runtime, created) == 1 and website.submissions == 1


def test_exhausted_cancel_propagation_enqueues_continuation_for_new_worker(
    workflow: Any,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, client, _, tick, created = _campaign(workflow)
    response = client.post(f"/api/v1/wisersone-downloads/{created['id']}/cancel")
    assert response.status_code == 200 and response.json()["status"] == "cancelling"
    cancellation_id = UUID(response.json()["job_id"])
    with runtime.database.engine.begin() as connection:
        connection.execute(update(jobs).where(jobs.c.id == cancellation_id).values(max_attempts=1))
    original = PostgresHistoricalCancellationRepository.request_job_cancellations

    def propagation_error(self: Any, campaign_id: UUID) -> None:
        if str(campaign_id) == created["campaign_id"]:
            raise OSError("模拟取消传播阶段的暂时存储故障")
        original(self, campaign_id)

    monkeypatch.setattr(
        PostgresHistoricalCancellationRepository, "request_job_cancellations", propagation_error
    )
    _run_job(workflow, created, cancellation_id)
    monkeypatch.setattr(
        PostgresHistoricalCancellationRepository, "request_job_cancellations", original
    )
    failed = next(job for job in _owned_jobs(runtime, created) if job["id"] == cancellation_id)
    assert failed["status"] == "failed" and failed["attempt"] == failed["max_attempts"] == 1
    current = client.get(f"/api/v1/wisersone-downloads/{created['id']}").json()
    assert current["status"] == "cancelling" and current["finished_at"] is None
    continuation = next(
        job for job in _owned_jobs(runtime, created) if str(job["id"]) == current["job_id"]
    )
    assert continuation["id"] != cancellation_id and continuation["status"] == "queued"
    assert continuation["payload"]["operation"] == "cancel"
    assert continuation["payload"]["step"] == failed["payload"]["step"] + 1
    result = _settle(client, workflow[4](), tick, created["id"], {"cancelled"})
    assert result["campaign_id"] == created["campaign_id"]
    assert _campaign_row(runtime, created)["status"] == "cancelled"
    assert _scoped_created_count(runtime, created) == 0 and workflow[1].submissions == 1


def _failed_import(workflow: Any, monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    runtime, _, client, _, new_worker, tick = workflow
    created = _ready(workflow)
    _run_job(workflow, created, _queued(runtime, created, WISERSONE_JOB_TYPE))
    chunk_job = _queued(runtime, created, HISTORICAL_IMPORT_CHUNK_JOB_TYPE)
    with runtime.database.engine.begin() as connection:
        connection.execute(update(jobs).where(jobs.c.id == chunk_job).values(max_attempts=1))
    canonical = _canonical_state(runtime, created)
    original = CanonicalArtifactReader.read_for_bounded_staging

    def artifact_read_error(self: Any, artifact: Any) -> Any:
        if artifact.id in canonical:
            raise OSError("模拟已冻结 Canonical 文件的暂时读取故障")
        return original(self, artifact)

    monkeypatch.setattr(CanonicalArtifactReader, "read_for_bounded_staging", artifact_read_error)
    _run_job(workflow, created, chunk_job)
    monkeypatch.setattr(CanonicalArtifactReader, "read_for_bounded_staging", original)
    _settle(client, new_worker(), tick, created["id"], {"failed", "partial_failed"})
    assert _campaign_row(runtime, created)["status"] in {"failed", "partial_failed"}
    return created


def test_claimed_ttl_blocks_both_retry_entries_before_campaign_activation(
    workflow: Any,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, _, client = workflow[:3]
    created = _failed_import(workflow, monkeypatch)
    canonical = _canonical_state(runtime, created)
    campaign = _campaign_row(runtime, created)
    before_jobs = [job["id"] for job in _owned_jobs(runtime, created)]
    published = runtime.settings.wisersone_input_dir / created["id"] / "wisersone_last24h.xlsx"
    entered, release = Event(), Event()
    original = Path.unlink

    def pause_unlink(path: Path, *args: Any, **kwargs: Any) -> None:
        if path == published:
            entered.set()
            assert release.wait(15), "认领后的文件清理未获得释放"
        original(path, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", pause_unlink)
    with ThreadPoolExecutor(max_workers=1) as pool:
        cleaning = pool.submit(
            cleanup_wisersone_files, runtime, now=beijing_now() + timedelta(days=8)
        )
        try:
            assert entered.wait(15)
            assert published.exists()
            with runtime.database.new_session() as session:
                pending = session.scalar(
                    select(downloads.c.file_delete_pending_at).where(
                        downloads.c.id == UUID(created["id"])
                    )
                )
                assert pending is not None
            for path in (
                f"/api/v1/wisersone-downloads/{created['id']}/retry",
                f"/api/v1/data-import-campaigns/{created['campaign_id']}/retry-failed",
            ):
                response = client.post(path)
                assert response.status_code == 409, response.text
                assert _campaign_row(runtime, created) == campaign
                assert [job["id"] for job in _owned_jobs(runtime, created)] == before_jobs
            assert _canonical_state(runtime, created) == canonical
        finally:
            release.set()
        assert cleaning.result(timeout=15) == 1
    assert not published.exists() and _canonical_state(runtime, created) == canonical


def test_existing_campaign_retry_holds_download_admission_against_ttl(
    workflow: Any,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime, _, client, _, new_worker, tick = workflow
    created = _failed_import(workflow, monkeypatch)
    canonical = _canonical_state(runtime, created)
    entered, release = Event(), Event()
    original = PostgresHistoricalImportRepository.prepare_failed_retry

    def admitted(self: Any, campaign_id: UUID) -> Any:
        if str(campaign_id) == created["campaign_id"]:
            entered.set()
            assert release.wait(15), "已获准的 Campaign 恢复未获得释放"
        return original(self, campaign_id)

    monkeypatch.setattr(PostgresHistoricalImportRepository, "prepare_failed_retry", admitted)
    with ThreadPoolExecutor(max_workers=2) as pool:
        retrying = pool.submit(
            client.post, f"/api/v1/data-import-campaigns/{created['campaign_id']}/retry-failed"
        )
        try:
            assert entered.wait(15)
            cleaning = pool.submit(
                cleanup_wisersone_files, runtime, now=beijing_now() + timedelta(days=8)
            )
            assert cleaning.result(timeout=15) == 0
        finally:
            release.set()
        response = retrying.result(timeout=15)
        assert response.status_code == 200, response.text
    published = runtime.settings.wisersone_input_dir / created["id"] / "wisersone_last24h.xlsx"
    assert published.exists() and _canonical_state(runtime, created) == canonical
    result = _settle(client, new_worker(), tick, created["id"], {"succeeded"})
    assert result["campaign_id"] == created["campaign_id"] and workflow[1].submissions == 1
