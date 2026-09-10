"""Stage 12B 真实文件系统、API、Job、Artifact 与 PostgreSQL Golden Path。"""

from __future__ import annotations

from datetime import timedelta
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from aima_ugc.adapters.persistence.postgres.brand_vehicle import PostgresBrandVehicleRepository
from aima_ugc.adapters.persistence.postgres.content_queries import (
    PostgresContentQueryRepository,
)
from aima_ugc.adapters.persistence.postgres.jobs import PostgresJobRepository
from aima_ugc.adapters.persistence.postgres.vehicles import PostgresVehicleCatalogRepository
from aima_ugc.bootstrap import historical_import_worker as historical_worker_module
from aima_ugc.bootstrap.administration_http import PostgresAdministrationHttpService
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.brand_vehicle_http import PostgresBrandVehicleHttpService
from aima_ugc.bootstrap.historical_import_http import PostgresHistoricalImportHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.runtime import PlatformRuntime
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.administration import (
    VehicleModelCreateRequest,
    VehicleModelUpdateRequest,
)
from aima_ugc.contracts.brand_vehicle import BrandCreateRequest
from aima_ugc.contracts.http import ContentFilterSnapshot
from aima_ugc.modules.content.query import ContentReadQuery
from aima_ugc.modules.content.tables import content_metric_observations_table, contents_table
from aima_ugc.modules.identity import Principal
from aima_ugc.modules.ingestion.historical_jobs import HISTORICAL_IMPORT_CHUNK_JOB_TYPE
from aima_ugc.modules.ingestion.historical_tables import (
    historical_import_campaign_items_table,
    historical_import_campaigns_table,
    processing_import_batch_identities_table,
    processing_import_batch_item_conflicts_table,
    processing_import_batch_items_table,
)
from aima_ugc.modules.ingestion.tables import processing_import_batches_table
from aima_ugc.modules.vehicles.tables import (
    content_brand_evidence_table,
    content_brand_review_locks_table,
    content_vehicle_evidence_table,
    content_vehicle_review_locks_table,
)
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs.tables import jobs_table
from aima_ugc.platform.storage.tables import artifacts_table
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, insert, select, update
from sqlalchemy.exc import IntegrityError


def _xlsx(*, title: str, text: str | None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(
        [
            "媒体名称（中文）",
            "标题",
            "内文",
            "作者",
            "出版日期",
            "原文链接",
        ]
    )
    sheet.append(
        [
            "小红书",
            title,
            text,
            "官方账号",
            "2025-01-02 10:00:00",
            "https://www.xiaohongshu.com/explore/stage12-shared-content",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _xlsx_with_cross_chunk_duplicates(*, rows: int) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(
        [
            "媒体名称（中文）",
            "标题",
            "内文",
            "作者",
            "出版日期",
            "原文链接",
        ]
    )
    for index in range(rows):
        sheet.append(
            [
                "小红书",
                "爱玛稳定首行" if index == 0 else f"爱玛后续重复行 {index}",
                "跨 Chunk 重复身份",
                "官方账号",
                "2025-01-02 10:00:00",
                "https://www.xiaohongshu.com/explore/stage12-cross-chunk-duplicate",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _principal() -> Principal:
    return Principal(
        principal_id="stage3-historical-admin",
        display_name="Stage3 Historical 管理员",
        role="administrator",
        source="development",
    )


def _brand(runtime: PlatformRuntime) -> str:
    brand = PostgresBrandVehicleHttpService(runtime).create_brand(
        BrandCreateRequest(
            code=f"AIMA-STAGE3-{uuid4()}",
            display_name="爱玛",
            role="owned",
            aliases=("爱玛",),
        ),
        principal=_principal(),
        request_id="stage3-historical-brand",
    )
    return str(brand.id)


def _drain(worker, *, maximum: int = 20) -> int:
    executed = 0
    for _ in range(maximum):
        if not worker.run_once():
            break
        executed += 1
    return executed


def test_source_manifest_identity_is_database_unique_when_ordinal_is_null(
    tmp_path: Path,
) -> None:
    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "unique.xlsx").write_bytes(
        _xlsx(title="爱玛源文件唯一身份", text="不可重复建立源文件 Item")
    )
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-source-unique-{uuid4()}",
                "relative_paths": ["unique.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        assert created.status_code == 202
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-source-unique-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker, maximum=1) == 1

        with runtime.database.engine.begin() as connection:
            source_item = dict(
                connection.execute(
                    select(historical_import_campaign_items_table).where(
                        historical_import_campaign_items_table.c.item_kind == "source_file"
                    )
                )
                .mappings()
                .one()
            )
        source_item["id"] = uuid4()
        source_item["job_id"] = None
        with pytest.raises(IntegrityError):
            with runtime.database.engine.begin() as connection:
                connection.execute(insert(historical_import_campaign_items_table), source_item)
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_historical_campaign_preflights_before_fill_only_import(tmp_path: Path) -> None:
    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "history.xlsx").write_bytes(
        _xlsx(title="爱玛历史冲突标题", text="历史正文补空")
    )
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        import_service = PostgresImportHttpService(runtime)
        historical_service = PostgresHistoricalImportHttpService(runtime)
        client = TestClient(
            create_app(
                import_service=import_service,
                historical_import_service=historical_service,
            )
        )
        brand_id = _brand(runtime)
        baseline = client.post(
            "/api/v1/import-batches",
            files=[
                (
                    "file",
                    (
                        "current.xlsx",
                        _xlsx(title="爱玛在线标题", text=None),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
                ("brand_ids", (None, brand_id)),
            ],
        )
        assert baseline.status_code == 202
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-golden-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker) == 1

        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-{uuid4()}",
                "relative_paths": ["history.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        assert created.status_code == 202
        campaign_id = created.json()["campaign_id"]
        before_preflight = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert before_preflight["status"] == "discovering"
        assert before_preflight["can_start"] is False
        assert before_preflight["progress"] == {
            "preflight_completed_file_count": 0,
            "preflight_percent": 0,
            "migration_completed_row_count": 0,
            "migration_percent": 0,
        }

        assert _drain(worker) == 2
        ready = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert ready["status"] == "ready"
        assert ready["can_start"] is True
        assert ready["total_rows"] == 1
        assert ready["progress"]["preflight_completed_file_count"] == 1
        assert ready["progress"]["preflight_percent"] == 100
        assert ready["progress"]["migration_percent"] == 0
        listed = client.get("/api/v1/historical-import-campaigns").json()["items"]
        listed_campaign = next(item for item in listed if item["id"] == campaign_id)
        assert listed_campaign["progress"] == ready["progress"]

        started = client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start")
        assert started.status_code == 200
        assert started.json()["status"] == "queued"
        assert _drain(worker) == 1

        completed = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert completed["status"] == "succeeded"
        assert completed["stats"]["filled"] == 1
        assert completed["progress"]["migration_completed_row_count"] == 1
        assert completed["progress"]["migration_percent"] == 100
        items = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}/items").json()[
            "items"
        ]
        assert [item["status"] for item in items] == ["succeeded", "succeeded"]
        conflict_response = client.get(
            f"/api/v1/historical-import-campaigns/{campaign_id}/conflicts"
        )
        assert conflict_response.status_code == 200
        assert [item["field_name"] for item in conflict_response.json()["items"]] == ["title"]

        with runtime.database.engine.begin() as connection:
            content = connection.execute(select(contents_table)).mappings().one()
            outcomes = (
                connection.execute(select(processing_import_batch_items_table.c.outcome))
                .scalars()
                .all()
            )
            conflicts = (
                connection.execute(
                    select(processing_import_batch_item_conflicts_table.c.field_name)
                )
                .scalars()
                .all()
            )
            metric_count = connection.scalar(
                select(func.count()).select_from(content_metric_observations_table)
            )
            artifact_count = connection.scalar(
                select(func.count()).select_from(historical_import_campaign_items_table)
            )
            analysis_job_count = connection.scalar(
                select(func.count())
                .select_from(jobs_table)
                .where(jobs_table.c.job_type == "analysis.content-label.v1")
            )
        assert content["title"] == "爱玛在线标题"
        assert content["text"] == "历史正文补空"
        assert outcomes == ["filled"]
        assert conflicts == ["title"]
        assert metric_count == 1
        assert artifact_count == 2
        assert analysis_job_count == 0
        query_session = runtime.database.new_session()
        try:
            with query_session.begin():
                campaign_records = PostgresContentQueryRepository(
                    query_session,
                    analysis_identity=None,
                ).list_contents(
                    ContentReadQuery(
                        filters=ContentFilterSnapshot(source_identifier=UUID(campaign_id)),
                        position=None,
                        limit=20,
                    )
                )
            assert [record.id for record in campaign_records] == [content["id"]]
        finally:
            query_session.close()
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_local_campaign_uploads_immutable_artifact_before_common_preflight(
    tmp_path: Path,
) -> None:
    payload = _xlsx(title="爱玛本地统一导入", text="本地文件正文")
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": None,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        local_request = {
            "client_idempotency_key": f"stage12-local-{uuid4()}",
            "files": [{"relative_path": "picked/folder/local.xlsx", "byte_size": len(payload)}],
            "brand_ids": [brand_id],
            "ingestion_policy": "standard_observation",
        }
        created = client.post("/api/v1/data-import-campaigns/local", json=local_request)
        assert created.status_code == 201
        repeated_create = client.post(
            "/api/v1/data-import-campaigns/local",
            json=local_request,
        )
        assert repeated_create.status_code == 201
        assert repeated_create.json() == created.json()
        campaign_id = created.json()["campaign_id"]
        item_id = created.json()["upload_items"][0]["item_id"]
        staging = client.get(f"/api/v1/data-import-campaigns/{campaign_id}").json()
        assert staging["status"] == "uploading"
        assert staging["source_kind"] == "local_upload"
        assert staging["ingestion_policy"] == "standard_observation"
        assert staging["declared_file_count"] == 1
        assert staging["can_start"] is False
        assert (
            client.post(f"/api/v1/data-import-campaigns/{campaign_id}/finalize").status_code == 409
        )

        uploaded = client.put(
            f"/api/v1/data-import-campaigns/{campaign_id}/items/{item_id}/content",
            files={
                "file": (
                    "local.xlsx",
                    payload,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert uploaded.status_code == 200
        assert uploaded.json()["byte_size"] == len(payload)
        repeated_upload = client.put(
            f"/api/v1/data-import-campaigns/{campaign_id}/items/{item_id}/content",
            files={
                "file": (
                    "local.xlsx",
                    payload,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert repeated_upload.status_code == 200
        assert repeated_upload.json()["artifact_id"] == uploaded.json()["artifact_id"]
        altered_payload = payload[:-1] + bytes([payload[-1] ^ 1])
        rejected_replay = client.put(
            f"/api/v1/data-import-campaigns/{campaign_id}/items/{item_id}/content",
            files={
                "file": (
                    "local.xlsx",
                    altered_payload,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert rejected_replay.status_code == 422
        with runtime.database.engine.connect() as connection:
            assert (
                connection.scalar(
                    select(artifacts_table.c.storage_status).where(
                        artifacts_table.c.id == UUID(uploaded.json()["artifact_id"])
                    )
                )
                == "linked"
            )

        finalized = client.post(f"/api/v1/data-import-campaigns/{campaign_id}/finalize")
        assert finalized.status_code == 202
        assert finalized.json()["status"] == "snapshotting"

        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-local-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker) == 1
        ready = client.get(f"/api/v1/data-import-campaigns/{campaign_id}").json()
        assert ready["status"] == "ready"
        assert ready["progress"]["preflight_percent"] == 100

        started = client.post(f"/api/v1/data-import-campaigns/{campaign_id}/start")
        assert started.status_code == 200
        assert _drain(worker) == 1
        completed = client.get(f"/api/v1/data-import-campaigns/{campaign_id}").json()
        assert completed["status"] == "succeeded"
        assert completed["stats"]["created"] == 1

        updated_payload = _xlsx(title="爱玛本地统一导入已更新", text="本地文件正文")
        updated_created = client.post(
            "/api/v1/data-import-campaigns/local",
            json={
                "client_idempotency_key": f"stage12-local-update-{uuid4()}",
                "files": [
                    {
                        "relative_path": "picked/updated.xlsx",
                        "byte_size": len(updated_payload),
                    }
                ],
                "brand_ids": [brand_id],
                "ingestion_policy": "standard_observation",
            },
        ).json()
        updated_campaign_id = updated_created["campaign_id"]
        updated_item_id = updated_created["upload_items"][0]["item_id"]
        assert (
            client.put(
                f"/api/v1/data-import-campaigns/{updated_campaign_id}/items/"
                f"{updated_item_id}/content",
                files={
                    "file": (
                        "updated.xlsx",
                        updated_payload,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            ).status_code
            == 200
        )
        assert (
            client.post(f"/api/v1/data-import-campaigns/{updated_campaign_id}/finalize").status_code
            == 202
        )
        assert _drain(worker) == 1
        assert (
            client.post(f"/api/v1/data-import-campaigns/{updated_campaign_id}/start").status_code
            == 200
        )
        assert _drain(worker) == 1
        updated = client.get(f"/api/v1/data-import-campaigns/{updated_campaign_id}").json()
        assert updated["status"] == "succeeded"
        assert updated["stats"]["updated"] == 1
        with runtime.database.engine.begin() as connection:
            content = connection.execute(select(contents_table)).mappings().one()
        assert content["title"] == "爱玛本地统一导入已更新"
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_local_campaign_can_be_cancelled_while_uploading(tmp_path: Path) -> None:
    """本地上传失败后允许管理员把尚未进入预检的 Campaign 收敛到终态。"""
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": None,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        created = client.post(
            "/api/v1/data-import-campaigns/local",
            json={
                "client_idempotency_key": f"stage12-local-cancel-{uuid4()}",
                "files": [
                    {
                        "relative_path": "picked/interrupted.xlsx",
                        "byte_size": 1024,
                    }
                ],
                "brand_ids": [brand_id],
                "ingestion_policy": "standard_observation",
            },
        )
        assert created.status_code == 201
        campaign_id = created.json()["campaign_id"]

        cancelled = client.post(f"/api/v1/data-import-campaigns/{campaign_id}/cancel")

        assert cancelled.status_code == 200
        assert cancelled.json()["status"] == "cancelled"
        assert cancelled.json()["finished_at"] is not None
        items = client.get(f"/api/v1/data-import-campaigns/{campaign_id}/items").json()["items"]
        assert len(items) == 1
        assert items[0]["item_kind"] == "source_file"
        assert items[0]["status"] == "cancelled"
        assert items[0]["finished_at"] is not None
        late_upload = client.put(
            f"/api/v1/data-import-campaigns/{campaign_id}/items/"
            f"{created.json()['upload_items'][0]['item_id']}/content",
            files={
                "file": (
                    "interrupted.xlsx",
                    b"x" * 1024,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert late_upload.status_code == 409
        with runtime.database.engine.connect() as connection:
            assert connection.scalar(select(func.count()).select_from(artifacts_table)) == 0
        assert (
            client.post(f"/api/v1/data-import-campaigns/{campaign_id}/finalize").status_code == 409
        )
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_historical_snapshot_fails_closed_when_source_changes_after_discovery(
    tmp_path: Path,
) -> None:
    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    source = historical_root / "changing.xlsx"
    source.write_bytes(_xlsx(title="爱玛发现时标题", text="发现时正文"))
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-change-{uuid4()}",
                "relative_paths": ["changing.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        assert created.status_code == 202
        campaign_id = created.json()["campaign_id"]
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-source-change-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )

        assert worker.run_once() is True
        source.write_bytes(
            _xlsx(
                title="爱玛发现后被修改且长度明显不同的标题",
                text="发现后正文也被修改，不能混入不可变快照",
            )
        )
        assert worker.run_once() is True

        campaign = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        items = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}/items").json()[
            "items"
        ]
        assert campaign["status"] == "failed"
        assert campaign["can_start"] is False
        assert items[0]["status"] == "failed"
        assert items[0]["error_code"] == "historical_source_changed"
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_historical_snapshot_technical_retry_reuses_bound_source_artifact(
    tmp_path: Path,
    monkeypatch,
) -> None:
    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "retry.xlsx").write_bytes(
        _xlsx(title="爱玛快照重试", text="技术重试必须复用不可变源 Artifact")
    )
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    original_convert = historical_worker_module.convert_historical_excel_to_chunks
    convert_attempts = 0

    def fail_first_conversion(**kwargs):
        nonlocal convert_attempts
        convert_attempts += 1
        if convert_attempts == 1:
            raise OSError("simulated transient conversion I/O failure")
        return original_convert(**kwargs)

    monkeypatch.setattr(
        historical_worker_module,
        "convert_historical_excel_to_chunks",
        fail_first_conversion,
    )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-snapshot-retry-{uuid4()}",
                "relative_paths": ["retry.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        campaign_id = created.json()["campaign_id"]
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-snapshot-retry-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )

        assert worker.run_once() is True
        assert worker.run_once() is True
        with runtime.database.engine.begin() as connection:
            source_item = (
                connection.execute(
                    select(historical_import_campaign_items_table).where(
                        historical_import_campaign_items_table.c.campaign_id == UUID(campaign_id),
                        historical_import_campaign_items_table.c.item_kind == "source_file",
                    )
                )
                .mappings()
                .one()
            )
            source_artifact_count = connection.scalar(
                select(func.count())
                .select_from(artifacts_table)
                .where(artifacts_table.c.kind == "historical-import.source")
            )
        assert source_item["status"] == "snapshotting"
        assert source_item["artifact_id"] is not None
        assert source_item["sha256"] is not None
        assert source_artifact_count == 1

        assert worker.run_once() is True
        ready = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert ready["status"] == "ready"
        assert ready["total_rows"] == 1
        with runtime.database.engine.begin() as connection:
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(artifacts_table)
                    .where(artifacts_table.c.kind == "historical-import.source")
                )
                == 1
            )
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_historical_single_source_schedules_chunks_in_order_for_stable_first_row(
    tmp_path: Path,
) -> None:
    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "duplicates.xlsx").write_bytes(_xlsx_with_cross_chunk_duplicates(rows=101))
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 2,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-stable-first-{uuid4()}",
                "relative_paths": ["duplicates.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        assert created.status_code == 202
        campaign_id = created.json()["campaign_id"]
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-stable-first-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker) == 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )

        chunks = [
            item
            for item in client.get(
                f"/api/v1/historical-import-campaigns/{campaign_id}/items"
            ).json()["items"]
            if item["item_kind"] == "chunk"
        ]
        assert [item["status"] for item in chunks] == ["queued", "ready"]

        assert worker.run_once() is True
        in_progress = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert in_progress["status"] == "running"
        assert in_progress["progress"]["migration_completed_row_count"] == 100
        assert in_progress["progress"]["migration_percent"] == 99
        assert worker.run_once() is True
        completed = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert completed["progress"]["migration_completed_row_count"] == 101
        assert completed["progress"]["migration_percent"] == 100
        with runtime.database.engine.begin() as connection:
            content = connection.execute(select(contents_table)).mappings().one()
            outcomes = tuple(
                connection.execute(
                    select(processing_import_batch_items_table.c.outcome).order_by(
                        processing_import_batch_items_table.c.source_row_ordinal
                    )
                ).scalars()
            )
        assert content["title"] == "爱玛稳定首行"
        assert outcomes[0] == "created"
        assert outcomes[1:] == ("duplicate",) * 100
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_historical_failed_chunk_range_is_included_in_campaign_accounting(
    tmp_path: Path,
    monkeypatch,
) -> None:
    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "failed-range.xlsx").write_bytes(_xlsx_with_cross_chunk_duplicates(rows=101))
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-failed-range-{uuid4()}",
                "relative_paths": ["failed-range.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        campaign_id = created.json()["campaign_id"]
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-failed-range-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker) == 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )
        assert worker.run_once() is True

        def fail_chunk_read(*args, **kwargs):
            raise ValueError("simulated structural chunk failure")

        monkeypatch.setattr(
            historical_worker_module,
            "read_historical_chunk",
            fail_chunk_read,
        )
        assert worker.run_once() is True

        campaign = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert campaign["status"] == "partial_failed"
        assert campaign["stats"]["created"] == 1
        assert campaign["stats"]["duplicate"] == 99
        assert campaign["stats"]["failed"] == 1
        assert sum(campaign["stats"].values()) == campaign["total_rows"] == 101
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_historical_failed_retry_preserves_cross_chunk_duplicate_identity(
    tmp_path: Path,
    monkeypatch,
) -> None:
    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "retry-duplicate.xlsx").write_bytes(
        _xlsx_with_cross_chunk_duplicates(rows=101)
    )
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    original_read_chunk = historical_worker_module.read_historical_chunk
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-retry-duplicate-{uuid4()}",
                "relative_paths": ["retry-duplicate.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        campaign_id = created.json()["campaign_id"]
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-retry-duplicate-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker) == 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )
        assert worker.run_once() is True

        def fail_chunk_read(*args, **kwargs):
            raise ValueError("simulated structural chunk failure")

        monkeypatch.setattr(
            historical_worker_module,
            "read_historical_chunk",
            fail_chunk_read,
        )
        assert worker.run_once() is True
        monkeypatch.setattr(
            historical_worker_module,
            "read_historical_chunk",
            original_read_chunk,
        )

        retried = client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/retry-failed")
        assert retried.status_code == 200
        with runtime.database.engine.begin() as connection:
            batches = tuple(
                connection.execute(
                    select(processing_import_batches_table.c.id).order_by(
                        processing_import_batches_table.c.created_at
                    )
                ).scalars()
            )
            retry_identity_count = connection.scalar(
                select(func.count())
                .select_from(processing_import_batch_identities_table)
                .where(processing_import_batch_identities_table.c.batch_id == batches[1])
            )
        assert len(batches) == 2
        assert retry_identity_count == 1

        assert worker.run_once() is True
        campaign = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert campaign["status"] == "succeeded"
        assert campaign["stats"]["created"] == 1
        assert campaign["stats"]["duplicate"] == 100
        assert sum(campaign["stats"].values()) == campaign["total_rows"] == 101
        with runtime.database.engine.begin() as connection:
            retry_outcomes = tuple(
                connection.execute(
                    select(processing_import_batch_items_table.c.outcome).where(
                        processing_import_batch_items_table.c.batch_id == batches[1]
                    )
                ).scalars()
            )
        assert retry_outcomes == ("duplicate",)
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_historical_queued_cancel_reaches_terminal_without_business_writes(
    tmp_path: Path,
) -> None:
    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "cancel.xlsx").write_bytes(_xlsx_with_cross_chunk_duplicates(rows=101))
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 2,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-cancel-{uuid4()}",
                "relative_paths": ["cancel.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        campaign_id = created.json()["campaign_id"]
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-cancel-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker) == 2
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )

        cancelled = client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/cancel")
        assert cancelled.status_code == 200
        assert cancelled.json()["status"] == "cancelled"
        assert cancelled.json()["stats"]["failed"] == 101
        assert sum(cancelled.json()["stats"].values()) == 101
        items = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}/items").json()[
            "items"
        ]
        assert all(item["status"] == "cancelled" for item in items)
        with runtime.database.engine.begin() as connection:
            content_count = connection.scalar(select(func.count()).select_from(contents_table))
            ledger_count = connection.scalar(
                select(func.count()).select_from(processing_import_batch_items_table)
            )
            batch = (
                connection.execute(
                    select(
                        processing_import_batches_table.c.status,
                        processing_import_batches_table.c.stats,
                    )
                )
                .mappings()
                .one()
            )
        assert content_count == 0
        assert ledger_count == 0
        assert batch["status"] == "failed"
        assert batch["stats"]["failed"] == 101
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def test_historical_chunk_resumes_after_lease_takeover_without_duplicate_outcome(
    tmp_path: Path,
) -> None:
    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "takeover.xlsx").write_bytes(
        _xlsx(title="爱玛 Lease 接管", text="接管后只写入一次")
    )
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id = _brand(runtime)
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage12-takeover-{uuid4()}",
                "relative_paths": ["takeover.xlsx"],
                "recursive": False,
                "brand_ids": [brand_id],
            },
        )
        assert created.status_code == 202
        campaign_id = created.json()["campaign_id"]
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-preflight-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert _drain(worker) == 2
        assert (
            client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()["status"]
            == "ready"
        )
        assert (
            client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start").status_code
            == 200
        )

        crashed_session = runtime.database.new_session()
        try:
            with crashed_session.begin():
                claimed = PostgresJobRepository(crashed_session).claim_next(
                    supported_job_types=(HISTORICAL_IMPORT_CHUNK_JOB_TYPE,),
                    worker_id="stage12-crashed-worker",
                    lease_seconds=120,
                )
            assert claimed is not None
            assert claimed.lease_token is not None
        finally:
            crashed_session.close()

        with runtime.database.engine.begin() as connection:
            connection.execute(
                update(jobs_table)
                .where(jobs_table.c.id == claimed.id)
                .values(lease_expires_at=func.clock_timestamp() - timedelta(seconds=1))
            )

        takeover_worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage12-takeover-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert takeover_worker.run_once() is True

        completed = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}").json()
        assert completed["status"] == "succeeded"
        assert completed["stats"]["created"] == 1
        with runtime.database.engine.begin() as connection:
            takeover_count = connection.scalar(
                select(jobs_table.c.lease_takeover_count).where(jobs_table.c.id == claimed.id)
            )
            outcome_count = connection.scalar(
                select(func.count()).select_from(processing_import_batch_items_table)
            )
            content_count = connection.scalar(select(func.count()).select_from(contents_table))
        assert takeover_count == 1
        assert outcome_count == 1
        assert content_count == 1
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def _stage3_historical_evidence_catalog(runtime: PlatformRuntime) -> tuple[UUID, UUID]:
    """通过正式 Stage 2 Application Service 建立 Historical 验收目录。"""

    brand = PostgresBrandVehicleHttpService(runtime).create_brand(
        BrandCreateRequest(
            code="AIMA-STAGE3-HISTORICAL-EVIDENCE",
            display_name="爱玛",
            role="owned",
            aliases=("爱玛",),
        ),
        principal=_principal(),
        request_id="stage3-historical-evidence-brand",
    )
    vehicle = PostgresAdministrationHttpService(runtime).create_vehicle_model(
        VehicleModelCreateRequest(
            code="AIMA-STAGE3-HISTORICAL-Q7",
            display_name="Q7",
            brand_id=brand.id,
            aliases=("Q7",),
        ),
        principal=_principal(),
        request_id="stage3-historical-evidence-vehicle",
    )
    return brand.id, vehicle.id


def _stage3_historical_update_alias(runtime: PlatformRuntime, vehicle_id: UUID, alias: str) -> int:
    """通过正式管理员 Service 修改 live alias，并返回新目录版本。"""

    updated = PostgresAdministrationHttpService(runtime).update_vehicle_model(
        vehicle_id,
        VehicleModelUpdateRequest(aliases=(alias,)),
        principal=_principal(),
        request_id=f"stage3-historical-evidence-alias-{alias}",
    )
    return updated.catalog_version


def _stage3_historical_lock_evidence(
    runtime: PlatformRuntime,
    *,
    content_id: UUID,
    content_version: int,
    brand_id: UUID,
    vehicle_id: UUID,
) -> None:
    """在同一 Content Version 上锁定 Brand/Vehicle 人工结论。"""

    session = runtime.database.new_session()
    try:
        with session.begin():
            PostgresBrandVehicleRepository(session).replace_manual_brand_evidence(
                content_id=content_id,
                content_version=content_version,
                brand_ids=(brand_id,),
                unlock_existing=False,
                actor_ref="stage3-historical-evidence-reviewer",
            )
            PostgresVehicleCatalogRepository(session).replace_manual_evidence(
                content_id=content_id,
                content_version=content_version,
                model_ids=(vehicle_id,),
                unlock_existing=False,
                actor_ref="stage3-historical-evidence-reviewer",
            )
    finally:
        session.close()


def _stage3_run_historical_campaign(
    client: TestClient,
    worker,
    *,
    campaign_id: str,
) -> dict[str, object]:  # type: ignore[no-untyped-def]
    """把一个单文件 Historical Campaign 从 discovering 驱动到终态。"""

    for _ in range(10):
        response = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}")
        assert response.status_code == 200
        payload = response.json()
        if payload["status"] == "ready":
            break
        assert payload["status"] in {"discovering", "snapshotting"}
        assert worker.run_once() is True
    else:
        pytest.fail("Historical Campaign 未在测试预算内进入 ready")

    started = client.post(f"/api/v1/historical-import-campaigns/{campaign_id}/start")
    assert started.status_code == 200
    assert started.json()["status"] == "queued"
    for _ in range(10):
        response = client.get(f"/api/v1/historical-import-campaigns/{campaign_id}")
        assert response.status_code == 200
        payload = response.json()
        if payload["status"] in {"succeeded", "partial_failed", "failed", "cancelled"}:
            return payload
        assert worker.run_once() is True
    pytest.fail("Historical Campaign 未在测试预算内进入终态")


def test_historical_stage3_freezes_catalog_and_preserves_manual_evidence(
    tmp_path: Path,
) -> None:
    """Historical 用冻结目录写 Evidence；同版本再次 Campaign 不覆盖人工锁。"""

    historical_root = tmp_path / "approved-history"
    historical_root.mkdir()
    (historical_root / "stage3-evidence.xlsx").write_bytes(
        _xlsx(title="爱玛 Q7 历史冻结目录验证", text="Historical Evidence 固定正文")
    )
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "historical_import_root": historical_root,
            "historical_chunk_rows": 100,
            "historical_max_in_flight_jobs": 1,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
            "accounts RESTART IDENTITY CASCADE"
        )
    try:
        client = TestClient(
            create_app(
                historical_import_service=PostgresHistoricalImportHttpService(runtime),
                import_service=PostgresImportHttpService(runtime),
            )
        )
        brand_id, vehicle_id = _stage3_historical_evidence_catalog(runtime)
        created = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage3-historical-evidence-{uuid4()}",
                "relative_paths": ["stage3-evidence.xlsx"],
                "recursive": False,
                "brand_ids": [str(brand_id)],
            },
        )
        assert created.status_code == 202
        campaign_id = created.json()["campaign_id"]
        with runtime.database.engine.begin() as connection:
            frozen_snapshot = connection.scalar(
                select(historical_import_campaigns_table.c.keyword_pack_snapshot).where(
                    historical_import_campaigns_table.c.id == UUID(campaign_id)
                )
            )
        assert frozen_snapshot is not None
        frozen_catalog_version = frozen_snapshot["catalog"]["catalog_version"]
        assert frozen_snapshot["schema_version"] == "brand-vehicle-filter.v1"

        changed_catalog_version = _stage3_historical_update_alias(runtime, vehicle_id, "Q7CHANGED")
        assert changed_catalog_version > frozen_catalog_version
        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage3-historical-evidence-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        completed = _stage3_run_historical_campaign(
            client,
            worker,
            campaign_id=campaign_id,
        )
        assert completed["status"] == "succeeded"

        with runtime.database.engine.begin() as connection:
            content_id, content_version = connection.execute(
                select(contents_table.c.id, contents_table.c.current_version)
            ).one()
            vehicle_rows = tuple(
                connection.execute(
                    select(content_vehicle_evidence_table).where(
                        content_vehicle_evidence_table.c.content_id == content_id,
                        content_vehicle_evidence_table.c.content_version == content_version,
                        content_vehicle_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
            brand_rows = tuple(
                connection.execute(
                    select(content_brand_evidence_table).where(
                        content_brand_evidence_table.c.content_id == content_id,
                        content_brand_evidence_table.c.content_version == content_version,
                        content_brand_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
        assert len(vehicle_rows) == 1
        assert vehicle_rows[0]["vehicle_model_id"] == vehicle_id
        assert vehicle_rows[0]["source"] == "import"
        assert vehicle_rows[0]["matched_text"] == "Q7"
        assert vehicle_rows[0]["catalog_version"] == frozen_catalog_version
        assert len(brand_rows) == 1
        assert brand_rows[0]["brand_id"] == brand_id
        assert brand_rows[0]["source"] == "vehicle_match"
        assert brand_rows[0]["derived_vehicle_model_id"] == vehicle_id
        assert brand_rows[0]["catalog_version"] == frozen_catalog_version

        _stage3_historical_lock_evidence(
            runtime,
            content_id=content_id,
            content_version=content_version,
            brand_id=brand_id,
            vehicle_id=vehicle_id,
        )
        restored_catalog_version = _stage3_historical_update_alias(runtime, vehicle_id, "Q7")
        assert restored_catalog_version > changed_catalog_version

        replay = client.post(
            "/api/v1/historical-import-campaigns",
            json={
                "client_idempotency_key": f"stage3-historical-evidence-replay-{uuid4()}",
                "relative_paths": ["stage3-evidence.xlsx"],
                "recursive": False,
                "brand_ids": [str(brand_id)],
            },
        )
        assert replay.status_code == 202
        replay_completed = _stage3_run_historical_campaign(
            client,
            worker,
            campaign_id=replay.json()["campaign_id"],
        )
        assert replay_completed["status"] == "succeeded"

        with runtime.database.engine.begin() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 1
            assert connection.scalar(select(contents_table.c.current_version)) == content_version
            active_vehicle_rows = tuple(
                connection.execute(
                    select(content_vehicle_evidence_table).where(
                        content_vehicle_evidence_table.c.content_id == content_id,
                        content_vehicle_evidence_table.c.content_version == content_version,
                        content_vehicle_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
            active_brand_rows = tuple(
                connection.execute(
                    select(content_brand_evidence_table).where(
                        content_brand_evidence_table.c.content_id == content_id,
                        content_brand_evidence_table.c.content_version == content_version,
                        content_brand_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
            vehicle_locked = connection.scalar(
                select(content_vehicle_review_locks_table.c.is_locked).where(
                    content_vehicle_review_locks_table.c.content_id == content_id,
                    content_vehicle_review_locks_table.c.content_version == content_version,
                )
            )
            brand_locked = connection.scalar(
                select(content_brand_review_locks_table.c.is_locked).where(
                    content_brand_review_locks_table.c.content_id == content_id,
                    content_brand_review_locks_table.c.content_version == content_version,
                )
            )
        assert [(row["source"], row["is_manual_locked"]) for row in active_vehicle_rows] == [
            ("manual_review", True)
        ]
        assert [(row["source"], row["is_manual_locked"]) for row in active_brand_rows] == [
            ("manual_review", True)
        ]
        assert vehicle_locked is True
        assert brand_locked is True
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, "
                "accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()
