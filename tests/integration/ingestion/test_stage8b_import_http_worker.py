from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from aima_ugc.adapters.persistence.postgres.brand_vehicle import PostgresBrandVehicleRepository
from aima_ugc.adapters.persistence.postgres.collection_targets import PostgresCollectionTargetReader
from aima_ugc.adapters.persistence.postgres.jobs import PostgresJobRepository
from aima_ugc.adapters.persistence.postgres.vehicles import PostgresVehicleCatalogRepository
from aima_ugc.bootstrap.administration_http import PostgresAdministrationHttpService
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.brand_vehicle_http import PostgresBrandVehicleHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.import_worker import PostgresImportJobExecutor
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.administration import (
    VehicleModelCreateRequest,
    VehicleModelUpdateRequest,
)
from aima_ugc.contracts.brand_vehicle import BrandCreateRequest
from aima_ugc.modules.content.tables import (
    content_metric_observations_table,
    content_versions_table,
    contents_table,
)
from aima_ugc.modules.identity import Principal
from aima_ugc.modules.ingestion.import_job import (
    IMPORT_JOB_TYPE,
    ImportJobPayload,
)
from aima_ugc.modules.ingestion.tables import processing_import_batches_table
from aima_ugc.modules.vehicles.tables import (
    content_brand_evidence_table,
    content_brand_review_locks_table,
    content_vehicle_evidence_table,
    content_vehicle_review_locks_table,
)
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import JobExecutionFence, LeaseLostError
from aima_ugc.platform.jobs.tables import jobs_table
from aima_ugc.platform.storage.canonical import CANONICAL_CONTENT_ARTIFACT_KIND
from aima_ugc.platform.storage.tables import artifacts_table, canonical_artifact_links_table
from fastapi.testclient import TestClient
from httpx import Response
from openpyxl import Workbook
from sqlalchemy import func, select, update


class _ExecutionContext:
    def __init__(self, fence: JobExecutionFence) -> None:
        self._fence = fence

    @property
    def fence(self) -> JobExecutionFence:
        return self._fence

    def heartbeat(self, *, progress: int) -> None:
        del progress

    def cancel_requested(self) -> bool:
        return False


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    sheet.append(
        [
            "小红书",
            "爱玛新品发布",
            "与品牌无关的补充正文",
            "官方账号",
            "2026-08-20 10:00:00",
            "https://www.xiaohongshu.com/explore/stage8b-content-1",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _principal() -> Principal:
    return Principal(
        principal_id="stage3-import-admin",
        display_name="Stage3 Import 管理员",
        role="administrator",
        source="development",
    )


def _configure_and_upload(client: TestClient, runtime) -> Response:  # type: ignore[no-untyped-def]
    brand = PostgresBrandVehicleHttpService(runtime).create_brand(
        BrandCreateRequest(
            code="AIMA-STAGE3-IMPORT",
            display_name="爱玛",
            role="owned",
            aliases=("爱玛",),
        ),
        principal=_principal(),
        request_id="stage3-import-brand",
    )
    created = client.post(
        "/api/v1/import-batches",
        files=[
            (
                "file",
                (
                    "stage8b.xlsx",
                    _xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("brand_ids", (None, str(brand.id))),
        ],
    )
    assert created.status_code == 202
    return created


def _truncate(runtime) -> None:  # type: ignore[no-untyped-def]
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, vehicle_brands, accounts "
            "RESTART IDENTITY CASCADE"
        )


def test_http_upload_worker_and_status_query_use_stage3_brand_filter(tmp_path) -> None:
    settings = load_settings().model_copy(
        update={"data_dir": tmp_path / "data", "log_dir": tmp_path / "logs"}
    )
    runtime = create_worker_runtime(settings=settings)
    _truncate(runtime)
    try:
        service = PostgresImportHttpService(runtime)
        client = TestClient(create_app(import_service=service))
        created = _configure_and_upload(client, runtime)

        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage3-import-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert worker.run_once() is True
        assert worker.run_once() is False

        batch = client.get(f"/api/v1/import-batches/{created.json()['batch_id']}")
        job = client.get(f"/api/v1/jobs/{created.json()['job_id']}")
        assert batch.status_code == job.status_code == 200
        assert batch.json()["status"] == batch.json()["stage"] == "succeeded"
        assert batch.json()["stats"]["rows_seen"] == 1
        assert batch.json()["stats"]["rows_ingested"] == 1
        assert job.json()["status"] == "succeeded"
        assert job.json()["attempt"] == 1

        session = runtime.database.new_session()
        try:
            with session.begin():
                supplement_targets = PostgresCollectionTargetReader(session).list_batch_targets(
                    batch_id=UUID(created.json()["batch_id"]),
                    platforms=("xiaohongshu",),
                )
        finally:
            session.close()
        assert len(supplement_targets) == 1
        assert supplement_targets[0].external_content_id == "stage8b-content-1"

        with runtime.database.engine.begin() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 1
            assert connection.scalar(select(func.count()).select_from(content_versions_table)) == 1
            assert (
                connection.scalar(
                    select(func.count()).select_from(content_metric_observations_table)
                )
                == 1
            )
            persisted_batch = (
                connection.execute(select(processing_import_batches_table)).mappings().one()
            )
            persisted_job = connection.execute(select(jobs_table)).mappings().one()
            persisted_artifacts = connection.execute(select(artifacts_table)).mappings().all()
            canonical_link = connection.execute(
                select(canonical_artifact_links_table)
            ).mappings().one()
        assert persisted_batch["job_id"] == persisted_job["id"]
        assert persisted_batch["status"] == "succeeded"
        assert len(persisted_artifacts) == 2
        assert {artifact["kind"] for artifact in persisted_artifacts} == {
            "file-import.raw",
            CANONICAL_CONTENT_ARTIFACT_KIND,
        }
        assert all(artifact["storage_status"] == "linked" for artifact in persisted_artifacts)
        assert canonical_link["processing_import_batch_id"] == persisted_batch["id"]
        snapshot = persisted_batch["stats"]["filter_snapshot"]
        assert snapshot["schema_version"] == "brand-vehicle-filter.v1"
        assert snapshot["search_semantics"] == "not_applicable"
        assert persisted_job["payload"]["filter_snapshot"] == snapshot
        assert persisted_job["job_type"] == IMPORT_JOB_TYPE
        assert "keyword_selection" not in persisted_job["payload"]
    finally:
        _truncate(runtime)
        runtime.close()


def test_unavailable_source_artifact_fails_job_and_batch_without_content(tmp_path) -> None:
    settings = load_settings().model_copy(
        update={"data_dir": tmp_path / "data", "log_dir": tmp_path / "logs"}
    )
    runtime = create_worker_runtime(settings=settings)
    _truncate(runtime)
    try:
        client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        created = _configure_and_upload(client, runtime)
        with runtime.database.engine.begin() as connection:
            connection.execute(update(artifacts_table).values(storage_status="error"))

        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage3-unavailable-artifact",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert worker.run_once() is True
        batch = client.get(f"/api/v1/import-batches/{created.json()['batch_id']}")
        job = client.get(f"/api/v1/jobs/{created.json()['job_id']}")
        assert batch.json()["status"] == batch.json()["stage"] == "failed"
        assert batch.json()["error_summary"] == "invalid_import"
        assert job.json()["status"] == "failed"
        assert job.json()["error_code"] == "invalid_import"
        with runtime.database.engine.begin() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 0
    finally:
        _truncate(runtime)
        runtime.close()


def test_import_retry_after_business_commit_is_fenced_and_does_not_duplicate_content(
    tmp_path: Path,
) -> None:
    settings = load_settings().model_copy(
        update={"data_dir": tmp_path / "data", "log_dir": tmp_path / "logs"}
    )
    runtime = create_worker_runtime(settings=settings)
    _truncate(runtime)
    try:
        client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        created = _configure_and_upload(client, runtime)
        session = runtime.database.new_session()
        try:
            with session.begin():
                first_job = PostgresJobRepository(session).claim_next(
                    supported_job_types=(IMPORT_JOB_TYPE,),
                    worker_id="stage3-first-attempt",
                    lease_seconds=120,
                )
        finally:
            session.close()
        assert first_job is not None and first_job.lease_token is not None
        first_fence = JobExecutionFence(job_id=first_job.id, lease_token=first_job.lease_token)
        first_result = PostgresImportJobExecutor(runtime).execute(
            payload=ImportJobPayload.model_validate(first_job.payload),
            fence=first_fence,
            context=_ExecutionContext(first_fence),
        )
        assert first_result.outcome == "succeeded"

        with runtime.database.engine.begin() as connection:
            connection.execute(
                update(jobs_table)
                .where(jobs_table.c.id == first_job.id)
                .values(lease_expires_at=datetime.now(UTC) - timedelta(seconds=1))
            )
        session = runtime.database.new_session()
        try:
            with session.begin():
                retry_job = PostgresJobRepository(session).claim_next(
                    supported_job_types=(IMPORT_JOB_TYPE,),
                    worker_id="stage3-retry-attempt",
                    lease_seconds=120,
                )
        finally:
            session.close()
        assert retry_job is not None and retry_job.lease_token is not None
        retry_fence = JobExecutionFence(job_id=retry_job.id, lease_token=retry_job.lease_token)
        with pytest.raises(LeaseLostError):
            PostgresImportJobExecutor(runtime).execute(
                payload=ImportJobPayload.model_validate(first_job.payload),
                fence=first_fence,
                context=_ExecutionContext(first_fence),
            )
        retry_result = PostgresImportJobExecutor(runtime).execute(
            payload=ImportJobPayload.model_validate(retry_job.payload),
            fence=retry_fence,
            context=_ExecutionContext(retry_fence),
        )
        assert retry_result.outcome == "succeeded"

        session = runtime.database.new_session()
        try:
            with session.begin():
                PostgresJobRepository(session).succeed(
                    job_id=retry_job.id,
                    lease_token=retry_job.lease_token,
                    result=retry_result.result,
                )
        finally:
            session.close()
        with runtime.database.engine.begin() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 1
            assert connection.scalar(select(func.count()).select_from(content_versions_table)) == 1
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(artifacts_table)
                    .where(artifacts_table.c.kind == CANONICAL_CONTENT_ARTIFACT_KIND)
                )
                == 1
            )
            assert (
                connection.scalar(select(func.count()).select_from(canonical_artifact_links_table))
                == 1
            )
            batch = connection.execute(select(processing_import_batches_table)).mappings().one()
            assert str(batch["id"]) == created.json()["batch_id"]
            assert batch["status"] == "succeeded"
    finally:
        _truncate(runtime)
        runtime.close()


def _stage3_evidence_workbook() -> bytes:
    """构造同时命中品牌与 Q7 车型的 Stage 3 固定输入。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    sheet.append(
        [
            "小红书",
            "爱玛 Q7 冻结目录验证",
            "同一内容用于验证自动证据不会覆盖人工锁。",
            "官方账号",
            "2026-08-20 12:00:00",
            "https://www.xiaohongshu.com/explore/stage3-evidence-content",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _stage3_evidence_catalog(runtime) -> tuple[UUID, UUID]:  # type: ignore[no-untyped-def]
    """通过正式 Stage 2 Application Service 建立爱玛 + Q7 目录。"""

    brand = PostgresBrandVehicleHttpService(runtime).create_brand(
        BrandCreateRequest(
            code="AIMA-STAGE3-EVIDENCE",
            display_name="爱玛",
            role="owned",
            aliases=("爱玛",),
        ),
        principal=_principal(),
        request_id="stage3-evidence-brand",
    )
    vehicle = PostgresAdministrationHttpService(runtime).create_vehicle_model(
        VehicleModelCreateRequest(
            code="AIMA-STAGE3-EVIDENCE-Q7",
            display_name="Q7",
            brand_id=brand.id,
            aliases=("Q7",),
        ),
        principal=_principal(),
        request_id="stage3-evidence-vehicle",
    )
    return brand.id, vehicle.id


def _stage3_upload_evidence_batch(client: TestClient, *, brand_id: UUID) -> Response:
    """从正式公开入口创建一个冻结 Brand/Vehicle Snapshot 的单文件任务。"""

    created = client.post(
        "/api/v1/import-batches",
        files=[
            (
                "file",
                (
                    "stage3-evidence.xlsx",
                    _stage3_evidence_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("brand_ids", (None, str(brand_id))),
        ],
    )
    assert created.status_code == 202
    return created


def _stage3_update_vehicle(
    runtime,
    vehicle_id: UUID,
    alias: str,
    *,
    brand_id: UUID,
) -> int:  # type: ignore[no-untyped-def]
    """通过正式管理员 Service 修改 live alias/归属，并返回新目录版本。"""

    updated = PostgresAdministrationHttpService(runtime).update_vehicle_model(
        vehicle_id,
        VehicleModelUpdateRequest(aliases=(alias,), brand_id=brand_id),
        principal=_principal(),
        request_id=f"stage3-evidence-alias-{alias}",
    )
    return updated.catalog_version


def _stage3_lock_manual_evidence(
    runtime,
    *,
    content_id: UUID,
    content_version: int,
    brand_id: UUID,
    vehicle_id: UUID,
) -> None:  # type: ignore[no-untyped-def]
    """用正式 Owner Repository 对同一 Content Version 建立 Brand/Vehicle 人工锁。"""

    session = runtime.database.new_session()
    try:
        with session.begin():
            PostgresBrandVehicleRepository(session).replace_manual_brand_evidence(
                content_id=content_id,
                content_version=content_version,
                brand_ids=(brand_id,),
                unlock_existing=False,
                actor_ref="stage3-evidence-reviewer",
            )
            PostgresVehicleCatalogRepository(session).replace_manual_evidence(
                content_id=content_id,
                content_version=content_version,
                model_ids=(vehicle_id,),
                unlock_existing=False,
                actor_ref="stage3-evidence-reviewer",
            )
    finally:
        session.close()


def test_stage3_import_freezes_catalog_and_preserves_manual_evidence(tmp_path: Path) -> None:
    """live alias/归属漂移后仍按冻结 Snapshot 写证据，且不覆盖人工锁。"""

    settings = load_settings().model_copy(
        update={"data_dir": tmp_path / "data", "log_dir": tmp_path / "logs"}
    )
    runtime = create_worker_runtime(settings=settings)
    _truncate(runtime)
    try:
        client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        brand_id, vehicle_id = _stage3_evidence_catalog(runtime)
        created = _stage3_upload_evidence_batch(client, brand_id=brand_id)
        with runtime.database.engine.begin() as connection:
            payload = connection.scalar(
                select(jobs_table.c.payload).where(
                    jobs_table.c.id == UUID(created.json()["job_id"])
                )
            )
        assert payload is not None
        frozen_snapshot = payload["filter_snapshot"]
        frozen_catalog_version = frozen_snapshot["catalog"]["catalog_version"]
        assert frozen_snapshot["schema_version"] == "brand-vehicle-filter.v1"

        competitor = PostgresBrandVehicleHttpService(runtime).create_brand(
            BrandCreateRequest(
                code="COMPETITOR-STAGE3-EVIDENCE",
                display_name="竞品",
                role="competitor",
                aliases=("竞品",),
            ),
            principal=_principal(),
            request_id="stage3-evidence-competitor",
        )
        changed_catalog_version = _stage3_update_vehicle(
            runtime,
            vehicle_id,
            "Q7CHANGED",
            brand_id=competitor.id,
        )
        assert changed_catalog_version > frozen_catalog_version

        worker = create_job_worker(
            runtime=runtime,
            registry=create_collection_job_registry(runtime=runtime),
            worker_id="stage3-evidence-worker",
            lease_seconds=120,
            retry_delay_seconds=0,
        )
        assert worker.run_once() is True
        batch = client.get(f"/api/v1/import-batches/{created.json()['batch_id']}")
        assert batch.status_code == 200
        assert batch.json()["status"] == "succeeded"

        with runtime.database.engine.begin() as connection:
            content_id, content_version = connection.execute(
                select(contents_table.c.id, contents_table.c.current_version)
            ).one()
            vehicle_rows = tuple(
                connection.execute(
                    select(content_vehicle_evidence_table).where(
                        content_vehicle_evidence_table.c.content_id == content_id,
                        content_vehicle_evidence_table.c.content_version == content_version,
                        content_vehicle_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
            brand_rows = tuple(
                connection.execute(
                    select(content_brand_evidence_table).where(
                        content_brand_evidence_table.c.content_id == content_id,
                        content_brand_evidence_table.c.content_version == content_version,
                        content_brand_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
        assert len(vehicle_rows) == 1
        assert vehicle_rows[0]["vehicle_model_id"] == vehicle_id
        assert vehicle_rows[0]["source"] == "import"
        assert vehicle_rows[0]["matched_text"] == "Q7"
        assert vehicle_rows[0]["catalog_version"] == frozen_catalog_version
        assert len(brand_rows) == 1
        assert brand_rows[0]["brand_id"] == brand_id
        assert brand_rows[0]["source"] == "vehicle_match"
        assert brand_rows[0]["derived_vehicle_model_id"] == vehicle_id
        assert brand_rows[0]["catalog_version"] == frozen_catalog_version

        _stage3_lock_manual_evidence(
            runtime,
            content_id=content_id,
            content_version=content_version,
            brand_id=brand_id,
            vehicle_id=vehicle_id,
        )
        restored_catalog_version = _stage3_update_vehicle(
            runtime,
            vehicle_id,
            "Q7",
            brand_id=brand_id,
        )
        assert restored_catalog_version > changed_catalog_version

        replay = _stage3_upload_evidence_batch(client, brand_id=brand_id)
        assert worker.run_once() is True
        replay_batch = client.get(f"/api/v1/import-batches/{replay.json()['batch_id']}")
        assert replay_batch.status_code == 200
        assert replay_batch.json()["status"] == "succeeded"

        with runtime.database.engine.begin() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 1
            assert connection.scalar(select(contents_table.c.current_version)) == content_version
            active_vehicle_rows = tuple(
                connection.execute(
                    select(content_vehicle_evidence_table).where(
                        content_vehicle_evidence_table.c.content_id == content_id,
                        content_vehicle_evidence_table.c.content_version == content_version,
                        content_vehicle_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
            active_brand_rows = tuple(
                connection.execute(
                    select(content_brand_evidence_table).where(
                        content_brand_evidence_table.c.content_id == content_id,
                        content_brand_evidence_table.c.content_version == content_version,
                        content_brand_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
            vehicle_locked = connection.scalar(
                select(content_vehicle_review_locks_table.c.is_locked).where(
                    content_vehicle_review_locks_table.c.content_id == content_id,
                    content_vehicle_review_locks_table.c.content_version == content_version,
                )
            )
            brand_locked = connection.scalar(
                select(content_brand_review_locks_table.c.is_locked).where(
                    content_brand_review_locks_table.c.content_id == content_id,
                    content_brand_review_locks_table.c.content_version == content_version,
                )
            )
        assert [(row["source"], row["is_manual_locked"]) for row in active_vehicle_rows] == [
            ("manual_review", True)
        ]
        assert [(row["source"], row["is_manual_locked"]) for row in active_brand_rows] == [
            ("manual_review", True)
        ]
        assert vehicle_locked is True
        assert brand_locked is True
    finally:
        _truncate(runtime)
        runtime.close()
