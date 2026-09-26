from __future__ import annotations

import logging
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from threading import Barrier
from typing import cast
from uuid import UUID, uuid4

import pytest
from aima_ugc.adapters.persistence.postgres.brand_vehicle import PostgresBrandVehicleRepository
from aima_ugc.adapters.persistence.postgres.canonical_replay import (
    PostgresCanonicalReplayRepository,
)
from aima_ugc.adapters.persistence.postgres.content_complete import (
    PostgresCompleteContentRepository,
)
from aima_ugc.adapters.persistence.postgres.content_visibility import (
    content_has_active_source,
)
from aima_ugc.adapters.persistence.postgres.jobs import PostgresJobRepository
from aima_ugc.adapters.persistence.postgres.vehicles import PostgresVehicleCatalogRepository
from aima_ugc.bootstrap import canonical_replay_worker as canonical_replay_worker_module
from aima_ugc.bootstrap.administration_http import PostgresAdministrationHttpService
from aima_ugc.bootstrap.api import create_app
from aima_ugc.bootstrap.brand_vehicle_http import PostgresBrandVehicleHttpService
from aima_ugc.bootstrap.canonical_replay_http import PostgresCanonicalReplayHttpService
from aima_ugc.bootstrap.canonical_replay_reversal_worker import (
    PostgresCanonicalReplayReversalJobExecutor,
)
from aima_ugc.bootstrap.canonical_replay_worker import PostgresCanonicalReplayJobExecutor
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.runtime import PlatformRuntime
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.administration import VehicleModelCreateRequest
from aima_ugc.contracts.brand_vehicle import BrandAliasCreateRequest, BrandCreateRequest
from aima_ugc.contracts.canonical import CanonicalAuthorV1, CanonicalContentV1
from aima_ugc.contracts.http import CanonicalReplayCreateRequest
from aima_ugc.modules.content.contribution_tables import content_source_contributions_table
from aima_ugc.modules.content.extended_tables import content_external_ids_table
from aima_ugc.modules.content.read_model_tables import (
    voice_plaza_content_projection_table,
    voice_plaza_filter_catalog_entries_table,
)
from aima_ugc.modules.content.tables import (
    content_metric_observations_table,
    content_versions_table,
    contents_table,
)
from aima_ugc.modules.identity import Principal
from aima_ugc.modules.ingestion.canonical_replay import (
    CANONICAL_REPLAY_JOB_TYPE,
    CanonicalReplayCounters,
    CanonicalReplayJobPayload,
)
from aima_ugc.modules.ingestion.canonical_replay_tables import (
    canonical_replay_all_requests_table,
    canonical_replay_content_changes_table,
    canonical_replay_runs_table,
    canonical_replay_seen_content_table,
    canonical_replay_validation_proofs_table,
)
from aima_ugc.modules.ingestion.replay_shard_tables import canonical_replay_run_shards_table
from aima_ugc.modules.ingestion.reversal_shard_tables import reversal_shards_table
from aima_ugc.modules.ingestion.tables import processing_import_batches_table
from aima_ugc.modules.system.tables import audit_events_table
from aima_ugc.modules.vehicles.models import ContentVehicleEvidence
from aima_ugc.modules.vehicles.tables import (
    content_brand_evidence_table,
    content_brand_review_locks_table,
    content_vehicle_evidence_table,
)
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import JobExecutionFence, LeaseLostError
from aima_ugc.platform.jobs.tables import jobs_table
from aima_ugc.platform.storage.tables import artifacts_table, canonical_artifact_links_table
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import event, func, select, text, update


class _ExecutionContext:
    def __init__(
        self,
        fence: JobExecutionFence,
        *,
        lose_after_heartbeat: bool = False,
        cancel_after_checks: int | None = None,
    ) -> None:
        self.fence = fence
        self._lose_after_heartbeat = lose_after_heartbeat
        self._cancel_after_checks = cancel_after_checks
        self._cancel_checks = 0

    def heartbeat(self, *, progress=None):  # type: ignore[no-untyped-def]
        del progress
        if self._lose_after_heartbeat:
            raise LeaseLostError("模拟首批提交后的 Worker lease 丢失")

    def cancel_requested(self) -> bool:
        self._cancel_checks += 1
        return (
            self._cancel_after_checks is not None
            and self._cancel_checks > self._cancel_after_checks
        )


def _principal() -> Principal:
    return Principal(
        principal_id="canonical-replay-integration",
        display_name="Canonical Replay 集成测试",
        role="administrator",
        source="development",
    )


def _xlsx(*, rows: tuple[tuple[str, str], ...]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    for external_id, title in rows:
        sheet.append(
            [
                "小红书",
                title,
                "Replay 固定测试正文",
                "测试账号",
                "2026-09-11 08:00:00",
                f"https://www.xiaohongshu.com/explore/{external_id}",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _truncate(runtime: PlatformRuntime) -> None:
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE audit_events, jobs, artifacts, keyword_packs, vehicle_brands, accounts "
            "RESTART IDENTITY CASCADE"
        )


def _runtime(tmp_path: Path) -> PlatformRuntime:
    settings = load_settings().model_copy(
        update={"data_dir": tmp_path / "data", "log_dir": tmp_path / "logs"}
    )
    return create_worker_runtime(settings=settings)


def _worker(runtime: PlatformRuntime, *, suffix: str):  # type: ignore[no-untyped-def]
    return create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id=f"canonical-replay-{suffix}",
        lease_seconds=120,
        retry_delay_seconds=0,
    )


def _create_brand_without_matching_alias(runtime: PlatformRuntime) -> UUID:
    return _create_brand(runtime, alias="绝不命中的旧识别词")


def _create_brand(runtime: PlatformRuntime, *, alias: str) -> UUID:
    created = PostgresBrandVehicleHttpService(runtime).create_brand(
        BrandCreateRequest(
            display_name="Replay 测试品牌",
            role="owned",
            aliases=(alias,),
        ),
        principal=_principal(),
        request_id="canonical-replay-brand",
    )
    return created.id


def _add_replay_alias(
    runtime: PlatformRuntime,
    brand_id: UUID,
    *,
    text: str = "星曜",
) -> int:
    updated = PostgresBrandVehicleHttpService(runtime).add_alias(
        brand_id,
        BrandAliasCreateRequest(text=text),
        principal=_principal(),
        request_id="canonical-replay-alias",
    )
    return updated.catalog_version


def _create_replay_vehicle(runtime: PlatformRuntime, *, brand_id: UUID) -> UUID:
    created = PostgresAdministrationHttpService(runtime).create_vehicle_model(
        VehicleModelCreateRequest(
            display_name="Replay 星曜车型",
            brand_id=brand_id,
            aliases=("星曜",),
        ),
        principal=_principal(),
        request_id="canonical-replay-vehicle",
    )
    return created.id


def _import_canonical(
    client: TestClient,
    runtime: PlatformRuntime,
    *,
    filename: str,
    rows: tuple[tuple[str, str], ...],
    brand_ids: tuple[UUID, ...],
    expected_rows_ingested: int = 0,
) -> UUID:
    files: list[tuple[str, tuple[str | None, object, str | None]]] = [
        (
            "file",
            (
                filename,
                _xlsx(rows=rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        )
    ]
    files.extend(("brand_ids", (None, str(brand_id), None)) for brand_id in brand_ids)
    created = client.post(
        "/api/v1/import-batches",
        files=files,
    )
    assert created.status_code == 202
    assert _worker(runtime, suffix=f"import-{filename}").run_once() is True
    batch = client.get(f"/api/v1/import-batches/{created.json()['batch_id']}")
    assert batch.status_code == 200
    assert batch.json()["status"] == "succeeded"
    assert batch.json()["stats"]["rows_ingested"] == expected_rows_ingested
    with runtime.database.engine.connect() as connection:
        artifact_id = connection.scalar(
            select(canonical_artifact_links_table.c.artifact_id).where(
                canonical_artifact_links_table.c.processing_import_batch_id
                == UUID(created.json()["batch_id"])
            )
        )
    assert isinstance(artifact_id, UUID)
    return artifact_id


def _create_replay(
    client: TestClient,
    *,
    artifact_ids: tuple[UUID, ...],
    brand_id: UUID,
    idempotency_key: str,
    batch_size: int = 1,
) -> dict[str, object]:
    response = client.post(
        "/api/v1/canonical-replays",
        json={
            "idempotency_key": idempotency_key,
            "artifact_ids": [str(item) for item in artifact_ids],
            "brand_ids": [str(brand_id)],
            "batch_size": batch_size,
        },
    )
    assert response.status_code == 202
    return response.json()


def _create_all_replay(client: TestClient, *, idempotency_key: str) -> dict[str, object]:
    response = client.post(
        "/api/v1/canonical-replays/all",
        json={"idempotency_key": idempotency_key},
    )
    assert response.status_code == 202
    return response.json()


def test_large_replay_reversal_uses_durable_content_shards_and_completion_barrier(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """同一次撤回由正式 jobs/两个 Worker 结清，父任务不能先于子任务完成。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        row_count = 10_000
        rows = tuple(
            (f"durable-reversal-{uuid4()}-{index}", "星曜重筛") for index in range(row_count)
        )
        _import_canonical(
            client, runtime, filename="durable-reversal.xlsx", rows=rows, brand_ids=(brand_id,)
        )
        _add_replay_alias(runtime, brand_id)
        created = _create_all_replay(client, idempotency_key=f"durable-reversal-{uuid4()}")
        assert _worker(runtime, suffix="durable-ingest").run_once()
        request_id = UUID(str(created["request_id"]))
        queued = client.post(f"/api/v1/canonical-replays/all/{request_id}/revoke")
        assert queued.status_code == 202
        with runtime.database.engine.connect() as connection:
            parent_id = connection.scalar(
                select(canonical_replay_all_requests_table.c.reversal_job_id).where(
                    canonical_replay_all_requests_table.c.id == request_id
                )
            )
        assert isinstance(parent_id, UUID)
        original_process = PostgresCanonicalReplayReversalJobExecutor.process_shard
        fault_injected = False

        def fail_after_child_commit(self, shard_id, *, fence, context):  # type: ignore[no-untyped-def]
            nonlocal fault_injected
            processed = original_process(self, shard_id, fence=fence, context=context)
            if fence.job_id != parent_id and not fault_injected:
                with runtime.database.engine.connect() as connection:
                    assert (
                        connection.scalar(
                            select(canonical_replay_all_requests_table.c.lifecycle_status).where(
                                canonical_replay_all_requests_table.c.id == request_id
                            )
                        )
                        == "reverting"
                    )
                    assert (
                        connection.scalar(
                            select(func.count())
                            .select_from(voice_plaza_content_projection_table)
                            .where(voice_plaza_content_projection_table.c.is_visible.is_(False))
                        )
                        >= processed
                    )
                fault_injected = True
                raise ValueError("模拟分片业务已提交但子 Job 终态失败")
            return processed

        monkeypatch.setattr(
            PostgresCanonicalReplayReversalJobExecutor, "process_shard", fail_after_child_commit
        )
        parent = _worker(runtime, suffix="durable-parent")
        child = _worker(runtime, suffix="durable-child")
        with ThreadPoolExecutor(max_workers=1) as pool:
            parent_result = pool.submit(parent.run_once)
            deadline = time.monotonic() + 120
            while time.monotonic() < deadline:
                with runtime.database.engine.connect() as connection:
                    status = connection.scalar(
                        select(jobs_table.c.status).where(jobs_table.c.id == parent_id)
                    )
                if status == "running":
                    break
                time.sleep(0.01)
            else:
                pytest.fail("撤回父 Job 未进入运行状态")
            while not parent_result.done() and time.monotonic() < deadline:
                if not child.run_once():
                    time.sleep(0.02)
            assert parent_result.result(timeout=1) is True
        with runtime.database.engine.connect() as connection:
            units = (
                connection.execute(
                    select(reversal_shards_table).where(
                        reversal_shards_table.c.replay_request_id == request_id
                    )
                )
                .mappings()
                .all()
            )
            request = (
                connection.execute(
                    select(canonical_replay_all_requests_table).where(
                        canonical_replay_all_requests_table.c.id == request_id
                    )
                )
                .mappings()
                .one()
            )
            remaining = connection.scalar(
                select(func.count())
                .select_from(canonical_replay_content_changes_table)
                .where(
                    canonical_replay_content_changes_table.c.all_request_id == request_id,
                    canonical_replay_content_changes_table.c.reverted_at.is_(None),
                )
            )
            child_jobs = connection.scalar(
                select(func.count())
                .select_from(jobs_table)
                .where(jobs_table.c.job_type == "ingestion.reversal-shard.v1")
            )
        assert len(units) == 4
        assert all(unit["status"] == "succeeded" for unit in units)
        assert sum(unit["processed_content_count"] for unit in units) == row_count
        assert child_jobs >= 1
        assert fault_injected
        assert request["lifecycle_status"] == "reverted"
        assert request["reverted_content_count"] == row_count
        assert remaining == 0
    finally:
        _truncate(runtime)
        runtime.close()


def test_replay_reversal_failed_parent_retries_from_persisted_shard_checkpoint(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """用户重试失败撤回时，新父 Job 接管旧片断点与已结清账本。"""

    monkeypatch.setattr(
        "aima_ugc.bootstrap.canonical_replay_reversal_worker.REVERSAL_MIN_PARALLEL_CONTENTS",
        1,
    )
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        row_count = 3_000
        _import_canonical(
            client,
            runtime,
            filename="reversal-retry.xlsx",
            rows=tuple(
                (f"reversal-retry-{uuid4()}-{index}", "星曜重筛") for index in range(row_count)
            ),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        created = _create_all_replay(client, idempotency_key=f"reversal-retry-{uuid4()}")
        assert _worker(runtime, suffix="reversal-retry-ingest").run_once()
        request_id = UUID(str(created["request_id"]))
        assert client.post(f"/api/v1/canonical-replays/all/{request_id}/revoke").status_code == 202
        with runtime.database.engine.connect() as connection:
            old_job_id = connection.scalar(
                select(canonical_replay_all_requests_table.c.reversal_job_id).where(
                    canonical_replay_all_requests_table.c.id == request_id
                )
            )
        assert isinstance(old_job_id, UUID)
        original_process = PostgresCanonicalReplayReversalJobExecutor.process_shard
        fault_injected = False

        def fail_parent_after_commit(self, shard_id, *, fence, context):  # type: ignore[no-untyped-def]
            nonlocal fault_injected
            processed = original_process(self, shard_id, fence=fence, context=context)
            if fence.job_id == old_job_id and not fault_injected:
                fault_injected = True
                raise ValueError("模拟撤回父 Job 首片提交后终态失败")
            return processed

        monkeypatch.setattr(
            PostgresCanonicalReplayReversalJobExecutor, "process_shard", fail_parent_after_commit
        )
        worker = _worker(runtime, suffix="reversal-retry-parent")
        assert worker.run_once()
        assert fault_injected
        _worker(runtime, suffix="reversal-retry-old-child").run_once()
        retried = client.post(f"/api/v1/canonical-replays/all/{request_id}/revoke")
        assert retried.status_code == 202, retried.text
        with runtime.database.engine.connect() as connection:
            new_job_id = connection.scalar(
                select(canonical_replay_all_requests_table.c.reversal_job_id).where(
                    canonical_replay_all_requests_table.c.id == request_id
                )
            )
        assert isinstance(new_job_id, UUID) and new_job_id != old_job_id
        assert worker.run_once()
        with runtime.database.engine.connect() as connection:
            request = (
                connection.execute(
                    select(canonical_replay_all_requests_table).where(
                        canonical_replay_all_requests_table.c.id == request_id
                    )
                )
                .mappings()
                .one()
            )
            shards = (
                connection.execute(
                    select(reversal_shards_table).where(
                        reversal_shards_table.c.replay_request_id == request_id
                    )
                )
                .mappings()
                .all()
            )
            remaining = connection.scalar(
                select(func.count())
                .select_from(canonical_replay_content_changes_table)
                .where(
                    canonical_replay_content_changes_table.c.all_request_id == request_id,
                    canonical_replay_content_changes_table.c.reverted_at.is_(None),
                )
            )
        assert request["lifecycle_status"] == "reverted"
        assert request["reverted_content_count"] == row_count
        assert sum(int(shard["processed_content_count"]) for shard in shards) == row_count
        assert all(shard["status"] == "succeeded" for shard in shards)
        assert remaining == 0
    finally:
        _truncate(runtime)
        runtime.close()


def test_replay_run_shards_preserve_cross_artifact_deduplication_and_row_counts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """同一身份跨 Artifact 重复时，分片只能让原始顺序的首条入库。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        shared = f"sharded-shared-{uuid4()}"
        first_artifact = _import_canonical(
            client,
            runtime,
            filename="shard-first.xlsx",
            rows=((shared, "星曜首条"), (f"sharded-first-{uuid4()}", "星曜独立一")),
            brand_ids=(brand_id,),
        )
        second_artifact = _import_canonical(
            client,
            runtime,
            filename="shard-second.xlsx",
            rows=((shared, "星曜后条"), (f"sharded-second-{uuid4()}", "星曜独立二")),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        monkeypatch.setattr(PostgresCanonicalReplayJobExecutor, "_shard_count", lambda *_: 4)
        created = _create_replay(
            client,
            artifact_ids=(first_artifact, second_artifact),
            brand_id=brand_id,
            idempotency_key=f"shard-dedup-{uuid4()}",
            batch_size=2,
        )
        run_id = UUID(str(created["run_id"]))
        parent = _worker(runtime, suffix="shard-ingest-parent")
        child = _worker(runtime, suffix="shard-ingest-child")
        with ThreadPoolExecutor(max_workers=1) as pool:
            parent_result = pool.submit(parent.run_once)
            deadline = time.monotonic() + 30
            while time.monotonic() < deadline:
                with runtime.database.engine.connect() as connection:
                    shard_count = connection.scalar(
                        select(func.count())
                        .select_from(canonical_replay_run_shards_table)
                        .where(canonical_replay_run_shards_table.c.run_id == run_id)
                    )
                if shard_count == 4:
                    break
                time.sleep(0.01)
            else:
                pytest.fail("Replay 子工作单元未创建")
            while not parent_result.done() and time.monotonic() < deadline:
                if not child.run_once():
                    time.sleep(0.02)
            assert parent_result.result(timeout=1) is True
        result = client.get(f"/api/v1/canonical-replays/{run_id}")
        assert result.status_code == 200
        assert result.json()["job"]["status"] == "succeeded"
        assert result.json()["stats"] == {
            "rows_seen": 4,
            "rows_matched": 4,
            "rows_filtered_out": 0,
            "duplicates_removed": 1,
            "rows_ingested": 3,
            "existing_convergence": 0,
            "invalid_artifact_rows": 0,
        }
        with runtime.database.engine.connect() as connection:
            saved = (
                connection.execute(
                    select(contents_table).where(contents_table.c.external_content_id == shared)
                )
                .mappings()
                .one()
            )
            shards = (
                connection.execute(
                    select(canonical_replay_run_shards_table).where(
                        canonical_replay_run_shards_table.c.run_id == run_id
                    )
                )
                .mappings()
                .all()
            )
        assert saved["title"] == "星曜首条"
        assert len(shards) == 4
        assert all(item["status"] == "succeeded" for item in shards)
        assert sum(item["rows_seen"] for item in shards) == 4
    finally:
        _truncate(runtime)
        runtime.close()


@pytest.mark.parametrize("batch_size", [1, 2])
def test_new_alias_replay_deduplicates_and_converges_through_content_owner(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    batch_size: int,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-duplicate.xlsx",
            rows=(
                ("canonical-replay-same", "星曜首条"),
                ("canonical-replay-same", "星曜重复记录"),
            ),
            brand_ids=(brand_id,),
        )
        catalog_version = _add_replay_alias(runtime, brand_id)

        def reject_row_claim(*args: object, **kwargs: object) -> None:
            raise AssertionError("Replay 不应逐行声明 Content identity")

        monkeypatch.setattr(
            PostgresCanonicalReplayRepository,
            "claim_content_identity",
            reject_row_claim,
        )

        first = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-first-{uuid4()}",
            batch_size=batch_size,
        )
        assert _worker(runtime, suffix="first").run_once() is True
        first_run = client.get(f"/api/v1/canonical-replays/{first['run_id']}")

        assert first_run.status_code == 200
        assert first_run.json()["job"]["status"] == "succeeded"
        assert first_run.json()["catalog_version"] == catalog_version
        assert first_run.json()["stats"] == {
            "rows_seen": 2,
            "rows_matched": 2,
            "rows_filtered_out": 0,
            "duplicates_removed": 1,
            "rows_ingested": 1,
            "existing_convergence": 0,
            "invalid_artifact_rows": 0,
        }

        replay_again = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-existing-{uuid4()}",
            batch_size=batch_size,
        )
        assert _worker(runtime, suffix="existing").run_once() is True
        second_run = client.get(f"/api/v1/canonical-replays/{replay_again['run_id']}")

        assert second_run.json()["stats"] == {
            "rows_seen": 2,
            "rows_matched": 2,
            "rows_filtered_out": 0,
            "duplicates_removed": 1,
            "rows_ingested": 0,
            "existing_convergence": 1,
            "invalid_artifact_rows": 0,
        }
        with runtime.database.engine.connect() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 1
            assert (
                connection.scalar(
                    select(func.count()).select_from(canonical_replay_seen_content_table)
                )
                == 2
            )
            evidence = connection.execute(select(content_brand_evidence_table)).mappings().all()
        assert {row["brand_id"] for row in evidence} == {brand_id}
        assert {row["catalog_version"] for row in evidence} == {catalog_version}
    finally:
        _truncate(runtime)
        runtime.close()


def test_replay_batches_stable_authors_without_scalar_content_writes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """真实 Replay 循环中的稳定作者进入集合路径，并输出可核验的回退计数。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-stable-authors.xlsx",
            rows=tuple(
                (f"replay-stable-author-{index}", f"星曜稳定作者 {index}") for index in range(3)
            ),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)

        original_lineage = PostgresCanonicalReplayJobExecutor._content_with_lineage

        def inject_stable_author(
            session: object,
            content: CanonicalContentV1,
            **kwargs: object,
        ) -> CanonicalContentV1:
            mapped = original_lineage(session, content, **kwargs)  # type: ignore[arg-type]
            index = mapped.external_content_id.rsplit("-", 1)[-1]
            return mapped.model_copy(
                update={
                    "author": CanonicalAuthorV1(
                        external_account_id=f"stable-account-{index}",
                        alternate_ids={"red_id": f"stable-red-{index}"},
                        display_name=f"稳定作者 {index}",
                    ),
                    "observed_fields": [
                        *mapped.observed_fields,
                        "author.external_account_id",
                        "author.alternate_ids",
                        "author.display_name",
                    ],
                }
            )

        def reject_scalar_content_write(*_args: object, **_kwargs: object) -> None:
            raise AssertionError("稳定作者 Replay 不应进入逐行 Content 写入")

        monkeypatch.setattr(
            PostgresCanonicalReplayJobExecutor,
            "_content_with_lineage",
            staticmethod(inject_stable_author),
        )
        monkeypatch.setattr(
            PostgresCompleteContentRepository,
            "_ingest_content",
            reject_scalar_content_write,
        )
        runtime.logger.setLevel(logging.DEBUG)
        for handler in runtime.logger.handlers:
            handler.setLevel(logging.DEBUG)

        created = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-stable-authors-{uuid4()}",
            batch_size=100,
        )
        assert _worker(runtime, suffix="stable-authors").run_once() is True
        detail = client.get(f"/api/v1/canonical-replays/{created['run_id']}")
        assert detail.status_code == 200
        assert detail.json()["stats"]["rows_ingested"] == 3

        worker_log = (runtime.settings.log_dir / "worker.log").read_text(encoding="utf-8")
        completed_line = next(
            line
            for line in worker_log.splitlines()
            if "event=canonical_replay.batch_completed" in line
        )
        assert "stable_author_count=3" in completed_line
        assert "batched_remainder_count=3" in completed_line
        assert "scalar_fallback_count=0" in completed_line
    finally:
        _truncate(runtime)
        runtime.close()


def test_all_replay_revoke_hides_replay_only_content_and_preserves_history(
    tmp_path: Path,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        _import_canonical(
            client,
            runtime,
            filename="replay-reversible.xlsx",
            rows=(("canonical-replay-reversible", "星曜仅由重筛入库"),),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)

        created = _create_all_replay(
            client,
            idempotency_key=f"all-replay-reversible-{uuid4()}",
        )
        request_id = UUID(cast(str, created["request_id"]))
        assert _worker(runtime, suffix="reversible-ingest").run_once() is True

        with runtime.database.engine.connect() as connection:
            content = (
                connection.execute(
                    select(contents_table).where(
                        contents_table.c.external_content_id == "canonical-replay-reversible"
                    )
                )
                .mappings()
                .one()
            )
            content_id = cast(UUID, content["id"])
            assert content["replay_visibility_owner_id"] == request_id
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(canonical_replay_content_changes_table)
                    .where(canonical_replay_content_changes_table.c.all_request_id == request_id)
                )
                == 1
            )
            assert (
                connection.scalar(
                    select(content_has_active_source(contents_table.c.id)).where(
                        contents_table.c.id == content_id
                    )
                )
                is True
            )
            assert (
                connection.scalar(
                    select(voice_plaza_content_projection_table.c.is_visible).where(
                        voice_plaza_content_projection_table.c.content_id == content_id
                    )
                )
                is True
            )

        requested = client.post(f"/api/v1/canonical-replays/all/{request_id}/revoke")
        assert requested.status_code == 202
        assert requested.json()["lifecycle_status"] == "reverting"
        with runtime.database.engine.connect() as connection:
            audit = (
                connection.execute(
                    select(audit_events_table).where(
                        audit_events_table.c.event_type == "canonical_replay_revoke_requested",
                        audit_events_table.c.object_id == str(request_id),
                    )
                )
                .mappings()
                .one()
            )
            assert audit["actor_ref"] == "local-administrator"
        assert _worker(runtime, suffix="reversible-revoke").run_once() is True

        with runtime.database.engine.connect() as connection:
            request = (
                connection.execute(
                    select(canonical_replay_all_requests_table).where(
                        canonical_replay_all_requests_table.c.id == request_id
                    )
                )
                .mappings()
                .one()
            )
            content = (
                connection.execute(select(contents_table).where(contents_table.c.id == content_id))
                .mappings()
                .one()
            )
            assert request["lifecycle_status"] == "reverted"
            assert request["hidden_content_count"] == 1
            assert content["replay_visibility_owner_id"] == request_id
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(content_versions_table)
                    .where(content_versions_table.c.content_id == content_id)
                )
                == 2
            )
            assert (
                connection.scalar(
                    select(content_has_active_source(contents_table.c.id)).where(
                        contents_table.c.id == content_id
                    )
                )
                is False
            )
            assert (
                connection.scalar(
                    select(voice_plaza_content_projection_table.c.is_visible).where(
                        voice_plaza_content_projection_table.c.content_id == content_id
                    )
                )
                is False
            )
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(voice_plaza_filter_catalog_entries_table)
                    .where(voice_plaza_filter_catalog_entries_table.c.content_id == content_id)
                )
                == 0
            )
    finally:
        _truncate(runtime)
        runtime.close()


def test_new_content_batch_persists_vehicle_and_derived_brand_evidence(
    tmp_path: Path,
) -> None:
    """车型匹配的新内容快路径必须同时保存车型证据、派生品牌证据和 Replay ledger。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        _import_canonical(
            client,
            runtime,
            filename="replay-vehicle-evidence.xlsx",
            rows=(("canonical-replay-vehicle", "星曜车型命中"),),
            brand_ids=(brand_id,),
        )
        vehicle_id = _create_replay_vehicle(runtime, brand_id=brand_id)
        created = _create_all_replay(
            client,
            idempotency_key=f"all-replay-vehicle-evidence-{uuid4()}",
        )
        request_id = UUID(cast(str, created["request_id"]))
        assert _worker(runtime, suffix="vehicle-evidence").run_once() is True

        with runtime.database.engine.connect() as connection:
            vehicle = connection.execute(select(content_vehicle_evidence_table)).mappings().one()
            brand = connection.execute(select(content_brand_evidence_table)).mappings().one()
            ledger = (
                connection.execute(
                    select(canonical_replay_content_changes_table).where(
                        canonical_replay_content_changes_table.c.all_request_id == request_id
                    )
                )
                .mappings()
                .one()
            )
        assert vehicle["vehicle_model_id"] == vehicle_id
        assert vehicle["source"] == "alias_match"
        assert brand["brand_id"] == brand_id
        assert brand["source"] == "vehicle_match"
        assert brand["derived_vehicle_model_id"] == vehicle_id
        assert len(ledger["vehicle_evidence_after"]) == 1
        assert len(ledger["brand_evidence_after"]) == 1
    finally:
        _truncate(runtime)
        runtime.close()


def test_vehicle_evidence_owner_deduplicates_same_identity_before_upsert(
    tmp_path: Path,
) -> None:
    """Owner 合批写入不能让同一唯一键在单条 SQL 内更新两次。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        brand_id = _create_brand(runtime, alias="星曜")
        _create_replay_vehicle(runtime, brand_id=brand_id)
        _import_canonical(
            client,
            runtime,
            filename="vehicle-owner-duplicate.xlsx",
            rows=(("vehicle-owner-duplicate", "星曜车型"),),
            brand_ids=(brand_id,),
            expected_rows_ingested=1,
        )
        session = runtime.database.new_session()
        try:
            with session.begin():
                row = session.execute(select(content_vehicle_evidence_table)).mappings().one()
                first = replace(
                    ContentVehicleEvidence(**dict(row)),
                    id=uuid4(),
                    source="alias_match",
                    is_manual_locked=False,
                )
                second = replace(first, id=uuid4(), matched_text="同车型另一别名")
                written, locked = PostgresVehicleCatalogRepository(
                    session
                ).replace_automatic_alias_evidence_batch(
                    entries=((first.content_id, first.content_version, (first, second)),)
                )
                assert (written, locked) == (1, 0)
        finally:
            session.close()
        with runtime.database.engine.connect() as connection:
            rows = (
                connection.execute(
                    select(content_vehicle_evidence_table).where(
                        content_vehicle_evidence_table.c.source == "alias_match"
                    )
                )
                .mappings()
                .all()
            )
        assert len(rows) == 1
        assert rows[0]["is_active"] is True
        assert rows[0]["matched_text"] == first.matched_text
    finally:
        _truncate(runtime)
        runtime.close()


def test_existing_content_replay_with_two_aliases_for_same_vehicle_finishes(
    tmp_path: Path,
) -> None:
    """已有内容重筛同时命中同车型两别名时应一次成功并只留下代表证据。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand(runtime, alias="星曜")
        _import_canonical(
            client,
            runtime,
            filename="replay-existing-two-aliases.xlsx",
            rows=(("existing-two-aliases", "星曜车型"),),
            brand_ids=(brand_id,),
            expected_rows_ingested=1,
        )
        vehicle = PostgresAdministrationHttpService(runtime).create_vehicle_model(
            VehicleModelCreateRequest(
                display_name="Replay 星曜车型",
                brand_id=brand_id,
                aliases=("星曜", "星曜车型"),
            ),
            principal=_principal(),
            request_id="replay-existing-two-aliases-vehicle",
        )
        created = _create_all_replay(
            client,
            idempotency_key=f"replay-existing-two-aliases-{uuid4()}",
        )
        assert _worker(runtime, suffix="existing-two-aliases").run_once() is True
        with runtime.database.engine.connect() as connection:
            job = connection.execute(
                select(jobs_table.c.status, jobs_table.c.attempt)
                .join(
                    canonical_replay_runs_table,
                    canonical_replay_runs_table.c.job_id == jobs_table.c.id,
                )
                .where(
                    canonical_replay_runs_table.c.all_request_id
                    == UUID(cast(str, created["request_id"]))
                )
            ).one()
            evidence = (
                connection.execute(
                    select(content_vehicle_evidence_table).where(
                        content_vehicle_evidence_table.c.source == "alias_match",
                        content_vehicle_evidence_table.c.is_active.is_(True),
                    )
                )
                .mappings()
                .all()
            )
        assert job == ("succeeded", 1)
        assert len(evidence) == 1
        assert evidence[0]["vehicle_model_id"] == vehicle.id
        assert evidence[0]["matched_text"] == "星曜车型"
    finally:
        _truncate(runtime)
        runtime.close()


@pytest.mark.parametrize("row_count", [2, 101])
def test_all_replay_batches_contribution_ledger_writes(tmp_path: Path, row_count: int) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        _import_canonical(
            client,
            runtime,
            filename="replay-ledger-batch.xlsx",
            rows=tuple(
                (f"replay-ledger-{index}", f"星曜第 {index} 条") for index in range(row_count)
            ),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        created = _create_all_replay(client, idempotency_key=f"replay-ledger-{uuid4()}")
        ledger_inserts: list[str] = []
        content_updates: list[str] = []
        source_pair_reads: list[str] = []
        content_reads: list[str] = []
        statement_count = 0

        def count_ledger_insert(
            connection: object,
            cursor: object,
            statement: str,
            parameters: object,
            context: object,
            executemany: bool,
        ) -> None:
            nonlocal statement_count
            del connection, cursor, parameters, context, executemany
            statement_count += 1
            if statement.lstrip().startswith("INSERT INTO canonical_replay_content_changes"):
                ledger_inserts.append(statement)
            if statement.lstrip().startswith("UPDATE contents"):
                content_updates.append(statement)
            if statement.lstrip().startswith("SELECT provider_request_attempts.raw_artifact_id"):
                source_pair_reads.append(statement)
            if statement.lstrip().startswith("SELECT contents."):
                content_reads.append(statement)

        event.listen(runtime.database.engine, "before_cursor_execute", count_ledger_insert)
        try:
            assert _worker(runtime, suffix="ledger-batch").run_once() is True
        finally:
            event.remove(runtime.database.engine, "before_cursor_execute", count_ledger_insert)

        worker_log = (runtime.settings.log_dir / "worker.log").read_text(encoding="utf-8")
        assert "event=canonical_replay.preflight_completed" in worker_log
        assert "event=canonical_replay.ingestion_completed" in worker_log

        with runtime.database.engine.connect() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == row_count
            for table in (
                content_versions_table,
                content_metric_observations_table,
                content_external_ids_table,
                content_source_contributions_table,
                content_brand_evidence_table,
            ):
                assert connection.scalar(select(func.count()).select_from(table)) == row_count
            assert (
                connection.scalar(
                    select(func.count()).select_from(canonical_replay_content_changes_table)
                )
                == row_count
            )
            deltas = connection.scalars(
                select(canonical_replay_content_changes_table.c.delta)
            ).all()
            assert len(deltas) == row_count
            assert all(
                isinstance(delta, dict)
                and delta.get("schema_version") == "content-source-contribution.v1"
                and delta.get("created_content") is True
                for delta in deltas
            )
        assert len(ledger_inserts) == (row_count + 999) // 1000
        assert len(content_updates) == 0
        assert len(source_pair_reads) <= 1
        assert len(content_reads) == 0
        if row_count == 101:
            # 新内容主路径必须保持集合式 SQL；该上限同时防止 Content、来源贡献、
            # 自动证据或 Replay ledger 中任一环节重新退化为逐行往返。
            assert statement_count < 200
            request_id = UUID(cast(str, created["request_id"]))
            assert (
                client.post(f"/api/v1/canonical-replays/all/{request_id}/revoke").status_code == 202
            )
            remaining_scans = 0
            reversal_statement_count = 0

            def count_remaining_scan(
                connection: object,
                cursor: object,
                statement: str,
                parameters: object,
                context: object,
                executemany: bool,
            ) -> None:
                nonlocal remaining_scans, reversal_statement_count
                del connection, cursor, parameters, context, executemany
                reversal_statement_count += 1
                if (
                    "count(distinct(canonical_replay_content_changes.content_id))"
                    in statement.lower()
                ):
                    remaining_scans += 1

            event.listen(runtime.database.engine, "before_cursor_execute", count_remaining_scan)
            try:
                assert _worker(runtime, suffix="ledger-batch-reversal").run_once() is True
            finally:
                event.remove(
                    runtime.database.engine,
                    "before_cursor_execute",
                    count_remaining_scan,
                )
            assert remaining_scans == 1
            assert reversal_statement_count < 200
            with runtime.database.engine.connect() as connection:
                assert (
                    connection.scalar(
                        select(canonical_replay_all_requests_table.c.lifecycle_status).where(
                            canonical_replay_all_requests_table.c.id == request_id
                        )
                    )
                    == "reverted"
                )
    finally:
        _truncate(runtime)
        runtime.close()


def test_all_replay_batches_existing_convergence_without_per_row_sql(tmp_path: Path) -> None:
    """101 条既有内容必须保持集合收敛，防止 Current/Evidence/贡献退回逐行 SQL。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        _import_canonical(
            client,
            runtime,
            filename="replay-existing-batch.xlsx",
            rows=tuple(
                (f"replay-existing-{index}", f"星曜已有第 {index} 条") for index in range(101)
            ),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        first = _create_all_replay(client, idempotency_key=f"replay-existing-first-{uuid4()}")
        assert _worker(runtime, suffix="existing-first").run_once() is True

        second = _create_all_replay(client, idempotency_key=f"replay-existing-second-{uuid4()}")
        second_request_id = UUID(cast(str, second["request_id"]))
        statement_count = 0

        def count_sql(
            connection: object,
            cursor: object,
            statement: str,
            parameters: object,
            context: object,
            executemany: bool,
        ) -> None:
            nonlocal statement_count
            del connection, cursor, statement, parameters, context, executemany
            statement_count += 1

        event.listen(runtime.database.engine, "before_cursor_execute", count_sql)
        try:
            assert _worker(runtime, suffix="existing-second").run_once() is True
        finally:
            event.remove(runtime.database.engine, "before_cursor_execute", count_sql)

        with runtime.database.engine.connect() as connection:
            run = (
                connection.execute(
                    select(canonical_replay_runs_table).where(
                        canonical_replay_runs_table.c.all_request_id == second_request_id
                    )
                )
                .mappings()
                .one()
            )
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 101
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(canonical_replay_content_changes_table)
                    .where(
                        canonical_replay_content_changes_table.c.all_request_id == second_request_id
                    )
                )
                == 101
            )
        assert run["rows_seen"] == 101
        assert run["rows_ingested"] == 0
        assert run["existing_convergence"] == 101
        assert statement_count < 150

        # 第二次撤回必须恢复第一次 Replay 的可见性归属，不能隐藏仍有前次贡献的内容。
        first_request_id = UUID(cast(str, first["request_id"]))
        assert (
            client.post(f"/api/v1/canonical-replays/all/{second_request_id}/revoke").status_code
            == 202
        )
        assert _worker(runtime, suffix="existing-second-reversal").run_once() is True
        with runtime.database.engine.connect() as connection:
            assert set(connection.scalars(select(contents_table.c.replay_visibility_owner_id))) == {
                first_request_id
            }
    finally:
        _truncate(runtime)
        runtime.close()


def test_all_replay_revoke_keeps_version_when_only_evidence_converged(
    tmp_path: Path,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        _add_replay_alias(runtime, brand_id)
        _import_canonical(
            client,
            runtime,
            filename="replay-evidence-only.xlsx",
            rows=(("canonical-replay-evidence-only", "星曜证据幂等收敛"),),
            brand_ids=(brand_id,),
            expected_rows_ingested=1,
        )

        snapshot_session = runtime.database.new_session()
        try:
            content = snapshot_session.execute(select(contents_table)).mappings().one()
            content_id = cast(UUID, content["id"])
            original_version = cast(int, content["current_version"])
            evidence_before = PostgresBrandVehicleRepository(
                snapshot_session
            ).snapshot_automatic_brand_evidence(
                content_id=content_id,
                content_version=original_version,
            )
        finally:
            snapshot_session.close()

        created = _create_all_replay(
            client,
            idempotency_key=f"all-replay-evidence-only-{uuid4()}",
        )
        request_id = UUID(cast(str, created["request_id"]))
        assert _worker(runtime, suffix="evidence-only-replay").run_once() is True

        with runtime.database.engine.connect() as connection:
            content = (
                connection.execute(select(contents_table).where(contents_table.c.id == content_id))
                .mappings()
                .one()
            )
            change = (
                connection.execute(
                    select(canonical_replay_content_changes_table).where(
                        canonical_replay_content_changes_table.c.all_request_id == request_id
                    )
                )
                .mappings()
                .one()
            )
        assert content["current_version"] == original_version
        assert content["replay_visibility_owner_id"] == request_id
        assert change["version_before"] == change["version_after"] == original_version
        assert change["delta"] == {
            "schema_version": "content-source-contribution.v1",
            "created_content": False,
            "content_fields": {},
            "author_snapshot": None,
            "collections": {},
            "account": None,
        }

        requested = client.post(f"/api/v1/canonical-replays/all/{request_id}/revoke")
        assert requested.status_code == 202
        assert _worker(runtime, suffix="evidence-only-revoke").run_once() is True

        verify_session = runtime.database.new_session()
        try:
            content = (
                verify_session.execute(
                    select(contents_table).where(contents_table.c.id == content_id)
                )
                .mappings()
                .one()
            )
            evidence_after = PostgresBrandVehicleRepository(
                verify_session
            ).snapshot_automatic_brand_evidence(
                content_id=content_id,
                content_version=original_version,
            )
            version_count = verify_session.scalar(
                select(func.count())
                .select_from(content_versions_table)
                .where(content_versions_table.c.content_id == content_id)
            )
        finally:
            verify_session.close()
        assert content["current_version"] == original_version
        assert content["replay_visibility_owner_id"] is None
        assert version_count == 1
        assert evidence_after == evidence_before
    finally:
        _truncate(runtime)
        runtime.close()


def test_all_replay_revoke_batches_changed_evidence_without_per_content_sql(
    tmp_path: Path,
) -> None:
    """多条既有内容的自动证据撤回应集合执行，且恢复原快照。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        _add_replay_alias(runtime, brand_id)
        _import_canonical(
            client,
            runtime,
            filename="replay-evidence-batch.xlsx",
            rows=tuple(
                (f"replay-evidence-batch-{index}", f"星曜证据集合样本 {index}")
                for index in range(101)
            ),
            brand_ids=(brand_id,),
            expected_rows_ingested=101,
        )
        _add_replay_alias(runtime, brand_id, text="证据集合")
        created = _create_all_replay(
            client,
            idempotency_key=f"all-replay-evidence-batch-{uuid4()}",
        )
        request_id = UUID(cast(str, created["request_id"]))
        assert _worker(runtime, suffix="evidence-batch-replay").run_once() is True
        with runtime.database.engine.connect() as connection:
            changes = (
                connection.execute(
                    select(canonical_replay_content_changes_table).where(
                        canonical_replay_content_changes_table.c.all_request_id == request_id
                    )
                )
                .mappings()
                .all()
            )
            assert len(changes) == 101
            assert all(
                row["brand_evidence_before"] != row["brand_evidence_after"]
                and row["version_before"] == row["version_after"]
                for row in changes
            )

        assert client.post(f"/api/v1/canonical-replays/all/{request_id}/revoke").status_code == 202
        statement_count = 0

        def count_sql(
            connection: object,
            cursor: object,
            statement: str,
            parameters: object,
            context: object,
            executemany: bool,
        ) -> None:
            nonlocal statement_count
            del connection, cursor, statement, parameters, context, executemany
            statement_count += 1

        event.listen(runtime.database.engine, "before_cursor_execute", count_sql)
        try:
            assert _worker(runtime, suffix="evidence-batch-revoke").run_once() is True
        finally:
            event.remove(runtime.database.engine, "before_cursor_execute", count_sql)
        assert statement_count < 100
        with runtime.database.engine.connect() as connection:
            request = (
                connection.execute(
                    select(canonical_replay_all_requests_table).where(
                        canonical_replay_all_requests_table.c.id == request_id
                    )
                )
                .mappings()
                .one()
            )
            assert request["lifecycle_status"] == "reverted"
            assert request["retained_content_count"] == 101
            assert request["restored_evidence_count"] == 202
            assert (
                connection.scalar(select(func.count()).select_from(content_brand_evidence_table))
                == 101
            )
            assert set(connection.scalars(select(content_brand_evidence_table.c.matched_text))) == {
                "星曜"
            }
    finally:
        _truncate(runtime)
        runtime.close()


def test_all_replay_revoke_preserves_content_claimed_by_later_normal_import(
    tmp_path: Path,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        rows = tuple(
            (f"canonical-replay-later-import-{index}", f"星曜后续导入保护 {index}")
            for index in range(101)
        )
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        _import_canonical(
            client,
            runtime,
            filename="replay-before-later-import.xlsx",
            rows=rows,
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)

        created = _create_all_replay(
            client,
            idempotency_key=f"all-replay-later-import-{uuid4()}",
        )
        request_id = UUID(cast(str, created["request_id"]))
        assert _worker(runtime, suffix="later-import-replay").run_once() is True

        _import_canonical(
            client,
            runtime,
            filename="normal-import-after-replay.xlsx",
            rows=rows,
            brand_ids=(brand_id,),
            expected_rows_ingested=len(rows),
        )
        with runtime.database.engine.connect() as connection:
            content = (
                connection.execute(
                    select(contents_table).where(contents_table.c.external_content_id == rows[0][0])
                )
                .mappings()
                .one()
            )
            content_id = cast(UUID, content["id"])
            version_before_revoke = cast(int, content["current_version"])
            assert content["replay_visibility_owner_id"] is None

        requested = client.post(f"/api/v1/canonical-replays/all/{request_id}/revoke")
        assert requested.status_code == 202
        statement_count = 0

        def count_sql(
            connection: object,
            cursor: object,
            statement: str,
            parameters: object,
            context: object,
            executemany: bool,
        ) -> None:
            nonlocal statement_count
            del connection, cursor, statement, parameters, context, executemany
            statement_count += 1

        event.listen(runtime.database.engine, "before_cursor_execute", count_sql)
        try:
            assert _worker(runtime, suffix="later-import-revoke").run_once() is True
        finally:
            event.remove(runtime.database.engine, "before_cursor_execute", count_sql)
        assert statement_count < 100

        with runtime.database.engine.connect() as connection:
            request = (
                connection.execute(
                    select(canonical_replay_all_requests_table).where(
                        canonical_replay_all_requests_table.c.id == request_id
                    )
                )
                .mappings()
                .one()
            )
            content = (
                connection.execute(select(contents_table).where(contents_table.c.id == content_id))
                .mappings()
                .one()
            )
            assert request["lifecycle_status"] == "reverted"
            assert request["retained_content_count"] == len(rows)
            assert request["skipped_content_count"] == len(rows)
            assert request["hidden_content_count"] == 0
            assert content["current_version"] == version_before_revoke
            assert content["replay_visibility_owner_id"] is None
            assert (
                connection.scalar(
                    select(content_has_active_source(contents_table.c.id)).where(
                        contents_table.c.id == content_id
                    )
                )
                is True
            )
    finally:
        _truncate(runtime)
        runtime.close()


def test_all_replay_revoke_carries_manual_brand_lock_to_reversal_version(
    tmp_path: Path,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        selected_brand = _create_brand_without_matching_alias(runtime)
        manual_brand = _create_brand_without_matching_alias(runtime)
        _import_canonical(
            client,
            runtime,
            filename="replay-before-manual-lock.xlsx",
            rows=(("canonical-replay-manual-lock", "星曜人工锁保护"),),
            brand_ids=(selected_brand,),
        )
        _add_replay_alias(runtime, selected_brand)
        created = _create_all_replay(
            client,
            idempotency_key=f"all-replay-manual-lock-{uuid4()}",
        )
        request_id = UUID(cast(str, created["request_id"]))
        assert _worker(runtime, suffix="manual-lock-replay").run_once() is True

        with runtime.database.engine.connect() as connection:
            content = connection.execute(select(contents_table)).mappings().one()
        content_id = cast(UUID, content["id"])
        replay_version = cast(int, content["current_version"])
        session = runtime.database.new_session()
        try:
            with session.begin():
                PostgresBrandVehicleRepository(session).replace_manual_brand_evidence(
                    content_id=content_id,
                    content_version=replay_version,
                    brand_ids=(manual_brand,),
                    unlock_existing=False,
                    actor_ref="canonical-replay-manual-reviewer",
                )
        finally:
            session.close()

        requested = client.post(f"/api/v1/canonical-replays/all/{request_id}/revoke")
        assert requested.status_code == 202
        assert _worker(runtime, suffix="manual-lock-revoke").run_once() is True

        with runtime.database.engine.connect() as connection:
            current_version = cast(
                int,
                connection.scalar(
                    select(contents_table.c.current_version).where(
                        contents_table.c.id == content_id
                    )
                ),
            )
            lock = (
                connection.execute(
                    select(content_brand_review_locks_table).where(
                        content_brand_review_locks_table.c.content_id == content_id,
                        content_brand_review_locks_table.c.content_version == current_version,
                    )
                )
                .mappings()
                .one()
            )
            active = tuple(
                connection.execute(
                    select(content_brand_evidence_table).where(
                        content_brand_evidence_table.c.content_id == content_id,
                        content_brand_evidence_table.c.content_version == current_version,
                        content_brand_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
        assert current_version == replay_version + 1
        assert lock["is_locked"] is True
        assert lock["actor_ref"] == "canonical-replay-manual-reviewer"
        assert [(row["brand_id"], row["source"], row["is_manual_locked"]) for row in active] == [
            (manual_brand, "manual_review", True)
        ]
    finally:
        _truncate(runtime)
        runtime.close()


def test_all_artifacts_are_preflighted_before_first_content_write(tmp_path: Path) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        first_artifact = _import_canonical(
            client,
            runtime,
            filename="replay-valid.xlsx",
            rows=(("canonical-replay-valid", "星曜合法记录"),),
            brand_ids=(brand_id,),
        )
        second_artifact = _import_canonical(
            client,
            runtime,
            filename="replay-corrupt.xlsx",
            rows=(("canonical-replay-corrupt", "星曜损坏记录"),),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        with runtime.database.engine.connect() as connection:
            storage_key = connection.scalar(
                select(artifacts_table.c.storage_key).where(artifacts_table.c.id == second_artifact)
            )
        assert isinstance(storage_key, str)
        target = runtime.artifact_store.root.joinpath(*storage_key.split("/"))
        target.write_bytes(b"corrupted-after-link")

        replay = _create_replay(
            client,
            artifact_ids=(first_artifact, second_artifact),
            brand_id=brand_id,
            idempotency_key=f"replay-preflight-{uuid4()}",
        )
        assert _worker(runtime, suffix="preflight").run_once() is True
        failed = client.get(f"/api/v1/canonical-replays/{replay['run_id']}")

        assert failed.status_code == 200
        assert failed.json()["job"]["status"] == "failed"
        assert failed.json()["job"]["error_code"] == "canonical_replay_artifact_invalid"
        assert failed.json()["stats"]["rows_seen"] == 0
        with runtime.database.engine.connect() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 0
            assert (
                connection.scalar(
                    select(func.count()).select_from(canonical_replay_seen_content_table)
                )
                == 0
            )
    finally:
        _truncate(runtime)
        runtime.close()


def test_later_import_parent_failure_is_found_before_first_content_write(tmp_path: Path) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        first_artifact = _import_canonical(
            client,
            runtime,
            filename="replay-parent-first.xlsx",
            rows=(("replay-parent-first", "星曜第一条"),),
            brand_ids=(brand_id,),
        )
        second_artifact = _import_canonical(
            client,
            runtime,
            filename="replay-parent-second.xlsx",
            rows=(("replay-parent-second", "星曜第二条"),),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        replay = _create_replay(
            client,
            artifact_ids=(first_artifact, second_artifact),
            brand_id=brand_id,
            idempotency_key=f"replay-parent-failure-{uuid4()}",
        )
        with runtime.database.engine.begin() as connection:
            input_artifact_id = connection.scalar(
                select(processing_import_batches_table.c.input_artifact_id)
                .join(
                    canonical_artifact_links_table,
                    canonical_artifact_links_table.c.processing_import_batch_id
                    == processing_import_batches_table.c.id,
                )
                .where(canonical_artifact_links_table.c.artifact_id == second_artifact)
            )
            assert isinstance(input_artifact_id, UUID)
            connection.execute(
                update(artifacts_table)
                .where(artifacts_table.c.id == input_artifact_id)
                .values(storage_status="error")
            )
        assert _worker(runtime, suffix="parent-failure").run_once() is True
        detail = client.get(f"/api/v1/canonical-replays/{replay['run_id']}").json()
        assert detail["job"]["status"] == "failed"
        assert detail["stats"]["rows_seen"] == 0
        with runtime.database.engine.connect() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 0
    finally:
        _truncate(runtime)
        runtime.close()


def test_replay_enqueue_is_idempotent_and_rejects_parameter_drift(tmp_path: Path) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-idempotency.xlsx",
            rows=(("canonical-replay-idempotency", "星曜幂等记录"),),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        key = f"replay-idempotency-{uuid4()}"

        first = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=key,
        )
        repeated = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=key,
        )
        drift = client.post(
            "/api/v1/canonical-replays",
            json={
                "idempotency_key": key,
                "artifact_ids": [str(artifact_id)],
                "brand_ids": [str(brand_id)],
                "batch_size": 2,
            },
        )

        assert repeated == first
        assert drift.status_code == 409
        assert drift.json()["errors"][0]["code"] == "canonical_replay_conflict"
    finally:
        _truncate(runtime)
        runtime.close()


def test_concurrent_same_idempotency_key_returns_one_run_and_job(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(create_app(import_service=PostgresImportHttpService(runtime)))
        brand_id = _create_brand_without_matching_alias(runtime)
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-concurrent.xlsx",
            rows=(("canonical-replay-concurrent", "并发幂等记录"),),
            brand_ids=(brand_id,),
        )
        body = CanonicalReplayCreateRequest(
            idempotency_key=f"replay-concurrent-{uuid4()}",
            artifact_ids=(artifact_id,),
            brand_ids=(brand_id,),
            batch_size=1,
        )
        barrier = Barrier(2, timeout=5)
        original = PostgresCanonicalReplayRepository._lock_idempotency_key

        def synchronized_lock(repository, key):  # type: ignore[no-untyped-def]
            barrier.wait()
            original(repository, key)

        monkeypatch.setattr(
            PostgresCanonicalReplayRepository,
            "_lock_idempotency_key",
            synchronized_lock,
        )

        def create(index: int):  # type: ignore[no-untyped-def]
            return PostgresCanonicalReplayHttpService(runtime).create_replay(
                body,
                actor_ref=_principal().principal_id,
                request_id=f"canonical-replay-concurrent-{index}",
            )

        with ThreadPoolExecutor(max_workers=2) as executor:
            responses = tuple(executor.map(create, range(2)))

        assert responses[0] == responses[1]
        with runtime.database.engine.connect() as connection:
            assert (
                connection.scalar(
                    select(func.count())
                    .select_from(jobs_table)
                    .where(jobs_table.c.job_type == CANONICAL_REPLAY_JOB_TYPE)
                )
                == 1
            )
    finally:
        _truncate(runtime)
        runtime.close()


def test_small_batches_open_each_artifact_once_after_preflight(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand(runtime, alias="星曜")
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-linear-read.xlsx",
            rows=tuple(
                (f"canonical-replay-linear-{index}", f"星曜线性读取 {index}") for index in range(5)
            ),
            brand_ids=(brand_id,),
            expected_rows_ingested=5,
        )
        preflight_calls = 0
        ingest_calls = 0
        original_preflight = (
            canonical_replay_worker_module.CanonicalArtifactReader.read_for_preflight
        )
        original_ingest = canonical_replay_worker_module.CanonicalArtifactReader.read_preflighted

        def counted_preflight(reader, artifact):  # type: ignore[no-untyped-def]
            """记录全输入预检读取次数。"""

            nonlocal preflight_calls
            preflight_calls += 1
            yield from original_preflight(reader, artifact)

        def counted_ingest(reader, artifact):  # type: ignore[no-untyped-def]
            """记录每个文件的业务读取次数。"""

            nonlocal ingest_calls
            ingest_calls += 1
            yield from original_ingest(reader, artifact)

        monkeypatch.setattr(
            canonical_replay_worker_module.CanonicalArtifactReader,
            "read_for_preflight",
            counted_preflight,
        )
        monkeypatch.setattr(
            canonical_replay_worker_module.CanonicalArtifactReader,
            "read_preflighted",
            counted_ingest,
        )
        _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-linear-{uuid4()}",
            batch_size=1,
        )

        assert _worker(runtime, suffix="linear-read").run_once() is True
        assert preflight_calls == 1
        assert ingest_calls == 1
    finally:
        _truncate(runtime)
        runtime.close()


def test_validation_proof_reuses_only_matching_contract_and_verified_bytes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-proof.xlsx",
            rows=(("canonical-replay-proof", "星曜证明记录"),),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        first = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-proof-first-{uuid4()}",
        )
        assert _worker(runtime, suffix="proof-first").run_once() is True
        assert (
            client.get(f"/api/v1/canonical-replays/{first['run_id']}").json()["job"]["status"]
            == "succeeded"
        )
        with runtime.database.engine.connect() as connection:
            proof = connection.execute(
                select(canonical_replay_validation_proofs_table).where(
                    canonical_replay_validation_proofs_table.c.artifact_id == artifact_id
                )
            ).one()
            storage_key = connection.scalar(
                select(artifacts_table.c.storage_key).where(artifacts_table.c.id == artifact_id)
            )
        assert proof.sha256 and proof.validation_version
        assert isinstance(storage_key, str)

        parsed = 0
        bytes_verified = 0
        original_parse = canonical_replay_worker_module.CanonicalArtifactReader.read_for_preflight
        original_verify = (
            canonical_replay_worker_module.CanonicalArtifactReader.verify_bytes_for_preflight
        )

        def counted_parse(reader, artifact):  # type: ignore[no-untyped-def]
            nonlocal parsed
            parsed += 1
            yield from original_parse(reader, artifact)

        def counted_verify(reader, artifact):  # type: ignore[no-untyped-def]
            nonlocal bytes_verified
            bytes_verified += 1
            return original_verify(reader, artifact)

        monkeypatch.setattr(
            canonical_replay_worker_module.CanonicalArtifactReader,
            "read_for_preflight",
            counted_parse,
        )
        monkeypatch.setattr(
            canonical_replay_worker_module.CanonicalArtifactReader,
            "verify_bytes_for_preflight",
            counted_verify,
        )
        second = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-proof-second-{uuid4()}",
        )
        assert _worker(runtime, suffix="proof-second").run_once() is True
        assert (
            client.get(f"/api/v1/canonical-replays/{second['run_id']}").json()["job"]["status"]
            == "succeeded"
        )
        assert parsed == 0
        assert bytes_verified == 1

        with runtime.database.engine.begin() as connection:
            connection.execute(
                update(canonical_replay_validation_proofs_table)
                .where(canonical_replay_validation_proofs_table.c.artifact_id == artifact_id)
                .values(validation_version="obsolete-validator")
            )
        third = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-proof-third-{uuid4()}",
        )
        assert _worker(runtime, suffix="proof-third").run_once() is True
        assert (
            client.get(f"/api/v1/canonical-replays/{third['run_id']}").json()["job"]["status"]
            == "succeeded"
        )
        assert parsed == 1

        target = runtime.artifact_store.root.joinpath(*storage_key.split("/"))
        target.write_bytes(b"tampered-after-proof")
        fourth = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-proof-fourth-{uuid4()}",
        )
        assert _worker(runtime, suffix="proof-fourth").run_once() is True
        response = client.get(f"/api/v1/canonical-replays/{fourth['run_id']}").json()
        assert response["job"]["status"] == "failed"
        assert response["stats"]["rows_seen"] == 0
        assert parsed == 1
        assert bytes_verified == 2
    finally:
        _truncate(runtime)
        runtime.close()


def test_selected_replay_preserves_existing_out_of_scope_brand_evidence(
    tmp_path: Path,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        selected_brand = _create_brand(runtime, alias="星曜")
        out_of_scope_brand = _create_brand(runtime, alias="月影")
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-preserve-evidence.xlsx",
            rows=(("canonical-replay-preserve-evidence", "星曜与月影联名"),),
            brand_ids=(selected_brand, out_of_scope_brand),
            expected_rows_ingested=1,
        )
        with runtime.database.engine.connect() as connection:
            before = set(
                connection.scalars(
                    select(content_brand_evidence_table.c.brand_id).where(
                        content_brand_evidence_table.c.is_active.is_(True)
                    )
                )
            )
        assert before == {selected_brand, out_of_scope_brand}

        replay = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=selected_brand,
            idempotency_key=f"replay-preserve-evidence-{uuid4()}",
        )
        assert _worker(runtime, suffix="preserve-evidence").run_once() is True
        response = client.get(f"/api/v1/canonical-replays/{replay['run_id']}")
        assert response.json()["stats"]["existing_convergence"] == 1
        with runtime.database.engine.connect() as connection:
            after = set(
                connection.scalars(
                    select(content_brand_evidence_table.c.brand_id).where(
                        content_brand_evidence_table.c.is_active.is_(True)
                    )
                )
            )
        assert after == {selected_brand, out_of_scope_brand}
    finally:
        _truncate(runtime)
        runtime.close()


def test_replay_replaces_same_brand_automatic_evidence_and_preserves_manual_lock(
    tmp_path: Path,
) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        selected_brand = _create_brand_without_matching_alias(runtime)
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-evidence-merge.xlsx",
            rows=(("canonical-replay-evidence-merge", "星曜证据合并"),),
            brand_ids=(selected_brand,),
        )
        first_catalog_version = _add_replay_alias(runtime, selected_brand)
        _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=selected_brand,
            idempotency_key=f"replay-evidence-first-{uuid4()}",
        )
        assert _worker(runtime, suffix="evidence-first").run_once() is True

        second_catalog_version = _add_replay_alias(
            runtime,
            selected_brand,
            text="目录版本推进别名",
        )
        assert second_catalog_version > first_catalog_version
        _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=selected_brand,
            idempotency_key=f"replay-evidence-second-{uuid4()}",
        )
        assert _worker(runtime, suffix="evidence-second").run_once() is True

        with runtime.database.engine.connect() as connection:
            content = connection.execute(
                select(contents_table.c.id, contents_table.c.current_version)
            ).one()
            selected_rows = tuple(
                connection.execute(
                    select(content_brand_evidence_table).where(
                        content_brand_evidence_table.c.content_id == content.id,
                        content_brand_evidence_table.c.brand_id == selected_brand,
                    )
                ).mappings()
            )
        assert [row["catalog_version"] for row in selected_rows if row["is_active"]] == [
            second_catalog_version
        ]
        assert any(
            row["catalog_version"] == first_catalog_version and not row["is_active"]
            for row in selected_rows
        )

        manual_brand = _create_brand_without_matching_alias(runtime)
        session = runtime.database.new_session()
        try:
            with session.begin():
                PostgresBrandVehicleRepository(session).replace_manual_brand_evidence(
                    content_id=content.id,
                    content_version=content.current_version,
                    brand_ids=(manual_brand,),
                    unlock_existing=False,
                    actor_ref="canonical-replay-manual-reviewer",
                )
        finally:
            session.close()

        _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=selected_brand,
            idempotency_key=f"replay-evidence-manual-lock-{uuid4()}",
        )
        assert _worker(runtime, suffix="evidence-manual-lock").run_once() is True
        with runtime.database.engine.connect() as connection:
            active = tuple(
                connection.execute(
                    select(content_brand_evidence_table).where(
                        content_brand_evidence_table.c.content_id == content.id,
                        content_brand_evidence_table.c.content_version == content.current_version,
                        content_brand_evidence_table.c.is_active.is_(True),
                    )
                ).mappings()
            )
        assert [(row["brand_id"], row["source"], row["is_manual_locked"]) for row in active] == [
            (manual_brand, "manual_review", True)
        ]
    finally:
        _truncate(runtime)
        runtime.close()


def test_repository_snapshot_is_exactly_the_requested_active_brand(tmp_path: Path) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        selected = _create_brand_without_matching_alias(runtime)
        ignored = _create_brand_without_matching_alias(runtime)
        session = runtime.database.new_session()
        try:
            with session.begin():
                snapshot = PostgresBrandVehicleRepository(session).snapshot(brand_ids=(selected,))
        finally:
            session.close()

        assert snapshot.filter_scope == "selected"
        assert snapshot.selected_brand_ids == (selected,)
        assert {item.id for item in snapshot.brands} == {selected}
        assert ignored not in {item.id for item in snapshot.brands}
    finally:
        _truncate(runtime)
        runtime.close()


def test_running_replay_cancels_between_committed_batches(tmp_path: Path) -> None:
    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-running-cancel.xlsx",
            rows=(
                ("canonical-replay-cancel-1", "星曜取消第一条"),
                ("canonical-replay-cancel-2", "星曜取消第二条"),
            ),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        created = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-running-cancel-{uuid4()}",
            batch_size=1,
        )
        run_id = UUID(str(created["run_id"]))
        job_id = UUID(str(created["job_id"]))

        claim_session = runtime.database.new_session()
        try:
            with claim_session.begin():
                claim = PostgresJobRepository(claim_session).claim_next(
                    supported_job_types=(CANONICAL_REPLAY_JOB_TYPE,),
                    worker_id="replay-cancelling-worker",
                    lease_seconds=30,
                )
                assert claim is not None and claim.lease_token is not None
        finally:
            claim_session.close()
        fence = JobExecutionFence(job_id=job_id, lease_token=claim.lease_token)

        result = PostgresCanonicalReplayJobExecutor(runtime).execute(
            payload=CanonicalReplayJobPayload(run_id=run_id),
            fence=fence,
            context=_ExecutionContext(fence, cancel_after_checks=3),
        )

        assert result.outcome == "cancelled"
        session = runtime.database.new_session()
        try:
            with session.begin():
                run = PostgresCanonicalReplayRepository(session).get(run_id)
                assert run is not None
                assert run.checkpoint_artifact_ordinal == 0
                assert run.checkpoint_row_number == 1
                assert run.rows_seen == 1
        finally:
            session.close()
        with runtime.database.engine.connect() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 1
    finally:
        _truncate(runtime)
        runtime.close()


def test_cancellation_before_fenced_checkpoint_rolls_back_batch(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """批次写入期间收到取消时，提交前 Fence 复核必须回滚整批。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-fenced-cancel.xlsx",
            rows=(
                ("canonical-replay-fenced-cancel-1", "星曜提交前取消一"),
                ("canonical-replay-fenced-cancel-2", "星曜提交前取消二"),
            ),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        created = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-fenced-cancel-{uuid4()}",
            batch_size=2,
        )
        run_id = UUID(str(created["run_id"]))
        job_id = UUID(str(created["job_id"]))
        claim_session = runtime.database.new_session()
        try:
            with claim_session.begin():
                claim = PostgresJobRepository(claim_session).claim_next(
                    supported_job_types=(CANONICAL_REPLAY_JOB_TYPE,),
                    worker_id="replay-fenced-cancel-worker",
                    lease_seconds=30,
                )
                assert claim is not None and claim.lease_token is not None
        finally:
            claim_session.close()
        fence = JobExecutionFence(job_id=job_id, lease_token=claim.lease_token)
        original_advance = PostgresCanonicalReplayRepository.advance
        cancellation_sent = False

        def cancel_before_advance(
            repository: PostgresCanonicalReplayRepository,
            **kwargs: object,
        ):
            nonlocal cancellation_sent
            if not cancellation_sent:
                cancel_session = runtime.database.new_session()
                try:
                    with cancel_session.begin():
                        PostgresJobRepository(cancel_session).request_cancel(job_id)
                finally:
                    cancel_session.close()
                cancellation_sent = True
            return original_advance(repository, **kwargs)  # type: ignore[arg-type]

        monkeypatch.setattr(
            PostgresCanonicalReplayRepository,
            "advance",
            cancel_before_advance,
        )
        with pytest.raises(LeaseLostError):
            PostgresCanonicalReplayJobExecutor(runtime).execute(
                payload=CanonicalReplayJobPayload(run_id=run_id),
                fence=fence,
                context=_ExecutionContext(fence),
            )

        session = runtime.database.new_session()
        try:
            run = PostgresCanonicalReplayRepository(session).get(run_id)
            assert run is not None
            assert run.checkpoint_artifact_ordinal == 0
            assert run.checkpoint_row_number == 0
            assert run.rows_seen == 0
        finally:
            session.close()
        with runtime.database.engine.connect() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 0
            assert (
                connection.scalar(
                    select(func.count()).select_from(canonical_replay_content_changes_table)
                )
                == 0
            )
            assert (
                connection.scalar(
                    select(func.count()).select_from(canonical_replay_seen_content_table)
                )
                == 0
            )
    finally:
        _truncate(runtime)
        runtime.close()


def test_replay_takeover_resumes_checkpoint_and_rejects_stale_fence(tmp_path: Path) -> None:
    """首批提交后接管从检查点续跑，旧 Token 不能再推进统计。"""

    runtime = _runtime(tmp_path)
    _truncate(runtime)
    try:
        client = TestClient(
            create_app(
                import_service=PostgresImportHttpService(runtime),
                canonical_replay_service=PostgresCanonicalReplayHttpService(runtime),
            )
        )
        brand_id = _create_brand_without_matching_alias(runtime)
        artifact_id = _import_canonical(
            client,
            runtime,
            filename="replay-takeover.xlsx",
            rows=(
                ("canonical-replay-takeover-1", "星曜接管第一条"),
                ("canonical-replay-takeover-2", "星曜接管第二条"),
            ),
            brand_ids=(brand_id,),
        )
        _add_replay_alias(runtime, brand_id)
        created = _create_replay(
            client,
            artifact_ids=(artifact_id,),
            brand_id=brand_id,
            idempotency_key=f"replay-takeover-{uuid4()}",
            batch_size=1,
        )
        run_id = UUID(str(created["run_id"]))
        job_id = UUID(str(created["job_id"]))

        first_session = runtime.database.new_session()
        try:
            with first_session.begin():
                first_claim = PostgresJobRepository(first_session).claim_next(
                    supported_job_types=(CANONICAL_REPLAY_JOB_TYPE,),
                    worker_id="replay-old-worker",
                    lease_seconds=30,
                )
                assert first_claim is not None
                assert first_claim.lease_token is not None
        finally:
            first_session.close()
        old_fence = JobExecutionFence(job_id=job_id, lease_token=first_claim.lease_token)

        with pytest.raises(LeaseLostError):
            PostgresCanonicalReplayJobExecutor(runtime).execute(
                payload=CanonicalReplayJobPayload(run_id=run_id),
                fence=old_fence,
                context=_ExecutionContext(old_fence, lose_after_heartbeat=True),
            )

        with runtime.database.engine.begin() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 1
            connection.execute(
                text(
                    "UPDATE jobs SET lease_expires_at = clock_timestamp() - interval '1 second' "
                    "WHERE id = :job_id"
                ),
                {"job_id": job_id},
            )

        takeover_session = runtime.database.new_session()
        try:
            with takeover_session.begin():
                takeover = PostgresJobRepository(takeover_session).claim_next(
                    supported_job_types=(CANONICAL_REPLAY_JOB_TYPE,),
                    worker_id="replay-new-worker",
                    lease_seconds=30,
                )
                assert takeover is not None
                assert takeover.lease_token is not None
                assert takeover.lease_token != first_claim.lease_token
        finally:
            takeover_session.close()
        takeover_fence = JobExecutionFence(job_id=job_id, lease_token=takeover.lease_token)

        stale_session = runtime.database.new_session()
        try:
            with pytest.raises(LeaseLostError), stale_session.begin():
                PostgresCanonicalReplayRepository(stale_session).advance(
                    run_id=run_id,
                    expected_artifact_ordinal=0,
                    expected_row_number=1,
                    next_artifact_ordinal=0,
                    next_row_number=2,
                    counters=CanonicalReplayCounters(1, 1, 0, 0, 1, 0),
                    fence=old_fence,
                )
        finally:
            stale_session.close()

        result = PostgresCanonicalReplayJobExecutor(runtime).execute(
            payload=CanonicalReplayJobPayload(run_id=run_id),
            fence=takeover_fence,
            context=_ExecutionContext(takeover_fence),
        )

        assert result.outcome == "succeeded"
        assert result.result == {
            "run_id": str(run_id),
            "artifact_count": 1,
            "rows_seen": 2,
            "rows_matched": 2,
            "rows_filtered_out": 0,
            "duplicates_removed": 0,
            "rows_ingested": 2,
            "existing_convergence": 0,
            "invalid_artifact_rows": 0,
        }
        with runtime.database.engine.connect() as connection:
            assert connection.scalar(select(func.count()).select_from(contents_table)) == 2
    finally:
        _truncate(runtime)
        runtime.close()
