"""公开 HTTP → 持久 Job → 文件/Canonical → PostgreSQL 的真实接线。"""

from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pytest
from aima_ugc.adapters.providers.wisersone.models import (
    ExportProgress,
    ExportTask,
    SubmissionUnknown,
)
from aima_ugc.bootstrap import api as api_module
from aima_ugc.bootstrap import worker as worker_module
from aima_ugc.bootstrap.brand_vehicle_http import PostgresBrandVehicleHttpService
from aima_ugc.bootstrap.scheduler import run_scheduler_once
from aima_ugc.bootstrap.wisersone_cleanup import cleanup_wisersone_files
from aima_ugc.bootstrap.wisersone_worker import PostgresWisersOneJobExecutor
from aima_ugc.contracts.brand_vehicle import BrandCreateRequest
from aima_ugc.modules.collection.tables import collection_schedule_occurrences_table
from aima_ugc.modules.content.tables import contents_table
from aima_ugc.modules.identity import Principal
from aima_ugc.modules.ingestion.historical_tables import historical_import_campaigns_table
from aima_ugc.modules.ingestion.wisersone_jobs import WISERSONE_JOB_TYPE
from aima_ugc.modules.ingestion.wisersone_tables import wisersone_downloads_table
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs.tables import jobs_table
from aima_ugc.platform.storage.tables import canonical_artifact_links_table
from aima_ugc.platform.time import beijing_now
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select, update


def _xlsx() -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "文章"
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    sheet.append(
        [
            "小红书",
            "爱玛测试内容",
            "爱玛",
            "测试",
            "2026-10-03 10:00:00",
            "https://www.xiaohongshu.com/explore/wisersone-workflow",
        ]
    )
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class _Website:
    """只替代外部网站；保留生产 Job、Reader、Mapper、Artifact 和数据库。"""

    def __init__(self) -> None:
        self.submissions = 0
        self.waits = 0
        self.unknown = False
        self.receipts: dict[UUID, ExportTask] = {}

    def saved_task(self, run_id: UUID) -> ExportTask | None:
        return self.receipts.get(run_id)

    def submit(self, run_id: UUID, **callbacks: Any) -> ExportTask:
        callbacks["before_submit"]()
        self.submissions += 1
        if self.unknown:
            raise SubmissionUnknown
        task = ExportTask(f"website-{run_id}")
        self.receipts[run_id] = task
        callbacks["on_submitted"](task)
        return task

    def poll(self, task: ExportTask, destination: Path, **options: Any) -> ExportProgress:
        assert task in self.receipts.values()
        if self.waits:
            self.waits -= 1
            return ExportProgress(False, 25, None)
        if not options.get("download", True):
            return ExportProgress(True, 100, None)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(_xlsx())
        return ExportProgress(True, 100, destination)


@pytest.fixture
def workflow(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    approved, managed = tmp_path / "approved", tmp_path / "managed"
    approved.mkdir()
    managed.mkdir()
    runtime = worker_module.create_worker_runtime(
        settings=load_settings().model_copy(
            update={
                "data_dir": tmp_path / "data",
                "log_dir": tmp_path / "logs",
                "historical_import_root": approved,
                "wisersone_input_dir": managed,
                "wisersone_auth_dir": tmp_path / "auth",
            }
        )
    )
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE collection_plans, jobs, artifacts, keyword_packs, "
            "vehicle_brands, accounts CASCADE"
        )
    website = _Website()
    monkeypatch.setattr(api_module, "create_platform_runtime", lambda *_: runtime)
    monkeypatch.setattr(
        worker_module,
        "PostgresWisersOneJobExecutor",
        lambda r: PostgresWisersOneJobExecutor(r, exporter=website),
    )
    client = TestClient(api_module.create_app())
    brand = PostgresBrandVehicleHttpService(runtime).create_brand(
        BrandCreateRequest(display_name="爱玛", role="owned", aliases=("爱玛",)),
        principal=Principal(
            principal_id="wisersone-test",
            display_name="测试",
            role="administrator",
            source="development",
        ),
        request_id="wisersone-test",
    )

    def new_worker():
        return worker_module.create_job_worker(
            runtime=runtime,
            registry=worker_module.create_collection_job_registry(runtime=runtime),
            worker_id=f"wisersone-test-{uuid4()}",
            lease_seconds=120,
            retry_delay_seconds=0,
        )

    def tick(worker):
        # 只加速测试调度，不替代生产续跑/事务/领取实现。
        with runtime.database.engine.begin() as connection:
            connection.execute(
                update(jobs_table)
                .where(jobs_table.c.status == "queued")
                .values(available_at=beijing_now() - timedelta(seconds=1))
            )
        assert worker.run_once()

    try:
        yield runtime, website, client, str(brand.id), new_worker, tick
    finally:
        client.close()
        runtime.close()


def _create(client: TestClient, brand: str, key: str | None = None):
    response = client.post(
        "/api/v1/wisersone-downloads",
        json={
            "client_idempotency_key": key or str(uuid4()),
            "brand_ids": [brand],
        },
    )
    assert response.status_code == 202, response.text
    return response.json()


def test_download_resume_beyond_thirty_minutes_and_import_then_cleanup(workflow) -> None:
    runtime, website, client, brand, new_worker, tick = workflow
    created = _create(client, brand)
    run_id = created["id"]
    website.waits = 65
    worker = new_worker()
    tick(worker)
    for index in range(65):
        if index == 30:
            worker = new_worker()
        tick(worker)
        result = client.get(f"/api/v1/wisersone-downloads/{run_id}").json()
        assert result["status"] == "waiting"
    # 65 个正常 30 秒续跑超过旧 30 分钟限制；每个阶段只执行一次，无重新提交。
    with runtime.database.engine.connect() as connection:
        jobs = (
            connection.execute(
                select(jobs_table).where(jobs_table.c.job_type == WISERSONE_JOB_TYPE)
            )
            .mappings()
            .all()
        )
        assert len(jobs) == 67
        assert all(job["attempt"] <= 1 for job in jobs)
    for _ in range(30):
        tick(worker)
        result = client.get(f"/api/v1/wisersone-downloads/{run_id}").json()
        if result["status"] in {"succeeded", "failed", "partial_failed"}:
            break
    assert result["status"] == "succeeded", result
    assert result["campaign_id"] and result["sha256"]
    assert website.submissions == 1
    published = runtime.settings.wisersone_input_dir / run_id / "wisersone_last24h.xlsx"
    assert published.is_file()
    assert cleanup_wisersone_files(runtime) == 0
    with runtime.database.engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(contents_table)) == 1
        canonical_count = connection.scalar(
            select(func.count()).select_from(canonical_artifact_links_table)
        )
        assert canonical_count
    past = beijing_now() - timedelta(days=8)
    with runtime.database.engine.begin() as connection:
        connection.execute(
            update(wisersone_downloads_table)
            .where(wisersone_downloads_table.c.id == UUID(run_id))
            .values(finished_at=past)
        )
        connection.execute(
            update(historical_import_campaigns_table)
            .where(historical_import_campaigns_table.c.id == UUID(result["campaign_id"]))
            .values(finished_at=past)
        )
    assert cleanup_wisersone_files(runtime) == 1
    assert not published.exists()
    with runtime.database.engine.connect() as connection:
        assert (
            connection.scalar(select(func.count()).select_from(canonical_artifact_links_table))
            == canonical_count
        )
    assert client.post(f"/api/v1/wisersone-downloads/{run_id}/retry").status_code == 409


def test_parallel_duplicate_request_creates_one_job_and_cancel_stops_send(workflow) -> None:
    runtime, website, client, brand, new_worker, tick = workflow
    key = str(uuid4())
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: _create(client, brand, key), range(2)))
    assert results[0]["id"] == results[1]["id"]
    run_id = results[0]["id"]
    assert (
        client.post(f"/api/v1/wisersone-downloads/{run_id}/cancel").json()["status"] == "cancelled"
    )
    # 请求取消 queued Job 会立即终结，领取器不执行外部调用。
    assert new_worker().run_once() is False
    assert website.submissions == 0
    with runtime.database.engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(wisersone_downloads_table)) == 1
    response = client.post(
        "/api/v1/wisersone-downloads", json={"client_idempotency_key": key, "brand_ids": []}
    )
    assert response.status_code == 409
    assert "request_id" in response.json()


def test_unknown_submission_is_not_automatically_repeated(workflow) -> None:
    _, website, client, brand, new_worker, tick = workflow
    website.unknown = True
    created = _create(client, brand)
    tick(new_worker())
    result = client.get(f"/api/v1/wisersone-downloads/{created['id']}").json()
    assert result["status"] == "attention" and result["send_state"] == "unknown"
    assert client.post(f"/api/v1/wisersone-downloads/{created['id']}/retry").status_code == 409
    assert website.submissions == 1


def test_plan_schedulers_enqueue_only_latest_slot_once_then_automatically_import(workflow) -> None:
    runtime, website, client, brand, new_worker, tick = workflow
    response = client.post(
        "/api/v1/collection-plans",
        json={
            "plan_type": "wisersone",
            "name": "网站自动导入",
            "schedule_expr": "* * * * *",
            "brand_ids": [brand],
        },
    )
    assert response.status_code == 201, response.text
    plan = response.json()
    assert plan["plan_type"] == "wisersone" and plan["comment_policy"] is None
    first = beijing_now().replace(second=0, microsecond=0)
    assert run_scheduler_once(runtime, now=first).initialized == 1
    due = first + timedelta(minutes=5)
    with ThreadPoolExecutor(max_workers=2) as pool:
        outcomes = list(pool.map(lambda _: run_scheduler_once(runtime, now=due), range(2)))
    assert sum(result.enqueued for result in outcomes) == 1
    assert sum(result.skipped for result in outcomes) == 4
    downloads = client.get("/api/v1/wisersone-downloads").json()["items"]
    assert len(downloads) == 1 and downloads[0]["occurrence_id"]
    with runtime.database.engine.connect() as connection:
        assert (
            connection.scalar(
                select(func.count()).select_from(collection_schedule_occurrences_table)
            )
            == 5
        )
    worker = new_worker()
    for _ in range(25):
        tick(worker)
        result = client.get(f"/api/v1/wisersone-downloads/{downloads[0]['id']}").json()
        if result["status"] in {"succeeded", "failed", "partial_failed"}:
            break
    assert result["status"] == "succeeded", result
    assert website.submissions == 1
