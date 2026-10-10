"""真实 PostgreSQL Session → Export 冻结/Worker/XLSX 与个人字段配置隔离。"""

from __future__ import annotations

from collections.abc import Iterator
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from threading import Barrier
from typing import Any
from uuid import UUID, uuid4

import pytest
from aima_ugc.adapters.persistence.postgres.analysis_schemes import (
    PostgresAnalysisSchemeRepository,
)
from aima_ugc.adapters.persistence.postgres.export_column_defaults import (
    PostgresExportColumnDefaultRepository,
)
from aima_ugc.bootstrap.content_http import PostgresContentHttpService
from aima_ugc.bootstrap.export_worker import PostgresDataExportJobExecutor
from aima_ugc.bootstrap.feishu_auth_http import (
    FeishuAuthRoutes,
    FeishuAuthSettings,
    FeishuLoginRequiredResolver,
    install_feishu_auth_routes,
)
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.product_http import PostgresProductHttpService
from aima_ugc.bootstrap.reporting_http import PostgresReportingHttpService
from aima_ugc.bootstrap.route_authorization import install_route_authorization
from aima_ugc.bootstrap.runtime import PlatformRuntime
from aima_ugc.bootstrap.workbench_http import PostgresWorkbenchHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.entrypoints.api_main import create_app
from aima_ugc.modules.analysis.scheme_tables import (
    analysis_scheme_versions_table,
    analysis_schemes_table,
)
from aima_ugc.modules.identity import DevelopmentIdentityResolver
from aima_ugc.modules.identity.feishu import SESSION_COOKIE_NAME, PrincipalStore, SessionStore
from aima_ugc.modules.identity.tables import identity_sessions_table
from aima_ugc.modules.reporting.column_catalog import EXPORT_COLUMN_CATALOG_VERSION
from aima_ugc.modules.reporting.data_export_job import (
    DataExportJobHandler,
    register_data_export_job,
)
from aima_ugc.modules.reporting.http import ExportColumnDefaultConflict
from aima_ugc.modules.reporting.tables import (
    reporting_data_exports_table,
    reporting_user_export_column_defaults_table,
)
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs import JobRegistry
from aima_ugc.platform.jobs.tables import jobs_table
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, event, func, select, text, update

from tests.integration.content.test_stage8d_voice_plaza_runtime import _seed_import

_ORIGIN = "http://testserver"
_DEFAULT_PATH = "/api/v1/me/export-column-default"


@pytest.fixture
def role_runtime(tmp_path: Path) -> Iterator[PlatformRuntime]:
    """复用仓库隔离 PostgreSQL 生命周期；不连接外部 Provider。"""
    settings = load_settings().model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts RESTART IDENTITY CASCADE"
        )
    try:
        yield runtime
    finally:
        with runtime.database.engine.begin() as connection:
            connection.exec_driver_sql(
                "TRUNCATE TABLE jobs, artifacts, keyword_packs, accounts RESTART IDENTITY CASCADE"
            )
        runtime.close()


def _clients(runtime: PlatformRuntime) -> tuple[dict[str, TestClient], dict[str, str]]:
    """三个账号走真实 PrincipalStore、SessionStore、Cookie 和生产 Resolver。"""
    settings = FeishuAuthSettings(
        app_id=f"cli_reporting_roles_{uuid4().hex}",
        admin_group_id="grp_admin",
        user_group_id="grp_user",
        redirect_uri=f"{_ORIGIN}/api/v1/auth/feishu/callback",
        cookie_secure=False,
    )
    routes = FeishuAuthRoutes(
        auth_settings=settings,
        session_factory=runtime.database.new_session,
    )
    resolver = FeishuLoginRequiredResolver(routes)
    application = create_app(
        identity_resolver=resolver,
        content_service=PostgresContentHttpService(runtime),
        product_service=PostgresProductHttpService(runtime),
        reporting_service=PostgresReportingHttpService(runtime),
    )
    install_feishu_auth_routes(application, auth_routes=routes)
    install_route_authorization(application, identity_resolver=resolver)
    clients: dict[str, TestClient] = {}
    principal_ids: dict[str, str] = {}
    with runtime.database.new_session() as session, session.begin():
        for key, role in (("admin", "administrator"), ("a", "user"), ("b", "user")):
            principal = PrincipalStore(session).resolve_or_create(
                connector_id=settings.connector_id,
                provider="feishu",
                provider_subject=f"reporting-role:{key}",
                display_name=f"导出账号{key}",
            )
            token = SessionStore(session).issue(
                principal_id=principal.principal_id,
                display_name=principal.display_name,
                role=role,
            )
            client = TestClient(application, headers={"Origin": _ORIGIN})
            client.cookies.set(SESSION_COOKIE_NAME, token)
            clients[key] = client
            principal_ids[key] = principal.principal_id
    return clients, principal_ids


def _seed_contents(runtime: PlatformRuntime) -> list[str]:
    """使用正式上传/导入 Worker，保留真实来源与内容版本。"""
    client = TestClient(
        create_app(
            identity_resolver=DevelopmentIdentityResolver(),
            import_service=PostgresImportHttpService(runtime),
            content_service=PostgresContentHttpService(runtime),
        )
    )
    batch_id = _seed_import(client, runtime)
    worker = create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id="role-export-import",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    for _ in range(10):
        status = client.get(f"/api/v1/import-batches/{batch_id}").json()["status"]
        if status == "succeeded":
            break
        assert status in {"queued", "running"}
        assert worker.run_once() is True
    else:
        pytest.fail("正式导入没有在测试预算内完成")
    result = client.get("/api/v1/contents")
    assert result.status_code == 200
    ids = [item["id"] for item in result.json()["items"]]
    assert len(ids) == 2
    return ids


def _create_export(
    client: TestClient, content_ids: list[str], columns: list[str]
) -> dict[str, Any]:
    response = client.post(
        "/api/v1/data-exports",
        json={
            "targets": {"scope": "selected", "content_ids": content_ids},
            "columns": columns,
        },
    )
    assert response.status_code == 202, response.text
    return response.json()


def _save_default(
    client: TestClient,
    *,
    revision: int,
    columns: list[str] | None,
) -> Any:
    return client.put(
        _DEFAULT_PATH,
        json={
            "revision": revision,
            "catalog_version": EXPORT_COLUMN_CATALOG_VERSION,
            "columns": columns,
        },
    )


def test_real_sessions_isolate_export_history_detail_download_and_frozen_xlsx(
    role_runtime: PlatformRuntime,
) -> None:
    """用户 A/B 不能交叉读取，管理员可读全部；真实 XLSX 保留冻结顺序。"""
    runtime = role_runtime
    content_ids = _seed_contents(runtime)
    clients, principal_ids = _clients(runtime)
    admin, first, second = clients["admin"], clients["a"], clients["b"]
    saved = _save_default(first, revision=0, columns=["text", "title"])
    assert saved.status_code == 200
    assert saved.json()["revision"] == 1
    export_a = _create_export(first, content_ids, ["title", "author_display_name"])
    export_b = _create_export(second, content_ids, ["text"])

    with runtime.database.new_session() as session:
        owner_rows = (
            session.execute(
                select(
                    reporting_data_exports_table.c.id,
                    reporting_data_exports_table.c.created_by,
                    reporting_data_exports_table.c.request_snapshot,
                )
            )
            .mappings()
            .all()
        )
    owners = {str(row["id"]): row for row in owner_rows}
    assert owners[export_a["export_id"]]["created_by"] == principal_ids["a"]
    assert owners[export_b["export_id"]]["created_by"] == principal_ids["b"]
    assert owners[export_a["export_id"]]["request_snapshot"]["requested_by"] == principal_ids["a"]
    for caller, other in ((first, export_b), (second, export_a)):
        assert caller.get(f"/api/v1/data-exports/{other['export_id']}").status_code == 404
        assert caller.get(f"/api/v1/data-exports/{other['export_id']}/download").status_code == 404
        assert caller.get(f"/api/v1/jobs/{other['job_id']}").status_code == 403
        assert caller.get(f"/api/v1/content-analysis-jobs/{other['job_id']}").status_code == 403
    for export in (export_a, export_b):
        assert admin.get(f"/api/v1/data-exports/{export['export_id']}").status_code == 200

    registry = JobRegistry()
    register_data_export_job(registry, DataExportJobHandler(PostgresDataExportJobExecutor(runtime)))
    worker = create_job_worker(
        runtime=runtime,
        registry=registry,
        worker_id="role-export-xlsx",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    for _ in range(10):
        result = first.get(f"/api/v1/data-exports/{export_a['export_id']}").json()
        if result["job"]["status"] == "succeeded":
            break
        assert worker.run_once() is True
    else:
        pytest.fail("真实导出 Worker 没有完成")

    # Job 受理后更改个人默认，不能反向改变已冻结的文件字段。
    assert _save_default(first, revision=1, columns=["sentiment"]).status_code == 200
    download = first.get(f"/api/v1/data-exports/{export_a['export_id']}/download")
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
    try:
        rows = list(workbook["内容"].iter_rows(values_only=True))
        assert rows[0] == ("标题", "作者")
        assert {row[0] for row in rows[1:]} == {"爱玛 Q7 续航体验", "爱玛门店服务体验"}
        assert {row[1] for row in rows[1:]} == {"用户甲", "用户乙"}
        for sheet in workbook:
            headers = next(sheet.iter_rows(values_only=True))
            assert "Raw/来源定位" not in headers
            assert "来源项ID" not in headers
    finally:
        workbook.close()
    assert first.get(_DEFAULT_PATH).json()["columns"] == ["sentiment"]
    assert second.get(_DEFAULT_PATH).json()["columns"] is None

    # 旧目录文件的评论表可能仍有 Raw，内容列安全也不足以批准下载。
    with runtime.database.new_session() as session, session.begin():
        session.execute(
            update(reporting_data_exports_table)
            .where(
                reporting_data_exports_table.c.id == UUID(export_a["export_id"]),
            )
            .values(column_catalog_version=2)
        )
    assert first.get(f"/api/v1/data-exports/{export_a['export_id']}/download").status_code == 403
    assert admin.get(f"/api/v1/data-exports/{export_a['export_id']}/download").status_code == 200
    with runtime.database.new_session() as session, session.begin():
        session.execute(
            update(reporting_data_exports_table)
            .where(
                reporting_data_exports_table.c.id == UUID(export_a["export_id"]),
            )
            .values(column_catalog_version=EXPORT_COLUMN_CATALOG_VERSION)
        )

    # 别人的大量更新记录不能挤掉当前用户最近的合法导出。
    for _ in range(25):
        _create_export(second, content_ids, ["title"])
    assert [item["id"] for item in first.get("/api/v1/data-exports").json()["items"]] == [
        export_a["export_id"],
    ]
    second_history = second.get("/api/v1/data-exports").json()["items"]
    assert len(second_history) == 20
    assert export_a["export_id"] not in {item["id"] for item in second_history}
    assert len(admin.get("/api/v1/data-exports").json()["items"]) == 20

    # 历史快照即使提及当前用户，没有可靠 created_by 也不能认领。
    with runtime.database.new_session() as session, session.begin():
        session.execute(
            update(reporting_data_exports_table)
            .where(
                reporting_data_exports_table.c.id == UUID(export_a["export_id"]),
            )
            .values(created_by=None)
        )
    assert first.get(f"/api/v1/data-exports/{export_a['export_id']}").status_code == 404
    assert first.get(f"/api/v1/data-exports/{export_a['export_id']}/download").status_code == 404
    assert admin.get(f"/api/v1/data-exports/{export_a['export_id']}/download").status_code == 200


def test_personal_defaults_reset_revision_conflict_and_another_device_restore(
    role_runtime: PlatformRuntime,
) -> None:
    """数据库配置独立于浏览器，重置不删除修订事实。"""
    runtime = role_runtime
    clients, principal_ids = _clients(runtime)
    first, second = clients["a"], clients["b"]
    assert first.get(_DEFAULT_PATH).json() == {
        "revision": 0,
        "columns": None,
        "saved_catalog_version": None,
        "updated_at": None,
    }
    assert _save_default(first, revision=0, columns=["vehicles", "title"]).status_code == 200
    assert _save_default(second, revision=0, columns=["author_display_name"]).status_code == 200
    stale = _save_default(first, revision=0, columns=["text"])
    assert stale.status_code == 409
    assert stale.json()["errors"][0]["code"] == "export_column_default_conflict"
    assert first.get(_DEFAULT_PATH).json()["columns"] == ["vehicles", "title"]
    assert second.get(_DEFAULT_PATH).json()["columns"] == ["author_display_name"]

    another_device = TestClient(first.app, headers={"Origin": _ORIGIN})
    with runtime.database.new_session() as session, session.begin():
        token = SessionStore(session).issue(
            principal_id=principal_ids["a"],
            display_name="另一设备的用户甲",
            role="user",
        )
    another_device.cookies.set(SESSION_COOKIE_NAME, token)
    assert another_device.get(_DEFAULT_PATH).json() == first.get(_DEFAULT_PATH).json()
    reset = _save_default(another_device, revision=1, columns=None)
    assert reset.status_code == 200
    assert reset.json()["revision"] == 2
    assert reset.json()["columns"] is None
    assert _save_default(first, revision=1, columns=["text"]).status_code == 409
    with runtime.database.new_session() as session:
        row = PostgresExportColumnDefaultRepository(session).get(principal_ids["a"])
        assert row is not None
        assert row["revision"] == 2
        assert row["selected_columns"] is None


@pytest.mark.parametrize(
    "saved_columns,expected",
    [
        (("removed-column", "raw_locator"), []),
        (("removed-column", "title", "raw_locator"), ["title"]),
    ],
)
def test_legacy_default_is_filtered_without_adding_new_columns_or_writing_original(
    role_runtime: PlatformRuntime,
    saved_columns: tuple[str, ...],
    expected: list[str],
) -> None:
    """目录升级/降权只投影有效有权字段；全部失效由弹窗安全回退。"""
    runtime = role_runtime
    clients, principal_ids = _clients(runtime)
    with runtime.database.new_session() as session, session.begin():
        repository = PostgresExportColumnDefaultRepository(session)
        before = dict(
            repository.save(
                principal_ids["a"],
                revision=0,
                columns=saved_columns,
                catalog_version=1,
            )
        )

    response = clients["a"].get(_DEFAULT_PATH)

    assert response.status_code == 200
    assert response.json()["columns"] == expected
    assert response.json()["revision"] == 1
    assert response.json()["saved_catalog_version"] == 1
    with runtime.database.new_session() as session:
        original = PostgresExportColumnDefaultRepository(session).get(principal_ids["a"])
        assert original is not None
        assert dict(original) == before


@pytest.mark.parametrize("initial_revision", [0, 1])
def test_personal_default_compare_and_swap_has_one_winner_under_real_concurrency(
    role_runtime: PlatformRuntime,
    initial_revision: int,
) -> None:
    """用两个数据库连接真实并发，分别覆盖初次 INSERT 与后续 UPDATE。"""
    runtime = role_runtime
    principal_id = f"development:default-race:{uuid4().hex}"
    barrier = Barrier(2)
    if initial_revision:
        with runtime.database.new_session() as session, session.begin():
            PostgresExportColumnDefaultRepository(session).save(
                principal_id,
                revision=0,
                columns=("title",),
                catalog_version=EXPORT_COLUMN_CATALOG_VERSION,
            )

    def save(columns: tuple[str, ...]) -> str:
        try:
            with runtime.database.new_session() as session, session.begin():
                barrier.wait(timeout=10)
                PostgresExportColumnDefaultRepository(session).save(
                    principal_id,
                    revision=initial_revision,
                    columns=columns,
                    catalog_version=EXPORT_COLUMN_CATALOG_VERSION,
                )
            return "saved"
        except ExportColumnDefaultConflict:
            return "conflict"

    try:
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(save, (("text", "title"), ("author_display_name",))))
        assert sorted(results) == ["conflict", "saved"]
        with runtime.database.new_session() as session:
            saved = PostgresExportColumnDefaultRepository(session).get(principal_id)
            assert saved is not None
            assert saved["revision"] == initial_revision + 1
            assert saved["selected_columns"] in (["text", "title"], ["author_display_name"])
    finally:
        with runtime.database.new_session() as session, session.begin():
            session.execute(
                delete(reporting_user_export_column_defaults_table).where(
                    reporting_user_export_column_defaults_table.c.principal_id == principal_id,
                )
            )


def test_role_downgrade_filters_defaults_without_rewriting_and_blocks_old_file(
    role_runtime: PlatformRuntime,
) -> None:
    """同一稳定账号降权后，旧配置与历史文件都不能恢复内部定位字段权限。"""
    runtime = role_runtime
    content_ids = _seed_contents(runtime)
    clients, principal_ids = _clients(runtime)
    admin = clients["admin"]
    saved = _save_default(admin, revision=0, columns=["raw_locator", "title", "source_item_id"])
    assert saved.status_code == 200
    export = _create_export(admin, content_ids, ["raw_locator", "title", "source_item_id"])
    with runtime.database.new_session() as session:
        before = dict(PostgresExportColumnDefaultRepository(session).get(principal_ids["admin"]))
    registry = JobRegistry()
    register_data_export_job(registry, DataExportJobHandler(PostgresDataExportJobExecutor(runtime)))
    worker = create_job_worker(
        runtime=runtime,
        registry=registry,
        worker_id="role-downgrade-file",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    for _ in range(10):
        if (
            admin.get(f"/api/v1/data-exports/{export['export_id']}").json()["job"]["status"]
            == "succeeded"
        ):
            break
        assert worker.run_once() is True
    else:
        pytest.fail("管理员导出 Worker 未完成")
    assert admin.get(f"/api/v1/data-exports/{export['export_id']}/download").status_code == 200
    with runtime.database.new_session() as session, session.begin():
        token = SessionStore(session).issue(
            principal_id=principal_ids["admin"],
            display_name="已降权账号",
            role="user",
        )
    downgraded = TestClient(admin.app, headers={"Origin": _ORIGIN})
    downgraded.cookies.set(SESSION_COOKIE_NAME, token)
    effective = downgraded.get(_DEFAULT_PATH)
    assert effective.status_code == 200
    assert effective.json()["columns"] == ["title"]
    assert effective.json()["revision"] == 1
    assert downgraded.get(f"/api/v1/data-exports/{export['export_id']}").status_code == 200
    assert downgraded.get(f"/api/v1/data-exports/{export['export_id']}/download").status_code == 403
    assert _save_default(downgraded, revision=1, columns=["raw_locator"]).status_code == 422
    with runtime.database.new_session() as session:
        after = dict(PostgresExportColumnDefaultRepository(session).get(principal_ids["admin"]))
    assert after == before
    # 当前 Session 角色继续遵循快照；本地撤销才立即失效。
    assert admin.get("/api/v1/principal").json()["role"] == "administrator"
    with runtime.database.new_session() as session, session.begin():
        session.execute(
            update(identity_sessions_table)
            .where(
                identity_sessions_table.c.principal_id == principal_ids["admin"],
            )
            .values(revoked_at=before["updated_at"])
        )
    assert admin.get(_DEFAULT_PATH).status_code == 401
    assert downgraded.get(_DEFAULT_PATH).status_code == 401


def test_rejected_columns_do_not_create_jobs_or_exports_in_real_database(
    role_runtime: PlatformRuntime,
) -> None:
    """状态码之外核对业务事实，拒绝字段不产生任何 Job/Export。"""
    runtime = role_runtime
    clients, _ = _clients(runtime)
    with runtime.database.new_session() as session:
        before = (
            session.scalar(select(func.count()).select_from(jobs_table)),
            session.scalar(select(func.count()).select_from(reporting_data_exports_table)),
        )
    response = clients["a"].post(
        "/api/v1/data-exports",
        json={
            "targets": {"scope": "selected", "content_ids": [str(uuid4())]},
            "columns": ["raw_locator"],
        },
    )
    assert response.status_code == 422
    with runtime.database.new_session() as session:
        after = (
            session.scalar(select(func.count()).select_from(jobs_table)),
            session.scalar(select(func.count()).select_from(reporting_data_exports_table)),
        )
    assert after == before


def test_read_only_content_revision_changes_with_data_and_active_scheme_without_jobs(
    role_runtime: PlatformRuntime,
) -> None:
    """普通用户只取安全修订摘要，数据或方案变化后无需管理任务权限即可刷新。"""
    runtime = role_runtime
    _seed_contents(runtime)
    client = TestClient(
        create_app(
            identity_resolver=DevelopmentIdentityResolver(role="user"),
            workbench_service=PostgresWorkbenchHttpService(runtime),
        )
    )
    statements: list[str] = []

    def capture(
        _connection: object,
        _cursor: object,
        statement: str,
        _parameters: object,
        _context: object,
        _executemany: bool,
    ) -> None:
        statements.append(statement)

    with runtime.database.new_session() as session:
        jobs_before = session.scalar(select(func.count()).select_from(jobs_table))
    event.listen(runtime.database.engine, "before_cursor_execute", capture)
    try:
        first = client.get("/api/v1/content-data-revision")
        second = client.get("/api/v1/content-data-revision")
    finally:
        event.remove(runtime.database.engine, "before_cursor_execute", capture)
    assert first.status_code == second.status_code == 200
    assert first.json() == second.json()
    assert set(first.json()) == {"revision"}
    assert len(first.json()["revision"]) == 64
    assert statements
    assert all(
        statement.lstrip().split()[0].upper()
        not in {
            "INSERT",
            "UPDATE",
            "DELETE",
            "CREATE",
            "ALTER",
            "DROP",
        }
        for statement in statements
    )
    with runtime.database.new_session() as session:
        assert session.scalar(select(func.count()).select_from(jobs_table)) == jobs_before
    with runtime.database.engine.begin() as connection:
        connection.scalar(text("SELECT nextval('workbench_data_revision_seq')"))
    changed_data = client.get("/api/v1/content-data-revision")
    assert changed_data.status_code == 200
    assert changed_data.json()["revision"] != first.json()["revision"]
    assert client.get("/api/v1/analysis/content-runs").status_code == 403

    draft = None
    with runtime.database.new_session() as session:
        active = PostgresAnalysisSchemeRepository(session).get_active_version()
    assert active is not None
    try:
        with runtime.database.new_session() as session, session.begin():
            repository = PostgresAnalysisSchemeRepository(session)
            draft = repository.create_draft(
                name=f"安全修订只读测试-{uuid4().hex}",
                description="只改变统计身份",
                definition=active.definition,
                actor_ref="role-revision-test",
            )
            repository.activate_version(draft.id, expected_version=draft.version)
        changed_scheme = client.get("/api/v1/content-data-revision")
        assert changed_scheme.status_code == 200
        assert changed_scheme.json()["revision"] != changed_data.json()["revision"]
        with runtime.database.new_session() as session:
            assert session.scalar(select(func.count()).select_from(jobs_table)) == jobs_before
    finally:
        if draft is not None:
            with runtime.database.new_session() as session, session.begin():
                repository = PostgresAnalysisSchemeRepository(session)
                repository.activate_version(active.id, expected_version=active.version)
                session.execute(
                    update(analysis_schemes_table)
                    .where(
                        analysis_schemes_table.c.id == draft.scheme_id,
                    )
                    .values(active_version_id=None)
                )
                session.execute(
                    delete(analysis_scheme_versions_table).where(
                        analysis_scheme_versions_table.c.scheme_id == draft.scheme_id,
                    )
                )
                session.execute(
                    delete(analysis_schemes_table).where(
                        analysis_schemes_table.c.id == draft.scheme_id,
                    )
                )
