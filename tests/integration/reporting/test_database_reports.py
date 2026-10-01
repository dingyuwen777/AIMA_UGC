"""真实 PostgreSQL 下报告快照、模型退避、下载与独立发布的工作流。"""

from __future__ import annotations

import json
import os
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from threading import Event
from uuid import UUID, uuid4
from zipfile import ZipFile

import aima_ugc.bootstrap.report_runs_worker as report_worker
import httpx
import pytest
from aima_ugc.adapters.feishu import FeishuApiError
from aima_ugc.adapters.llm import OpenAICompatibleContentLabelingLLM
from aima_ugc.adapters.persistence.postgres.provider_lifecycle import (
    PostgresProviderConfigLifecycleRepository,
)
from aima_ugc.adapters.persistence.postgres.report_runs import PostgresReportRepository
from aima_ugc.bootstrap.administration_http import PostgresAdministrationHttpService
from aima_ugc.bootstrap.content_http import PostgresContentHttpService
from aima_ugc.bootstrap.import_http import PostgresImportHttpService
from aima_ugc.bootstrap.report_runs_http import PostgresReportHttpService
from aima_ugc.bootstrap.worker import (
    create_collection_job_registry,
    create_job_worker,
    create_worker_runtime,
)
from aima_ugc.contracts.administration import ProviderConfigCreateRequest
from aima_ugc.contracts.http import ContentAnalysisSubmitRequest, ContentTargetSelection
from aima_ugc.contracts.reports import ReportSubmitRequest
from aima_ugc.entrypoints.api_main import create_app
from aima_ugc.modules.content.tables import accounts_table, contents_table
from aima_ugc.modules.identity import Principal
from aima_ugc.modules.reporting.report_tables import report_items_table
from aima_ugc.modules.system.tables import provider_configs_table
from aima_ugc.platform.config import load_settings
from aima_ugc.platform.jobs.tables import jobs_table
from aima_ugc.platform.time import beijing_now
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, select, text, update
from sqlalchemy.exc import OperationalError

from tests.integration.content.test_stage8d_voice_plaza_runtime import (
    _analysis_registry,
    _relevant_response,
    _seed_import,
)


@pytest.fixture
def report_system(tmp_path: Path):  # type: ignore[no-untyped-def]
    """禁止对运行库清表；测试只能显式连接 report 专用数据库。"""
    settings = load_settings()
    if settings.db_port != 55437 or settings.db_name != "aima_report_test":
        pytest.skip("只允许在本任务隔离数据库中运行此破坏性 fixture")
    settings = settings.model_copy(
        update={
            "data_dir": tmp_path / "data",
            "log_dir": tmp_path / "logs",
            "feishu_dry_run": False,
        }
    )
    runtime = create_worker_runtime(settings=settings)
    with runtime.database.engine.begin() as connection:
        connection.exec_driver_sql(
            "TRUNCATE TABLE jobs, artifacts, accounts, vehicle_brands, provider_configs "
            "RESTART IDENTITY CASCADE"
        )
    admin = Principal(
        principal_id="report-test",
        display_name="报告测试管理员",
        role="administrator",
        source="development",
    )
    provider = PostgresAdministrationHttpService(runtime).create_provider_config(
        ProviderConfigCreateRequest(
            provider_kind="llm",
            provider="fake",
            display_name="报告模型",
            base_url="https://example.test/v1",
            model="fake-content-labeler-v1",
            api_key="local-test-only",
            is_default=True,
        ),
        principal=admin,
        request_id="report-model",
    )
    service = PostgresReportHttpService(runtime)
    client = TestClient(
        create_app(
            import_service=PostgresImportHttpService(runtime),
            administration_service=PostgresAdministrationHttpService(runtime),
            report_runs_service=service,
        )
    )
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["媒体名称（中文）", "标题", "内文", "作者", "出版日期", "原文链接"])
    for index, day in enumerate((20, 20, 19, 18)):
        sheet.append(
            [
                "小红书",
                f"爱玛报告体验{index}",
                "爱玛续航体验很好，门店服务值得改进",
                f"真实用户{index}",
                f"2026-08-{day:02} 10:00:00",
                f"https://www.xiaohongshu.com/explore/report-{index}",
            ]
        )
    output = BytesIO()
    book.save(output)
    book.close()
    _seed_import(client, runtime, workbook=output.getvalue())
    worker = create_job_worker(
        runtime=runtime,
        registry=create_collection_job_registry(runtime=runtime),
        worker_id="report-integration",
        lease_seconds=120,
        retry_delay_seconds=2,
    )
    assert worker.run_once()
    with runtime.database.new_session() as session:
        ids = tuple(
            session.scalars(
                select(contents_table.c.id).order_by(contents_table.c.external_content_id)
            )
        )
    with runtime.database.engine.begin() as connection:
        account_id = uuid4()
        connection.execute(
            accounts_table.insert().values(
                id=account_id,
                platform="xiaohongshu",
                external_account_id="report-author",
                current_follower_count=50000,
                first_seen_at=beijing_now(),
                last_seen_at=beijing_now(),
                updated_at=beijing_now(),
            )
        )
        connection.execute(update(contents_table).values(author_account_id=account_id))
    PostgresContentHttpService(runtime).create_analysis(
        ContentAnalysisSubmitRequest(
            targets=ContentTargetSelection(scope="selected", content_ids=(ids[0],))
        ),
        request_id="report-analysis",
    )
    analysis_worker = create_job_worker(
        runtime=runtime,
        registry=_analysis_registry(runtime, _relevant_response()),
        worker_id="report-seed-analysis",
        lease_seconds=120,
        retry_delay_seconds=0,
    )
    for _ in range(4):
        if not analysis_worker.run_once():
            break
    from tests.integration.stage3_brand_support import stage3_filter_brand_id

    body = {
        "brand_id": stage3_filter_brand_id(runtime, alias="爱玛"),
        "start_date": "2026-08-20",
        "end_date": "2026-08-20",
    }
    try:
        yield runtime, client, worker, body, provider
    finally:
        client.close()
        runtime.close()


def test_report_history_prevents_deleting_its_frozen_provider(report_system) -> None:  # type: ignore[no-untyped-def]
    """报告使用独立的新 Provider，避免既有 AI 历史守卫掩盖报告引用遗漏。"""
    runtime, client, _, body, _ = report_system
    principal = Principal(
        principal_id="report-test",
        display_name="报告测试管理员",
        role="administrator",
        source="development",
    )
    provider = PostgresAdministrationHttpService(runtime).create_provider_config(
        ProviderConfigCreateRequest(
            provider_kind="llm",
            provider="fake",
            display_name="仅供报告使用的模型",
            base_url="https://example.test/v1",
            model="report-model-v1",
            api_key="local-test-only",
            is_default=True,
        ),
        principal=principal,
        request_id="report-only-model",
    )
    response = client.post("/api/v1/reports", json=body)
    assert response.status_code == 202
    assert response.json()["provider_config_id"] == str(provider.id)
    with runtime.database.new_session() as session, session.begin():
        lifecycle = PostgresProviderConfigLifecycleRepository(session)
        assert lifecycle.archive(provider.id, archived_at=beijing_now()) is not None
        assert lifecycle.delete_blockers(provider.id) == ("报告历史引用了该 Provider",)
        with pytest.raises(RuntimeError, match="报告历史引用"):
            lifecycle.delete_archived(provider.id)


def test_report_freeze_holds_provider_lock_until_history_is_visible(
    report_system, monkeypatch: pytest.MonkeyPatch
) -> None:  # type: ignore[no-untyped-def]
    """报告未提交时，配置删除事务不能越过 System Owner 的 Provider 行锁。"""
    runtime, client, _, body, provider = report_system
    freezing = Event()
    release = Event()
    original_freeze = PostgresReportRepository.freeze

    def pause_freeze(self, report_id, period, targets):  # type: ignore[no-untyped-def]
        if period == "current":
            freezing.set()
            assert release.wait(10), "测试没有释放报告冻结事务"
        return original_freeze(self, report_id, period, targets)

    monkeypatch.setattr(PostgresReportRepository, "freeze", pause_freeze)
    with ThreadPoolExecutor(max_workers=1) as pool:
        creating = pool.submit(client.post, "/api/v1/reports", json=body)
        try:
            assert freezing.wait(10), "报告创建没有进入冻结阶段"
            with runtime.database.new_session() as session:
                with pytest.raises(OperationalError) as error, session.begin():
                    session.execute(text("SET LOCAL lock_timeout = '200ms'"))
                    PostgresProviderConfigLifecycleRepository(session).get_for_update(provider.id)
                assert error.value.orig.sqlstate == "55P03"
        finally:
            release.set()
        response = creating.result(timeout=10)
    assert response.status_code == 202
    with runtime.database.new_session() as session, session.begin():
        lifecycle = PostgresProviderConfigLifecycleRepository(session)
        assert lifecycle.get_for_update(provider.id) is not None
        assert "报告历史引用了该 Provider" in lifecycle.delete_blockers(provider.id)


def test_report_keeps_legal_environment_provider_fallback(
    report_system, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:  # type: ignore[no-untyped-def]
    """环境 Provider 没有数据库行，不应被新增历史锁误判为删除中的配置。"""
    runtime, _, _, body, _ = report_system
    with runtime.database.engine.begin() as connection:
        connection.execute(delete(provider_configs_table))
    environment_secrets = tmp_path / "environment-secrets"
    environment_secrets.mkdir()
    key_path = environment_secrets / "llm_api_key"
    key_path.write_text("local-test-only", encoding="utf-8")
    settings = runtime.settings.model_copy(
        update={
            "llm_base_url": "https://example.test/v1",
            "llm_model": "environment-report-model",
            "external_secret_dir": environment_secrets,
        }
    )
    environment_runtime = create_worker_runtime(settings=settings)
    calls, clients = _fake_model(monkeypatch)
    try:
        response = PostgresReportHttpService(environment_runtime).create(
            ReportSubmitRequest.model_validate(body),
            actor_ref="report-test",
            request_id="environment-report",
        )
        assert response.model == "environment-report-model"
        assert response.generation_job.status == "queued"
        with environment_runtime.database.new_session() as session:
            assert session.scalar(select(provider_configs_table.c.id)) is None
        worker = create_job_worker(
            runtime=environment_runtime,
            registry=create_collection_job_registry(runtime=environment_runtime),
            worker_id="environment-report",
            lease_seconds=120,
            retry_delay_seconds=1,
        )
        assert worker.run_once()
        assert PostgresReportHttpService(environment_runtime).get(response.id).status == "generated"
        assert calls == ["selection", "advice"]
    finally:
        for model_client in clients:
            model_client.close()
        environment_runtime.close()


def _fake_model(
    monkeypatch: pytest.MonkeyPatch, *, overload_advice: bool = False, timeout_advice: bool = False
):  # type: ignore[no-untyped-def]
    """保留生产 HTTP Adapter 与输出校验，只替换外部 HTTP 传输。"""
    calls: list[str] = []

    def respond(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        prompt = body["messages"][0]["content"]
        is_selection = "selected_item_nos" in prompt
        calls.append("selection" if is_selection else "advice")
        if overload_advice and calls.count("advice") == 1 and not is_selection:
            if timeout_advice:
                raise httpx.ReadTimeout("controlled model timeout", request=request)
            return httpx.Response(503, json={"error": {"message": "overloaded"}})
        raw = (
            {"selected_item_nos": [1]}
            if is_selection
            else {
                "items": [
                    {
                        "item_no": 1,
                        "action_advice": "建议主动联系用户核实续航使用环境，并落实门店服务改进。",
                    }
                ]
            }
        )
        return httpx.Response(
            200, json={"choices": [{"message": {"content": json.dumps(raw, ensure_ascii=False)}}]}
        )

    transport = httpx.MockTransport(respond)
    clients = []

    def factory(**kwargs):  # type: ignore[no-untyped-def]
        assert 0 < kwargs["timeout_seconds"] <= 120
        assert kwargs["max_connections"] == 1
        client = httpx.Client(base_url=kwargs["base_url"] + "/", transport=transport)
        clients.append(client)
        return OpenAICompatibleContentLabelingLLM(**kwargs, client=client)

    monkeypatch.setattr(report_worker, "OpenAICompatibleContentLabelingLLM", factory)
    return calls, clients


def _make_due(runtime, job_id: str) -> None:  # type: ignore[no-untyped-def]
    """模拟时间前进，避免测试实际等待退避。"""
    with runtime.database.engine.begin() as connection:
        connection.execute(
            update(jobs_table)
            .where(jobs_table.c.id == UUID(job_id))
            .values(available_at=beijing_now() - timedelta(seconds=1))
        )


@pytest.mark.parametrize("timeout_advice", [False, True])
def test_database_report_snapshot_retry_download_and_publish(
    report_system, monkeypatch: pytest.MonkeyPatch, timeout_advice: bool
) -> None:  # type: ignore[no-untyped-def]
    """从正式 API 到真实 Worker/文件；发布失败与恢复不再调用 LLM。"""
    runtime, client, worker, body, provider = report_system
    calls, clients = _fake_model(monkeypatch, overload_advice=True, timeout_advice=timeout_advice)
    checked = client.post("/api/v1/reports/preflight", json=body)
    assert checked.status_code == 200, checked.text
    assert checked.json()["content_count"] == 2
    assert checked.json()["previous_content_count"] == 1
    assert checked.json()["analyzed_count"] == 1
    with runtime.database.engine.begin() as connection:
        connection.execute(update(accounts_table).values(current_follower_count=50000))
    created = client.post("/api/v1/reports", json=body)
    assert created.status_code == 202, created.text
    assert created.json()["provider_config_id"] == str(provider.id)
    assert created.json()["provider_revision"] == provider.revision
    assert created.json()["analysis_bases"][0]["period"] == "current"
    assert created.json()["analysis_bases"][0]["content_count"] == 1
    assert len(created.json()["prompt_sha256"]) == 64
    assert len(created.json()["selection_prompt_sha256"]) == 64
    report_id = created.json()["id"]
    path = f"/api/v1/reports/{report_id}"
    with runtime.database.new_session() as session:
        frozen = PostgresReportRepository(session).load_records(UUID(report_id), "current")
        basis = list(
            session.scalars(
                select(report_items_table.c.analysis_basis).where(
                    report_items_table.c.report_run_id == UUID(report_id)
                )
            )
        )
    assert len(frozen) == 2
    assert all(record.content.author_follower_count == 50000 for record in frozen)
    assert any(value.get("analysis_scheme_version_id") for value in basis)
    # 创建后的内容指标修改不能进入既有报告。
    with runtime.database.engine.begin() as connection:
        connection.execute(update(contents_table).values(current_like_count=999))
        connection.execute(update(accounts_table).values(current_follower_count=999999))
    assert worker.run_once()
    retried = client.get(path).json()
    assert retried["generation_job"]["status"] == "queued", retried
    assert retried["generation_job"]["error_code"] == ("timeout" if timeout_advice else "http_503")
    assert retried["generation_job"]["available_at"] > retried["generation_job"]["created_at"]
    assert retried["files"] == []
    assert not worker.run_once()
    _make_due(runtime, retried["generation_job"]["id"])
    assert worker.run_once()
    generated = client.get(path).json()
    assert generated["status"] == "generated", generated
    assert calls == ["selection", "advice", "advice"]
    with runtime.database.new_session() as session:
        assert PostgresReportRepository(session).load_records(UUID(report_id), "current") == frozen
    downloaded = {}
    for file in generated["files"]:
        response = client.get(file["download_url"])
        assert response.status_code == 200
        assert len(response.content) == file["byte_size"]
        downloaded[file["filename"]] = response.content
    assert {
        "report.docx",
        "report.md",
        "report-data.xlsx",
        "report-charts.xlsx",
    } <= downloaded.keys()
    with ZipFile(BytesIO(downloaded["report.docx"])) as package:
        assert package.testzip() is None
        assert any(name.startswith("word/charts/chart") for name in package.namelist())
    data = load_workbook(BytesIO(downloaded["report-data.xlsx"]))
    assert data["内容"].max_row == 3
    data.close()
    markdown = downloaded["report.md"].decode()
    assert "| 内容声量 | 2 |" in markdown
    publication_calls = []

    def publish(**kwargs):  # type: ignore[no-untyped-def]
        publication_calls.append(kwargs["checkpoint"].get("test_uploaded"))
        assert kwargs["report"].word_path.is_file()
        assert kwargs["report"].source_excel_path.is_file()
        kwargs["checkpoint"].set("test_uploaded", "remote-file-confirmed")
        if len(publication_calls) == 1:
            raise FeishuApiError("temporary", status_code=503, retriable=True)
        return {
            "native_document_url": "https://example.test/document",
            "editable_chart_sheet_url": "https://example.test/sheet",
            "representative_table_url": "https://example.test/table",
        }

    monkeypatch.setattr(report_worker, "publish_prepared_report_to_feishu", publish)
    publishing = client.post(path + "/publish")
    assert publishing.status_code == 200, publishing.text
    assert worker.run_once()
    failed_publish = client.get(path).json()
    assert failed_publish["status"] == "generated"
    assert failed_publish["publication_job"]["status"] == "queued"
    assert client.get(generated["files"][0]["download_url"]).status_code == 200
    _make_due(runtime, failed_publish["publication_job"]["id"])
    assert worker.run_once()
    published = client.get(path).json()
    assert published["status"] == "published"
    assert published["native_document_url"] == "https://example.test/document"
    assert publication_calls == [None, "remote-file-confirmed"]
    assert calls == ["selection", "advice", "advice"]
    for client_instance in clients:
        client_instance.close()

    # 到期复用正式清理器；历史快照保留，API 明确拒绝下载。
    import aima_ugc.bootstrap.report_runs_http as report_http
    from aima_ugc.bootstrap.artifact_cleanup import run_artifact_cleanup_once

    future = beijing_now() + timedelta(days=61)
    monkeypatch.setattr(report_http, "beijing_now", lambda: future)
    assert client.get(generated["files"][0]["download_url"]).status_code == 410
    cleaned = run_artifact_cleanup_once(runtime, now=future, include_capacity=False)
    assert cleaned.failed == 0
    assert cleaned.deleted >= len(generated["files"])
    expired = client.get(path).json()
    assert expired["status"] == "expired"
    assert expired["files"] == []
    assert client.post(path + "/publish").status_code == 410
    assert client.get("/api/v1/reports").json()["items"][0]["id"] == report_id


def test_empty_scope_catalog_cancel_and_manual_retry(report_system) -> None:  # type: ignore[no-untyped-def]
    """业务输入、原子失败及 queued 取消/恢复由真实 Job 状态证明。"""
    runtime, client, worker, body, _ = report_system
    assert (
        client.post("/api/v1/reports", json={**body, "brand_id": str(uuid4())}).status_code == 422
    )
    assert (
        client.post(
            "/api/v1/reports", json={**body, "vehicle_model_ids": [str(uuid4())]}
        ).status_code
        == 422
    )
    empty = client.post(
        "/api/v1/reports", json={**body, "start_date": "2025-01-01", "end_date": "2025-01-01"}
    )
    assert empty.status_code == 422
    assert client.get("/api/v1/reports").json()["items"] == []
    created = client.post("/api/v1/reports", json=body).json()
    path = f"/api/v1/reports/{created['id']}"
    cancelled = client.post(path + "/cancel").json()
    assert cancelled["status"] == "cancelled"
    assert not worker.run_once()
    restored = client.post(path + "/retry").json()
    assert restored["status"] == "queued"
    assert restored["generation_job"]["id"] != created["generation_job"]["id"]
    assert (
        client.post(path + "/retry").json()["generation_job"]["id"]
        == restored["generation_job"]["id"]
    )
    assert client.post(path + "/publish").status_code == 409


def test_model_overload_exhaustion_and_cancel_fencing(
    report_system,
    monkeypatch: pytest.MonkeyPatch,
) -> None:  # type: ignore[no-untyped-def]
    """连续过载有限次终止；取消后的 Worker 不能提交任何断点。"""
    from aima_ugc.adapters.persistence.postgres.jobs import PostgresJobRepository
    from aima_ugc.modules.reporting.report_jobs import REPORT_GENERATION_JOB
    from aima_ugc.platform.jobs import JobExecutionFence, LeaseLostError

    runtime, client, worker, body, _ = report_system
    clients = []
    calls = []

    def factory(**kwargs):  # type: ignore[no-untyped-def]
        def respond(request: httpx.Request) -> httpx.Response:
            calls.append(request)
            return httpx.Response(429, json={"error": {"message": "overloaded"}})

        adapter_client = httpx.Client(
            base_url=kwargs["base_url"], transport=httpx.MockTransport(respond)
        )
        clients.append(adapter_client)
        return OpenAICompatibleContentLabelingLLM(**kwargs, client=adapter_client)

    monkeypatch.setattr(report_worker, "OpenAICompatibleContentLabelingLLM", factory)
    created = client.post("/api/v1/reports", json=body).json()
    path = f"/api/v1/reports/{created['id']}"
    for attempt in range(created["generation_job"]["max_attempts"]):
        if attempt:
            _make_due(runtime, created["generation_job"]["id"])
        assert worker.run_once()
    failed = client.get(path).json()
    assert failed["status"] == "failed"
    assert failed["generation_job"]["error_code"] == "http_429"
    assert failed["generation_job"]["attempt"] == failed["generation_job"]["max_attempts"]
    assert len(calls) == failed["generation_job"]["max_attempts"]
    assert not worker.run_once()
    assert failed["files"] == []
    # 使用同一生产服务的安全模式配置，拒绝发布必须发生在排队和外部 I/O 前。
    original_settings = runtime.settings
    object.__setattr__(
        runtime, "settings", original_settings.model_copy(update={"feishu_dry_run": True})
    )
    guarded = client.post(path + "/publish")
    assert guarded.status_code == 409
    assert guarded.json()["errors"][0]["code"] == "report_publication_dry_run"
    object.__setattr__(runtime, "settings", original_settings)
    restored = client.post(path + "/retry").json()
    with runtime.database.new_session() as session, session.begin():
        job = PostgresJobRepository(session).claim_next(
            supported_job_types=(REPORT_GENERATION_JOB,), worker_id="stale-test", lease_seconds=120
        )
    assert job is not None and job.lease_token is not None
    assert str(job.id) == restored["generation_job"]["id"]
    assert client.post(path + "/cancel").status_code == 200
    with runtime.database.new_session() as session, session.begin(), pytest.raises(LeaseLostError):
        PostgresReportRepository(session).checkpoint(
            UUID(created["id"]),
            JobExecutionFence(job_id=job.id, lease_token=job.lease_token),
            kind="generation",
            values={"forbidden": "stale"},
        )
    with runtime.database.new_session() as session:
        assert (
            PostgresReportRepository(session).get(UUID(created["id"]))["generation_checkpoint"]
            == {}
        )
    for adapter_client in clients:
        adapter_client.close()


@pytest.mark.skipif(
    os.environ.get("AIMA_REPORT_BROWSER_ACCEPTANCE") != "1", reason="显式启动专用前端55440后运行"
)
def test_browser_to_real_report_api_worker_and_download(
    report_system, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:  # type: ignore[no-untyped-def]
    """真实浏览器、HTTP、PostgreSQL、Worker和下载；模型仅替换外部传输。"""
    from threading import Thread
    from time import monotonic, sleep

    import uvicorn
    from playwright.sync_api import expect, sync_playwright

    runtime, client, worker, body, _ = report_system
    calls, clients = _fake_model(monkeypatch)
    app = client.app
    server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=55439, log_level="warning"))
    serving = Thread(target=server.run, daemon=True)
    serving.start()
    deadline = monotonic() + 15
    while not server.started and monotonic() < deadline:
        sleep(0.05)
    assert server.started
    try:
        with httpx.Client(base_url="http://127.0.0.1:55439") as live_client:
            catalogue = live_client.get("/api/v1/vehicle-brands?offset=0&limit=200")
            assert catalogue.status_code == 200, catalogue.text
            assert any(item["id"] == body["brand_id"] for item in catalogue.json()["items"]), (
                catalogue.text
            )
            vehicles = live_client.get("/api/v1/vehicle-models?offset=0&limit=200")
            assert vehicles.status_code == 200, vehicles.text
        with sync_playwright() as browser_api:
            browser = browser_api.chromium.launch(channel="chrome", headless=True)
            page = browser.new_page(viewport={"width": 1440, "height": 900})
            errors = []
            page.on("pageerror", lambda error: errors.append(error))
            page.goto("http://127.0.0.1:55440/admin/configuration")
            page.get_by_role("button", name="报告生成", exact=True).click()
            expect(page.get_by_label("报告品牌").locator("option")).not_to_have_count(1)
            page.get_by_label("报告品牌").select_option(body["brand_id"])
            page.get_by_label("报告开始日期").fill(body["start_date"])
            page.get_by_label("报告结束日期").fill(body["end_date"])
            page.get_by_role("button", name="检查数据与模型").click()
            expect(page.get_by_role("region", name="报告预检")).to_contain_text(
                "本期 2 条 · 上期 1 条"
            )
            with page.expect_response(
                lambda response: (
                    response.url.endswith("/api/v1/reports") and response.request.method == "POST"
                )
            ) as created_response:
                page.get_by_role("region", name="数据库报告生成").get_by_role(
                    "button", name="生成报告", exact=True
                ).click()
            response = created_response.value
            assert response.status == 202
            assert worker.run_once()
            expect(page.get_by_role("link", name="下载 Word")).to_be_visible(timeout=15000)
            with page.expect_download() as downloaded:
                page.get_by_role("link", name="下载 Word").click()
            output = tmp_path / "browser-report.docx"
            downloaded.value.save_as(output)
            with ZipFile(output) as archive:
                assert archive.testzip() is None
                assert "word/document.xml" in archive.namelist()
            assert calls == ["selection", "advice"]
            assert errors == []
            page.screenshot(path=str(tmp_path / "report-page.png"), full_page=True)
            browser.close()
    finally:
        server.should_exit = True
        serving.join(timeout=10)
        assert not serving.is_alive()
        for adapter_client in clients:
            adapter_client.close()


def test_partial_report_files_are_cleaned_without_removing_completed_report(
    report_system,
    monkeypatch: pytest.MonkeyPatch,
) -> None:  # type: ignore[no-untyped-def]
    """文件集合中途失败产生孤儿，恢复提交后的报告按60天期限保留。"""
    from aima_ugc.bootstrap.artifact_cleanup import run_artifact_cleanup_once
    from aima_ugc.platform.storage import ArtifactService
    from aima_ugc.platform.storage.tables import artifacts_table

    runtime, client, worker, body, _ = report_system
    calls, clients = _fake_model(monkeypatch)
    original_store = ArtifactService.store_stream
    writes = []

    def store(self, **kwargs):  # type: ignore[no-untyped-def]
        if writes:
            raise OSError("controlled disk failure")
        artifact = original_store(self, **kwargs)
        writes.append(artifact.id)
        return artifact

    monkeypatch.setattr(ArtifactService, "store_stream", store)
    created = client.post("/api/v1/reports", json=body).json()
    assert worker.run_once()
    assert len(writes) == 1
    assert client.get(f"/api/v1/reports/{created['id']}").json()["files"] == []
    cleaned = run_artifact_cleanup_once(
        runtime, now=beijing_now() + timedelta(days=2), include_capacity=False
    )
    assert cleaned.failed == 0
    with runtime.database.new_session() as session:
        assert (
            session.scalar(
                select(artifacts_table.c.storage_status).where(artifacts_table.c.id == writes[0])
            )
            == "deleted"
        )
    monkeypatch.setattr(ArtifactService, "store_stream", original_store)
    _make_due(runtime, created["generation_job"]["id"])
    assert worker.run_once()
    generated = client.get(f"/api/v1/reports/{created['id']}").json()
    assert generated["status"] == "generated"
    assert calls == ["selection", "advice"]
    report_files = [UUID(file["artifact_id"]) for file in generated["files"]]
    run_artifact_cleanup_once(
        runtime, now=beijing_now() + timedelta(days=2), include_capacity=False
    )
    with runtime.database.new_session() as session:
        assert set(
            session.scalars(
                select(artifacts_table.c.storage_status).where(
                    artifacts_table.c.id.in_(report_files)
                )
            )
        ) == {"linked"}
    for adapter_client in clients:
        adapter_client.close()
